"""Excel collection templates and additive, row-validated CRM import."""
import zipfile
from backend.crm_region_identity import region_identity, normalize_business_regions
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

CRM_IMPORT_CONTEXT = threading.local()


def crm_trace_export(raw, caller, path, filename, filters=None):
    """Stamp the delivered workbook and audit its exact bytes before release."""
    book = load_workbook(io.BytesIO(raw))
    timestamp = crm_now()
    actor = crm_text(caller.get('username'))
    for sheet in book.worksheets:
        sheet.oddFooter.center.text = f'导出人：{actor}  时间：{timestamp}'
        sheet.evenFooter.center.text = sheet.oddFooter.center.text
        sheet.firstFooter.center.text = sheet.oddFooter.center.text
    book.properties.creator = actor
    book.properties.lastModifiedBy = actor
    output = io.BytesIO()
    book.save(output)
    book.close()
    delivered = output.getvalue()
    crm_audit(actor, 'export', path, 200, filename, {
        'filters': filters or {}, 'bytes': len(delivered),
        'sha256': crm_hashlib.sha256(delivered).hexdigest(), 'exported_at': timestamp,
    })
    return delivered

CRM_EXCEL_SHEETS = {
    '单位': ['单位名称*', '单位类型*', '城市', '线网规模', '招标偏好', '历史合作', '历史中标率'],
    '联系人': ['单位名称*', '姓名*', '职务', '决策角色', '权力半径', '影响力', '立场', '职业轨迹', '轮岗窗口'],
    '关系': ['单位名称*', '姓名*', '我方对接人', '关系来源', '引荐路径', '单点风险'],
    '动作': ['单位名称*', '姓名*', '我方对接人', '动作类型*', '发生日期*', '关联决策节点', '情报产出', '投入成本', '有效动作', '客户主动', '项目编号', '下一步事项', '负责人', '跟进日期'],
    '项目决策链': ['项目编号*', '决策节点*', '单位名称*', '姓名*', '我方对接人'],
}


def _crm_sheet_dimensions(sheet):
    """Return safe worksheet dimensions for files created by Excel/WPS alike.

    A few third-party writers omit the worksheet dimension metadata.  openpyxl
    then exposes ``max_row`` or ``max_column`` as ``None`` on a read-only
    worksheet, so comparing it directly with the import limits would abort the
    whole upload with ``NoneType > int`` before row validation can run.
    """
    try:
        rows = int(getattr(sheet, 'max_row', 0) or 0)
        columns = int(getattr(sheet, 'max_column', 0) or 0)
    except (TypeError, ValueError):
        return 0, 0
    return max(0, rows), max(0, columns)


def _crm_normalize_legacy_business_values(values, headers):
    """Keep old collection files importable without reintroducing retired regions."""
    if values.get('区域') in ('南区', '华南区', '广东'):
        # Historical collection files used 南区 for the Guangdong team. The live
        # product now uses 广东区, so migrate only during import (new UI entries
        # still reject retired region names).
        values['区域'] = '广东区'
        # A few early exported sheets carried a serial number in the final
        # “我方姓名” column. Never create a staff profile named "1" from it.
    if headers in (CRM_BUSINESS_CONTACT_LEGACY_HEADERS, CRM_BUSINESS_OLD_HEADERS, CRM_BUSINESS_LEGACY_HEADERS) and str(values.get('我方姓名', '')).strip().isdigit():
        values['我方姓名'] = ''
    return values


# 当前区域汇总表以公司全称为客户组织根节点，不再收集“子公司全称”。
# 保留旧 22 列表头，仅为让已下载或历史归档文件继续可导入。
CRM_BUSINESS_EXTERNAL_LEGACY_HEADERS = ['区域', '公司全称', '单位*', '部门', '职务', '姓名*', '手机', '办公电话', '邮箱', '微信 / 其他联系', '重要程度', '支持程度', '跟进轨迹', '上次互动时间', '本周互动计划', '计划跟进日期', '单位类型', '省份', '城市', '维系周期（天）', '我方姓名']
CRM_BUSINESS_EXTERNAL_LEGACY_SUBSIDIARY_HEADERS = ['区域', '公司全称', '子公司全称', '单位*', '部门', '职务', '姓名*', '手机', '办公电话', '邮箱', '微信 / 其他联系', '重要程度', '支持程度', '跟进轨迹', '上次互动时间', '本周互动计划', '计划跟进日期', '单位类型', '省份', '城市', '维系周期（天）', '我方姓名']
# 新模板将我方人员的组织信息拆成独立列；姓名仍是唯一必填项。
# 保留两种旧表头，以兼容已下载及历史归档的收集表。
CRM_OUR_STAFF_HEADERS = ['我方公司', '我方部门', '我方岗位']
CRM_BUSINESS_EXTERNAL_HEADERS = CRM_BUSINESS_EXTERNAL_LEGACY_HEADERS + CRM_OUR_STAFF_HEADERS
CRM_BUSINESS_EXTERNAL_SUBSIDIARY_HEADERS = CRM_BUSINESS_EXTERNAL_LEGACY_SUBSIDIARY_HEADERS + CRM_OUR_STAFF_HEADERS
CRM_BUSINESS_EXTERNAL_LAYOUTS = (
    CRM_BUSINESS_EXTERNAL_HEADERS,
    CRM_BUSINESS_EXTERNAL_SUBSIDIARY_HEADERS,
    CRM_BUSINESS_EXTERNAL_LEGACY_HEADERS,
    CRM_BUSINESS_EXTERNAL_LEGACY_SUBSIDIARY_HEADERS,
)
CRM_IMPORT_COMPLETION_REQUIRED_FIELDS = ('公司全称', '姓名*', '单位类型')


# ---------------------------------------------------------------------------
# 经营分析：按“2026联指经营数据·标准导入模板”导入多子表，整份数据集为联指级共享。
# ---------------------------------------------------------------------------
def crm_business_analysis_path():
    # 旧版单文件（迁移前）；新导入按日期归档到 crm_business_analysis_dir()。
    return Path(USER_DATA_DIR) / '_crm_business_analysis.json'


def crm_business_analysis_dir():
    return Path(USER_DATA_DIR) / '_crm_business_analysis'


def crm_business_analysis_template_path():
    return Path(BASE_DIR) / 'backend' / 'assets' / 'business_analysis_template.xlsx'


CRM_TASK_ORDER_SHEETS = {
    '订货任务令': ['确认主要责任人', '区域', '客户经理', '子项目编码', '机会点名称', '任务令作战部署的目标/动作', '赢率', '可参与份额-总额(万元)', '合同预计签订时间', '毛利率', '管理订货(万元)', '毛利(万元)', '月份', '任务令标签', '是否已签合同', '合同签订金额(万元)', '合同签订时间'],
    '回款任务令': ['商务专员', '主要责任人', '区域', '销售子项目编码', '销售合同编号', '销售子项目名称', '客户单位', '回款性质', '回款条件', '计划回款金额(元)', '任务令标签', '月份', '是否已回款', '回款时间'],
    '财务收入任务令': ['区域', '项目经理', '项目类型', '机会点/子项目编码', '机会点/子项目名称', '不含税金额(元)', '收入金额(含税/元)', '任务令标签', '月份', '是否已确认收入', '确认收入时间'],
}

CRM_TASK_PROGRESS_KEYS = {'任务令总表说明', '任务令公司进度', '任务令区域进度', '任务令责任人进度', '任务令进度总表原表'}


def _crm_task_progress_sheets(book):
    source_name = next((name for name in book.sheetnames if name.endswith('任务进度总表')), '')
    if not source_name:
        return []
    values = [list(row) for row in book[source_name].iter_rows(values_only=True)]
    report_title = crm_text(values[0][0] if values and values[0] else '', 200)
    cutoff = (re.search(r'数据截止[：:]\s*([^|｜]+)', report_title) or [None, ''])[1].strip()
    unit = next((crm_text(value, 30).split('：', 1)[-1] for row in values[:5] for value in row if isinstance(value, str) and value.startswith('单位：')), '万元')
    company_rows, region_rows, owner_rows = [], [], []
    section = ''
    for index, row in enumerate(values):
        title = crm_text(row[0] if row else '', 120)
        if re.match(r'^[一二三四五六七八九十]+、', title):
            section = title
            continue
        if index == 0 or not title or title in {'区域', '指标', '指标\n（不含公司其他）'}:
            continue
        if section.startswith('一、公司整体'):
            if len(row) > 10 and isinstance(row[2], (int, float)):
                company_rows.append([title, *[row[pos] if pos < len(row) else None for pos in range(2, 11)]])
            continue
        if not re.match(r'^[二三四五六七八九十]+、', section) or len(row) <= 10 or not isinstance(row[2], (int, float)):
            continue
        metric_match = re.search(r'\d+月(.+?)(?:冲刺)?任务', section)
        metric = crm_text(metric_match.group(1) if metric_match else '', 30).strip()
        record = [
            '区域' if section.startswith('二、区域') else section.split('、', 1)[-1].split(' - ', 1)[0].strip(),
            metric, title, crm_text(row[1], 80),
            *[row[pos] if pos < len(row) else None for pos in range(2, 11)],
        ]
        (region_rows if section.startswith('二、区域') else owner_rows).append(record)
    source = f'{source_name} · {report_title}' if report_title else source_name
    result = [{'key': '任务令总表说明', 'title': '任务令总表说明', 'source': source,
               'headers': ['报表标题', '数据截止', '单位'], 'rows': [[report_title, cutoff, unit]]}]
    if company_rows:
        result.append({'key': '任务令公司进度', 'title': '任务令公司进度', 'source': source,
                       'headers': ['指标', '年度目标', '1-10月完成', '年度完成率', '年度缺口', '本月任务令', '本月完成', '本月完成率', '本月排名', '本月缺口'], 'rows': company_rows})
    progress_headers = ['层级', '指标', '区域', '责任人', '任务令', '第一周完成', '第二周完成', '第三周完成', '第四周完成', '合计完成', '完成率', '排名', '缺口']
    if region_rows:
        result.append({'key': '任务令区域进度', 'title': '任务令区域进度', 'source': source, 'headers': progress_headers, 'rows': region_rows})
    if owner_rows:
        result.append({'key': '任务令责任人进度', 'title': '任务令责任人进度', 'source': source, 'headers': ['角色', *progress_headers[1:]], 'rows': owner_rows})
    width = max((len(row) for row in values), default=0)
    raw_rows = []
    for row_number, row in enumerate(values, 1):
        if not any(value not in (None, '') for value in row):
            continue
        cleaned = [value if isinstance(value, (int, float, bool)) or value is None else str(value) for value in row[:width]]
        raw_rows.append([row_number, *cleaned, *([None] * (width - len(cleaned)))])
    result.append({'key': '任务令进度总表原表', 'title': '任务令进度总表原表', 'source': source,
                   'headers': ['行号', *[f'第{index}列' for index in range(1, width + 1)]], 'rows': raw_rows})
    return result


def crm_task_order_template():
    book = Workbook()
    guide = book.active
    guide.title = '填写说明'
    lines = [
        ['联指经营助手 · 任务令数据导入模板'],
        ['只填写“订货任务令、回款任务令、财务收入任务令”三个工作表，不修改表头和工作表名。'],
        ['一行一项任务；金额单位严格按表头填写；月份填写1-12；任务令标签填写“保底、冲刺或推进”。'],
        ['重复导入会生成新的经营分析历史版本，并替换上一版本中的三张任务令表，其他经营数据保持不变。'],
    ]
    for line in lines:
        guide.append(line)
    guide.column_dimensions['A'].width = 105
    for row in guide:
        row[0].alignment = Alignment(wrap_text=True, vertical='center')
        guide.row_dimensions[row[0].row].height = 34
    guide['A1'].font = Font(size=16, bold=True, color='15499A')
    for name, headers in CRM_TASK_ORDER_SHEETS.items():
        sheet = book.create_sheet(name)
        sheet.append(headers)
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='174B91')
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            sheet.column_dimensions[cell.column_letter].width = 18
        sheet.column_dimensions['E'].width = 38
        sheet.column_dimensions['F'].width = 42
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = f'A1:{get_column_letter(len(headers))}5001'
        sheet.row_dimensions[1].height = 36
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def crm_task_order_import(payload, username):
    caller = next((u for u in read_config().get('users', []) if u.get('username') == username), {'username': username})
    if not (crm_capabilities(caller).get('write') and crm_capabilities(caller).get('data_scope') == 'all'):
        return {'ok': False, 'error': '需要管理级全公司数据范围及编辑权限才能导入任务令数据'}
    encoded = payload.get('data', '') if isinstance(payload, dict) else ''
    try:
        raw = base64.b64decode(encoded, validate=True)
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        return {'ok': False, 'error': '无法读取任务令 Excel 文件'}
    imported = []
    for name, expected in CRM_TASK_ORDER_SHEETS.items():
        source_name = next((title for title in book.sheetnames if title == name or title.endswith('-' + name)), '')
        if not source_name:
            continue
        sheet = book[source_name]
        values = [list(row) for row in sheet.iter_rows(values_only=True)]
        if not values:
            continue
        headers = [str(value or '').strip() for value in values[0]]
        while headers and not headers[-1]:
            headers.pop()
        if len(headers) < 5:
            book.close()
            return {'ok': False, 'error': f'{name}表头不完整，请使用任务令标准模板'}
        rows = []
        for source in values[1:]:
            row = []
            for index in range(len(headers)):
                value = source[index] if index < len(source) else None
                row.append(value if isinstance(value, (int, float, bool)) else str(value or ''))
            if any(value not in ('', None) for value in row):
                rows.append(row)
        imported.append({'key': name, 'title': name, 'source': '任务令专用导入模板', 'headers': headers, 'rows': rows})
    imported.extend(_crm_task_progress_sheets(book))
    book.close()
    if not imported:
        return {'ok': False, 'error': '未找到任务令工作表，请使用任务令标准模板'}
    current = crm_business_analysis(caller).get('data') or {}
    task_names = set(CRM_TASK_ORDER_SHEETS) | CRM_TASK_PROGRESS_KEYS
    sheets = [sheet for sheet in current.get('sheets', []) if sheet.get('key') not in task_names] + imported
    try:
        dashboard = crm_compute_dashboard_data(sheets)
    except Exception:
        dashboard = current.get('dashboard')
    result = {'imported_at': crm_now(), 'imported_by': username, 'filename': crm_text(payload.get('filename'), 200), 'sheets': sheets, 'dashboard': dashboard}
    base = ''.join(char for char in result['imported_at'] if char.isdigit()) or 'v'
    version = base
    suffix = 1
    if crmdb_enabled():
        taken = set(crmdb_blob_keys('business_analysis:'))
        while f'business_analysis:{version}' in taken:
            version = f'{base}-{suffix}'; suffix += 1
        crmdb_blob_put(f'business_analysis:{version}', result)
    else:
        directory = crm_business_analysis_dir(); directory.mkdir(parents=True, exist_ok=True)
        while (directory / f'{version}.json').exists():
            version = f'{base}-{suffix}'; suffix += 1
        (directory / f'{version}.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    return {'ok': True, 'sheets': len(imported), 'rows': sum(len(sheet['rows']) for sheet in imported), 'version': version, 'imported_at': result['imported_at']}


def crm_business_dashboard_html(user, version=None):
    """Render the 10-page executive dashboard HTML with the chosen snapshot's data inlined."""
    if not crm_capabilities(user).get('read'):
        return None
    res = crm_business_analysis(user, version)
    data = (res or {}).get('data')
    dashboard = (data or {}).get('dashboard')
    tpl_path = Path(BASE_DIR) / 'backend' / 'assets' / 'dashboard_template.html'
    try:
        tpl = tpl_path.read_text(encoding='utf-8')
    except Exception:
        return None
    payload = json.dumps(dashboard or {}, ensure_ascii=False)
    return tpl.replace('__DASHBOARD_DATA__', payload)


def _crm_business_versions():
    """List every archived business-analysis snapshot, newest first (legacy file included)."""
    if crmdb_enabled():
        versions = [v for v in crmdb_business_versions()]
        versions.sort(key=lambda v: (v.get('imported_at', ''), v.get('version', '')), reverse=True)
        return versions
    versions = []
    directory = crm_business_analysis_dir()
    if directory.exists():
        for f in directory.glob('*.json'):
            try:
                meta = json.loads(f.read_text(encoding='utf-8'))
            except Exception:
                continue
            versions.append({'version': f.stem, 'imported_at': meta.get('imported_at', ''),
                             'imported_by': meta.get('imported_by', ''), 'filename': meta.get('filename', ''),
                             'sheets': len(meta.get('sheets', []))})
    legacy = crm_business_analysis_path()
    if legacy.exists():
        try:
            meta = json.loads(legacy.read_text(encoding='utf-8'))
            versions.append({'version': 'legacy', 'imported_at': meta.get('imported_at', ''),
                             'imported_by': meta.get('imported_by', ''), 'filename': meta.get('filename', ''),
                             'sheets': len(meta.get('sheets', []))})
        except Exception:
            pass
    versions.sort(key=lambda v: (v.get('imported_at', ''), v.get('version', '')), reverse=True)
    return versions


def crm_business_scope_data(data, user, preview=None):
    caps = crm_capabilities(user)
    if preview is not None:
        if user.get('role') != 'superadmin':
            raise ValueError('仅超级管理员可以切换经营视图')
        level, target = preview
        caps = dict(caps, data_scope={'company': 'all', 'region': 'regions', 'personal': 'self'}[level],
                    regions=[target] if level == 'region' else [])
        if level == 'personal':
            user = next(item for item in read_config().get('users', []) if item.get('username') == target)
    level = 'company' if caps.get('data_scope') == 'all' else ('region' if caps.get('data_scope') == 'regions' else 'personal')
    regions = caps.get('regions', []) if level == 'region' else []
    scope = {'level': level, 'label': {'company': '管理级 · 全公司', 'region': '区域级 · ' + '、'.join(regions), 'personal': '个人级 · 本人负责业务'}[level],
             'regions': regions, 'can_import': level == 'company' and bool(caps.get('write'))}
    if not data:
        return data, scope
    result = dict(data)
    result['scope'] = scope
    if level == 'company':
        return result, scope
    def region_key(value):
        return region_identity(value)[0]
    allowed_regions = {region_key(value) for value in regions}
    accounts = read_config().get('users', [])
    account = next((item for item in accounts if item.get('username') == user.get('username')), user)
    identities = {str(account.get(key) or '').strip() for key in ('username', 'employee_id')} - {''}
    name = str(account.get('name') or '').strip()
    if name and sum(str(item.get('name') or '').strip() == name for item in accounts) == 1:
        identities.add(name)
    region_headers = ('区域', '项目归属', '*项目归属', '区域归属', '区域组织', '26年区域')
    owner_headers = ('客户经理', '*客户经理', '责任人', '确认主要责任人', '主要责任人', '项目经理', '26年客户经理')
    permitted = {
        '02_作战平台经营结果', '06_费用明细', '07_费用率', '11_现金流构成',
        '16_机会点清单', '17_滚动冲刺机会点', '18_年初全预算承接', '19_已完成订单明细',
        '任务令区域进度', '任务令责任人进度', '订货任务令', '回款任务令', '财务收入任务令',
        '实际回款明细', '本次区域经营汇总',
        '区域经营_年度指标', '区域经营_月度实绩', '区域经营_回款计划', '区域经营_应收款结构',
        '区域经营_期间指标', '区域经营_业务明细', '区域经营_历史回款计划',
        '区域经营_原表期间', '区域经营_原表业务', '区域经营_月度拆解', '区域经营_应收原表',
    }
    filtered = []
    for sheet in data.get('sheets', []):
        key = sheet.get('key', '')
        headers = sheet.get('headers', [])
        if level == 'region' and key in ('03_作战平台目标达成', '09_订单质量'):
            indices = [index for index, header in enumerate(headers) if header == '指标' or region_key(header) in allowed_regions]
            if len(indices) > 1:
                filtered.append({'key': key, 'title': key, 'source': scope['label'], 'headers': [headers[index] for index in indices],
                                 'rows': [[row[index] if index < len(row) else None for index in indices] for row in sheet.get('rows', [])]})
            continue
        if key not in permitted:
            continue
        if key in ('订货任务令', '回款任务令', '财务收入任务令', '任务令责任人进度'):
            personal_headers = ('责任人', '确认主要责任人', '主要责任人', '项目经理', '客户经理', '*客户经理')
        else:
            personal_headers = owner_headers
        candidates = region_headers if level == 'region' else personal_headers
        header = next((field for field in candidates if field in headers), None)
        if not header:
            continue
        index = headers.index(header)
        selected = [list(row) for row in sheet.get('rows', []) if index < len(row) and
                    (region_key(row[index]) in allowed_regions if level == 'region' else str(row[index] or '').strip() in identities)]
        if selected:
            filtered.append({'key': key, 'title': key, 'source': scope['label'], 'headers': list(headers), 'rows': selected})
    result['sheets'] = filtered
    result['dashboard'] = crm_compute_dashboard_data(filtered)
    result['dashboard']['_meta']['口径'] = scope['label'] + '；仅展示明确归属的数据，不包含公司汇总及无法归属的记录。'
    result.pop('filename', None)
    result.pop('imported_by', None)
    return result, scope


def crm_business_analysis(user, version=None, view=None, region=None, manager=None):
    """Read one archived snapshot (default: newest) plus the list of all versions."""
    if not crm_capabilities(user).get('read'):
        return {'ok': False, 'error': '无查看权限'}
    preview = None
    view_options = None
    if any((view, region, manager)) and user.get('role') != 'superadmin':
        return {'ok': False, 'error': '仅超级管理员可以切换经营视图'}
    if user.get('role') == 'superadmin':
        accounts = read_config().get('users', [])
        view_options = {'regions': list(CRM_BUSINESS_REGIONS),
                        'managers': [{'username': item['username'], 'name': item.get('name') or item['username']}
                                     for item in accounts if item.get('username')]}
        view = view or 'company'
        if view not in ('company', 'region', 'personal'):
            return {'ok': False, 'error': '无效的经营视图'}
        if view == 'region' and region not in view_options['regions']:
            return {'ok': False, 'error': '请选择有效区域'}
        if view == 'personal' and manager not in {item['username'] for item in view_options['managers']}:
            return {'ok': False, 'error': '请选择有效客户经理'}
        preview = (view, region if view == 'region' else manager)
    versions = _crm_business_versions()
    if not versions:
        _, scope = crm_business_scope_data(None, user, preview)
        return {'ok': True, 'data': None, 'versions': [], 'scope': scope, 'view_options': view_options}
    chosen = next((v for v in versions if v['version'] == version), None) if version else None
    if not chosen:
        chosen = versions[0]
    try:
        if crmdb_enabled():
            data = crmdb_blob_get(f"business_analysis:{chosen['version']}")
            if not isinstance(data, dict):
                raise ValueError('快照缺失')
        else:
            path = crm_business_analysis_path() if chosen['version'] == 'legacy' else crm_business_analysis_dir() / f"{chosen['version']}.json"
            data = json.loads(path.read_text(encoding='utf-8'))
        data['version'] = chosen['version']
        # 早期版本没有 dashboard，用已存子表即时补算并回写，历史版本也能直接出图。
        if not data.get('dashboard') and data.get('sheets'):
            try:
                data['dashboard'] = crm_compute_dashboard_data(data['sheets'])
                if chosen['version'] != 'legacy':
                    archived = {k: v for k, v in data.items() if k != 'version'}
                    if crmdb_enabled():
                        crmdb_blob_put(f"business_analysis:{chosen['version']}", archived)
                    else:
                        path.write_text(json.dumps(archived, ensure_ascii=False, indent=2), encoding='utf-8')
            except Exception:
                pass
    except Exception:
        data = None
    if data and data.get('sheets'):
        data = {**data, 'sheets': normalize_business_regions(data['sheets'])}
        previous_meta = (data.get('dashboard') or {}).get('_meta')
        data['dashboard'] = crm_compute_dashboard_data(data['sheets'])
        if previous_meta:
            data['dashboard']['_meta'] = previous_meta
    data, scope = crm_business_scope_data(data, user, preview)
    if scope['level'] != 'company':
        versions = [{key: item.get(key) for key in ('version', 'imported_at')} for item in versions]
    return {'ok': True, 'data': data, 'versions': versions, 'scope': scope, 'view_options': view_options}


CRM_ANALYSIS_BLOCKS = {
    'overall': ('公司整体', ['01_', '10_', '19_', '公司月度订货', '公司月度回款', '公司月度现金流', '实际回款明细']),
    'taskOrder': ('任务令', ['任务令', '订货任务令', '回款任务令', '财务收入任务令']),
    'structure': ('经营结构', ['06_']),
    'region': ('区域经营', ['区域经营_', '02_', '03_', '05_', '06_', '07_', '09_', '11_', '16_', '本次区域经营汇总']),
    'industry': ('产业经营', ['04_', '08_', '13_', '14_', '15_', '16_']),
    'foundation': ('经营基础', ['11_', '12_', '13_', '14_']),
    'projects': ('项目监测', ['16_', '17_', '18_', '19_']),
}


def crm_analysis_template_export(user, version=None, blocks=None):
    if not crm_capabilities(user).get('read') or crm_capabilities(user).get('data_scope') != 'all':
        raise ValueError('仅管理级全公司数据范围可下载当前数据模板')
    selected = blocks or list(CRM_ANALYSIS_BLOCKS)
    if any(key not in CRM_ANALYSIS_BLOCKS for key in selected):
        raise ValueError('无效的经营板块')
    result = crm_business_analysis(user, version=version)
    data = result.get('data')
    if not result.get('ok') or not data:
        raise ValueError('当前没有可导出的经营数据')
    if version and data.get('version') != version:
        raise ValueError('所选版本不存在，请刷新后重试')
    sheets = data.get('sheets', [])

    def workbook(items, label):
        book = Workbook()
        meta = book.active
        meta.title = '_板块导入配置'
        meta.append(['business-block-v1', data.get('version', ''), label])
        meta.append(['按同名表更新，未包含的表保持不变；共享表会同时更新多个页面。'])
        meta.append(['保留表名和第3行表头，第4行起填写数据；金额与比率单位沿用原表。'])
        meta.append(['填写规则', '说明'])
        meta.append(['模板版本', '2026-10 页面匹配版；携带所选版本实际数据，不生成示例业绩。'])
        meta.append(['更新范围', '同名工作表整表替换，不是追加行；删除数据行会删除该表对应记录，未上传表保持不变。'])
        meta.append(['缺失与零值', '未提供的数据留空；真实为零才填0。请填写数值，不使用公式，避免未缓存公式被读取为空。'])
        meta.append(['区域范围', '东区、西区、中区、广东区；东区包含浙北、浙南、江苏办事处。不要将办事处与东区汇总重复相加。'])
        meta.append(['期间指标', '按原表年度、期间填写：月份、季度、全年分别维护；目标与实际必须同期间。应收存量不按月份分摊。'])
        meta.append(['订货与管理订货', '管理订货为折算后口径，不与订货相加；保留原始折算结果，缺失时不要直接复制订货。'])
        meta.append(['利润与费用', '订货毛利、收入毛利及营业利润分别填写；公司口径与利润核算角度不相加；负费用保留负号。'])
        meta.append(['回款与计划', '实际回款与预计回款分开；明细金额元与汇总万元按各表表头或来源说明填写，不混用。'])
        meta.append(['月度与TOP10', '月度汇总支撑趋势；项目明细支撑TOP10，需完整项目名称、日期和金额，不能仅填汇总。'])
        meta.append(['展示序号', '页面01、02等是展示排序，无需写入数据；图表、完成率等展示以页面实际计算口径为准。'])
        meta.append(['共享来源', '同一表可能影响多个页面；各板块下载文件内的共享表只维护一份最新内容，避免旧副本覆盖。'])
        meta.append([])
        meta.append(['工作表', '影响页面', '数据行数', '来源/单位说明', '填写字段（保持顺序）'])
        for item in items:
            related = [name for name, prefixes in CRM_ANALYSIS_BLOCKS.values() if any(item['key'].startswith(prefix) for prefix in prefixes)]
            meta.append([item['key'], '、'.join(related) or '完整数据保留表', len(item.get('rows', [])), item.get('source', ''), ' / '.join(item.get('headers', []))])
        meta.freeze_panes = 'A4'
        for row in meta.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        for column, width in zip('ABCDE', [30, 90, 14, 60, 100]):
            meta.column_dimensions[column].width = width
        for cell in meta[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='245A81')
        for item in items:
            sheet = book.create_sheet(item['key'])
            sheet.append([item.get('title') or item['key']])
            sheet.append([item.get('source') or '来源：当前经营数据版本'])
            sheet.append(item.get('headers', []))
            for values in item.get('rows', []):
                sheet.append(values)
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
            sheet.freeze_panes = 'A4'
            sheet.auto_filter.ref = f'A3:{get_column_letter(max(1, sheet.max_column))}{max(3, sheet.max_row)}'
            for cell in sheet[3]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='245A81')
            for column in range(1, sheet.max_column + 1):
                header = str(sheet.cell(3, column).value or '')
                sheet.column_dimensions[get_column_letter(column)].width = 42 if any(word in header for word in ('说明', '名称', '备注', '口径')) else 24 if any(word in header for word in ('日期', '时间', '编号')) else 18
                for row in range(4, sheet.max_row + 1):
                    cell = sheet.cell(row, column)
                    if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                        if any(word in header for word in ('金额', '目标', '实际', '毛利', '费用', '回款', '收入', '差额', '应收', '订货')) and not any(word in header for word in ('率', '比', '年', '月', '数', '排名')):
                            cell.number_format = '#,##0.00;[Red]-#,##0.00'
            sheet.sheet_view.showGridLines = False
        output = io.BytesIO()
        book.save(output)
        book.close()
        return output.getvalue()

    if len(selected) == 1:
        label, prefixes = CRM_ANALYSIS_BLOCKS[selected[0]]
        if selected[0] == 'region' and any(sheet['key'] == '区域经营_年度指标' for sheet in sheets):
            prefixes = ['区域经营_']
        items = [sheet for sheet in sheets if any(sheet['key'].startswith(prefix) for prefix in prefixes)]
        if not items:
            raise ValueError('当前版本没有该板块的数据表')
        return workbook(items, label), f'经营分析_{label}_当前数据.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
        for key in dict.fromkeys(selected):
            label, prefixes = CRM_ANALYSIS_BLOCKS[key]
            if key == 'region' and any(sheet['key'] == '区域经营_年度指标' for sheet in sheets):
                prefixes = ['区域经营_']
            items = [sheet for sheet in sheets if any(sheet['key'].startswith(prefix) for prefix in prefixes)]
            if items:
                archive.writestr(f'{label}_当前数据.xlsx', workbook(items, label))
        if set(selected) == set(CRM_ANALYSIS_BLOCKS):
            archive.writestr('全部板块_完整当前数据.xlsx', workbook(sheets, '全部板块'))
        archive.writestr('使用说明.txt', '先阅读各文件的_板块导入配置：包含填写规则、页面来源对应和字段清单。\n按板块修改后，在经营分析点击重新导入逐个上传XLSX，不直接上传ZIP。\n全部板块_完整当前数据.xlsx可一次更新全部来源表。\n同名表整表替换，不是追加；未包含的表保留。\n共享表存在于多个板块模板中，以最后导入内容为准，请勿反复导入旧副本。\n模板包含所选版本全部实际行，不是随机样例；缺失来源不能凭空补数。\n空值不等于0；金额单位沿用原表；不要使用公式。')
    return output.getvalue(), '经营分析_分板块当前数据模板.zip', 'application/zip'


def crm_business_import(payload, username):
    """Parse every data sheet of the business-analysis workbook into title/headers/rows.

    The template is intentionally generic (20 子表): row 0 is the sheet title, an optional
    source note follows, then a header row and data rows. We detect the header row rather
    than hard-coding offsets, so future template tweaks keep importing cleanly.
    """
    caller = next((u for u in read_config().get('users', []) if u.get('username') == username), {'username': username})
    if not (crm_capabilities(caller).get('write') and crm_capabilities(caller).get('data_scope') == 'all'):
        return {'ok': False, 'error': '需要管理级全公司数据范围及编辑权限才能导入经营数据'}
    encoded = payload.get('data', '') if isinstance(payload, dict) else ''
    if not isinstance(encoded, str) or not encoded:
        return {'ok': False, 'error': '请选择经营数据 Excel 文件'}
    try:
        raw = base64.b64decode(encoded, validate=True)
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        return {'ok': False, 'error': '无法读取经营数据 Excel 文件'}

    def cell(v):
        if v is None:
            return ''
        if isinstance(v, bool):
            return v
        if isinstance(v, (int, float)):
            return v
        return str(v)

    block_import = '_板块导入配置' in book.sheetnames and book['_板块导入配置'].cell(1, 1).value == 'business-block-v1'
    sheets = []
    for ws in book.worksheets:
        if ws.title == '_板块导入配置':
            continue
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if not rows:
            continue
        title = str((rows[0][0] if rows[0] else '') or ws.title).strip()
        header_idx = 2 if block_import else None
        for i in range(1, min(len(rows), 8)):
            if block_import:
                break
            nonempty = sum(1 for c in rows[i] if c not in (None, ''))
            if nonempty >= 3:
                header_idx = i
                break
        if header_idx is None:
            header_idx = 1 if len(rows) > 1 else 0
        source = ''
        if header_idx >= 2:
            prev = rows[header_idx - 1]
            source = str((prev[0] if prev else '') or '').strip()
        headers = [str(c).strip() if c not in (None, '') else '' for c in rows[header_idx]]
        while headers and headers[-1] == '':
            headers.pop()
        width = len(headers)
        if width == 0:
            continue
        data_rows = []
        for r in rows[header_idx + 1:]:
            vals = [cell(r[j]) if j < len(r) else '' for j in range(width)]
            if all(v == '' for v in vals):
                continue
            data_rows.append(vals)
        sheets.append({'key': ws.title, 'title': title, 'source': source, 'headers': headers, 'rows': data_rows})
    book.close()
    if not sheets:
        return {'ok': False, 'error': '未解析到任何工作表，请使用标准模板'}
    if block_import:
        current = crm_business_analysis(caller)
        if not current.get('ok'):
            return {'ok': False, 'error': '无法读取当前版本，未导入'}
        previous = (current.get('data') or {}).get('sheets', [])
        incoming = {sheet['key']: sheet for sheet in sheets}
        for sheet in previous:
            if sheet['key'] in incoming and sheet.get('headers') != incoming[sheet['key']].get('headers'):
                return {'ok': False, 'error': f"{sheet['key']}表头与当前版本不一致，请重新下载模板"}
        sheets = [incoming.pop(sheet['key'], sheet) for sheet in previous] + list(incoming.values())
    try:
        dashboard = crm_compute_dashboard_data(sheets)
    except Exception:
        dashboard = None
    result = {'imported_at': crm_now(), 'imported_by': username,
              'filename': crm_text(payload.get('filename'), 200), 'sheets': sheets, 'dashboard': dashboard}
    # 按导入时间归档为独立快照，保留历史版本而非覆盖；同秒多次导入追加序号避免冲突。
    base = ''.join(ch for ch in result['imported_at'] if ch.isdigit()) or 'v'
    version = base
    n = 1
    if crmdb_enabled():
        taken = set(crmdb_blob_keys('business_analysis:'))
        while f'business_analysis:{version}' in taken:
            version = f'{base}-{n}'
            n += 1
        crmdb_blob_put(f'business_analysis:{version}', result)
    else:
        directory = crm_business_analysis_dir()
        directory.mkdir(parents=True, exist_ok=True)
        while (directory / f'{version}.json').exists():
            version = f'{base}-{n}'
            n += 1
        (directory / f'{version}.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    return {'ok': True, 'sheets': len(sheets), 'imported_at': result['imported_at'], 'version': version}


CRM_OPPORTUNITY_IMPORT_HEADERS = [
    '机会点编码', '机会点名称', '业务日期', '业务组织', '业主单位', '机会点阶段', '数据状态',
    '业务状态', '作废状态', '项目归属', '投资规模(万元)', '可参与份额-总额(万元)',
    '可参与份额-集成产业(万元)', '可参与份额-服务1.0(含配套件)(万元)', '项目预计招标时间',
    '合同预计签订时间', '创建人', '创建日期', '修改人', '修改日期', '客户经理', '备注',
    '作废原因', '审核人', '审核日期', '总包or二代单位', '涉及产品', '赢率',
    '生成销售项目次数', '作废人', '作废日期', '品牌',
]


def crm_opportunity_template():
    """Build the opportunity import template: header row + examples + a guidance sheet.

    Header names match crm_opportunity_import exactly so a filled template imports cleanly;
    the guidance sheet is skipped on import (it lacks the required header trio).
    """
    book = Workbook()
    sheet = book.active
    sheet.title = '机会点导入'
    sheet.append(CRM_OPPORTUNITY_IMPORT_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='15499C')
        cell.alignment = Alignment(horizontal='center')
    examples = [
        {
            '机会点编码': 'OPP-2026-001', '机会点名称': '某地铁6号线信号系统采购', '业务日期': '2026-09-01',
            '业务组织': '湖南承希科技有限公司', '业主单位': '某某地铁集团有限公司', '机会点阶段': '8.招投标阶段',
            '数据状态': '已审核', '业务状态': '正常', '作废状态': '未作废', '项目归属': '中区作战平台',
            '投资规模(万元)': 3000, '可参与份额-总额(万元)': 1200, '可参与份额-集成产业(万元)': 900,
            '可参与份额-服务1.0(含配套件)(万元)': 300, '项目预计招标时间': '2026-10-15',
            '合同预计签订时间': '2026-12-20', '创建人': '张三', '创建日期': '2026-09-01',
            '修改人': '张三', '修改日期': '2026-09-01', '客户经理': '张三', '总包or二代单位': '某某设计院',
            '涉及产品': '通信系统', '赢率': '60%', '生成销售项目次数': 0, '品牌': '华为',
        },
        {
            '机会点编码': 'OPP-2025-133', '机会点名称': '某旧城光纤入户项目（示例作废）',
            '业务日期': '2025-08-01', '业务组织': '湖南承希科技有限公司', '业主单位': '某某建设公司',
            '机会点阶段': '2.项目概算阶段', '数据状态': '已审核', '业务状态': '终止', '作废状态': '已作废',
            '项目归属': '西区作战平台', '投资规模(万元)': 500, '可参与份额-总额(万元)': 0,
            '项目预计招标时间': '2025-08-01', '创建人': '王五', '创建日期': '2025-07-01',
            '修改人': '王五', '修改日期': '2025-08-01', '客户经理': '王五', '作废原因': '业主预算取消',
            '赢率': '0%', '作废人': '王五', '作废日期': '2025-08-01',
        },
    ]
    for example in examples:
        sheet.append([example.get(header, '') for header in CRM_OPPORTUNITY_IMPORT_HEADERS])
    widths = [18, 34, 14, 24, 30, 20, 12, 12, 12, 18, 18, 22, 24, 30, 20, 20,
              14, 16, 14, 16, 14, 28, 24, 14, 16, 26, 24, 10, 18, 14, 16, 18]
    for i, w in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = w
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions

    guide = book.create_sheet('填写说明')
    notes = [
        ['商机（机会点）导入模板 · 填写说明'],
        [''],
        ['1. 请在「机会点导入」工作表中，从第 2 行开始逐行填写；首行表头请勿修改或删除。'],
        ['2. 必填列：机会点编码、机会点名称、业主单位；同一编码重复导入会更新原记录（保留已维护的阶段进展/关系人/材料）。'],
        ['3. 模板字段及顺序与机会点业务系统导出文件保持一致；系统当前读取核心字段生成商机看板，其余扩展列不参与当前看板计算。'],
        ['4. 机会点阶段：可填标准六阶段（立项阶段 / 初设阶段 / 招标（深化）设计 / 招标执行 / 设计联络阶段 / 合同签订），'],
        ['   也可填带编号的细分环节（如“11.合同签订”“9.中标谈判”），系统会自动归并到六大生命周期阶段，原文保留为阶段备注。'],
        ['5. 可参与份额-总额(万元)：填数字（万元）；赢率：可填“60%”或“60”。'],
        ['6. 时间列建议统一填写 YYYY-MM-DD；导入后可在商机详情“联网核对时间点”用招标/中标公告确认。'],
        ['7. 作废状态填“已作废”即归档；“未作废”或留空均视为有效商机。'],
        ['8. 示例行仅供参考，正式导入前请删除示例数据。'],
    ]
    for row in notes:
        guide.append(row)
    guide['A1'].font = Font(bold=True, size=13, color='15499C')
    guide.column_dimensions['A'].width = 96

    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def crm_compute_dashboard_data(sheets):
    """Port of build_data.py: derive the 10-page dashboard dataset from parsed sheets.

    Mirrors the reference build_data.py exactly (row 3 header / row 4+ data), operating on
    our already-parsed {headers, rows} instead of re-opening the workbook.
    """
    import collections
    by = {s['key']: s for s in sheets}

    def rows(name):
        s = by.get(name)
        if not s:
            return []
        hdr = s['headers']
        out = []
        for row in s['rows']:
            d = {}
            for i, h in enumerate(hdr):
                if not h:
                    continue
                v = row[i] if i < len(row) else None
                d[h] = None if v == '' else v
            if any(v is not None for v in d.values()):
                out.append(d)
        return out

    def n(v):
        if not isinstance(v, (int, float)) or isinstance(v, bool):
            return None
        v = float(v)
        return round(v, 4) if abs(v) < 10 else round(v, 2)

    def f(v):
        try:
            return float(v)
        except (TypeError, ValueError):
            return 0.0

    D = {}
    core = rows('01_核心指标总览')
    idx = {r.get('指标'): r for r in core}

    def light(q3, reverse=False):
        if q3 is None:
            return 'gray'
        if reverse:
            return 'red' if q3 > 0.95 else ('amber' if q3 > 0.8 else 'green')
        return 'green' if q3 >= 0.9 else ('amber' if q3 >= 0.6 else 'red')

    K = []
    for name, unit, rev in [('订货', '万元', False), ('管理订货', '万元', False), ('财务收入', '万元', False),
                            ('回款', '万元', False), ('毛利', '万元', False), ('刚性费用', '万元', True)]:
        r = idx.get(name)
        if not r:
            continue
        K.append({'指标': name, '单位': unit, '全年目标': n(r.get('全年目标')), 'Q3目标': n(r.get('Q3季度目标')),
                  '实际': n(r.get('实际完成')), '全年完成率': n(r.get('全年完成率')), 'Q3完成率': n(r.get('Q3完成率')),
                  '全年缺口': n(r.get('全年缺口')), '10月预估': n(r.get('10月预估')), '截至10月累计': n(r.get('截至10月累计')),
                  '全预算': n(r.get('全预算')), '对比同期预算偏差': n(r.get('对比同期预算偏差')),
                  '灯': light(r.get('Q3完成率'), rev), '逆向指标': rev})
    D['P01_红绿灯'] = K
    if idx.get('毛利率') and idx.get('费效比') and idx.get('营业利润'):
        D['P01_辅助'] = {'毛利率实际': n(idx['毛利率'].get('实际完成')), '毛利率目标': n(idx['毛利率'].get('全年目标')),
                        '费效比实际': n(idx['费效比'].get('实际完成')), '费效比目标': n(idx['费效比'].get('全年目标')),
                        '营业利润': n(idx['营业利润'].get('实际完成')), '统计口径': '联指-不含其他 / 截至2026年8月 Q3累计'}

    plat = [r for r in rows('02_作战平台经营结果')
            if region_identity(r.get('区域'))[0]]
    D['P02_平台'] = [{'平台': r.get('区域'), '区总': r.get('区总'), '订货': n(r.get('订货')), '收入': n(r.get('收入')),
                     '毛利': n(r.get('订货毛利')), '毛利率': n(r.get('订货毛利率')), '销售费用': n(r.get('销售费用(含薪酬)')),
                     '贡献毛利': n(r.get('贡献毛利')), '分摊费用': n(r.get('分摊费用')), '最终利润': n(r.get('最终利润'))} for r in plat]

    tgt = rows('03_作战平台目标达成')
    T = {r.get('指标'): r for r in tgt}
    cols = ['联指', '东区作战平台', '西区作战平台', '广东作战平台', '中区作战平台']

    def tv(metric, col):
        return n(T.get(metric, {}).get(col)) if T.get(metric) else None
    D['P04_目标达成'] = [{'平台': region_identity(c)[0] or c,
                        'Q3订货任务令': tv('Q3管理订货任务令', c), '实际订货': tv('实际订货完成', c), '管理订货': tv('管理订货完成', c),
                        'Q3收入任务令': tv('Q3财务收入任务令', c), '已完成收入': tv('Q3已完成收入', c),
                        'Q3回款任务令': tv('Q3回款任务令', c), '已完成回款': tv('Q3已完成回款', c),
                        'Q3毛利任务令': tv('Q3毛利任务令', c), '已完成毛利': tv('毛利完成(含变更)', c),
                        'Q3费用目标': tv('Q3费用目标', c), '费用已使用': tv('费用已使用', c), '费效比': tv('区域费效比', c)} for c in cols]

    pl = {r.get('项目'): r for r in rows('10_公司损益')}

    def plv(item):
        return n(pl.get(item, {}).get('公司')) if pl.get(item) else None
    if pl:
        D['P04_利润桥'] = [
            {'科目': '财务收入毛利', '值': plv('财务收入毛利'), '类型': 'start'},
            {'科目': '战略补贴', '值': plv('战略补贴（含研发+人才）'), '类型': 'plus'},
            {'科目': '销售费用', '值': -(plv('销售费用') or 0), '类型': 'minus'},
            {'科目': '销售平台费用', '值': -(plv('销售平台费用') or 0), '类型': 'minus'},
            {'科目': '管理费用', '值': -(plv('管理费用') or 0), '类型': 'minus'},
            {'科目': '财务费用', '值': -(plv('财务费用') or 0), '类型': 'minus'},
            {'科目': '营业利润', '值': plv('营业利润'), '类型': 'end'}]
    D['P04_损益全表'] = [{'科目': r.get('项目'), '公司': n(r.get('公司')), '利润核算角度': n(r.get('公司-利润核算角度'))} for r in rows('10_公司损益')]

    ind = rows('04_产业结构')
    top = [r for r in ind if r.get('子产业') in ('集成', '服务1.0', '服务2.0', '场研', '其他', '合计')]
    D['P05_产业'] = [{'产业': r.get('子产业'), '订货目标': n(r.get('产业订货目标')), '订货': n(r.get('订货')), '未税订货': n(r.get('未税订货')),
                     '解决方案毛利': n(r.get('解决方案毛利')), '解决方案毛利率': n(r.get('解决方案毛利率')),
                     '毛利': n(r.get('毛利')), '毛利率': n(r.get('毛利率')), '专项预算': n(r.get('专项预算')),
                     '订货完成率': n(r.get('完成率'))} for r in top]
    D['P05_毛利流失'] = []
    for r in top:
        if r.get('子产业') != '合计' and r.get('解决方案毛利') in (None, 0):
            continue
        sm = r.get('解决方案毛利') or 0
        gm = r.get('毛利') or 0
        D['P05_毛利流失'].append({'产业': r.get('子产业'), '解决方案毛利': n(sm), '公司留存毛利': n(gm),
                                '流失金额': n(sm - gm), '流失率': round((sm - gm) / sm, 4) if sm else None})
    sub = [r for r in ind if r.get('大产业') in ('集成', '服务1.0', '服务2.0')]
    D['P05_子产业'] = [{'子产业': r.get('子产业'), '大产业': r.get('大产业'), '订货': n(r.get('订货')), '毛利': n(r.get('毛利')), '毛利率': n(r.get('毛利率'))} for r in sub]

    exp = rows('06_费用明细')
    E = {r.get('区域'): r for r in exp}
    items = ['招待费-日常', '差旅费', '市内交通费', '办事处费用', '薪酬类', '其他']
    if E.get('公司合计'):
        D['P06_费用构成'] = [{'科目': k, '金额': n(E['公司合计'].get(k))} for k in items]
    D['P06_平台费用'] = [{'平台': r.get('区域'), '区总': r.get('区总'),
                        **{k: n(r.get(k)) for k in items}, '合计': n(r.get('合计')), '订货毛利': n(r.get('订货毛利')), '费效比': n(r.get('费效比')),
                        '管理平台分摊': n(r.get('管理平台分摊')), '销售平台分摊': n(r.get('销售平台分摊'))}
                       for r in exp if region_identity(r.get('区域'))[0]]
    D['P06_预算执行'] = [{'平台': region_identity(c)[0] or c, 'Q3费用目标': tv('Q3费用目标', c), '已使用': tv('费用已使用', c),
                        '剩余': tv('费用剩余', c),
                        '使用率': round(T['费用已使用'][c] / T['Q3费用目标'][c], 4) if (T.get('费用已使用') and T.get('Q3费用目标') and T['Q3费用目标'].get(c)) else None,
                        '费效比': tv('区域费效比', c)} for c in cols]
    D['P06_产业费效'] = [{'产业': r.get('产业'), '合计费用': n(r.get('合计')), '产业毛利': n(r.get('产业毛利')),
                        '产业贡毛': n(r.get('产业贡毛')), '费效比': n(r.get('产业费效比'))} for r in rows('08_产业费用构成')]

    pre = [r for r in rows('14_预付条款分析') if r.get('分组') == '合计(剔公司-其他)']
    D['P07_预付趋势'] = [{'年份': int(r['签约年']), '项目数': r.get('项目数'), '合同额万元': round((r.get('合同额(元)') or 0) / 1e4, 1),
                        '有预付数量占比': n(r.get('有预付条款数量占比')), '有预付金额占比': n(r.get('有预付条款金额占比')),
                        '预付≥30%数量占比': n(r.get('预付≥30%数量占比')), '预付≥30%金额占比': n(r.get('预付≥30%金额占比'))} for r in pre if r.get('签约年') is not None]
    D['P07_预付分产业'] = [{'产业': r.get('分组'), '年份': int(r['签约年']), '预付≥30%金额占比': n(r.get('预付≥30%金额占比'))}
                        for r in rows('14_预付条款分析') if r.get('分组') != '合计(剔公司-其他)' and r.get('签约年') is not None]
    D['P07_账龄'] = [{'产业': r.get('产业'), '年份': int(r['签约年份']), 'M1': n(r.get('1个月回款率')),
                     'M3': n(r.get('3个月回款率')), 'M6': n(r.get('6个月回款率'))} for r in rows('13_回款转化率') if r.get('签约年份') is not None]
    cf = rows('11_现金流构成')
    agg = collections.OrderedDict()
    for r in cf:
        if r.get('区域') == '小计':
            agg[r.get('项目年份')] = {'经营现金流': round((r.get('现金流') or 0) / 1e4, 1),
                                    '累计项目现金流': round((r.get('累计项目现金流') or 0) / 1e4, 1),
                                    '回款': round((r.get('回款') or 0) / 1e4, 1), '合同付款': round((r.get('合同付款') or 0) / 1e4, 1)}
    if not agg:  # 模板缺“小计”行时，按项目年份聚合各区域行作为兜底
        tmp = collections.OrderedDict()
        for r in cf:
            y = r.get('项目年份')
            if not y or r.get('区域') == '小计':
                continue
            cur = tmp.setdefault(y, {'现金流': 0.0, '累计项目现金流': 0.0, '回款': 0.0, '合同付款': 0.0})
            for src, key in (('现金流', '现金流'), ('累计项目现金流', '累计项目现金流'), ('回款', '回款'), ('合同付款', '合同付款')):
                cur[key] += (r.get(src) or 0)
        for y, v in tmp.items():
            agg[y] = {'经营现金流': round(v['现金流'] / 1e4, 1), '累计项目现金流': round(v['累计项目现金流'] / 1e4, 1),
                      '回款': round(v['回款'] / 1e4, 1), '合同付款': round(v['合同付款'] / 1e4, 1)}
    D['P07_现金流年份'] = [{'项目年份': k, **v} for k, v in agg.items()]
    D['P07_现金流区域'] = [{'区域': r.get('区域'), '回款': round((r.get('回款') or 0) / 1e4, 1),
                        '合同付款': round((r.get('合同付款') or 0) / 1e4, 1), '经营现金流': round((r.get('现金流') or 0) / 1e4, 1),
                        '累计项目现金流': round((r.get('累计项目现金流') or 0) / 1e4, 1)}
                       for r in cf if r.get('项目年份') == '2026年项目' and r.get('区域') != '小计']

    prov = [r for r in rows('05_省份分布') if r.get('省份') not in ('合计',)]
    for _r in prov:
        if str(_r.get('省份')).strip() == '/':
            _r['省份'] = '公司-其他'
    D['P08_省份'] = [{'省份': r.get('省份'), '订货': n(r.get('订货')), '毛利': n(r.get('毛利')), '毛利率': n(r.get('毛利率')),
                     '现金流': n(r.get('现金流')), '回款金额': n(r.get('回款金额')), '回款转化率': n(r.get('回款转化率')),
                     '收入金额': n(r.get('收入金额')), '收入转化率': n(r.get('收入转化率'))} for r in prov]
    opp = rows('16_机会点清单')
    bands = [('≥90% 高赢率', 0.9, 1.01), ('75%–90%', 0.75, 0.9), ('50%–75%', 0.5, 0.75), ('<50%', 0.0, 0.5)]
    fun = []
    for lb, a, b in bands:
        sel = [r for r in opp if a <= f(r.get('赢率')) < b]
        fun.append({'赢率区间': lb, '机会点数': len(sel),
                    '可参与份额': round(sum(f(r.get('可参与份额-总额(万元)')) for r in sel), 1),
                    '预估毛利': round(sum(f(r.get('预估毛利')) for r in sel), 1)})
    D['P08_漏斗'] = fun
    st = collections.Counter()
    sa = collections.Counter()
    for r in opp:
        k = str(r.get('机会点阶段') or '未标注')
        st[k] += 1
        sa[k] += f(r.get('可参与份额-总额(万元)'))
    D['P08_阶段'] = [{'阶段': k, '机会点数': v, '可参与份额': round(sa[k], 1)} for k, v in st.most_common()]
    rg = collections.Counter()
    rga = collections.Counter()
    for r in opp:
        k = region_identity(r.get('项目归属'))[0]
        if not k:
            continue
        rg[k] += 1
        rga[k] += f(r.get('可参与份额-总额(万元)'))
    D['P08_机会点按平台'] = [{'平台': k, '机会点数': v, '可参与份额': round(rga[k], 1)} for k, v in rg.most_common()]

    attr = rows('15_项目属性分布')
    D['P09_项目属性'] = [{'产业': r.get('产业'), '属性': r.get('属性'), '订货': n(r.get('订货')), '毛利': n(r.get('毛利')),
                        '毛利率': n(r.get('毛利率')), '现金流': n(r.get('现金流')), '回款转化率': n(r.get('回款转化率')),
                        '收入转化率': n(r.get('收入转化率'))} for r in attr if r.get('属性') in ('合作', '自行拓展', '合计')]
    Q = {r.get('指标'): r for r in rows('09_订单质量')}

    def qv(metric, col):
        return Q.get(metric, {}).get(col) if Q.get(metric) else None
    D['P09_订单质量'] = [{'平台': region_identity(c)[0] or c, '整体订单量': qv('整体订单量', c), '达标订单量': qv('达标订单量', c),
                        '达标数量占比': n(qv('达标项目数量占比', c)), '达标金额占比': n(qv('达标项目金额占比', c)),
                        '目标毛利率': n(qv('目标毛利率', c)), '实际毛利率': n(qv('实际毛利率', c))} for c in cols]

    D['P10_责任人'] = [{'平台': r.get('区域'), '区总': r.get('区总'), '订货缺口': None, '最终利润': n(r.get('最终利润'))}
                     for r in plat if r.get('区域') != '上海作战平台']
    for x in D['P10_责任人']:
        t = [y for y in D['P04_目标达成'] if y['平台'] == x['平台']]
        if t:
            x['订货缺口'] = round((t[0]['Q3订货任务令'] or 0) - (t[0]['管理订货'] or 0), 1)

    D['_meta'] = {'数据源': '2026联指经营数据_标准导入模板.xlsx', '口径': '联指（不含其他），截至2026年8月 / Q3累计', '单位': '除注明外均为万元'}
    return D


def crm_opportunity_import(payload, username):
    """Import opportunity workbooks; cancelled rows are retained as archive only."""
    encoded = payload.get('data', '') if isinstance(payload, dict) else ''
    if not isinstance(encoded, str) or not encoded:
        return {'ok': False, 'error': '请选择机会点 Excel 文件'}
    try:
        raw = base64.b64decode(encoded, validate=True)
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        return {'ok': False, 'error': '无法读取机会点 Excel 文件'}
    required = {'机会点编码', '机会点名称', '业主单位'}
    rows = []
    for sheet in book.worksheets:
        values = list(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        if not values: continue
        headers = [str(v or '').strip() for v in values[0]]
        if not required.issubset(headers): continue
        for cells in sheet.iter_rows(min_row=2, values_only=True):
            source = dict(zip(headers, cells))
            code = crm_text(source.get('机会点编码'), 80)
            if not code: continue
            cancelled = crm_text(source.get('作废状态'), 40) == '已作废'
            raw_stage = crm_text(source.get('机会点阶段'))
            def money(header):
                try:
                    return round(float(source.get(header) or 0), 4)
                except (TypeError, ValueError):
                    return 0
            try:
                sales_projects = int(float(source.get('生成销售项目次数') or 0))
            except (TypeError, ValueError):
                sales_projects = 0
            # 投资规模是项目总投资；可参与份额才是我方可争取的部分（= 集成产业 + 服务1.0），两者不可混用。
            rows.append({'id': code, 'code': code, 'name': crm_text(source.get('机会点名称')), 'owner_unit': crm_text(source.get('业主单位')), 'stage': raw_stage, 'sales_stage': crm_map_opportunity_stage(raw_stage), 'win_rate': crm_text(source.get('赢率'), 20), 'status': crm_text(source.get('业务状态'), 40), 'region': crm_text(source.get('项目归属'), 80), 'manager': crm_text(source.get('客户经理'), 80), 'bid_date': crm_text(source.get('项目预计招标时间'), 40), 'contract_date': crm_text(source.get('合同预计签订时间'), 40),
                         'amount': money('可参与份额-总额(万元)'), 'investment': money('投资规模(万元)'),
                         'share_integration': money('可参与份额-集成产业(万元)'), 'share_service': money('可参与份额-服务1.0(含配套件)(万元)'),
                         'sales_projects': sales_projects, 'business_date': crm_text(source.get('业务日期'), 40),
                         'contractor': crm_text(source.get('总包or二代单位'), 120), 'products': crm_text(source.get('涉及产品'), 120),
                         'brand': crm_text(source.get('品牌'), 80), 'remark': crm_text(source.get('备注'), 500),
                         'cancelled': cancelled, 'cancel_reason': crm_text(source.get('作废原因')), 'updated_at': crm_now()})
    book.close()
    if not rows: return {'ok': False, 'error': '未找到包含机会点编码、机会点名称、业主单位的工作表'}
    mode = crm_text(payload.get('mode'), 20) or 'update'
    if mode not in ('preview', 'skip', 'update'):
        return {'ok': False, 'error': '无效的重复商机处理方式'}
    stored = crm_load(username, 'opportunities')
    source_count = len(rows)
    unique_rows = {}
    source_duplicates = set()
    for row in rows:
        if row['id'] in unique_rows:
            source_duplicates.add(row['id'])
        unique_rows[row['id']] = row
    rows = list(unique_rows.values())
    existing_codes = {row['id'] for row in rows if row['id'] in stored}
    # 同一编码若已由其他账号导入，共享视图只会算一条；导入前一并提示，去重导入时跳过。
    elsewhere = crm_opportunity_codes_elsewhere(username)
    foreign_codes = {row['id'] for row in rows if row['id'] not in stored and row['id'] in elsewhere}
    duplicate_codes = existing_codes | source_duplicates | foreign_codes
    duplicates = []
    for code in sorted(duplicate_codes)[:100]:
        incoming = unique_rows[code]
        previous = stored.get(code, {})
        reasons = []
        if code in existing_codes: reasons.append('系统中已存在')
        if code in foreign_codes: reasons.append(f"其他账号已导入（{'、'.join(elsewhere[code])}）")
        if code in source_duplicates: reasons.append('导入文件内重复')
        duplicates.append({'code': code, 'name': incoming.get('name', ''),
                           'owner_unit': incoming.get('owner_unit', ''),
                           'existing_name': previous.get('name', ''), 'reason': '、'.join(reasons)})
    if mode == 'preview':
        return {'ok': True, 'preview': True, 'total': source_count, 'unique_count': len(rows),
                'new_count': sum(row['id'] not in stored and row['id'] not in foreign_codes for row in rows),
                'duplicate_count': len(duplicate_codes), 'duplicates': duplicates,
                'duplicates_truncated': len(duplicate_codes) > len(duplicates)}
    skipped = 0
    if mode == 'skip':
        skipped = len(existing_codes | foreign_codes)
        rows = [row for row in rows if row['id'] not in stored and row['id'] not in foreign_codes]
    created = sum(row['id'] not in stored for row in rows)
    updated = len(rows) - created
    for row in rows:
        previous = stored.get(row['id'], {})
        for key in ('owner_unit_id', 'outcome', 'outcome_source', 'outcome_changed_at', 'stage_progress', 'duplicate_of'):
            if key in previous:
                row[key] = previous[key]
        # 已人工确认的销售阶段优先保留；否则采用导入映射结果。
        if previous.get('sales_stage'):
            row['sales_stage'] = previous['sales_stage']
        if previous.get('duplicate_of'):
            row['cancelled'] = True
            row['outcome'] = 'cancelled'
            row['cancel_reason'] = previous.get('cancel_reason') or f"重复商机，已归并至 {previous['duplicate_of']}"
        elif row['cancelled']:
            row['outcome'] = 'cancelled'
        elif previous.get('outcome') in ('won', 'lost') and previous.get('outcome_source') != 'import':
            pass  # 人工确认的成交结果优先于导入推断
        else:
            # 已生成销售项目 = 已签约；业务状态丢标 = 输单；其余为在谈。
            row['outcome'] = 'won' if row['sales_projects'] >= 1 else 'lost' if '丢标' in row['status'] else 'open'
            row['outcome_source'] = 'import'
        if row.get('outcome') != previous.get('outcome'):
            row['outcome_changed_at'] = row['updated_at']
        if previous.get('owner_unit') and previous['owner_unit'] != row['owner_unit']:
            row['owner_unit_id'] = ''
        row['created_at'] = previous.get('created_at', row['updated_at'])
        row['version'] = int(previous.get('version', 0)) + 1
        stored[row['id']] = row
    crm_store(username, 'opportunities', stored); crm_on_change(username)
    return {'ok': True, 'imported': len(rows), 'created': created, 'updated': updated, 'skipped': skipped,
            'active': sum(r.get('outcome') == 'open' for r in rows), 'won': sum(r.get('outcome') == 'won' for r in rows),
            'lost': sum(r.get('outcome') == 'lost' for r in rows), 'archived': sum(r['cancelled'] for r in rows)}


def crm_sales_contract_import(payload, username):
    encoded = payload.get('data', '') if isinstance(payload, dict) else ''
    if not isinstance(encoded, str) or not encoded:
        return {'ok': False, 'error': '请选择销售合同 Excel 文件'}
    try:
        raw = base64.b64decode(encoded, validate=True)
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        return {'ok': False, 'error': '无法读取销售合同 Excel 文件'}
    required = {'销售项目编码', '销售子项目编码', '销售合同编码', '合同金额'}
    contracts = {}
    for sheet in book.worksheets:
        first = list(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        if not first:
            continue
        headers = [str(value or '').strip() for value in first[0]]
        if not required.issubset(headers):
            continue
        for cells in sheet.iter_rows(min_row=2, values_only=True):
            source = dict(zip(headers, cells))
            project_code = crm_text(source.get('销售项目编码'), 80)
            subproject_code = crm_text(source.get('销售子项目编码'), 80)
            contract_code = crm_text(source.get('销售合同编码'), 80)
            if not project_code or not contract_code:
                continue
            try:
                amount = round(float(source.get('合同金额') or 0), 2)
            except (TypeError, ValueError):
                amount = 0
            key = subproject_code or contract_code
            contracts[key] = {
                'id': key, 'project_code': project_code, 'project_name': crm_text(source.get('销售项目名称'), 200),
                'subproject_code': subproject_code, 'subproject_name': crm_text(source.get('销售子项目名称'), 200),
                'contract_code': contract_code, 'customer': crm_text(source.get('客户'), 200),
                'owner_unit': crm_text(source.get('业主单位'), 200), 'signed_at': crm_text(source.get('合同签订日期'), 40),
                'amount_yuan': amount, 'project_manager': crm_text(source.get('项目经理'), 80),
                'account_manager': crm_text(source.get('客户经理'), 80), 'archived': crm_text(source.get('归档'), 20) == '是',
                'industry': crm_text(source.get('产业'), 80), 'payment_terms': crm_text(source.get('收款条件'), 500),
                'expected_margin_yuan': source.get('预计毛利(不含税)') if isinstance(source.get('预计毛利(不含税)'), (int, float)) else None,
                'imported_at': crm_now(), 'imported_by': username,
            }
    book.close()
    if not contracts:
        return {'ok': False, 'error': '未找到销售合同数据，请确认表头包含销售项目编码、销售子项目编码、销售合同编码和合同金额'}
    opportunities = crm_load(username, 'opportunities')
    grouped = {}
    for contract in contracts.values():
        grouped.setdefault(contract['project_code'], []).append(contract)
    matched = 0
    now = crm_now()
    for code, items in grouped.items():
        opportunity = opportunities.get(code)
        if not opportunity:
            continue
        signed_dates = sorted(item['signed_at'] for item in items if item['signed_at'])
        opportunity['outcome'] = 'won'
        opportunity['outcome_source'] = 'sales_contract'
        opportunity['outcome_changed_at'] = now
        opportunity['cancelled'] = False
        opportunity['cancel_reason'] = ''
        opportunity['sales_stage'] = '合同签订'
        opportunity['sales_projects'] = len({item['subproject_code'] or item['contract_code'] for item in items})
        opportunity['contract_count'] = len(items)
        opportunity['contract_amount_yuan'] = round(sum(item['amount_yuan'] for item in items), 2)
        opportunity['contract_codes'] = [item['contract_code'] for item in items]
        if signed_dates:
            opportunity['contract_date'] = signed_dates[0]
        opportunity['updated_at'] = now
        opportunity['version'] = int(opportunity.get('version', 0)) + 1
        matched += 1
    crm_store(username, 'sales_contracts', contracts)
    crm_store(username, 'opportunities', opportunities)
    crm_on_change(username)
    return {'ok': True, 'contracts': len(contracts), 'contract_projects': len(grouped),
            'matched_opportunities': matched, 'unmatched_projects': len(grouped) - matched,
            'won_opportunities': sum(item.get('outcome') == 'won' for item in opportunities.values())}


CRM_OPPORTUNITY_EXPORT_HEADERS = [
    '机会点编码', '机会点名称', '业主单位', '项目归属', '机会点阶段', '赢率', '业务状态',
    '客户经理', '项目预计招标时间', '合同预计签订时间', '投资规模(万元)', '可参与份额-总额(万元)',
    '可参与份额-集成产业(万元)', '可参与份额-服务1.0(含配套件)(万元)', '生成销售项目次数',
    '作废状态', '作废原因', '最后更新时间', '关联客户ID', '销售阶段', '成交结果', '结果更新时间',
]


def crm_opportunity_export(caller, state='all', stage='', keyword='', region='', sales_stage=''):
    """Export the same opportunity subset currently shown in the workbench."""
    if not crm_capabilities(caller).get('read'):
        raise PermissionError('无机会点查看权限')
    state = crm_text(state, 20) or 'all'
    if state not in ('all', 'active', 'archived'):
        raise ValueError('机会点状态筛选无效')
    stage = crm_text(stage, 80)
    keyword = crm_text(keyword, 160).lower()
    records = crm_read_opportunities(caller)['opportunities']
    rows = []
    for item in records:
        cancelled = bool(item.get('cancelled'))
        if state == 'active' and cancelled:
            continue
        if state == 'archived' and not cancelled:
            continue
        if any(value and (str(item.get(field) or '').strip() or '__missing__') != value for field, value in [('stage', stage), ('region', region), ('sales_stage', sales_stage)]):
            continue
        search_text = ' '.join(str(item.get(field, '')) for field in ('code', 'name', 'owner_unit', 'manager')).lower()
        if keyword and keyword not in search_text:
            continue
        rows.append([
            item.get('code', ''), item.get('name', ''), item.get('owner_unit', ''), item.get('region', ''),
            item.get('stage', ''), item.get('win_rate', ''), item.get('status', ''), item.get('manager', ''),
            item.get('bid_date', ''), item.get('contract_date', ''), item.get('investment', ''), item.get('amount', 0),
            item.get('share_integration', ''), item.get('share_service', ''), item.get('sales_projects', ''),
            '已作废' if cancelled else '', item.get('cancel_reason', ''), item.get('updated_at', ''),
            item.get('owner_unit_id', ''), item.get('sales_stage', ''),
            {'open': '进行中', 'won': '赢单', 'lost': '输单', 'cancelled': '取消'}.get(item.get('outcome'), '取消' if cancelled else '进行中'),
            item.get('outcome_changed_at', ''),
        ])
    book = Workbook()
    sheet = book.active
    sheet.title = '机会点清单'
    sheet.append(CRM_OPPORTUNITY_EXPORT_HEADERS)
    for row in rows:
        sheet.append(row)
    header_fill = PatternFill('solid', fgColor='1F4E78')
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {'A':18, 'B':34, 'C':30, 'D':16, 'E':14, 'F':10, 'G':14, 'H':14, 'I':18, 'J':18, 'K':22, 'L':12, 'M':26, 'N':22}.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue(), '机会点清单.xlsx'


def _crm_normalize_external_business_values(values, headers):
    """Map the approved regional collection layout onto the canonical business schema."""
    if headers not in CRM_BUSINESS_EXTERNAL_LAYOUTS:
        return values
    company = values.get('公司全称', '').strip()
    subsidiary = values.get('子公司全称', '').strip()
    short_name = values.get('单位*', '').strip()
    # “单位” is directly below company full name in the current template.
    # Legacy files can still keep a subsidiary between these two levels.
    values['单位*'] = short_name or subsidiary or company
    # Optional interaction dates often contain source-system placeholders or
    # multiple date notes.  They must not stop an otherwise complete contact
    # row from entering CRM; retain only a single, valid date for action data.
    for key in ('上次互动时间', '计划跟进日期'):
        value = values.get(key, '').strip()
        if not value or value in ('/', '-', '—', '--'):
            values[key] = ''
            continue
        try:
            parsed = datetime.strptime(value.replace('/', '-'), '%Y-%m-%d')
            if len(value.split('/')[0].split('-')[0]) != 4:
                raise ValueError()
            values[key] = parsed.strftime('%Y-%m-%d')
        except ValueError:
            values[key] = ''
    # The CRM has only two importance values; preserve uncertain free text as
    # blank instead of inventing a priority level during an automated import.
    if values.get('重要程度', '').strip() not in ('', '重要', '一般'):
        values['重要程度'] = ''
    normalized = {header: values.get(header, '') for header in CRM_BUSINESS_HEADERS}
    normalized.update({header: values.get(header, '') for header in CRM_OUR_STAFF_HEADERS})
    # Keep hierarchy-only metadata out of legacy field handling. The business
    # row builder uses it to create company full name as the top-level unit and
    # the subsidiary (when present) as its direct child.
    normalized['_company_full_name'] = company
    normalized['_subsidiary_full_name'] = subsidiary
    normalized['_unit_name'] = short_name
    # Keep user-facing source fields for correction/retry cards. Internal
    # business fields above remain the values consumed by crm_business_row.
    normalized['公司全称'] = company
    normalized['子公司全称'] = subsidiary
    normalized['省份'] = values.get('省份', '')
    return normalized


def _crm_import_completion_path():
    return USER_DATA_DIR / '.crm-import-completion' / 'tasks.json'


def _crm_import_completion_read():
    if crmdb_enabled():
        payload = crmdb_blob_get('import_tasks')
        return payload if isinstance(payload, dict) else {}
    path = _crm_import_completion_path()
    try:
        payload = json.loads(path.read_text(encoding='utf-8'))
        return payload if isinstance(payload, dict) else {}
    except (OSError, ValueError):
        return {}


def _crm_import_completion_write(tasks):
    if crmdb_enabled():
        crmdb_blob_put('import_tasks', tasks)
        return
    path = _crm_import_completion_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix('.tmp')
    temporary.write_text(json.dumps(tasks, ensure_ascii=False, indent=2), encoding='utf-8')
    temporary.replace(path)


def _crm_import_completion_load():
    # During a bulk import the completion queue is held in memory; rewriting the
    # whole tasks file per pending row is O(n^2) and dominates large imports.
    batch = getattr(CRM_IMPORT_CONTEXT, 'batch', None)
    if batch is not None:
        if 'completion' not in batch:
            batch['completion'] = _crm_import_completion_read()
        return batch['completion']
    return _crm_import_completion_read()


def _crm_import_completion_store(tasks):
    batch = getattr(CRM_IMPORT_CONTEXT, 'batch', None)
    if batch is not None:
        batch['completion'] = tasks
        batch['completion_dirty'] = True
        return
    _crm_import_completion_write(tasks)


def _crm_flush_completion_batch():
    batch = getattr(CRM_IMPORT_CONTEXT, 'batch', None)
    if batch and batch.get('completion_dirty'):
        _crm_import_completion_write(batch['completion'])
        batch['completion_dirty'] = False


def _crm_account_for_member_identity(member_identity, config=None):
    """Resolve a username, employee ID, or legacy display name only if unique.

    任务列表按行反复解析账号，config 由调用方读一次传入，避免每行重读配置。
    """
    identity = crm_text(member_identity, 80)
    candidates = [user.get('username', '') for user in (config or read_config()).get('users', [])
                  if user.get('username') and identity in {
                      crm_text(user.get('username'), 80),
                      crm_text(user.get('employee_id'), 80),
                      crm_text(user.get('name'), 80),
                  }]
    return candidates[0] if len(candidates) == 1 else ''


def _crm_account_for_member_name(member_name):
    return _crm_account_for_member_identity(member_name)


def _crm_user_display_name(username, config=None):
    """Return the platform user's Chinese display name for business-facing views."""
    username = crm_text(username, 80)
    user = next((item for item in (config or read_config()).get('users', []) if item.get('username') == username), {})
    return crm_text(user.get('name'), 80) or username


def _crm_import_completion_assignee(task, config=None):
    return (_crm_account_for_member_identity(task.get('assignee_username'), config) or
            _crm_account_for_member_identity(task.get('assignee_name'), config))


def _crm_import_missing_fields(raw_values, values, headers):
    # The three-field gate is defined for the external regional collection
    # workbook. Keep established system/legacy templates import-compatible;
    # those templates do not expose “公司全称” as an independent field.
    if headers not in CRM_BUSINESS_EXTERNAL_LAYOUTS:
        return []
    company = raw_values.get('公司全称', '')
    required_values = {'公司全称': company, '姓名*': values.get('姓名*', ''), '单位类型': values.get('单位类型', '')}
    return [field for field in CRM_IMPORT_COMPLETION_REQUIRED_FIELDS if not str(required_values[field] or '').strip()]


def _crm_import_completion_create(raw_values, values, headers, missing_fields, uploader_username, row_number):
    """Persist an incomplete row outside CRM records until its owner completes it."""
    member_name = crm_text(raw_values.get('我方姓名') or values.get('我方姓名'), 80)
    target_username = _crm_account_for_member_name(member_name)
    fingerprint = json.dumps({'source': raw_values, 'uploader': uploader_username}, ensure_ascii=False, sort_keys=True)
    tasks = _crm_import_completion_load()
    existing = next((task for task in tasks.values() if task.get('fingerprint') == fingerprint and task.get('status') == 'pending'), None)
    if existing:
        return existing
    task = {
        'id': crm_new_id('import-completion'), 'status': 'pending', 'fingerprint': fingerprint,
        'uploader_username': uploader_username, 'source_row': row_number,
        'assignee_name': member_name, 'assignee_username': target_username,
        'missing_fields': list(missing_fields), 'values': dict(values), 'original_values': dict(raw_values),
        'created_at': crm_now(), 'updated_at': crm_now(),
    }
    tasks[task['id']] = task
    _crm_import_completion_store(tasks)
    return task


def crm_import_completion_tasks(caller):
    username = crm_text((caller or {}).get('username'), 80)
    capabilities = crm_capabilities(caller or {})
    is_manager = bool(capabilities.get('manage') or (caller or {}).get('role') == 'superadmin')
    tasks = []
    config = read_config()
    for stored in _crm_import_completion_load().values():
        task = dict(stored)
        resolved_assignee = _crm_import_completion_assignee(task, config)
        if resolved_assignee:
            task['assignee_username'] = resolved_assignee
        task['uploader_name'] = _crm_user_display_name(task.get('uploader_username'), config)
        if task.get('assignee_username'):
            task['assignee_name'] = _crm_user_display_name(task['assignee_username'], config)
        tasks.append(task)
    if not is_manager:
        tasks = [task for task in tasks if task.get('assignee_username') == username]
    tasks.sort(key=lambda task: (task.get('status') != 'pending', task.get('created_at', '')), reverse=False)
    return {'ok': True, 'can_manage': is_manager, 'tasks': tasks,
            'pending': sum(task.get('status') == 'pending' for task in tasks)}


CRM_IMPORT_COMPLETION_EXPORT_HEADERS = [
    '任务编号', '区域', '客户经理', '接收账号', '公司全称', '姓名*', '单位类型',
    '单位*', '部门', '职务',
    '重要程度', '支持程度', '跟进轨迹', '上次互动时间', '本周互动计划', '计划跟进日期',
    '省份', '城市', '维系周期（天）', '我方姓名',
    # 个人联系方式统一置于最右，便于先完成客户主体和业务信息维护。
    '手机', '办公电话', '邮箱', '微信 / 其他联系',
]


def _crm_completion_visible_tasks(caller, region='', manager=''):
    """Return pending tasks visible to caller, optionally narrowed by manager filters."""
    username = crm_text((caller or {}).get('username'), 80)
    capabilities = crm_capabilities(caller or {})
    is_manager = bool(capabilities.get('manage') or (caller or {}).get('role') == 'superadmin')
    region, manager = crm_text(region, 40), crm_text(manager, 80)
    if (region or manager) and not is_manager:
        raise PermissionError('只有客户经营管理员可以按区域或客户经理导出资料补充')
    result = []
    config = read_config()
    for stored in _crm_import_completion_load().values():
        task = dict(stored)
        if task.get('status') != 'pending':
            continue
        resolved_assignee = _crm_import_completion_assignee(task, config)
        if resolved_assignee:
            task['assignee_username'] = resolved_assignee
        if not is_manager and task.get('assignee_username') != username:
            continue
        values = task.get('original_values') or task.get('values') or {}
        task_region = crm_text(values.get('区域') or values.get('经营区域'), 40)
        if region and task_region != region:
            continue
        if manager and task.get('assignee_username') != manager:
            continue
        result.append(task)
    return sorted(result, key=lambda task: task.get('created_at', ''))


def crm_import_completion_excel_export(caller, region='', manager=''):
    """Download selected pending completion tasks for offline correction."""
    tasks = _crm_completion_visible_tasks(caller, region, manager)
    book = Workbook()
    sheet = book.active
    sheet.title = '资料补充'
    sheet.append(CRM_IMPORT_COMPLETION_EXPORT_HEADERS)
    for index, cell in enumerate(sheet[1], 1):
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = PatternFill('solid', fgColor='174B91' if index <= 7 else '627891')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[1].height = 30
    widths = [25, 12, 16, 16, 34, 18, 16, 24, 18, 18, 14, 14, 42, 18, 42, 18, 14, 14, 16, 16, 16, 18, 28, 24]
    for cell, width in zip(sheet[1], widths):
        sheet.column_dimensions[cell.column_letter].width = width
    highlighted_fields = ('公司全称', '单位类型', '部门', '省份', '城市')
    missing_fill = PatternFill('solid', fgColor='FFF2CC')
    completed_fill = PatternFill('solid', fgColor='E2F0D9')
    for task in tasks:
        values = dict(task.get('values') or {})
        original = dict(task.get('original_values') or {})
        def value(key):
            return original.get(key) or values.get(key) or ''
        sheet.append([
            task.get('id', ''), value('区域') or value('经营区域'), task.get('assignee_name', ''), task.get('assignee_username', ''),
            value('公司全称'), value('姓名*'), value('单位类型'), value('单位*'), value('部门'), value('职务'), value('重要程度'),
            value('支持程度'), value('跟进轨迹'), value('上次互动时间'), value('本周互动计划'), value('计划跟进日期'),
            value('省份'), value('城市'), value('维系周期（天）'), value('我方姓名'),
            value('手机'), value('办公电话'), value('邮箱'), value('微信 / 其他联系'),
        ])
        row_number = sheet.max_row
        for field in highlighted_fields:
            cell = sheet.cell(row_number, CRM_IMPORT_COMPLETION_EXPORT_HEADERS.index(field) + 1)
            cell.fill = completed_fill if str(cell.value or '').strip() else missing_fill
            cell.alignment = Alignment(wrap_text=True, vertical='center')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:X{len(tasks) + 1}'
    guide = book.create_sheet('填写说明')
    for line in [
        '资料补充导出表：请补齐黄色标记的公司全称、单位类型、部门、省份、城市，并可补充姓名*和接收账号；勿修改任务编号和工作表名称。',
        '黄色单元格表示待补充的公司全称、单位类型、部门、省份或城市；绿色单元格表示该项已有资料，请核对后保留。',
        '公司全称、姓名*、单位类型为必填。接收账号仅在管理员需要重新指定客户经理时填写。',
        '上传后系统按任务编号回写；成功补全的任务会自动从待补充列表移除，失败行保留以便修正后再次上传。',
    ]:
        guide.append([line])
    guide.column_dimensions['A'].width = 120
    out = io.BytesIO(); book.save(out); book.close()
    suffix = region or manager or '全部'
    return out.getvalue(), f'资料补充-{suffix}.xlsx'


def crm_import_completion_excel_import(payload, caller):
    """Apply an edited completion workbook and leave only failed tasks pending."""
    if not str((payload or {}).get('filename', '')).lower().endswith('.xlsx'):
        raise ValueError('请上传资料补充导出的 .xlsx 文件')
    encoded = (payload or {}).get('data', '')
    if not isinstance(encoded, str) or len(encoded) > 7 * 1024 * 1024:
        raise ValueError('文件不能超过5MB')
    try:
        raw = base64.b64decode(encoded, validate=True)
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    except Exception:
        raise ValueError('无法读取工作簿，请上传资料补充导出的xlsx文件')
    try:
        if '资料补充' not in book.sheetnames:
            raise ValueError('未找到“资料补充”工作表，请使用导出文件')
        sheet = book['资料补充']
        iterator = sheet.iter_rows()
        headers = [cell.value for cell in next(iterator, ())]
        required = CRM_IMPORT_COMPLETION_EXPORT_HEADERS[:7]
        if headers[:len(required)] != required:
            raise ValueError('表头不匹配，请使用资料补充导出文件')
        result = {'ok': True, 'completed': 0, 'skipped': 0, 'errors': []}
        for row_number, cells in enumerate(iterator, 2):
            row = {header: (str(cell.value).strip() if cell.value is not None else '') for header, cell in zip(headers, cells)}
            if not any(row.values()):
                continue
            task_id = crm_text(row.get('任务编号'), 120)
            if not task_id:
                result['errors'].append({'row': row_number, 'error': '缺少任务编号'})
                continue
            try:
                completed = crm_import_completion_complete({
                    'id': task_id,
                    'target_username': row.get('接收账号', ''),
                    'values': {
                        '公司全称': row.get('公司全称', ''), '姓名*': row.get('姓名*', ''), '单位类型': row.get('单位类型', ''),
                        '部门': row.get('部门', ''), '省份': row.get('省份', ''), '城市': row.get('城市', ''),
                    },
                }, caller)
                if completed.get('ok'):
                    result['completed'] += 1
                else:
                    result['skipped'] += 1
            except (ValueError, PermissionError) as exc:
                result['errors'].append({'row': row_number, 'task_id': task_id, 'error': str(exc)})
        return result
    finally:
        book.close()


def crm_import_completion_complete(payload, caller):
    task_id = crm_text((payload or {}).get('id'), 120)
    tasks = _crm_import_completion_load()
    task = tasks.get(task_id)
    if not task or task.get('status') != 'pending':
        raise ValueError('待补充记录不存在或已完成')
    username = crm_text((caller or {}).get('username'), 80)
    capabilities = crm_capabilities(caller or {})
    is_manager = bool(capabilities.get('manage') or (caller or {}).get('role') == 'superadmin')
    assigned_username = _crm_import_completion_assignee(task)
    if not is_manager and assigned_username != username:
        raise PermissionError('只能补充分派给本人的记录')
    supplied = (payload or {}).get('values', {})
    if not isinstance(supplied, dict):
        raise ValueError('补充数据格式无效')
    company = crm_text(supplied.get('公司全称') or supplied.get('company_full_name'), 200)
    person_name = crm_text(supplied.get('姓名*') or supplied.get('name'), 80)
    unit_type = crm_text(supplied.get('单位类型') or supplied.get('unit_type'), 40)
    target_username = crm_text(
        ((payload or {}).get('target_username') if is_manager else '') or assigned_username,
        80,
    )
    if not target_username or not find_user(target_username):
        raise ValueError('请指定有效的接收账号')
    missing = [field for field, value in [('公司全称', company), ('姓名*', person_name), ('单位类型', unit_type)] if not value]
    if missing:
        raise ValueError('请补充：' + '、'.join(missing))
    if unit_type not in CRM_UNIT_TYPES:
        raise ValueError('单位类型无效')
    values = dict(task.get('values') or {})
    # Offline completion sheets can enrich these business fields alongside the
    # three entry gates. Keep original values when the exported cell is left blank.
    for field in ('部门', '省份', '城市'):
        supplied_value = crm_text(supplied.get(field), 120)
        if supplied_value:
            values[field] = supplied_value
    # A subsidiary remains the contact's unit. If none was supplied, the
    # completed company full name is also the operating unit.
    values['_company_full_name'] = company
    values.update({'单位*': values.get('_subsidiary_full_name') or company,
                   '姓名*': person_name, '单位类型': unit_type})
    with CRM_LOCK:
        CRM_IMPORT_CONTEXT.active = True
        try:
            status = crm_business_row(values, target_username)
        finally:
            CRM_IMPORT_CONTEXT.active = False
        if status == 'created':
            crm_on_change(target_username)
    task.update({'status': 'completed', 'assignee_username': target_username,
                 'completed_by': username, 'completed_at': crm_now(), 'updated_at': crm_now(),
                 'completion_values': {'公司全称': company, '姓名*': person_name, '单位类型': unit_type}})
    tasks[task_id] = task
    _crm_import_completion_store(tasks)
    return {'ok': True, 'status': status, 'task': task}


def crm_excel_template(advanced=False):
    if not advanced:
        return crm_excel_simple_template()
    book = Workbook()
    guide = book.active
    guide.title = '填写说明'
    for line in [
        '联指经营助手 · Excel 收集模板',
        '星号字段必填；请保留工作表名称和第一行表头。空行自动忽略。',
        '只填需要导入的工作表。处理顺序：单位→联系人→关系→动作→项目决策链。',
        '联系人用单位名称+姓名关联；动作需要先有关系。我方对接人可留空（不指派，之后可在系统中指定）。',
        '同名单位或同单位同名联系人无法唯一匹配时，该行报错，请先在系统中核实。',
        '导入只新增，不覆盖已有档案；已有单位、联系人、关系和决策链跳过。相同动作跳过。',
        '有效行自动录入，错误行不录入；查看导入结果修正后可重新上传。',
        '日期填写 YYYY-MM-DD；是非字段填 是/否；影响力为0到100整数；职业轨迹用 → 分隔。',
        '发生日期不可晚于今天。请勿填写 Excel 公式；文件最多5MB，仅支持.xlsx。',
        '单位类型：' + '、'.join(CRM_UNIT_TYPES),
        '决策节点：' + '、'.join(CRM_DECISION_NODES),
        '决策角色：' + '、'.join(CRM_DECISION_ROLES),
        '示例仅供阅读，不参与导入：单位=某轨道集团；姓名=张工；日期=2026-09-01。',
    ]: guide.append([line])
    guide.column_dimensions['A'].width = 115
    for row in guide: row[0].alignment = Alignment(wrap_text=True, vertical='center'); guide.row_dimensions[row[0].row].height = 32
    for name, headers in CRM_EXCEL_SHEETS.items():
        sheet = book.create_sheet(name); sheet.append(headers); sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = f'A1:{sheet.cell(1,len(headers)).column_letter}10001'
        for cell in sheet[1]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='174B91')
            sheet.column_dimensions[cell.column_letter].width = 23
        for i, header in enumerate(headers, 1):
            options = CRM_UNIT_TYPES if header == '单位类型*' else CRM_DECISION_NODES if header in ('决策节点*','关联决策节点') else CRM_DECISION_ROLES if header == '决策角色' else ['是','否'] if header in ('单点风险','有效动作','客户主动') else ['支持','中立','反对','未知'] if header == '立场' else None
            if options:
                validation = DataValidation(type='list', formula1='"'+','.join(options)+'"', allow_blank=True)
                validation.errorTitle='请使用下拉选项'; validation.error='请选择有效值'; validation.showErrorMessage=True
                sheet.add_data_validation(validation); validation.add(f'{sheet.cell(1,i).column_letter}2:{sheet.cell(1,i).column_letter}10001')
    output=io.BytesIO(); book.save(output); book.close(); return output.getvalue()


def crm_excel_import(payload, username):
    if not str(payload.get('filename','')).lower().endswith('.xlsx'):
        raise ValueError('请上传 .xlsx 格式的模板文件')
    encoded = payload.get('data','')
    if not isinstance(encoded,str) or len(encoded) > 7 * 1024 * 1024:
        raise ValueError('文件不能超过5MB')
    try: raw=base64.b64decode(encoded,validate=True)
    except Exception: raise ValueError('文件编码无效')
    if len(raw)>5*1024*1024: raise ValueError('文件不能超过5MB')
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(f.file_size for f in archive.infolist())>30*1024*1024 or len(archive.infolist())>200:
                raise ValueError('工作簿内容过大')
        book=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
    except ValueError: raise
    except Exception: raise ValueError('无法读取工作簿，请使用未加密的xlsx模板')
    schemas = {'经营收集表': CRM_BUSINESS_HEADERS} if '经营收集表' in book.sheetnames else ({'客户收集表': CRM_SIMPLE_HEADERS} if '客户收集表' in book.sheetnames else CRM_EXCEL_SHEETS)
    rows=[]
    try:
        if not any(name in book.sheetnames for name in schemas): raise ValueError('未找到模板工作表，请先下载收集模板')
        for name,headers in schemas.items():
            if name not in book.sheetnames: continue
            sheet=book[name]
            if name == '客户收集表' and sheet.cell(1,9).value == '经营区域':
                headers = headers + ['经营区域']
                schemas[name] = headers
            max_row, max_column = _crm_sheet_dimensions(sheet)
            if max_row > 10000 or max_column > 100: raise ValueError('工作表范围过大，请删除多余行列')
            iterator=sheet.iter_rows(); first=next(iterator,())
            actual_headers = [c.value for c in first]
            if name == '经营收集表':
                actual_headers = [v[:-4] if isinstance(v,str) and v in ('我方公司（选填）','我方部门（选填）','我方岗位（选填）') else v for v in actual_headers]
                if actual_headers[:len(CRM_BUSINESS_EXTERNAL_HEADERS)] == CRM_BUSINESS_EXTERNAL_HEADERS:
                    headers = CRM_BUSINESS_EXTERNAL_HEADERS
                elif actual_headers[:len(CRM_BUSINESS_EXTERNAL_SUBSIDIARY_HEADERS)] == CRM_BUSINESS_EXTERNAL_SUBSIDIARY_HEADERS:
                    headers = CRM_BUSINESS_EXTERNAL_SUBSIDIARY_HEADERS
                elif actual_headers[:len(CRM_BUSINESS_EXTERNAL_LEGACY_HEADERS)] == CRM_BUSINESS_EXTERNAL_LEGACY_HEADERS:
                    headers = CRM_BUSINESS_EXTERNAL_LEGACY_HEADERS
                elif actual_headers[:len(CRM_BUSINESS_EXTERNAL_LEGACY_SUBSIDIARY_HEADERS)] == CRM_BUSINESS_EXTERNAL_LEGACY_SUBSIDIARY_HEADERS:
                    headers = CRM_BUSINESS_EXTERNAL_LEGACY_SUBSIDIARY_HEADERS
                elif actual_headers[:len(CRM_BUSINESS_LEGACY_HEADERS)] == CRM_BUSINESS_LEGACY_HEADERS:
                    headers = CRM_BUSINESS_LEGACY_HEADERS
                elif actual_headers[:len(CRM_BUSINESS_HEADERS)] == CRM_BUSINESS_HEADERS:
                    headers = CRM_BUSINESS_HEADERS
                elif actual_headers[:len(CRM_BUSINESS_CONTACT_LEGACY_HEADERS)] == CRM_BUSINESS_CONTACT_LEGACY_HEADERS:
                    headers = CRM_BUSINESS_CONTACT_LEGACY_HEADERS
                elif actual_headers[:len(CRM_BUSINESS_OLD_HEADERS)] == CRM_BUSINESS_OLD_HEADERS:
                    headers = CRM_BUSINESS_OLD_HEADERS
                else:
                    raise ValueError(f'{name}表头不匹配，请使用下载的模板')
                schemas[name] = headers
            elif actual_headers[:len(headers)] != headers: raise ValueError(f'{name}表头不匹配，请使用下载的模板')
            for row_number,cells in enumerate(iterator,2):
                if all(c.value is None or str(c.value).strip()=='' for c in cells): continue
                rows.append((name,row_number,cells))
        result={'ok':True,'created':0,'skipped':0,'pending':0,'errors':[], 'sheets':{n:{'created':0,'skipped':0,'pending':0,'failed':0} for n in schemas}}
        with CRM_LOCK:
            CRM_IMPORT_CONTEXT.active=True
            CRM_IMPORT_CONTEXT.batch={'data':{},'original':{},'dirty':set()}
            try:
                for name,row_number,cells in rows:
                    values = {}
                    try:
                        if any(c.data_type=='f' for c in cells): raise ValueError('不支持公式，请粘贴为值后上传')
                        headers=schemas[name]
                        values={h: (c.value.strftime('%Y-%m-%d') if isinstance(c.value,datetime) else str(c.value).strip() if c.value is not None else '') for h,c in zip(headers,cells)}
                        if name == '经营收集表':
                            source_values = dict(values)
                            values = _crm_normalize_legacy_business_values(values, headers)
                            values = _crm_normalize_external_business_values(values, headers)
                        if any(len(v)>2000 for v in values.values()): raise ValueError('单元格文字不能超过2000字')
                        for h in headers:
                            # The external collection workbook routes incomplete
                            # contact rows to the owner-completion queue instead of
                            # rejecting them at parse time. Its `单位*` is a display
                            # shorthand, while company full name is the import gate.
                            if (name == '经营收集表' and headers in CRM_BUSINESS_EXTERNAL_LAYOUTS
                                    and h in ('单位*', '姓名*')):
                                continue
                            if h.endswith('*') and not values.get(h): raise ValueError(f'{h[:-1]}必填')
                        if name == '经营收集表':
                            missing_fields = _crm_import_missing_fields(source_values, values, headers)
                            if missing_fields:
                                _crm_import_completion_create(source_values, values, headers, missing_fields, username, row_number)
                                status = 'pending'
                            else:
                                status = crm_business_row(values, username)
                        else:
                            status = crm_excel_simple_row(values,username) if name == '客户收集表' else crm_excel_row(name,values,username)
                        result[status]+=1; result['sheets'][name][status]+=1
                    except (ValueError,TypeError) as exc:
                        result['errors'].append({'sheet':name,'row':row_number,'error':str(exc),
                                                 'headers':list(headers), 'values':values})
                        result['sheets'][name]['failed']+=1
            finally:
                try:
                    crm_flush_import_batch()
                    _crm_flush_completion_batch()
                finally:
                    CRM_IMPORT_CONTEXT.batch=None
                    CRM_IMPORT_CONTEXT.active=False
                if result['created']: crm_on_change(username)
        return result
    finally: book.close()


def crm_excel_retry_rows(payload, username):
    """Retry business collection rows edited in the import-result UI."""
    raw_rows = (payload or {}).get('rows', [])
    if not isinstance(raw_rows, list) or not raw_rows or len(raw_rows) > 10000:
        raise ValueError('请选择1到10000条需要修正的记录')
    # Overwrite lets the operator confirm that mismatched rows (e.g. a different unit
    # type or business region) should update the existing archive instead of failing.
    overwrite_all = bool((payload or {}).get('overwrite'))
    result = {'ok': True, 'created': 0, 'skipped': 0, 'pending': 0, 'errors': [],
              'sheets': {'经营收集表': {'created': 0, 'skipped': 0, 'pending': 0, 'failed': 0}}}
    with CRM_LOCK:
        CRM_IMPORT_CONTEXT.active = True
        CRM_IMPORT_CONTEXT.batch = {'data': {}, 'original': {}, 'dirty': set()}
        try:
            for item in raw_rows:
                row_number = item.get('row', '') if isinstance(item, dict) else ''
                supplied = item.get('values', {}) if isinstance(item, dict) else {}
                if not isinstance(supplied, dict):
                    supplied = {}
                headers = (CRM_BUSINESS_EXTERNAL_SUBSIDIARY_HEADERS if '子公司全称' in supplied
                           else CRM_BUSINESS_EXTERNAL_HEADERS if '公司全称' in supplied
                           else CRM_BUSINESS_HEADERS if '手机' in supplied else CRM_BUSINESS_OLD_HEADERS)
                values = {header: str(supplied.get(header, '') or '').strip()[:2000] for header in headers}
                source_values = dict(values)
                values = _crm_normalize_legacy_business_values(values, headers)
                values = _crm_normalize_external_business_values(values, headers)
                try:
                    for header in headers:
                        if headers in CRM_BUSINESS_EXTERNAL_LAYOUTS and header in ('单位*', '姓名*'):
                            continue
                        if header.endswith('*') and not values[header]:
                            raise ValueError(f'{header[:-1]}必填')
                    row_overwrite = overwrite_all or bool(item.get('overwrite')) if isinstance(item, dict) else overwrite_all
                    missing_fields = _crm_import_missing_fields(source_values, values, headers)
                    if missing_fields:
                        _crm_import_completion_create(source_values, values, headers, missing_fields, username, row_number)
                        status = 'pending'
                    else:
                        status = crm_business_row(values, username, overwrite=row_overwrite)
                    result[status] += 1
                    result['sheets']['经营收集表'][status] += 1
                except (ValueError, TypeError) as exc:
                    result['errors'].append({'sheet': '经营收集表', 'row': row_number, 'error': str(exc),
                                             'headers': list(headers), 'values': values})
                    result['sheets']['经营收集表']['failed'] += 1
        finally:
            try:
                crm_flush_import_batch()
                _crm_flush_completion_batch()
            finally:
                CRM_IMPORT_CONTEXT.batch = None
                CRM_IMPORT_CONTEXT.active = False
            if result['created']:
                crm_on_change(username)
    return result


def crm_excel_row(sheet, v, username):
    def unique(collection, match, label, required=True):
        scoped_unit_id = v.get('_unit_id', '') if collection == 'units' else ''
        found = ([crm_load(username, collection)[scoped_unit_id]]
                 if scoped_unit_id and scoped_unit_id in crm_load(username, collection)
                 else [r for r in crm_load(username,collection).values() if match(r)])
        if len(found)>1: raise ValueError(label+'存在同名记录，无法唯一关联')
        if not found and required: raise ValueError(label+'不存在，请先填写对应工作表或建档')
        return found[0] if found else None
    def boolean(key):
        raw=v.get(key,'')
        if raw not in ('','是','否','true','false','1','0'): raise ValueError(key+'只能填写是或否')
        return raw in ('是','true','1')
    def enum(key, options):
        value=v.get(key,'')
        if value and value not in options: raise ValueError(key+'不在允许选项内')
        return value
    def save(fn,payload):
        result=fn(payload,username)
        if not result.get('ok'): raise ValueError(result.get('error','保存失败'))
        return 'created'
    unit=unique('units',lambda r:r.get('name')==v['单位名称*'],'单位',sheet!='单位')
    if sheet=='单位':
        enum('单位类型*',CRM_UNIT_TYPES)
        if unit: return 'skipped'
        return save(crm_unit_save,dict(name=v['单位名称*'],type=v['单位类型*'],region=v['城市'],business_region=v.get('经营区域',''),line_scale=v['线网规模'],bid_preference=v['招标偏好'],history=v['历史合作'],win_rate=v['历史中标率'],source=v.get('采集来源') or 'Excel收集'))
    person=unique('people',lambda r:r.get('unit_id')==unit['id'] and r.get('name')==v['姓名*'],'联系人',sheet!='联系人')
    if sheet=='联系人':
        role=enum('决策角色',CRM_DECISION_ROLES); stance=enum('立场',['支持','中立','反对','未知'])
        influence=int(v['影响力'] or 50)
        if not 0<=influence<=100: raise ValueError('影响力须为0到100整数')
        if person: return 'skipped'
        return save(crm_person_save,dict(department=v.get('所属部门',''),manager_id=v.get('直接上级ID',''),responsibilities=v.get('负责事项',''),name=v['姓名*'],unit_id=unit['id'],title=v['职务'],decision_role=role,power_radius=v['权力半径'],influence=influence,stance=stance or '未知',career_track=v['职业轨迹'].split('→') if v['职业轨迹'] else [],rotation_window=v['轮岗窗口']))
    member=v.get('我方对接人','')
    if sheet=='项目决策链':
        node=enum('决策节点*',CRM_DECISION_NODES)
        old=unique('links',lambda r:r.get('project_id')==v['项目编号*'] and r.get('person_id')==person['id'] and r.get('decision_node')==node and r.get('our_member')==member,'项目节点',False)
        if old: return 'skipped'
        return save(crm_link_save,dict(project_id=v['项目编号*'],person_id=person['id'],decision_node=node,our_member=member))
    relation=unique('relations',lambda r:r.get('person_id')==person['id'] and r.get('our_member')==member,'关系',sheet!='关系')
    if sheet=='关系':
        risk=boolean('单点风险')
        if relation: return 'skipped'
        return save(crm_relation_save,dict(person_id=person['id'],our_member=member,source=v['关系来源'],intro_path=v['引荐路径'],single_point_risk=risk))
    node=enum('关联决策节点',CRM_DECISION_NODES)
    for key in ('发生日期*','跟进日期'):
        if v[key]:
            try: datetime.strptime(v[key],'%Y-%m-%d')
            except ValueError: raise ValueError(key+'应为YYYY-MM-DD')
    if v['发生日期*']>crm_now()[:10]: raise ValueError('发生日期不能晚于今天')
    data=dict(relation_id=relation['id'],type=v['动作类型*'],happened_at=v['发生日期*'],decision_node=node,intel=v['情报产出'],cost=v['投入成本'],effective=boolean('有效动作'),bidirectional=boolean('客户主动'),project_id=v['项目编号'],next_step=v['下一步事项'],owner=v['负责人'],due_date=v['跟进日期'])
    if any(all(r.get(k,'')==value for k,value in data.items()) for r in crm_load(username,'actions').values()): return 'skipped'
    return save(crm_action_add,data)


CRM_SIMPLE_HEADERS = ['单位名称*', '单位类型*', '姓名*', '城市', '职务', '决策角色', '我方对接人', '认识渠道']


def crm_excel_simple_template():
    from openpyxl.comments import Comment
    book = Workbook()
    sheet = book.active
    sheet.title = '客户收集表'
    sheet.append(CRM_SIMPLE_HEADERS + ['经营区域'])
    sheet.freeze_panes = 'D2'
    sheet.auto_filter.ref = 'A1:I10001'
    tips = [
        '必填。一行一位联系人，同一单位可重复填写。请使用单位全称。',
        '必填。从下拉列表选择单位类型。',
        '必填。填写联系人姓名。前三列填完即可上传，其他列可留空。',
        '选填，例如上海。', '选填，例如项目经理。', '选填，从下拉列表选择；不确定就留空。',
        '选填。可留空（不指派，之后可在系统中指定）。', '选填。例如技术交流、同事引荐。', '选填，东区、西区、中区、广东区；按公司口径选择，不按城市自动推断。',
    ]
    for cell, tip in zip(sheet[1], tips):
        cell.font = Font(color='FFFFFF',bold=True)
        cell.fill = PatternFill('solid',fgColor='174B91' if cell.column <= 3 else '627891')
        cell.comment = Comment(tip + ' 导入只新增、不覆盖；同单位同名人员会关联已有记录，存在歧义则报错。', '填写提示')
        cell.alignment = Alignment(vertical='center')
        sheet.column_dimensions[cell.column_letter].width = 23
    sheet.row_dimensions[1].height = 30
    for column, options in [('B',CRM_UNIT_TYPES),('F',CRM_DECISION_ROLES),('I',CRM_BUSINESS_REGIONS)]:
        validation = DataValidation(type='list',formula1='"'+','.join(options)+'"',allow_blank=True)
        validation.showErrorMessage = True
        validation.error = '请从下拉列表选择'
        sheet.add_data_validation(validation); validation.add(f'{column}2:{column}10001')
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.print_title_rows = '1:1'
    sheet.oddHeader.center.text = '客户收集表：一行一位联系人，前三列必填，其他可留空'
    raw=io.BytesIO();book.save(raw);book.close();return raw.getvalue()


def crm_excel_simple_row(v, username, overwrite=False):
    # Validate references before creating any record for this row.
    area = v.get('经营区域','')
    if area and area not in CRM_BUSINESS_REGIONS: raise ValueError('经营区域请选择东区、西区、中区或广东区')
    if v['单位类型*'] not in CRM_UNIT_TYPES:
        raise ValueError('单位类型请从下拉列表选择')
    if v['决策角色'] and v['决策角色'] not in CRM_DECISION_ROLES:
        raise ValueError('决策角色请从下拉列表选择，不确定可留空')
    scoped_unit_id = v.get('_unit_id', '')
    units = ([crm_load(username, 'units')[scoped_unit_id]]
             if scoped_unit_id and scoped_unit_id in crm_load(username, 'units')
             else [u for u in crm_load(username,'units').values() if u.get('name')==v['单位名称*']])
    if len(units)>1: raise ValueError('单位存在同名记录，请在系统中核实')
    overwritten=False
    if units:
        existing=units[0]
        # A customer unit may legitimately serve more than one business region.
        # Add the incoming region to its membership instead of treating it as a
        # conflicting overwrite of the legacy primary-region field.
        if area and area not in crm_unit_regions(existing):
            res=crm_unit_save({**existing, 'business_regions': [*crm_unit_regions(existing), area]},username)
            if not res.get('ok'): raise ValueError(res.get('error','新增经营区域失败'))
            existing=res['unit']; overwritten=True
        if existing.get('type') != v['单位类型*']:
            if overwrite:
                res=crm_unit_save({**existing,'type':v['单位类型*']},username)
                if not res.get('ok'): raise ValueError(res.get('error','覆盖单位类型失败'))
                existing=res['unit']; overwritten=True
            else:
                raise ValueError('单位类型与现有档案不一致，请核实后填写')
        people=[p for p in crm_load(username,'people').values() if p.get('unit_id')==existing['id'] and p.get('name')==v['姓名*']]
        if len(people)>1: raise ValueError('同单位存在同名联系人，请在系统中核实')
        if people:
            relations=[r for r in crm_load(username,'relations').values() if r.get('person_id')==people[0]['id'] and r.get('our_member')==v.get('我方对接人','')]
            if len(relations)>1: raise ValueError('已有多条相同对接关系，请在系统中核实')
    statuses=[]
    sheets = ([] if scoped_unit_id else [('单位', {'采集来源':v.get('采集来源','Excel收集'), '经营区域':area})])
    sheets += [('联系人', {key:v.get(key,'') for key in ('所属部门','直接上级ID','负责事项')}), ('关系', {'关系来源':v['认识渠道']})]
    for sheet, extra in sheets:
        values={header:v.get(header,'') for header in CRM_EXCEL_SHEETS[sheet]}
        values.update(extra)
        if scoped_unit_id:
            values['_unit_id'] = scoped_unit_id
        statuses.append(crm_excel_row(sheet,values,username))
    return 'created' if ('created' in statuses or overwritten) else 'skipped'


def crm_quick_create(payload, username):
    """Both capture versions share the corresponding workbook validation."""
    if payload.get("format") == "business":
        return crm_business_capture(payload, username)
    mapping = {'单位名称*':'unit_name', '单位类型*':'unit_type', '姓名*':'name', '城市':'city', '经营区域':'business_region',
               '职务':'title', '决策角色':'role', '我方对接人':'our_member', '认识渠道':'source', '所属部门':'department', '直接上级ID':'manager_id', '负责事项':'responsibilities'}
    values = {}
    for header, field in mapping.items():
        value = payload.get(field, '')
        if not isinstance(value, str) or len(value)>500:
            return {'ok':False, 'error':'字段须为500字以内的文本'}
        values[header] = value.strip()
        if header.endswith('*') and not values[header]:
            return {'ok':False, 'error':header[:-1]+'必填'}
    values['采集来源'] = '快速建档'
    with CRM_LOCK:
        CRM_IMPORT_CONTEXT.active = True
        try:
            if values.get('直接上级ID') and values['直接上级ID'] not in crm_load(username, 'people'):
                raise ValueError('直接上级不存在或不属于当前账号')
            status = crm_excel_simple_row(values, username)
        except (ValueError, TypeError) as exc:
            return {'ok':False, 'error':str(exc)}
        finally:
            CRM_IMPORT_CONTEXT.active = False
            crm_on_change(username)
    return {'ok':True, 'status':status}

# Company collection format: one contact per row, keeping the regional working order.
# 联系方式可选；页面与新模板均使用这一组字段。保留旧表头以兼容已下载模板。
CRM_BUSINESS_CONTACT_LEGACY_HEADERS = ['区域', '单位*', '部门', '职务', '姓名*', '重要程度', '支持程度', '跟进轨迹', '上次互动时间', '本周互动计划', '计划跟进日期', '单位类型', '城市', '维系周期（天）', '我方姓名']
CRM_BUSINESS_HEADERS = CRM_BUSINESS_CONTACT_LEGACY_HEADERS[:5] + ['手机', '办公电话', '邮箱', '微信 / 其他联系'] + CRM_BUSINESS_CONTACT_LEGACY_HEADERS[5:]
CRM_BUSINESS_LEGACY_HEADERS = CRM_BUSINESS_CONTACT_LEGACY_HEADERS[:11] + ['我方对接人'] + CRM_BUSINESS_CONTACT_LEGACY_HEADERS[11:] + ['我方公司', '我方部门', '我方岗位']
CRM_BUSINESS_OLD_HEADERS = CRM_BUSINESS_LEGACY_HEADERS[:15]


def crm_business_template():
    from openpyxl.comments import Comment
    book = Workbook()
    sheet = book.active; sheet.title = '经营收集表'
    # The downloadable form deliberately follows the external regional
    # collection schema. It exposes the three direct-import gate fields in
    # the first row so an uploader does not need to learn an internal alias.
    headers = CRM_BUSINESS_EXTERNAL_HEADERS
    sheet.append(headers)
    tips = [
        '东区、西区、中区、广东区；不清楚可留空，按公司口径选择。',
        '关键项：填写工商/业务使用的公司全称。公司全称、姓名、单位类型齐全时可直接入库；缺项会进入资料补充待办。',
        '选填，供识别简称使用；不替代公司全称。',
        '所属部门，例如设备部。', '实际职务，例如总经理。',
        '关键项：联系人姓名。缺项会推送至“我方姓名”对应账号补充。',
        '选填，例如 13800000000。', '选填，例如 020-12345678。', '选填，例如 name@example.com。', '选填，例如微信号、企业微信或其他方便联系的渠道。',
        '重要或一般；未填写维系周期时分别建议14天或21天，导入后可编辑。',
        '选择反对、中立、支持或积极支持；不确定留空，不转换为系统关系强度。',
        '可分行记录已发生的跟进，原文存入联系人历史备注；不自动拆分成多次有效动作。',
        '真实日期，YYYY-MM-DD；不可晚于今天。填写后生成一次历史互动记录，不标记有效或客户主动。可用 / 表示暂无。',
        '未来计划，导入后形成待办，不作为已发生动作计分。',
        '计划的实际日期 YYYY-MM-DD；可留空待定，可填过去日期以保留逾期待办。',
        '关键项：必须从下拉项选择。缺项会进入资料补充待办，补齐后才会写入客户库。',
        '选填，例如广东省。', '城市独立于经营区域，例如广州。',
        '选填1到365天；填写时优先于重要程度建议，达到周期预警，达到两倍周期告警。',
        '填写我方人员中文姓名。系统精确匹配其拼音账号并将缺项资料推送给该账号；无法匹配时由管理员分派。',
        '选填，我方人员所属公司或组织，例如联合指挥中心。',
        '选填，我方人员所属部门，例如东区作战平台。',
        '选填，我方人员岗位，例如泛交通BD。',
    ]
    widths = [12,32,20,18,18,16,16,18,28,24,14,14,46,20,46,20,18,14,16,18,16,22,22,18]
    for i,(cell,tip,width) in enumerate(zip(sheet[1],tips,widths),1):
        cell.fill=PatternFill('solid',fgColor='174B91' if i <= 4 else '267E82' if i <= 8 else '8B6725' if i <= 11 else '627891')
        cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True,vertical='center')
        cell.comment=Comment(tip,'填写说明');sheet.column_dimensions[cell.column_letter].width=width
    sheet.row_dimensions[1].height=36
    sheet.freeze_panes='G2';sheet.auto_filter.ref='A1:X10001'
    # Named ranges are recognised by both Excel and WPS. Inline list strings
    # occasionally lose their visible arrow in WPS after a browser download.
    options_sheet = book.create_sheet('下拉选项')
    option_sets = [('经营区域', CRM_BUSINESS_REGIONS), ('重要程度', ['重要','一般']), ('支持程度', CRM_SUPPORT_LABELS), ('单位类型', CRM_UNIT_TYPES)]
    for column, (name, options) in enumerate(option_sets, 1):
        options_sheet.cell(1, column, name)
        for row, option in enumerate(options, 2): options_sheet.cell(row, column, option)
        letter = options_sheet.cell(1, column).column_letter
        book.defined_names.add(DefinedName(f'CRM_{name}', attr_text=f"'下拉选项'!${letter}$2:${letter}${len(options)+1}"))
    options_sheet.sheet_state = 'hidden'
    for col,name in [('A','经营区域'),('K','重要程度'),('L','支持程度'),('Q','单位类型')]:
        validation=DataValidation(type='list',formula1=f'=CRM_{name}',allow_blank=True)
        validation.showDropDown = False
        validation.showErrorMessage=True;validation.error='请从下拉选项选择';sheet.add_data_validation(validation);validation.add(f'{col}2:{col}10001')
    days=DataValidation(type='whole',operator='between',formula1=1,formula2=365,allow_blank=True)
    days.showErrorMessage=True;days.error='填写1到365之间的整数';sheet.add_data_validation(days);days.add('T2:T10001')
    for row in range(2,502):
        for col in (14,16): sheet.cell(row,col).number_format='yyyy-mm-dd'
        for col in (13,15): sheet.cell(row,col).alignment=Alignment(wrap_text=True,vertical='top')
    guide=book.create_sheet('填写说明')
    instructions=['客户经营收集表 · 东区、西区、中区、广东区', '只填“经营收集表”；一行一位联系人。公司全称、姓名、单位类型三项齐全才会直接入库。',
        '公司全称、姓名、单位类型任一缺项时，记录不会进入客户库，而会生成“资料补充”待办；系统按“我方姓名”精确匹配中文姓名对应的拼音账户推送。',
        '未匹配到我方姓名对应账号的待办由客户经营管理员或超级管理员查看、补充并指定接收账号；补齐三项后才写入该账号客户库。',
        '蓝色：组织与人员；青色：关系判断和历史；金色：未来计划；灰色：补充资料。',
        '跟进轨迹保留原文，不猜日期、不虚构有效互动。支持程度使用中文选项，旧数字编码自动兼容。',
        '上次互动时间生成历史互动；本周互动计划生成独立待办，完成待办不会自动计成客户触达。',
        '重要默认14天、一般默认21天维系；填写具体周期则以该周期为准。未填重要程度和周期时沿用系统层级预警。',
        '重复导入跳过已存在内容；已有联系人只补空字段，已有非空内容不覆盖。需要修改时到客户档案编辑。',
        '同名单位类型/区域冲突先核实；歧义记录报错，不自动合并。有效行直接导入，错误行按结果修正后可重传。',
        '手机、办公电话、邮箱、微信 / 其他联系均为选填；仅支持未加密.xlsx、最多5MB；不使用公式，不改工作表名和表头。旧八列、九列、旧经营收集表与进阶模板仍支持。',
        '本表仅整理业务资料，不填写密码等秘密。仅由具备导入权限的账号上传。']
    instructions.append('我方姓名用于匹配已有账号；我方公司、我方部门、我方岗位可选填，并会分别保存为独立组织字段，不能合并写在姓名单元格中。旧版表格仍可导入。')
    for line in instructions:guide.append([line])
    guide.column_dimensions['A'].width=110
    for row in guide: row[0].alignment=Alignment(wrap_text=True,vertical='center');guide.row_dimensions[row[0].row].height=36
    sample=book.create_sheet('示例（不导入）');sample.append([c.value for c in sheet[1]])
    sample.append(['东区','示例轨道集团有限公司','示例轨道集团','设备部','技术负责人','张工','13800000000','020-12345678','zhang@example.com','张工微信','重要','支持','2026-01-10 技术交流，确认接口事项','2026-01-10','确认项目接口资料','','业主单位','广东省','广州',14,'李经理','联合指挥中心','东区作战平台','客户经理'])
    for i,width in enumerate(widths,1):sample.column_dimensions[sample.cell(1,i).column_letter].width=width
    from copy import copy
    for i in range(1, len(headers)+1):
        sample.cell(1,i)._style = copy(sheet.cell(1,i)._style)
        sample.cell(1,i).comment = copy(sheet.cell(1,i).comment)
        sample.cell(2,i).alignment = Alignment(wrap_text=True,vertical='top')
        sample.cell(2,i).fill = PatternFill('solid',fgColor='EAF4FF')
    sample.row_dimensions[1].height=36;sample.row_dimensions[2].height=72
    sample.freeze_panes='G2';sample.sheet_properties.tabColor='F5C242'
    sample.cell(4,1,'这是填写实例，不会导入。请切换到“经营收集表”填写真实记录；不确定的信息留空。')
    sample.merge_cells(start_row=4,start_column=1,end_row=4,end_column=len(headers))
    sample.cell(4,1).font=Font(bold=True,color='174B91')
    sample.cell(5,1,'公司全称、姓名、单位类型是直接入库关键项；缺项会推送到“我方姓名”中文匹配的账号补充，补齐后再入库。')
    sample.merge_cells(start_row=5,start_column=1,end_row=5,end_column=len(headers))
    book.active=book.sheetnames.index('示例（不导入）')
    out=io.BytesIO();book.save(out);book.close();return out.getvalue()


# 经营管理员的三级导出：联合指挥（全部）、区域指挥（区域）、客户经理（个人）。
# 每一级都产出同一张客户关系表，区别只在范围过滤，便于跨级核对。
CRM_EXPORT_LEVELS = {
    'all': {'name': '联合指挥', 'scope': '全部客户关系'},
    'region': {'name': '区域指挥', 'scope': '区域客户关系'},
    'manager': {'name': '客户经理', 'scope': '个人客户关系'},
}
CRM_RELATIONSHIP_EXPORT_HEADERS = [
    '经营区域', '客户单位全称', '单位类型', '客户联系人', '部门', '职务', '决策角色',
    '重要程度', '支持程度', '我方对接人', '我方公司', '我方部门', '我方岗位', '关系强度', '最近互动',
    '未互动天数', '单点风险', '待跟进事项', '跟进日期',
    # 与资料补充表保持一致：联系方式置于最右侧。
    '手机', '办公电话', '邮箱', '微信 / 其他联系',
]


def crm_relationship_export(caller, level='all', target='', region=''):
    """Build a customer-relationship workbook for the requested command tier.

    Only the operating administrator (联合指挥) may export; the same person can
    narrow the sheet to a region (区域指挥) or a single account manager
    (客户经理·个人). Importance and support level are always carried through so
    the exported table matches what every in-app customer panel now shows.
    """
    caps = crm_capabilities(caller)
    if not (caps.get('manage') or caller.get('role') == 'superadmin'):
        raise PermissionError('只有客户经营管理员或超级管理员可导出客户关系表')
    if level not in CRM_EXPORT_LEVELS:
        raise ValueError('导出层级无效，请选择联合指挥、区域指挥或客户经理')
    target = str(target or '').strip()
    region = str(region or '').strip()
    if level == 'region' and target not in CRM_BUSINESS_REGIONS:
        raise ValueError('请选择要导出的经营区域')
    if level == 'manager' and not target:
        raise ValueError('请选择要导出的客户经理')
    if region and region not in CRM_BUSINESS_REGIONS:
        raise ValueError('请选择有效的经营区域')
    # A regional export historically carries its region in target. Keep that
    # compatible while allowing personal exports to apply the same region gate.
    if level == 'region':
        region = target
    if target.startswith('liaison:'):
        target = target.split(':', 1)[1]

    units = {u['id']: u for u in crm_read_units(caller)['units']}
    people = {p['id']: p for p in crm_read_people(caller)['people']}
    relations = crm_read_relations(caller)['relations']
    plans_by_relation = {}
    for relation in relations:
        pending = [a for a in relation.get('actions', []) if a.get('next_step') and not a.get('completed_at')]
        pending.sort(key=lambda a: a.get('due_date') or '9999-12-31')
        if pending:
            plans_by_relation[relation['id']] = pending

    rows, scope_target_label = [], target
    for relation in relations:
        person = people.get(relation.get('person_id'), {})
        unit = units.get(person.get('unit_id'), {})
        regions = crm_unit_regions(unit) or []
        staff = relation.get('our_staff') or {}
        if region and region not in regions:
            continue
        if level == 'manager' and staff.get('username') != target and str(relation.get('our_member') or '').strip() != target:
            continue
        if level == 'manager' and staff.get('name'):
            scope_target_label = staff.get('name')
        plan = plans_by_relation.get(relation['id'], [{}])[0]
        rows.append([
            '、'.join(regions) or '未分区',
            person.get('company_full_name') or unit.get('company_full_name') or unit.get('name', ''),
            unit.get('type', ''),
            person.get('name') or relation.get('person_name', ''),
            person.get('department', ''),
            person.get('title', ''),
            person.get('decision_role', ''),
            person.get('importance', '') or '待补充',
            person.get('support_level', '') or '待补充',
            staff.get('name') or relation.get('our_member') or '尚未指派',
            staff.get('company', ''),
            staff.get('department', ''),
            staff.get('title', ''),
            f"L{relation.get('strength', 0)} {relation.get('strength_label', '')}".strip(),
            relation.get('last_touch') or '暂无记录',
            '暂无' if relation.get('days_since_touch') is None else relation.get('days_since_touch'),
            '单点风险' if relation.get('single_point_risk') else '',
            plan.get('next_step', ''),
            plan.get('due_date', ''),
            person.get('mobile', ''),
            person.get('office_phone', ''),
            person.get('email', ''),
            person.get('wechat', ''),
        ])
    rows.sort(key=lambda r: (r[0], r[1], r[3]))

    meta = CRM_EXPORT_LEVELS[level]
    book = Workbook()
    sheet = book.active
    sheet.title = '客户关系'
    scope_label = scope_target_label if level != 'region' else region
    subtitle = meta['scope'] if level == 'all' else f"{meta['scope']}：{scope_label}"
    if level == 'manager' and region:
        subtitle += f"（{region}）"
    title = f"客户关系导出 · {meta['name']}（{subtitle}）"
    sheet.append([title] + [''] * (len(CRM_RELATIONSHIP_EXPORT_HEADERS) - 1))
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(CRM_RELATIONSHIP_EXPORT_HEADERS))
    sheet.cell(1, 1).font = Font(bold=True, color='FFFFFF', size=13)
    sheet.cell(1, 1).fill = PatternFill('solid', fgColor='15499C')
    sheet.cell(1, 1).alignment = Alignment(vertical='center')
    sheet.row_dimensions[1].height = 30
    sheet.append(CRM_RELATIONSHIP_EXPORT_HEADERS)
    widths = [12, 34, 12, 14, 16, 18, 16, 10, 12, 18, 22, 22, 18, 16, 20, 12, 12, 40, 14, 16, 18, 28, 24]
    for column, (cell, width) in enumerate(zip(sheet[2], widths), 1):
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = PatternFill('solid', fgColor='174B91' if column <= 7 else '267E82' if column <= 9 else '627891')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.row_dimensions[2].height = 30
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = 'A3'
    sheet.auto_filter.ref = f'A2:{sheet.cell(2, len(CRM_RELATIONSHIP_EXPORT_HEADERS)).column_letter}{len(rows) + 2}'
    for offset in range(len(rows)):
        for column in (18,):
            sheet.cell(3 + offset, column).alignment = Alignment(wrap_text=True, vertical='top')
    if not rows:
        sheet.cell(3, 1, '当前范围暂无客户关系记录。')
    filename_target = scope_label if level != 'all' else ''
    if level == 'manager' and region:
        filename_target = f"{region}-{scope_target_label}"
    filename = f"客户关系-{meta['name']}{('' if level == 'all' else '-' + filename_target)}.xlsx"
    out = io.BytesIO(); book.save(out); book.close()
    return out.getvalue(), filename


def crm_customer_organization_export(caller):
    """Export the visible customer hierarchy and its contacts as separate sheets."""
    caps = crm_capabilities(caller)
    if not (caps.get('manage') or caller.get('role') == 'superadmin'):
        raise PermissionError('只有客户经营管理员或超级管理员可导出客户组织架构')

    snapshot = crm_read_snapshot(caller)
    units = snapshot['units']
    people = snapshot['people']
    unit_by_id = {unit.get('id'): unit for unit in units if unit.get('id')}

    def organization_chain(unit):
        chain, seen = [], set()
        current = unit
        while current and current.get('id') not in seen:
            seen.add(current.get('id'))
            chain.append(current)
            current = unit_by_id.get(current.get('parent_unit_id', ''))
        return list(reversed(chain))

    people_by_unit = {}
    for person in people:
        people_by_unit.setdefault(person.get('unit_id', ''), []).append(person)

    organization_rows = []
    for unit in units:
        chain = organization_chain(unit)
        path = ' / '.join(item.get('name', '') for item in chain if item.get('name')) or unit.get('name', '')
        parent = chain[-2].get('name', '') if len(chain) > 1 else ''
        organization_rows.append([
            len(chain) or 1, unit.get('name', ''), parent,
            unit.get('organization_level', '') or ('单位' if parent else '客户组织'),
            unit.get('type', ''), '、'.join(crm_unit_regions(unit)), unit.get('region', ''), path,
            len(people_by_unit.get(unit.get('id', ''), [])),
        ])
    organization_rows.sort(key=lambda row: (row[7], row[0], row[1]))

    contact_rows = []
    for person in people:
        unit = unit_by_id.get(person.get('unit_id'), {})
        chain = organization_chain(unit)
        path = ' / '.join(item.get('name', '') for item in chain if item.get('name')) or unit.get('name', '')
        contact_rows.append([
            unit.get('name', ''), path,
            unit.get('organization_level', '') or ('单位' if unit.get('parent_unit_id') else '客户组织'),
            person.get('name', ''), person.get('department', ''), person.get('title', ''),
            person.get('decision_role', ''), person.get('importance', ''), person.get('support_level', ''),
            person.get('mobile') or person.get('office_phone') or person.get('email') or person.get('wechat') or '',
        ])
    contact_rows.sort(key=lambda row: (row[1], row[3]))

    book = Workbook()

    def write_sheet(sheet, title, headers, rows, widths):
        sheet.append([title] + [''] * (len(headers) - 1))
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        sheet.cell(1, 1).font = Font(bold=True, color='FFFFFF', size=13)
        sheet.cell(1, 1).fill = PatternFill('solid', fgColor='15499C')
        sheet.cell(1, 1).alignment = Alignment(vertical='center')
        sheet.row_dimensions[1].height = 30
        sheet.append(headers)
        for column, (cell, width) in enumerate(zip(sheet[2], widths), 1):
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='174B91' if column <= 4 else '267E82')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            sheet.column_dimensions[cell.column_letter].width = width
        sheet.row_dimensions[2].height = 30
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = 'A3'
        sheet.auto_filter.ref = f'A2:{sheet.cell(2, len(headers)).column_letter}{len(rows) + 2}'
        for row_number in range(3, len(rows) + 3):
            for column in range(1, len(headers) + 1):
                sheet.cell(row_number, column).alignment = Alignment(wrap_text=True, vertical='top')
        if not rows:
            sheet.cell(3, 1, '当前范围暂无客户组织数据。')

    organization_headers = ['层级', '客户组织名称', '上级组织', '组织层级', '单位类型', '经营区域', '所在地', '组织路径', '联系人数量']
    write_sheet(book.active, '客户组织架构导出', organization_headers, organization_rows,
                [8, 34, 30, 14, 14, 16, 16, 60, 14])
    book.active.title = '客户组织架构'
    contacts = book.create_sheet('客户联系人')
    contact_headers = ['客户组织名称', '客户组织路径', '客户单位层级', '联系人', '部门', '职务', '决策角色', '重要程度', '支持程度', '联系方式']
    write_sheet(contacts, '客户组织联系人清单', contact_headers, contact_rows,
                [30, 60, 14, 16, 20, 20, 18, 12, 14, 28])
    out = io.BytesIO(); book.save(out); book.close()
    return out.getvalue(), '客户组织架构.xlsx'


def crm_business_row(v, username, overwrite=False):
    staff_payload = None
    member = v.get('我方对接人') or v.get('我方姓名') or ''
    v['我方对接人'] = member
    fields = {field:v.get(header,'') for header,field in [('我方姓名','name'),('我方公司','company'),('我方部门','department'),('我方岗位','title')]}
    if any(fields.values()):
        existing = crm_member_profile(username, member)
        if existing:
            for field,value in fields.items():
                if value and existing.get(field) and existing[field] != value:
                    raise ValueError('我方人员资料与档案冲突，请在我方组织核实姓名、公司、部门和岗位后修改')
        staff_payload = {**(existing or {}), 'member_key':member, **{k:v for k,v in fields.items() if v}}
        staff_payload.setdefault('name', member)
        check = crm_staff_save(staff_payload, username, dry_run=True)
        if not check['ok']: raise ValueError(check['error'])
        if existing and all(not value or existing.get(field) == value for field,value in fields.items()):
            staff_payload = None
    importance=v['重要程度']; support=crm_support_label(v['支持程度']); cadence=v['维系周期（天）']
    if importance not in ('','重要','一般'): raise ValueError('重要程度请选择重要或一般')
    if support not in ('', *CRM_SUPPORT_LABELS): raise ValueError('支持程度请选择反对、中立、支持或积极支持，不清楚可留空')
    if cadence and (not cadence.isdigit() or not 1 <= int(cadence) <= 365): raise ValueError('维系周期须为1到365天')
    for key in ('上次互动时间','计划跟进日期'):
        if v[key]:
            try:
                # Accept Excel dates and common fully specified dates; never infer short years.
                parsed=datetime.strptime(v[key].replace('/','-'),'%Y-%m-%d')
                if len(v[key].split('/')[0].split('-')[0]) != 4: raise ValueError()
                v[key]=parsed.strftime('%Y-%m-%d')
            except ValueError: raise ValueError(key+'请填写完整日期 YYYY-MM-DD')
    if v['上次互动时间'] > crm_now()[:10]: raise ValueError('上次互动时间不能晚于今天')
    if v['计划跟进日期'] and not v['本周互动计划']: raise ValueError('填写计划日期时请同时填写本周互动计划')
    # New regional collection sheets carry a customer organization path:
    # company full name → subsidiary → unit → person department.  Only the
    # first three are unit records; department stays on the person so all
    # existing relation/action references remain stable.
    company_full_name = str(v.get('_company_full_name', '') or '').strip()
    subsidiary_full_name = str(v.get('_subsidiary_full_name', '') or '').strip()
    declared_unit_name = str(v.get('_unit_name', '') or '').strip()
    hierarchy_created = False
    simple={h:'' for h in CRM_SIMPLE_HEADERS}
    simple.update({'单位名称*':v['单位*'],'单位类型*':v['单位类型'] or '业主单位','姓名*':v['姓名*'],
        '城市':v['城市'],'职务':v['职务'],'我方对接人':v['我方对接人'],'经营区域':v['区域'],'所属部门':v['部门'],
        '决策角色':v.get('决策角色',''),'认识渠道':v.get('认识渠道',''),'直接上级ID':v.get('直接上级ID',''),'负责事项':v.get('负责事项',''),'采集来源':v.get('采集来源','Excel收集')})
    if simple['直接上级ID'] and simple['直接上级ID'] not in crm_load(username,'people'): raise ValueError('直接上级不存在或不属于当前账号')
    if company_full_name:
        levels = [('company', company_full_name)]
        if subsidiary_full_name and subsidiary_full_name != levels[-1][1]:
            levels.append(('subsidiary', subsidiary_full_name))
        if declared_unit_name and declared_unit_name != levels[-1][1]:
            levels.append(('unit', declared_unit_name))
        parent_id = ''
        unit = None
        for index, (level, name) in enumerate(levels):
            hierarchy = {
                'parent_unit_id': parent_id, 'organization_level': level,
                'company_full_name': company_full_name,
                'subsidiary_full_name': subsidiary_full_name if level in ('subsidiary', 'unit') else '',
                'unit_name': declared_unit_name if level == 'unit' else '',
            }
            # City describes the concrete operating organization, not the
            # corporate root.  If the workbook only declares a company, that
            # single node remains the operating organization and keeps city.
            node_city = v['城市'] if index == len(levels) - 1 else ''
            # A short unit name is only unique below its parent organization.
            # For example, several metro companies can each have a “轨道集团”
            # unit without being merged into one customer record.
            matches = [item for item in crm_load(username, 'units').values()
                       if item.get('name') == name
                       and item.get('parent_unit_id', '') == parent_id
                       and item.get('organization_level', '') == level]
            if len(matches) > 1:
                raise ValueError(f'{name}在当前组织下存在重复档案，请在客户档案中核实')
            if matches:
                target = matches[0]
                if v['区域'] and v['区域'] not in crm_unit_regions(target):
                    updated = crm_unit_save({**target, 'business_regions': [*crm_unit_regions(target), v['区域']]}, username)
                    if not updated.get('ok'):
                        raise ValueError(updated.get('error', f'{level}层级更新失败'))
                    target = updated['unit']; hierarchy_created = True
            else:
                created = crm_unit_save({
                    'name': name, 'type': v['单位类型'] or '业主单位',
                    'region': node_city, 'business_region': v['区域'],
                    'source': v.get('采集来源', 'Excel收集'), **hierarchy,
                }, username)
                if not created.get('ok'):
                    raise ValueError(created.get('error', f'{level}层级建档失败'))
                target = created['unit']; hierarchy_created = True
            parent_id = target['id']
            unit = target
        simple['_unit_id'] = unit['id']
        # Keep one canonical unit type; all assigned business regions stay on
        # the same organization record in `business_regions`.
        simple['单位类型*'] = unit.get('type') or simple['单位类型*']
        simple['经营区域'] = unit.get('business_region') or simple['经营区域']
        status = crm_excel_simple_row(simple, username, overwrite=overwrite)
    else:
        status=crm_excel_simple_row(simple,username,overwrite=overwrite)
        unit=next(u for u in crm_load(username,'units').values() if u.get('name')==v['单位*'])
    person=next(p for p in crm_load(username,'people').values() if p.get('unit_id')==unit['id'] and p.get('name')==v['姓名*'])
    additions={'department':v['部门'],'title':v['职务'],'mobile':v.get('手机',''),'office_phone':v.get('办公电话',''),'email':v.get('邮箱',''),'wechat':v.get('微信 / 其他联系',''),'importance':importance,'support_level':support,
               'maintenance_days':int(cadence) if cadence else (14 if importance=='重要' else 21 if importance=='一般' else ''), 'imported_history':v['跟进轨迹']}
    updates={k:value for k,value in additions.items() if value != '' and person.get(k,'') == ''}
    if updates:
        result=crm_person_save({**person,**updates},username)
        if not result['ok']: raise ValueError(result['error'])
        status='created'
    member=v['我方对接人']
    relation=next(r for r in crm_load(username,'relations').values() if r.get('person_id')==person['id'] and r.get('our_member')==member)
    records=[]
    if v['上次互动时间']: records.append({'relation_id':relation['id'],'type':'历史互动（表格录入）','happened_at':v['上次互动时间'],'effective':False,'bidirectional':False})
    if v['本周互动计划']: records.append({'relation_id':relation['id'],'type':'跟进计划','record_kind':'plan','next_step':v['本周互动计划'],'due_date':v['计划跟进日期'],'owner':member})
    for data in records:
        if any(all(a.get(k,'')==value for k,value in data.items()) for a in crm_load(username,'actions').values()): continue
        result=crm_action_add(data,username)
        if not result['ok']: raise ValueError(result['error'])
        status='created'
    if staff_payload:
        result = crm_staff_save(staff_payload, username)
        if not result['ok']: raise ValueError(result['error'])
        status='created'
    crm_recompute_relation(username,relation['id'])
    return 'created' if hierarchy_created else status


def crm_business_capture(payload, username):
    mapping = dict(zip(CRM_BUSINESS_HEADERS, ['business_region','unit_name','department','title','name','mobile','office_phone','email','wechat','importance','support_level','history','last_interaction','weekly_plan','due_date','unit_type','city','maintenance_days','our_name']))
    mapping.update({'决策角色':'role','认识渠道':'source','直接上级ID':'manager_id','负责事项':'responsibilities'})
    values = {}
    for header,field in mapping.items():
        value = payload.get(field,'')
        if not isinstance(value,str) or len(value)>2000:
            return {'ok':False,'error':'字段须为2000字以内的文本'}
        values[header]=value.strip()
        if header.endswith('*') and not values[header]:
            return {'ok':False,'error':header[:-1]+'必填'}
    company_full_name = payload.get('company_full_name', '')
    subsidiary_full_name = payload.get('subsidiary_full_name', '')
    if not isinstance(company_full_name, str) or not isinstance(subsidiary_full_name, str):
        return {'ok':False,'error':'客户组织字段须为文本'}
    if len(company_full_name) > 200 or len(subsidiary_full_name) > 200:
        return {'ok':False,'error':'客户组织名称不能超过200字'}
    if company_full_name.strip():
        values['_company_full_name'] = company_full_name.strip()
        values['_subsidiary_full_name'] = subsidiary_full_name.strip()
        values['_unit_name'] = values['单位*']
    for header, field in [('我方对接人', 'our_member'), ('我方公司', 'our_company'), ('我方部门', 'our_department'), ('我方岗位', 'our_position')]:
        value = payload.get(field, '')
        if not isinstance(value, str) or len(value) > 2000:
            return {'ok':False,'error':'字段须为2000字以内的文本'}
        values[header] = value.strip()
    values['采集来源']='页面收集'
    with CRM_LOCK:
        CRM_IMPORT_CONTEXT.active=True
        try:
            status=crm_business_row(values,username)
        except (ValueError,TypeError) as exc:
            return {'ok':False,'error':str(exc)}
        finally:
            CRM_IMPORT_CONTEXT.active=False
            crm_on_change(username)
    return {'ok':True,'status':status}
