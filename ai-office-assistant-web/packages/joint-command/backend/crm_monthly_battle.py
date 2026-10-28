"""Monthly task deployment with append-only progress and observed stage timing."""
import threading
import hashlib

MONTHLY_BATTLE_LOCK = threading.RLock()
MONTHLY_BATTLE_STAGES = ['初步设计阶段', '项目概算阶段', '初步设计评审阶段', '用户需求书引导阶段', '控制价确认阶段', '招标时间确认阶段', '招标公示阶段', '招投标阶段', '中标谈判（第三方谈判）', '合同条款沟通阶段', '合同签订']


def monthly_battle_action_cycles(item, now=None):
    now = now or datetime.now()
    cadence = 2
    created_at = datetime.fromisoformat(item.get('created_at') or now.isoformat())
    action_times = [datetime.fromisoformat(event['at']) for event in item.get('events', [])
                    if event.get('at') and event.get('event_type') != 'stage_change']
    cycles = []
    for offset in range(6, -1, -1):
        end = now.replace(hour=0, minute=0, second=0, microsecond=0) + timedelta(days=1 - offset * cadence)
        start = end - timedelta(days=cadence)
        status = 'gray' if end <= created_at else ('green' if any(start <= action < end and action <= now for action in action_times) else 'red')
        display_end = end - timedelta(days=1)
        cycles.append({'status': status, 'label': f'{start.month:02d}/{start.day:02d}-{display_end.month:02d}/{display_end.day:02d}'})
    return cycles


def monthly_battle_template(user):
    if not crm_capabilities(user).get('read'):
        raise PermissionError('无查看权限')
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    book = Workbook()
    sheet = book.active
    sheet.title = '月度任务令'
    headers = ['项目归属', '客户经理', '子项目编码', '机会点名称', '任务令作战部署的目标/动作', '赢率', '可参与份额-总额(万元)', '合同预计签订时间', '毛利率', '管理订货', '毛利', '税率', '未税金额', '月度', '原：保底/冲刺/推进', '原冲刺月份', '现调整后-冲刺月份', '现：保底/冲刺/推进', '确认主要责任人', '备注']
    sheet.append(headers)
    sheet.freeze_panes = 'D2'
    sheet.auto_filter.ref = 'A1:T501'
    for index, header in enumerate(headers, 1):
        cell = sheet.cell(1, index)
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = PatternFill('solid', fgColor='245C99')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        sheet.column_dimensions[get_column_letter(index)].width = 42 if index in (4, 5, 20) else 23
    sheet.row_dimensions[1].height = 38
    stages = book.create_sheet('阶段字典')
    stages.append(['阶段编号', '阶段名称', '填写格式'])
    for index, label in enumerate(MONTHLY_BATTLE_STAGES, 1):
        stages.append([index, label, f'阶段{index}-{label}'])
    stages.column_dimensions['B'].width = 32
    stages.column_dimensions['C'].width = 40
    validation = DataValidation(type='list', formula1="'阶段字典'!$C$2:$C$12", allow_blank=True)
    sheet.add_data_validation(validation)
    validation.add('E2:E501')
    guide = book.create_sheet('填写说明')
    for line in ['先选择上传月份，再上传本表；任务归属月份以上传窗口为准。', '机会点名称、子项目编码必须填写；同月同编码更新原记录并保留推进历史。', '客户经理、确认主要责任人填写平台真实姓名或账号，用于关联个人权限。', '项目归属填写东区-浙北、东区-浙南、东区-江苏、西区、中区或广东。', '目标/动作可选阶段，也可填写含“阶段7”等标记的具体动作；原表全文保留。', '金额单位为万元；赢率、毛利率、税率可填写百分比；预计签订日期建议YYYY-MM-DD。', '本模板不含示例业务数据；请在月度任务令工作表第2行开始填写。']:
        guide.append([line])
    guide.column_dimensions['A'].width = 110
    output = io.BytesIO()
    book.save(output)
    book.close()
    return {'ok': True, 'name': '月度任务令导入模板.xlsx', 'data': base64.b64encode(output.getvalue()).decode()}


def monthly_battle_unit(region):
    text = str(region or '').replace('办事处', '').replace('作战平台', '')
    checks = (('east-zhejiang-north', ('浙北',)), ('east-zhejiang-south', ('浙南',)), ('east-jiangsu', ('江苏',)),
              ('west', ('西区',)), ('central', ('中区',)), ('guangdong', ('广东',)))
    return next((unit for unit, words in checks if any(word in text for word in words)), '')


def monthly_battle_store(value=None):
    if crmdb_enabled():
        if value is not None:
            crmdb_blob_put('monthly_battle:v1', value)
        return crmdb_blob_get('monthly_battle:v1') or {'items': []}
    path = USER_DATA_DIR / '_crm_monthly_battle.json'
    if value is not None:
        office_doc_write('platform', path, value)
    return office_doc_read('platform', path, {'items': []})


def monthly_battle_level(user):
    caps = crm_capabilities(user)
    if user.get('role') in ('admin', 'superadmin') or caps.get('role') in ('manager', 'admin', 'superadmin') or caps.get('manage'):
        return 'company'
    roster_level = crm_battle_roster_level(user.get('username', ''))
    if roster_level in ('company', 'region'):
        return roster_level
    return 'region' if caps.get('data_scope') == 'regions' else 'personal'


def monthly_battle_visible(item, user):
    caps = crm_capabilities(user)
    if not caps.get('read'):
        return False
    if monthly_battle_level(user) == 'company':
        return True
    if user.get('username') in item.get('assignees', []):
        return True
    if monthly_battle_level(user) == 'region':
        roster_units = crm_battle_roster_units(user.get('username', ''))
        if roster_units:
            return monthly_battle_unit(item.get('region', '')) in roster_units
        return region_identity(item.get('region', ''))[0] in {region_identity(value)[0] for value in caps.get('regions', [])}
    return user.get('username') in item.get('assignees', [])


def monthly_battle_read(user, view='', region='', manager=''):
    if not crm_capabilities(user).get('read'):
        raise PermissionError('无查看权限')
    superadmin = user.get('role') == 'superadmin'
    if any((view, region, manager)) and not superadmin:
        raise PermissionError('仅超级管理员可切换权限视图')
    level = view or monthly_battle_level(user)
    if level not in ('company', 'region', 'personal'):
        raise ValueError('无效视图')
    accounts = read_config().get('users', []) if superadmin else []
    roster_units = crm_battle_roster()['units']
    roster_region_names = {unit['name'] for unit in roster_units}
    selectable_regions = roster_region_names or set(CRM_BUSINESS_REGIONS)
    if superadmin and level == 'region' and region not in selectable_regions:
        raise ValueError('请选择区域')
    if superadmin and level == 'personal' and manager not in {account['username'] for account in accounts}:
        raise ValueError('请选择人员')
    now = datetime.now()
    with MONTHLY_BATTLE_LOCK:
        items = json.loads(json.dumps(monthly_battle_store()['items']))
    result = []
    for item in items:
        if not monthly_battle_visible(item, user):
            continue
        if superadmin and level == 'region':
            selected_unit = next((unit['id'] for unit in roster_units if unit['name'] == region), '')
            if selected_unit and monthly_battle_unit(item.get('region', '')) != selected_unit:
                continue
            if not selected_unit and region_identity(item.get('region', ''))[0] != region_identity(region)[0]:
                continue
        if superadmin and level == 'personal' and manager not in item.get('assignees', []):
            continue
        events = item.get('events', [])
        spans = {str(stage): 0 for stage in range(1, 12)}
        for index, event in enumerate(events):
            if not event.get('stage') or event.get('completed'):
                continue
            end = datetime.fromisoformat(events[index + 1]['at']) if index + 1 < len(events) else now
            spans[str(event['stage'])] += max(0, (end - datetime.fromisoformat(event['at'])).total_seconds())
        item['stage_days'] = {key: round(value / 86400, 2) for key, value in spans.items()}
        last = datetime.fromisoformat(events[-1]['at'] if events else item['created_at'])
        item['overdue'] = not item.get('completed', False) and (now - last).total_seconds() >= item.get('cadence', 2) * 86400
        item['last_at'] = last.isoformat(timespec='seconds')
        item['action_cycles'] = monthly_battle_action_cycles(item, now)
        item['can_comment'] = monthly_battle_level(user) in ('company', 'region')
        caps = crm_capabilities(user)
        item['can_update'] = bool(caps.get('register_action') or caps.get('write')) and (user.get('username') in item.get('assignees', []) or caps.get('data_scope') == 'all' or caps.get('supervise'))
        result.append(item)
    label = {'company': '管理视图 · 全部任务', 'region': '区总视图 · 本人及所属作战单元', 'personal': '个人视图 · 本人任务'}[level]
    if superadmin and level == 'region':
        label = '区域视图 · ' + region
    if superadmin and level == 'personal':
        label += ' · ' + next((account.get('name') or manager for account in accounts if account['username'] == manager), manager)
    options = {'regions': [unit['name'] for unit in roster_units] or list(CRM_BUSINESS_REGIONS), 'users': [{'username': account['username'], 'name': account.get('name') or account['username']} for account in accounts]} if superadmin else None
    return {'ok': True, 'items': result, 'stages': MONTHLY_BATTLE_STAGES, 'scope': {'level': level, 'label': label}, 'view_options': options, 'can_import': bool(level == 'company' and monthly_battle_level(user) == 'company' and crm_capabilities(user).get('write'))}


def monthly_battle_import(payload, user):
    caps = crm_capabilities(user)
    if not (caps.get('write') and monthly_battle_level(user) == 'company'):
        raise PermissionError('仅全公司数据管理人员可导入')
    month = str(payload.get('month', ''))
    if not re.fullmatch(r'20\d{2}-(0[1-9]|1[0-2])', month):
        raise ValueError('请选择任务月份')
    raw = base64.b64decode(payload.get('data', ''), validate=True)
    if len(raw) > 10 * 1024 * 1024:
        raise ValueError('文件不能超过10MB')
    book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    imported = []
    users = read_config().get('users', [])
    try:
        for sheet in book:
            rows = iter(sheet.values)
            headers = [str(value or '').strip().lstrip('*') for value in next(rows, [])]
            if '机会点名称' not in headers or '子项目编码' not in headers:
                continue
            for row_number, row in enumerate(rows, 2):
                def field(name):
                    index = headers.index(name) if name in headers else -1
                    return row[index] if 0 <= index < len(row) and row[index] is not None else ''
                code, name = str(field('子项目编码')).strip(), str(field('机会点名称')).strip()
                if not name:
                    continue
                missing_code = not code
                if missing_code:
                    identity = region_identity(str(field('项目归属')))[0] + '|' + name
                    code = '待补编码-' + hashlib.sha256(identity.encode()).hexdigest()[:12]
                owner, manager = str(field('确认主要责任人')).strip(), str(field('客户经理')).strip()
                assignees = []
                for person in {owner, manager} - {''}:
                    matches = [account['username'] for account in users if account.get('name') == person or account.get('username') == person]
                    if len(matches) == 1:
                        assignees.extend(matches)
                target = str(field('任务令作战部署的目标/动作'))
                match = re.search(r'阶段\s*(1[01]|[1-9])(?=\D|$)', target)
                target_stage = int(match.group(1)) if match else next((index + 1 for index, label in enumerate(MONTHLY_BATTLE_STAGES) if label in target), 0)
                values = {'code': code, 'name': name, 'owner': owner, 'manager': manager, 'assignees': sorted(set(assignees)), 'region': str(field('项目归属')), 'target': target, 'target_stage': target_stage, 'amount': field('可参与份额-总额(万元)'), 'tag': str(field('现：保底/冲刺/推进') or field('原：保底/冲刺/推进')), 'missing_code': missing_code}
                source = {'sheet': sheet.title, 'row': row_number, 'values': {header: str(row[index]) if row[index] is not None else '' for index, header in enumerate(headers) if header and index < len(row)}}
                previous = next((item for item in imported if item['code'] == code), None)
                if previous:
                    if previous['name'] != name:
                        raise ValueError(f'编码{code}对应不同名称，请核对原表')
                    previous['sources'].append(source)
                    for key in ('owner', 'manager', 'region', 'target', 'target_stage', 'amount', 'tag'):
                        if previous[key] in ('', 0) and values[key] not in ('', 0):
                            previous[key] = values[key]
                        elif values[key] not in ('', 0) and previous[key] != values[key]:
                            previous['conflicts'] = sorted(set(previous['conflicts'] + [key]))
                    previous['assignees'] = sorted(set(previous['assignees'] + values['assignees']))
                else:
                    imported.append({**values, 'sources': [source], 'conflicts': []})
    finally:
        book.close()
    if not imported:
        raise ValueError('未找到含机会点名称、子项目编码的任务令表')
    with MONTHLY_BATTLE_LOCK:
        data = monthly_battle_store()
        for values in imported:
            item = next((item for item in data['items'] if item['month'] == month and item['code'] == values['code']), None)
            if item is None and not values['missing_code']:
                item = next((item for item in data['items'] if item['month'] == month and item.get('missing_code') and item['name'] == values['name'] and region_identity(item['region'])[0] == region_identity(values['region'])[0]), None)
            if item is None:
                item = {'id': secrets.token_urlsafe(12), 'month': month, 'created_at': datetime.now().isoformat(timespec='seconds'), 'stage': values['target_stage'], 'stage_source': 'import', 'completed': False, 'cadence': 2, 'revision': 0, 'events': []}
                data['items'].append(item)
            elif not item.get('events'):
                item['stage'] = values['target_stage']
                item['stage_source'] = 'import'
            item.update(values)
            item['revision'] += 1
        monthly_battle_store(data)
    return {'ok': True, 'count': len(imported), 'source_rows': sum(len(item['sources']) for item in imported), 'unassigned': sum(not item['assignees'] for item in imported)}


def monthly_battle_comment(payload, user):
    content = str(payload.get('content', '')).strip()
    if not content or len(content) > 2000:
        raise ValueError('请填写1至2000字点评')
    with MONTHLY_BATTLE_LOCK:
        data = monthly_battle_store()
        item = next((item for item in data['items'] if item['id'] == payload.get('id')), None)
        if item is None or not monthly_battle_visible(item, user) or monthly_battle_level(user) not in ('company', 'region'):
            raise PermissionError('仅管理者及所属区域负责人可点评')
        event = next((event for event in item.get('events', []) if event['id'] == payload.get('event_id')), None)
        if event is None:
            raise ValueError('动作不存在')
        event.setdefault('comments', []).append({'id': secrets.token_urlsafe(12), 'at': datetime.now().isoformat(timespec='seconds'), 'author': user.get('name') or user['username'], 'username': user['username'], 'content': content})
        item['revision'] += 1
        monthly_battle_store(data)
    return {'ok': True}


def monthly_battle_progress(payload, user):
    stage_change = payload.get('event_type') == 'stage_change'
    content = str(payload.get('content', '')).strip()
    if not stage_change and (not content or len(content) > 10000):
        raise ValueError('请填写1至10000字的沟通对象、推进细节和结果')
    if not stage_change and not str(payload.get('next_action', '')).strip():
        raise ValueError('请填写下一步动作及计划日期')
    stage, cadence = payload.get('stage'), payload.get('cadence', 2)
    if type(stage) is not int or not 0 <= stage <= 11 or cadence not in (1, 2):
        raise ValueError('阶段须为1至11，推进频率为1或2天')
    completed = payload.get('completed') is True
    if completed and stage != 11:
        raise ValueError('仅合同签订阶段可确认完成')
    attachments = payload.get('attachments') or []
    if not isinstance(attachments, list) or len(attachments) > 9:
        raise ValueError('每次最多9个截图附件')
    for attachment in attachments:
        if not isinstance(attachment, dict) or not re.fullmatch(r'[A-Za-z0-9_-]+', str(attachment.get('id', ''))) or not forum_attachment_entry(attachment['id']):
            raise ValueError('附件不存在，请重新上传')
    attachments = [forum_public_attachment({**attachment, 'url': forum_attachment_url(attachment['id'], attachment.get('name', '截图'))}) for attachment in attachments]
    with MONTHLY_BATTLE_LOCK:
        data = monthly_battle_store()
        item = next((item for item in data['items'] if item['id'] == payload.get('id')), None)
        caps = crm_capabilities(user)
        if item is None or not monthly_battle_visible(item, user) or not (caps.get('register_action') or caps.get('write')) or not (user.get('username') in item.get('assignees', []) or caps.get('data_scope') == 'all' or caps.get('supervise')):
            raise PermissionError('无权更新该机会点')
        if payload.get('revision') != item['revision']:
            raise ValueError('记录已被更新，请刷新后重试')
        stage_reason = str(payload.get('stage_reason') or '').strip()
        if stage_change:
            if stage == item['stage']:
                raise ValueError('阶段未发生变化')
            cadence = item['cadence']
            completed = item.get('completed', False) if stage == 11 else False
            content = ''
            attachments = []
        if stage == 0 and item['stage'] != 0:
            raise ValueError('已确认阶段不能改回未确认')
        if stage != item['stage'] and not stage_reason:
            raise ValueError('首次确认或修改阶段必须填写阶段修改说明')
        if len(stage_reason) > 2000:
            raise ValueError('阶段修改说明不能超过2000字')
        event = {'id': secrets.token_urlsafe(12), 'at': datetime.now().isoformat(timespec='seconds'), 'author': user.get('name') or user['username'], 'username': user['username'], 'from_stage': item['stage'], 'stage': stage, 'completed': completed, 'content': content, 'next_action': str(payload.get('next_action', '')).strip()[:2000], 'attachments': attachments}
        event['stage_reason'] = stage_reason if stage != item['stage'] else ''
        event['event_type'] = 'stage_change' if stage_change else 'action'
        if stage_change:
            event['next_action'] = ''
        item['events'].append(event)
        item.update(stage=stage, cadence=cadence, completed=completed, revision=item['revision'] + 1)
        monthly_battle_store(data)
    return {'ok': True}
