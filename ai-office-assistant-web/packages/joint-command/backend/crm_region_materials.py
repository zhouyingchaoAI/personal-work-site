from pathlib import Path
from hashlib import sha256

from openpyxl import Workbook, load_workbook


def build_region_workbook(directory):
    directory = Path(directory)
    sources = {}
    for suffix, sheet in [('summary', '汇总'), ('payment', '回款-总'), ('region_monthly', '区域月度订收回-总')]:
        path = directory / f'lianzhi_2026_{suffix}_values.xlsx'
        sources[suffix] = (load_workbook(path, data_only=True)[sheet], path.name, sha256(path.read_bytes()).hexdigest())
    summary = sources['summary'][0]
    monthly = sources['region_monthly'][0]
    payment = sources['payment'][0]
    if summary['B4'].value != '区域名称' or '3月' not in str(payment['A15'].value):
        raise ValueError('材料结构与预期不一致，请核对表头')
    book = Workbook()
    marker = book.active
    marker.title = '_板块导入配置'
    marker.append(['business-block-v1'])

    def add(name, headers, rows, source):
        sheet = book.create_sheet('区域经营_' + name)
        sheet.append([name])
        sheet.append([source])
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = 'A4'

    regions = ['东区', '西区', '中区', '广东区']
    annual = []
    for offset, region in enumerate(regions):
        row = offset + 5
        for metric, column in [('订货', 4), ('财务收入', 9), ('回款', 14), ('毛利', 19), ('管理订货', 24)]:
            annual.append([region, summary.cell(row, 3).value, metric, *[summary.cell(row, column + index).value for index in range(5)]])
    add('年度指标', ['区域', '负责人', '指标', '目标', '实际', '完成率', '差额', '排名'], annual, sources['summary'][1] + ' / 汇总 B4:AB8；单位万元；差额=实际-目标')
    records = []
    for offset, region in enumerate(regions):
        for month in range(1, 13):
            records.append([region, 2026, f'{month}月', *[monthly.cell(start + month - 1, offset + 2).value for start in (39, 22, 56)]])
    add('月度实绩', ['区域', '年度', '月份', '订货', '财务收入', '回款'], records, sources['region_monthly'][1] + '；单位万元；保留原表零值，不推断尚未发生月份的完成情况')
    plans = [[region, '2026年3月', *[payment.cell(offset + 18, column).value / 10000 if isinstance(payment.cell(offset + 18, column).value, (int, float)) else None for column in (2, 4, 5, 6, 7)]] for offset, region in enumerate(regions)]
    add('回款计划', ['区域', '计划月份', '预计回款', '电汇', '银承', '往月应收款', '本月应收款'], plans, sources['payment'][1] + ' / A15:G22；元换算万元；仅计划，不计入实际完成')
    receivables = [[region, payment.cell(row, 1).value, payment.cell(row, offset + 4).value, payment.cell(row, offset + 11).value] for offset, region in enumerate(regions) for row in range(45, 79) if payment.cell(row, 1).value]
    add('应收款结构', ['区域', '性质', '应收款', '数量'], receivables, sources['payment'][1] + ' / A44:N79；金额单位原表未注明，待核验；数量单位笔')
    notes = [['全部', '历史上海目标', '不作为第五区域展示；未擅自并入东区目标。四区域订货目标53000，公司目标55000。'], ['全部', '回款计划', '回款总表3月为计划，与月度实际回款分开，不混加。'], ['全部', '应收款口径', '公司应收款总额与四区域之和不一致，差额不分摊；原表金额单位未注明。'], ['全部', '人员与办事处', '本批材料未提供个人及三个办事处拆分实绩，不虚构分配。']]
    for suffix, (_, filename, digest) in sources.items():
        notes.append(['全部', filename, 'SHA256:' + digest])
    add('来源说明', ['范围', '事项', '说明'], notes, '三份用户指定材料；仅替换区域经营展示，保留其他板块和历史版本')
    return book
