from openpyxl import Workbook, load_workbook


def convert_region_periods(path):
    source = load_workbook(path, data_only=True, read_only=True)
    sheet = source['数据源']
    rows = list(sheet.values)
    if list(rows[0][:3]) != ['年度', '期间', '区域']:
        raise ValueError('数据源表头不符合期间模板')
    book = Workbook()
    book.active.title = '_板块导入配置'
    book.active.append(['business-block-v1'])
    regions = {'东区': '东区', '西区': '西区', '中区': '中区', '广东': '广东区', '广东区': '广东区'}

    def add(name, headers, values, note=''):
        target = book.create_sheet('区域经营_' + name)
        target.append([name])
        target.append(['2026经营数据_逐页时间切换模板.xlsx / 数据源；' + note])
        target.append(headers)
        for value in values:
            target.append(value)

    periods = []
    for row in rows[1:]:
        if row[0] == 2026 and row[2] in regions:
            periods.append([row[0], row[1], regions[row[2]], *row[3:12]])
    if not periods:
        raise ValueError('未找到四区域期间指标')
    add('期间指标', ['年度', '期间', '区域', '负责人', '指标', '目标', '实际', '完成率', '差额', '统计时间范围', '指标口径说明', '目标说明'], periods, '金额万元；缺失目标保留空值')
    add('年度指标', ['区域', '负责人', '指标', '目标', '实际', '完成率', '差额', '排名'], [[row[2], row[3], row[4], *row[5:9], None] for row in periods if row[1] == '全年'])
    monthly = []
    for region in regions.values():
        if region == '广东区' and any(row[0] == region for row in monthly):
            continue
        for month in range(1, 13):
            selected = {row[4]: row[6] for row in periods if row[2] == region and row[1] == f'{month}月'}
            monthly.append([region, 2026, f'{month}月', *[selected.get(metric) for metric in ['订货', '财务收入', '回款']]])
    add('月度实绩', ['区域', '年度', '月份', '订货', '财务收入', '回款'], monthly)
    add('业务明细', ['年度', '月份', '季度', '区域', '指标', '金额（万元）', '业务日期原值', '合同或单据编号', '来源工作表', '来源行号', '来源列号'], [[*row[13:16], regions[row[16]], *row[17:24]] for row in rows[1:] if row[13] == 2026 and row[16] in regions])
    add('应收款结构', ['区域', '性质', '应收款', '数量'], [[regions[row[48]], row[47], row[49], row[50]] for row in rows[1:] if row[48] in regions and row[47]], '金额万元；存量时点未明确，不按期间分摊')
    add('历史回款计划', ['区块', '期间', '区域', '指标', '数值', '单位', '数据性质', '来源位置', '说明'], [[*row[37:39], regions.get(row[39], row[39]), *row[40:46]] for row in rows[1:] if row[37]], '历史计划，不计入实际')
    add('回款计划', ['区域', '计划月份', '预计回款', '电汇', '银承', '往月应收款', '本月应收款'], [], '旧版计划停止展示，参见历史回款计划')
    add('来源说明', ['范围', '事项', '说明'], [['全部', '期间口径', '采用新文件数据源A:L期间指标；金额万元；毛利为预计毛利，非审定毛利。'], ['全部', '范围', '仅四区域；历史上海与公司其他原值保存在原材料，不合并或重复计入四区。'], ['全部', '应收存量', '单位万元；时点未明确，不随期间分摊。']])
    for name, start, end, region_offset in [
        ('原表期间', 0, 12, 2),
        ('原表业务', 13, 24, 3),
        ('月度拆解', 25, 36, 5),
        ('应收原表', 47, 55, 1),
    ]:
        headers = [str(value or '') for value in rows[0][start:end]]
        headers[region_offset] = '原表范围'
        values = []
        for row in rows[1:]:
            if row[start] is None:
                continue
            values.append([regions.get(row[start + region_offset], ''), *row[start:end]])
        add(name, ['区域', *headers], values, '区域列用于权限筛选；原表范围保留历史上海、公司其他及合计，不混入四区合计；不同记录层级不可重复相加')
    notes = [list(row[86:91]) for row in rows[1:] if any(value is not None for value in row[86:91])]
    add('原表核对说明', ['项目', '说明或指标', '原值1', '原值2', '原值3'], notes, '保留原表口径说明及核对数值')
    source.close()
    return book
