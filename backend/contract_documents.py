"""User-owned contract extraction and delivery document drafts."""
from decimal import Decimal, InvalidOperation

CONTRACT_DOCUMENT_KINDS = {"送货单", "到货签收单", "验收单"}
CONTRACT_DOCUMENT_LOCK = threading.RLock()


def contract_directory(username):
    if not isinstance(username, str) or not username or safe_username(username) != username:
        raise PermissionError("无效用户身份")
    directory = user_root(username) / "contract_documents"
    if directory.resolve() != USER_DATA_DIR.resolve() / username / "contract_documents":
        raise PermissionError("用户目录不可用")
    return directory


def contract_path(username, contract_id):
    if not re.fullmatch(r"[a-f0-9]{32}", str(contract_id)):
        raise ValueError("合同编号无效")
    directory = contract_directory(username) / contract_id
    if directory.is_symlink():
        raise PermissionError("合同目录不可用")
    return directory


def contract_read(username, contract_id):
    path = contract_path(username, contract_id) / "record.json"
    if not path.is_file():
        raise ValueError("合同不存在或无访问权限")
    return json.loads(path.read_text(encoding="utf-8"))


def contract_write(username, record):
    path = contract_path(username, record["id"]) / "record.json"
    temporary = path.with_suffix(".tmp")
    temporary.write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def contract_number(value, positive=False):
    try:
        number = Decimal(str(value).replace(",", "").strip())
    except InvalidOperation as exc:
        raise ValueError("数量必须是有效数字") from exc
    if not number.is_finite() or number < 0 or (positive and number == 0) or number > Decimal("1000000000"):
        raise ValueError("数量超出允许范围")
    return number


CONTRACT_TEMPLATE_PROFILES = (
    ("sales_service_2026", "2026 服务销售合同", "服务销售合同", ("合同服务清单（不含设备软硬件供货）", "服务有效期", "服务交付方式")),
    ("sales_equipment_2025", "2025 设备销售合同", "设备销售合同", ("设备清单", "到货签收约定", "保修责任")),
    ("sales_integration_2025", "2025 集成销售合同", "集成销售合同", ("合同清单", "服务范围", "到货签收及服务交付方式")),
    ("sales_training_2025", "2025 培训销售合同", "培训", ("培训服务内容", "培训费用", "结算方式")),
    ("sales_software_2025", "2025 软件销售合同", "软件", ("软件内容、服务方式、要求", "签收及验收约定", "质量保证期")),
)


def contract_detect_template(text):
    compact = re.sub(r"\s+", "", text)
    candidates = []
    for identifier, name, title, markers in CONTRACT_TEMPLATE_PROFILES:
        evidence = [marker for marker in markers if marker in compact]
        if title in compact[:600] and len(evidence) == len(markers):
            candidates.append({"template_id": identifier, "template_name": name, "evidence": evidence})
    if len(candidates) == 1:
        return {**candidates[0], "mode": "template", "reason": "标题与全部模板特征匹配；仍需核对字段", "version": 1}
    return {"template_id": "", "template_name": "未匹配已登记模板", "evidence": [], "mode": "general", "reason": "多个模板存在冲突，已转通用识别" if candidates else "模板特征不足，已转通用识别", "version": 1}


def contract_extract(raw, suffix, with_route=False):
    tables = []
    if suffix == ".docx":
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 80 * 1024 * 1024:
                raise ValueError("合同解压后过大")
        from docx import Document
        document = Document(io.BytesIO(raw))
        text_parts = [paragraph.text for paragraph in document.paragraphs]
        for table_index, table in enumerate(document.tables):
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            tables.append((table_index + 1, rows))
            text_parts.extend(" | ".join(row) for row in rows)
        source_text = "\n".join(text_parts)
    else:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted or len(reader.pages) > 150:
            raise ValueError("请上传未加密且不超过150页的PDF")
        pages = [page.extract_text() or "" for page in reader.pages]
        source_text = "\n".join(f"第{index + 1}页\n{text}" for index, text in enumerate(pages) if text.strip())
    route = contract_detect_template(source_text)
    aliases = {"name": ["产品名称", "货物名称", "设备名称", "服务名称", "服务项目", "品名", "名称", "内容"], "spec": ["规格型号", "规格", "型号", "服务内容"], "unit": ["单位"], "quantity": ["数量"]}
    aliases["spec"].append("规格型号/参数")
    aliases.update(brand=["品牌"], unit_price=["单价(元)", "单价（元）", "单价"], amount=["合计(元)", "合计（元）", "金额"], remark=["备注"])
    if route["template_id"] == "sales_service_2026":
        aliases["name"] = ["服务项目", "服务名称", "服务内容", "名称", "内容"]
        aliases["quantity"] = ["数量", "服务次数", "次数"]
    elif route["template_id"] == "sales_training_2025":
        aliases["name"] = ["培训内容", "培训项目", "课程名称", "名称"]
        aliases["quantity"] = ["数量", "次数", "培训次数"]
    lines = []
    for table_index, rows in tables:
        mapping = {}
        for row_index, cells in enumerate(rows):
            found = {}
            for field, labels in aliases.items():
                for cell_index, cell in enumerate(cells):
                    if re.sub(r"\s+", "", cell) in labels:
                        found[field] = cell_index
                        break
            if "name" in found and "quantity" in found:
                mapping = found
                continue
            if not mapping:
                continue
            values = {field: cells[position] if position < len(cells) else "" for field, position in mapping.items()}
            if not values.get("name") or values["name"] in {"/", "—"} or re.match(r"^(?:合计|总计|小计)", values["name"]):
                continue
            if not values.get("quantity") and not values.get("unit"):
                continue
            combined = re.fullmatch(r"([\d,]+(?:\.\d+)?)\s*(台|套|项|次|个|件|批|人天)", values.get("quantity", ""))
            if combined and not values.get("unit"):
                values["quantity"], values["unit"] = combined.groups()
            try:
                quantity = str(contract_number(values.get("quantity", ""), positive=True))
            except ValueError:
                quantity = ""
            lines.append({"id": secrets.token_hex(8), "name": values["name"], "spec": values.get("spec", ""), "unit": values.get("unit", ""), "quantity": quantity, "source": f"表{table_index} 第{row_index + 1}行", **{key: values.get(key, "") for key in ("brand", "unit_price", "amount", "remark")}})
    result = (source_text[:200000], lines[:500])
    return (*result, route) if with_route else result


def contract_structure(text, route=None):
    """Conservative candidates with original evidence, never inferred deadlines."""
    def first(pattern):
        match = re.search(pattern, text, re.M)
        return match.group(1).strip() if match else ""
    def party(label):
        # Parties may be parallel cells or text on the same line, never consume a newline.
        return first(r"(?:^|\|)[ \t]*" + label + r"[ \t]*方(?:[（(]章[）)])?[：:][ \t]*(?:\|[ \t]*)?([^\n|]*)")
    def section(title):
        return first(r"^第[^\n：:]{1,12}条[：:][ \t]*" + title + r"[^\n]*\n([\s\S]*?)(?=^第[^\n：:]{1,12}条[：:]|\Z)")
    payments = []
    for line in text.splitlines():
        cells = [cell.strip() for cell in line.split("|")]
        if len(cells) == 5 and cells[0] in {"预付款", "验收款", "开通试运营款", "质保金", "进度款", "尾款"} and (cells[1] or cells[2]):
            payments.append(dict(zip(("name", "percentage", "amount", "days", "condition"), cells)))
    result = {
        "contract_no": first(r"合同编号[：:][ \t]*([^\s|]*)"),
        "customer": party("甲"),
        "supplier": party("乙"),
        "total_amount": first(r"含税总价为[￥¥\s]*([\d,]+(?:\.\d+)?)"),
        "tax_rate": first(r"税率为\s*([\d.]+%)"),
        "payment_terms": "\n".join(line for line in text.splitlines() if any(word in line for word in ("预付款：", "开通试运营款：", "付款前", "顺延")))[:5000],
        "acceptance_terms": "\n".join(line for line in text.splitlines() if any(word in line for word in ("如验收", "验收标准对", "质量保证服务期")))[:5000],
        "payments": payments,
    }
    result["payment_terms"] = section("付款方式") or result["payment_terms"]
    result["acceptance_terms"] = "\n".join(filter(None, (section("到货签收约定"), section("保修责任"), result["acceptance_terms"])))[:5000]
    if route and route["mode"] == "template":
        result["total_amount"] = first(r"合同总金额[（(]含税[）)][：:][ \t]*(?:人民币)?(?:小写[：:])?[￥¥ \t]*([\d,]+(?:\.\d+)?)") or result["total_amount"]
        result["tax_rate"] = first(r"增值税税率为[ \t]*([\d.]+[ \t]*%)") or result["tax_rate"]
    return result


def contract_extract_preview(payload, username):
    record = contract_read(username, payload.get("id"))
    if record.get("archived") or record["revision"] != payload.get("revision"):
        raise ValueError("合同已归档或版本已更新，请刷新")
    name, raw = contract_original(username, record["id"])
    text, lines, route = contract_extract(raw, Path(name).suffix.lower(), with_route=True)
    record.update(contract_structure(text, route))
    record.update(source_text=text, lines=lines, extraction=route)
    return {"ok": True, "contract": record}


def contract_upload(payload, username):
    contract_directory(username)
    name = Path(str(payload.get("name", ""))).name
    suffix = Path(name).suffix.lower()
    if suffix not in {".docx", ".pdf"}:
        raise ValueError("仅支持DOCX、PDF；旧版DOC请另存为DOCX")
    encoded = payload.get("content_base64", "")
    if not isinstance(encoded, str) or len(encoded) > 28 * 1024 * 1024:
        raise ValueError("文件不能超过20MB")
    raw = base64.b64decode(encoded, validate=True)
    if not raw or len(raw) > 20 * 1024 * 1024:
        raise ValueError("文件为空或超过20MB")
    text, lines, route = contract_extract(raw, suffix, with_route=True)
    if not text.strip():
        raise ValueError("未提取到文字：扫描件OCR尚未接入，请上传文本PDF或DOCX")
    record = {"id": secrets.token_hex(16), "filename": name, "source_text": text, "project_code": "", "contract_no": "", "customer": "", "address": "", "lines": lines, "revision": 1, "documents": [], "created_at": datetime.now().isoformat(timespec="seconds"), "warning": "规则提取结果须逐项核对；PDF清单暂需参照原文补录。所有导出均为待审核草稿。"}
    record.update(contract_structure(text, route))
    record["extraction"] = route
    directory = contract_path(username, record["id"])
    directory.mkdir(parents=True, exist_ok=False)
    (directory / ("original" + suffix)).write_bytes(raw)
    contract_write(username, record)
    return {"ok": True, "contract": record}


def contract_list(username):
    directory = contract_directory(username)
    records = []
    if directory.exists():
        for path in directory.glob("*/record.json"):
            record = contract_read(username, path.parent.name)
            records.append({**{key: record[key] for key in ("id", "filename", "created_at", "revision")}, "contract_no": record.get("contract_no", ""), "customer": record.get("customer", ""), "archived": record.get("archived", False), "document_count": len(record["documents"])})
    return {"ok": True, "contracts": sorted(records, key=lambda record: record["created_at"], reverse=True)}


def contract_save(payload, username):
    with CONTRACT_DOCUMENT_LOCK:
        record = contract_read(username, payload.get("id"))
        if record.get("archived"):
            raise ValueError("合同已归档，请恢复后编辑")
        if payload.get("revision") != record["revision"]:
            raise ValueError("合同已更新，请重新加载后编辑")
        lines = payload.get("lines")
        if not isinstance(lines, list) or not 1 <= len(lines) <= 500:
            raise ValueError("请填写1至500条合同明细")
        cleaned = []
        existing_ids = {line["id"] for line in record["lines"]}
        used_ids = set()
        for line in lines:
            if not isinstance(line, dict) or not str(line.get("name", "")).strip():
                raise ValueError("明细名称不能为空")
            identifier = line.get("id")
            identifier = identifier if identifier in existing_ids and identifier not in used_ids else secrets.token_hex(8)
            used_ids.add(identifier)
            cleaned.append({"id": identifier, **{field: str(line.get(field, ""))[:500] for field in ("name", "spec", "unit", "source", "brand", "unit_price", "amount", "remark")}, "quantity": str(contract_number(line.get("quantity", ""), positive=True))})
        record["lines"] = cleaned
        record["extraction"] = contract_detect_template(record["source_text"])
        for field in ("project_code", "contract_no", "customer", "address"):
            record[field] = str(payload.get(field, "")).strip()[:500]
        for field in ("supplier", "total_amount", "tax_rate", "payment_terms", "acceptance_terms"):
            record[field] = str(payload.get(field, record.get(field, ""))).strip()[:5000]
        payments = payload.get("payments", record.get("payments", []))
        if not isinstance(payments, list) or len(payments) > 100 or any(not isinstance(row, dict) for row in payments):
            raise ValueError("付款节点格式无效")
        record["payments"] = [{key: str(row.get(key, ""))[:2000] for key in ("name", "percentage", "amount", "days", "condition")} for row in payments]
        record["revision"] += 1
        contract_write(username, record)
        return {"ok": True, "contract": record}


def contract_generate(payload, username):
    from docx import Document
    with CONTRACT_DOCUMENT_LOCK:
        record = contract_read(username, payload.get("id"))
        if record.get("archived"):
            raise ValueError("合同已归档，请恢复后生成单据")
        if payload.get("revision") != record["revision"]:
            raise ValueError("请先保存并核对当前清单")
        kind = payload.get("kind")
        if kind not in CONTRACT_DOCUMENT_KINDS or payload.get("confirmed") is not True:
            raise ValueError("请选择单据类型并确认已核对合同")
        if not record["contract_no"] or not record["customer"]:
            raise ValueError("请先保存合同编号及客户名称")
        selections = payload.get("selections")
        if not isinstance(selections, list) or not selections or len(selections) > 500:
            raise ValueError("请勾选本批次明细")
        source = {line["id"]: line for line in record["lines"]}
        selected = []
        seen = set()
        for selection in selections:
            line_id = selection.get("id")
            if line_id not in source or line_id in seen:
                raise ValueError("清单项无效或重复")
            seen.add(line_id)
            quantity = contract_number(selection.get("quantity", ""), positive=True)
            if quantity > contract_number(source[line_id]["quantity"]):
                raise ValueError("本批次数量不能超过合同数量")
            selected.append({**source[line_id], "batch_quantity": str(quantity), "sn": str(selection.get("sn", ""))[:2000]})
        document_id = secrets.token_hex(16)
        document = Document()
        document.add_heading(kind + "（待审核草稿）", 0)
        for label, value in [("单据编号", document_id), ("合同编号", record["contract_no"]), ("项目编码", record["project_code"]), ("客户", record["customer"]), ("交付地址", record["address"])]:
            document.add_paragraph(f"{label}：{value}")
        table = document.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        for cell, label in zip(table.rows[0].cells, ["名称", "规格型号", "单位", "合同数量", "本批次数量", "SN/二维码"]):
            cell.text = label
        for line in selected:
            for cell, field in zip(table.add_row().cells, ["name", "spec", "unit", "quantity", "batch_quantity", "sn"]):
                cell.text = line[field]
        document.add_paragraph("说明：本文件使用通用草稿版式，非公司正式模板；生成不代表已发货、签收或验收。请审核后使用。")
        document.add_paragraph("经办人：____________    接收/验收人：____________    日期：____________")
        output = io.BytesIO()
        document.save(output)
        (contract_path(username, record["id"]) / (document_id + ".docx")).write_bytes(output.getvalue())
        record["documents"].append({"id": document_id, "kind": kind, "created_at": datetime.now().isoformat(timespec="seconds"), "revision": record["revision"], "lines": selected, "status": "draft"})
        contract_write(username, record)
        return {"ok": True, "contract": record}


def contract_download(username, contract_id, document_id):
    record = contract_read(username, contract_id)
    entry = next((entry for entry in record["documents"] if entry["id"] == document_id), None)
    if not entry:
        raise ValueError("单据不存在或无访问权限")
    return entry["kind"] + "-" + document_id + ".docx", (contract_path(username, contract_id) / (document_id + ".docx")).read_bytes()


def contract_original(username, contract_id):
    record = contract_read(username, contract_id)
    suffix = Path(record["filename"]).suffix.lower()
    if suffix not in {".docx", ".pdf"}:
        raise ValueError("原件格式无效")
    path = contract_path(username, contract_id) / ("original" + suffix)
    if path.is_symlink() or not path.is_file():
        raise ValueError("合同原件不可用")
    return record["filename"], path.read_bytes()


def contract_archive(payload, username):
    if not isinstance(payload.get("archived"), bool):
        raise ValueError("归档状态无效")
    with CONTRACT_DOCUMENT_LOCK:
        record = contract_read(username, payload.get("id"))
        if record["revision"] != payload.get("revision"):
            raise ValueError("合同已更新，请刷新后操作")
        record["archived"] = payload["archived"]
        record["revision"] += 1
        contract_write(username, record)
        return {"ok": True, "contract": record}
