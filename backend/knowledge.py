"""知识库：文档入库（解析/分块/向量化）、语义检索与 RAG 问答。

嵌入模型走 OpenAI 兼容 /v1/embeddings，可配置本地 Xinference 或网关。
向量直接存 JSON 文件 + 余弦相似度内存检索，适合个人/小团队规模。

This module is loaded by backend.runtime into one shared application namespace.
"""
import math

KB_DEFAULT_EMBEDDING_MODEL = "bge-m3"
KB_CHUNK_SIZE = 500
KB_CHUNK_OVERLAP = 80
KB_MAX_FILE_BYTES = 20 * 1024 * 1024
# RAGFlow 的 DeepDoc 按流式/分页解析，API 默认上限 1GB；这里取 100MB 兼顾后端内存（base64 JSON 体积约 1.37 倍）。
KB_RAGFLOW_MAX_FILE_BYTES = 100 * 1024 * 1024
KB_MAX_CHUNKS_PER_DOC = 800
KB_EMBED_BATCH = 16

KB_ASK_SYSTEM = (
    "你是企业知识库问答助手。只依据提供的资料片段回答问题，引用时标注片段编号（如 [1]）。\n"
    "资料里没有的信息明确说\"知识库中没有找到相关内容\"，不要编造。用中文回答，简洁准确。"
)


KB_DEFAULT_QDRANT_URL = "http://qdrant:6333"
KB_DEFAULT_RAGFLOW_URL = "http://host.docker.internal:9380"
KB_RAGFLOW_DATASET_NAME = "企业知识库"
KB_RAGFLOW_CHAT_NAME = "知识库助手"
KB_RAGFLOW_CHAT_OPENER = (
    "你好！我是企业知识库助手，可以基于你的个人知识库和公司共享知识库回答问题，"
    "回答会标注引用来源，点击角标可查看原文。试试问我文档里的内容吧～"
)
KB_RAGFLOW_CHAT_PROMPT = (
    "你是企业知识库问答助手。请只依据知识库片段回答用户问题，并在回答中标注引用。\n"
    "如果用户要求生成试卷、培训题、总结、清单、流程或方案，只要资料片段与主题相关，即使资料不完整，"
    "也要基于现有资料完成可用草稿，并明确说明适用范围和资料局限。\n"
    "资料片段相关但不完整时，不能回答未找到；应说明‘以下内容基于知识库中已有资料整理’。\n"
    "只有当资料片段全部无关时，才回答：The answer you are looking for is not found in the knowledge base!\n"
    "请用中文回答。\n"
    "Here is the knowledge base:\n{knowledge}\nThe above is the knowledge base."
)


def _ragflow_chat_prompt_payload() -> dict:
    return {"opener": KB_RAGFLOW_CHAT_OPENER, "show_quote": True, "prompt": KB_RAGFLOW_CHAT_PROMPT}


def kb_embedding_settings() -> dict:
    config = read_config()
    return {
        "provider": str(config.get("knowledge_provider") or "").strip() or "builtin",
        "api_url": str(config.get("knowledge_embedding_api_url") or "").strip(),
        "api_key": str(config.get("knowledge_embedding_api_key") or ""),
        "model": str(config.get("knowledge_embedding_model") or "").strip() or KB_DEFAULT_EMBEDDING_MODEL,
        "vector_store": str(config.get("knowledge_vector_store") or "").strip() or "local",
        "qdrant_url": str(config.get("knowledge_qdrant_url") or "").strip() or KB_DEFAULT_QDRANT_URL,
        "rerank_model": str(config.get("knowledge_rerank_model") or "").strip(),
        "ragflow_url": str(config.get("knowledge_ragflow_url") or "").strip() or KB_DEFAULT_RAGFLOW_URL,
        "ragflow_api_key": str(config.get("knowledge_ragflow_api_key") or ""),
        # 浏览器侧访问 RAGFlow 的地址（容器内地址浏览器打不开），留空时前端按当前主机:9380 推断。
        "ragflow_public_url": str(config.get("knowledge_ragflow_public_url") or "").strip(),
    }


def kb_config_api() -> dict:
    settings = kb_embedding_settings()
    settings["api_key_masked"] = "已配置" if settings.pop("api_key") else "未配置"
    settings["ragflow_api_key_masked"] = "已配置" if settings.pop("ragflow_api_key") else "未配置"
    return {"ok": True, **settings}


def save_kb_config(payload: dict) -> dict:
    payload = payload or {}
    config = read_config()
    provider = str(payload.get("provider", "") or "").strip()
    if provider and provider not in ("builtin", "ragflow"):
        raise ValueError("知识库引擎只支持 builtin 或 ragflow")
    config["knowledge_provider"] = provider or "builtin"
    config["knowledge_ragflow_url"] = str(payload.get("ragflow_url", "") or "").strip().rstrip("/") or KB_DEFAULT_RAGFLOW_URL
    if "ragflow_public_url" in payload:
        config["knowledge_ragflow_public_url"] = str(payload.get("ragflow_public_url", "") or "").strip().rstrip("/")
    if payload.get("ragflow_api_key_clear"):
        config["knowledge_ragflow_api_key"] = ""
    else:
        ragflow_key = str(payload.get("ragflow_api_key", "") or "").strip()
        if ragflow_key:
            config["knowledge_ragflow_api_key"] = ragflow_key
    config["knowledge_embedding_api_url"] = str(payload.get("api_url", "") or "").strip().rstrip("/")
    config["knowledge_embedding_model"] = str(payload.get("model", "") or "").strip()[:100] or KB_DEFAULT_EMBEDDING_MODEL
    store = str(payload.get("vector_store", "") or "").strip()
    if store and store not in ("local", "qdrant"):
        raise ValueError("向量存储只支持 local 或 qdrant")
    config["knowledge_vector_store"] = store or "local"
    config["knowledge_qdrant_url"] = str(payload.get("qdrant_url", "") or "").strip().rstrip("/") or KB_DEFAULT_QDRANT_URL
    # 重排模型留空表示关闭重排；重排接口复用嵌入 API 的地址和 Key。
    config["knowledge_rerank_model"] = str(payload.get("rerank_model", "") or "").strip()[:100]
    # Key 留空表示保持原值不变；本地 Xinference 等无鉴权服务用 api_key_clear 显式清空。
    if payload.get("api_key_clear"):
        config["knowledge_embedding_api_key"] = ""
    else:
        api_key = str(payload.get("api_key", "") or "").strip()
        if api_key:
            config["knowledge_embedding_api_key"] = api_key
    write_config(config)
    # FastGPT consumes the same embedding/rerank service. Synchronize at the
    # configuration boundary so its model selector never drifts from the
    # platform knowledge-base settings.
    result = kb_config_api()
    try:
        result["fastgpt_sync"] = sync_fastgpt_embedding_model()
    except Exception as exc:
        result["fastgpt_sync"] = {"ok": False, "error": str(exc)}
    return result


# ---------------------------------------------------------------------------
# RAGFlow 桥接（成熟 RAG 引擎：DeepDoc 解析 + 混合检索 + 引用问答）
# ---------------------------------------------------------------------------

def _ragflow_request(settings: dict, method: str, path: str, body=None, raw_body=None, content_type="application/json"):
    headers = {"Authorization": "Bearer " + settings["ragflow_api_key"]}
    data = None
    if raw_body is not None:
        data = raw_body
        headers["Content-Type"] = content_type
    elif body is not None:
        data = json.dumps(body, ensure_ascii=False).encode("utf-8")
        headers["Content-Type"] = "application/json"
    request = urllib.request.Request(
        settings["ragflow_url"].rstrip("/") + path, data=data, headers=headers, method=method
    )
    try:
        with urllib.request.urlopen(request, timeout=180) as response:
            result = json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", "ignore")[:300]
        raise ValueError(f"RAGFlow 返回 {exc.code}：{detail}") from exc
    except Exception as exc:
        raise ValueError(f"RAGFlow 连接失败（{settings['ragflow_url']}）：{exc}") from exc
    if result.get("code") not in (0, None):
        raise ValueError(f"RAGFlow 错误：{result.get('message', result)}")
    return result.get("data")


def _ragflow_ready(settings: dict):
    if not settings["ragflow_api_key"]:
        raise ValueError("尚未配置 RAGFlow API Key，请在设置页填写")


def _kb_mapping(filename: str) -> dict:
    data = office_doc_read("knowledge", _kb_dir() / filename, {})
    return data if isinstance(data, dict) else {}


def _kb_mapping_save(filename: str, mapping: dict):
    office_doc_write("knowledge", _kb_dir() / filename, mapping)


def _kb_is_admin(username: str) -> bool:
    user = find_user(username)
    return bool(user and user.get("role") in ("admin", "superadmin"))


KB_DATASET_INIT_LOCK = threading.RLock()


def _ragflow_named_dataset(settings: dict, name: str, cache_key: str, permission: str = "") -> str:
    with KB_DATASET_INIT_LOCK:
        return _ragflow_named_dataset_locked(settings, name, cache_key, permission)


def _ragflow_named_dataset_locked(settings: dict, name: str, cache_key: str, permission: str = "") -> str:
    """按名称取（必要时创建）数据集，ID 缓存在映射文件。"""
    mapping = _kb_mapping("ragflow_datasets.json")
    cached = str(mapping.get(cache_key) or "").strip()
    if cached:
        return cached
    # v0.21 按名称过滤查不到时会直接报错，视作不存在。
    try:
        data = _ragflow_request(settings, "GET", f"/api/v1/datasets?page=1&page_size=100&name={urllib.parse.quote(name)}") or []
    except ValueError:
        data = []
    dataset = next((item for item in data if item.get("name") == name), None)
    if not dataset:
        chunk = kb_ragflow_chunk_settings()
        body = {
            "name": name,
            "chunk_method": chunk["chunk_method"],
            "parser_config": {"chunk_token_num": chunk["chunk_token_num"], "layout_recognize": chunk["layout_recognize"]},
        }
        if permission:
            body["permission"] = permission
        dataset = _ragflow_request(settings, "POST", "/api/v1/datasets", body)
    mapping[cache_key] = dataset["id"]
    _kb_mapping_save("ragflow_datasets.json", mapping)
    return dataset["id"]


def _ragflow_shared_dataset_id(settings: dict) -> str:
    # 兼容旧的 config 缓存。permission=team 让加入服务账号团队的用户租户可读。
    config = read_config()
    cached = str(config.get("knowledge_ragflow_dataset_id") or "").strip()
    if cached:
        return cached
    return _ragflow_named_dataset(settings, KB_RAGFLOW_DATASET_NAME, "__shared__", permission="team")


def _ragflow_user_dataset_id(settings: dict, username: str) -> str:
    """用户租户内的个人知识库（settings 须带该用户自己的 API Key）。"""
    safe = str(username or "anonymous")
    return _ragflow_named_dataset(settings, "个人知识库", f"tenant::{safe}")


# ---------------------------------------------------------------------------
# 多租户：每个办公助手用户对应一个 RAGFlow 账号/租户。
# 个人知识库存在用户自己的租户里（控制台只见自己的数据）；共享知识库在服务账号
# 租户下且 permission=team，用户开通时自动加入服务账号团队获得读权限。
# ---------------------------------------------------------------------------

KB_RAGFLOW_USER_EMAIL_DOMAIN = "office.chencytech.com"


def _ragflow_user_email(username: str) -> str:
    slug = re.sub(r"[^0-9a-z]", "", str(username or "").lower())[:20]
    digest = hashlib.md5(str(username or "anonymous").encode("utf-8")).hexdigest()[:8]
    return f"{slug or 'user'}-{digest}@{KB_RAGFLOW_USER_EMAIL_DOMAIN}"


def _ragflow_admin_login(settings: dict) -> str:
    """登录托管的服务账号，返回 Authorization（JWT）。"""
    config = read_config()
    email = str(config.get("knowledge_ragflow_admin_email") or "").strip()
    password = str(config.get("knowledge_ragflow_admin_password") or "").strip()
    if not email or not password:
        raise ValueError("尚未托管 RAGFlow 服务账号，请先在知识库引擎设置执行「一键初始化」")
    data, headers = _ragflow_web_request(
        settings, "/v1/user/login", {"email": email, "password": _ragflow_encrypt_password(password)}
    )
    if data.get("code") != 0:
        raise ValueError(f"RAGFlow 服务账号登录失败：{data.get('message')}（可重新执行一键初始化）")
    auth = headers.get("Authorization", "")
    if not auth:
        raise ValueError("RAGFlow 未返回登录凭证")
    return auth


def _ragflow_tenant_sync_llm(settings: dict, auth: str) -> list:
    """把平台大模型同步进该登录租户并设为默认对话模型（幂等）。"""
    steps = []
    assistant = assistant_settings()
    if not (assistant["url"] and assistant["key"]):
        return ["同步对话模型：跳过（平台未配置大模型 API）"]
    try:
        data, _ = _ragflow_web_request(settings, "/v1/llm/add_llm", {
            "llm_factory": "OpenAI-API-Compatible",
            "llm_name": assistant["model"],
            "model_type": "chat",
            "api_base": assistant["url"].rstrip("/") + "/v1",
            "api_key": assistant["key"],
            "max_tokens": 16384,
        }, auth=auth)
        steps.append("同步对话模型：" + ("成功" if data.get("code") == 0 else str(data.get("message", ""))[:60]))
    except Exception as exc:
        steps.append(f"同步对话模型：{str(exc)[:60]}")
    try:
        data, _ = _ragflow_web_request(settings, "/v1/user/tenant_info", auth=auth)
        info = data.get("data") or {}
        data, _ = _ragflow_web_request(settings, "/v1/user/set_tenant_info", {
            "tenant_id": info.get("tenant_id"),
            "llm_id": f'{assistant["model"]}@OpenAI-API-Compatible',
            "embd_id": info.get("embd_id"),
            "asr_id": info.get("asr_id", ""),
            "img2txt_id": info.get("img2txt_id", ""),
            "rerank_id": info.get("rerank_id", ""),
        }, auth=auth)
        steps.append("设默认模型：" + ("成功" if data.get("code") == 0 else str(data.get("message", ""))[:60]))
        steps.append(f"内置嵌入模型：{info.get('embd_id', '未知')}")
    except Exception as exc:
        steps.append(f"设默认模型：{str(exc)[:60]}")
    # 统一模型服务（Xinference）已配置嵌入时，新租户同步挂载并设为默认。
    config = read_config()
    xf_url = str(config.get("xinference_api_url") or "").strip()
    xf_embd = str(config.get("xinference_embedding_model") or "").strip()
    if xf_url and xf_embd:
        steps.append(f"Xinference 嵌入（{xf_embd}）：" + _ragflow_tenant_apply_embedding(settings, auth, xf_url, xf_embd))
    return steps


def _ragflow_join_shared_team(settings: dict, email: str, user_auth: str):
    """把用户拉进服务账号团队（可读共享知识库）。失败只记日志不阻塞开通。"""
    try:
        admin_auth = _ragflow_admin_login(settings)
        data, _ = _ragflow_web_request(settings, "/v1/user/tenant_info", auth=admin_auth)
        tenant_id = (data.get("data") or {}).get("tenant_id", "")
        if not tenant_id:
            return
        try:
            _ragflow_web_request(settings, f"/v1/tenant/{tenant_id}/user", {"email": email}, auth=admin_auth)
        except Exception:
            pass  # 已在团队中等情况
        _ragflow_web_request(settings, f"/v1/tenant/agree/{tenant_id}", method="PUT", auth=user_auth)
        # 旧版共享数据集可能还是 permission=me，确保切成 team。
        try:
            _ragflow_request(settings, "PUT", f"/api/v1/datasets/{_ragflow_shared_dataset_id(settings)}", {"permission": "team"})
        except Exception:
            pass
    except Exception as exc:
        print(f"[knowledge] 加入 RAGFlow 共享团队失败（{email}）：{exc}")


KB_RAGFLOW_ACCOUNT_LOCK = threading.RLock()


def _ragflow_user_account(settings: dict, username: str) -> dict:
    entry = _kb_mapping("ragflow_accounts.json").get(str(username or "anonymous")) or {}
    if entry.get("api_key"):
        return entry
    with KB_RAGFLOW_ACCOUNT_LOCK:
        return _ragflow_user_account_locked(settings, username)


def _ragflow_user_account_locked(settings: dict, username: str) -> dict:
    """取/自动开通用户专属 RAGFlow 账号：注册 → API Key → 同步模型 → 加入共享团队。幂等。"""
    key = str(username or "anonymous")
    accounts = _kb_mapping("ragflow_accounts.json")
    entry = accounts.get(key) or {}
    if entry.get("api_key"):
        return entry
    email = entry.get("email") or _ragflow_user_email(key)
    password = entry.get("password") or "Rf@" + secrets.token_hex(8)
    accounts[key] = {**entry, "email": email, "password": password}
    _kb_mapping_save("ragflow_accounts.json", accounts)
    encrypted = _ragflow_encrypt_password(password)
    try:
        _ragflow_web_request(settings, "/v1/user/register", {"nickname": key[:30] or "user", "email": email, "password": encrypted})
    except Exception:
        pass  # 已注册等情况，以登录结果为准
    data, headers = _ragflow_web_request(settings, "/v1/user/login", {"email": email, "password": encrypted})
    if data.get("code") != 0:
        raise ValueError(f"RAGFlow 用户开通失败（{email}）：{data.get('message')}")
    auth = headers.get("Authorization", "")
    data, _ = _ragflow_web_request(settings, "/v1/system/new_token", {}, auth=auth)
    token = (data.get("data") or {}).get("token", "")
    if not token:
        raise ValueError(f"RAGFlow 用户 API Key 生成失败：{data.get('message')}")
    _ragflow_tenant_sync_llm(settings, auth)
    _ragflow_join_shared_team(settings, email, auth)
    entry = {"email": email, "password": password, "api_key": token}
    accounts[key] = entry
    _kb_mapping_save("ragflow_accounts.json", accounts)
    return entry


def _ragflow_user_settings(settings: dict, username: str) -> dict:
    """返回带用户自己 API Key 的 settings（按需自动开通账号）。"""
    account = _ragflow_user_account(settings, username)
    user_settings = dict(settings)
    user_settings["ragflow_api_key"] = account["api_key"]
    return user_settings


def _kb_normalize_scope(settings: dict, scope: str) -> str:
    """共享库的真实数据集 ID 归一化为 'shared'，确保管理员权限校验不被绕过。"""
    scope = str(scope or "mine").strip()
    if settings["provider"] == "ragflow" and scope not in ("mine", "shared"):
        try:
            if scope == _ragflow_shared_dataset_id(settings):
                return "shared"
        except ValueError:
            pass
    return scope


def _ragflow_scope_settings(settings: dict, scope: str, username: str):
    """按作用域返回 (settings, dataset_id)：共享库走服务账号，个人库走用户自己的租户。
    scope 可以是 'mine'（默认个人库）、'shared'，或用户自建知识库的数据集 ID。"""
    if scope == "shared":
        return settings, _ragflow_shared_dataset_id(settings)
    user_settings = _ragflow_user_settings(settings, username)
    if scope and scope != "mine":
        # 指定的自建知识库：归属校验由 RAGFlow 按用户自己的 Key 把关。
        return user_settings, scope
    return user_settings, _ragflow_user_dataset_id(user_settings, username)


def _ragflow_user_query_datasets(settings: dict, username: str):
    """检索/问答用的 (用户settings, 有内容的数据集列表)：覆盖用户租户全部知识库 +
    team 可见的共享库。空数据集会让 RAGFlow 检索报 No chunk found，必须过滤。"""
    return _ragflow_selected_query_datasets(settings, username, None)


def _kb_normalize_selected_kb_ids(value) -> list:
    """把前端传入的知识库选择规范成去重后的 ID 列表；空列表表示全部可见知识库。"""
    if value is None:
        return []
    if isinstance(value, str):
        raw = [item.strip() for item in value.split(",")]
    else:
        raw = [str(item or "").strip() for item in value]
    seen = set()
    result = []
    for item in raw:
        if item and item not in seen:
            seen.add(item)
            result.append(item)
    return result


def _ragflow_selected_query_datasets(settings: dict, username: str, selected_kb_ids=None):
    """返回当前问答范围内有内容的数据集；默认个人/自建库优先，共享库最后补充。"""
    user_settings = _ragflow_user_settings(settings, username)
    default_id = _ragflow_user_dataset_id(user_settings, username)  # 确保默认个人库存在
    try:
        shared_id = _ragflow_shared_dataset_id(settings)
    except ValueError:
        shared_id = ""
    data = _ragflow_request(user_settings, "GET", "/api/v1/datasets?page=1&page_size=200") or []
    selected = set(_kb_normalize_selected_kb_ids(selected_kb_ids))
    personal = []
    shared = []
    other = []
    for item in data:
        dataset_id = str(item.get("id") or "")
        if not dataset_id:
            continue
        if int(item.get("chunk_count", item.get("chunk_num", 0)) or 0) <= 0:
            continue
        if selected and dataset_id not in selected:
            continue
        if dataset_id == shared_id:
            shared.append(dataset_id)
        elif dataset_id == default_id:
            personal.insert(0, dataset_id)
        else:
            other.append(dataset_id)
    dataset_ids = [
        *personal,
        *[item for item in other if item not in personal],
        *[item for item in shared if item not in personal and item not in other],
    ]
    return user_settings, dataset_ids


def _ragflow_chat_cache_key(username: str, dataset_ids) -> str:
    """按知识库集合隔离 RAGFlow chat，避免手动选择不同范围时互相覆盖。"""
    key = username or "anonymous"
    ids = [str(item) for item in (dataset_ids or []) if item]
    if not ids:
        return key
    digest = hashlib.md5("|".join(sorted(ids)).encode("utf-8")).hexdigest()[:12]
    return f"{key}::kb::{digest}"


def _kb_session_selected_kb_ids(session: dict, payload: dict = None) -> list:
    payload = payload or {}
    incoming = _kb_normalize_selected_kb_ids(payload.get("kb_ids"))
    if incoming:
        return incoming
    metadata = (session or {}).get("metadata") or {}
    return _kb_normalize_selected_kb_ids(metadata.get("selected_kb_ids"))


def _kb_update_session_selected_kbs(session_id: str, selected_kb_ids: list):
    getter = globals().get("get_session")
    updater = globals().get("update_session")
    if not (callable(getter) and callable(updater)) or not session_id:
        return
    session = getter(session_id)
    if not session:
        return
    metadata = dict(session.get("metadata") or {})
    metadata["selected_kb_ids"] = _kb_normalize_selected_kb_ids(selected_kb_ids)
    updater(session_id, metadata=metadata)


def kb_datasets(username: str = "") -> dict:
    """当前用户可见的知识库列表：自己租户的全部库（可建多个）+ 全员共享库。"""
    settings = kb_embedding_settings()
    if settings["provider"] != "ragflow":
        return {"ok": True, "provider": "builtin", "kbs": [
            {"id": "mine", "name": "我的知识库", "shared": False, "default": True, "document_count": 0, "chunk_count": 0},
            {"id": "shared", "name": "共享知识库", "shared": True, "default": False, "document_count": 0, "chunk_count": 0},
        ]}
    _ragflow_ready(settings)
    user_settings = _ragflow_user_settings(settings, username)
    default_id = _ragflow_user_dataset_id(user_settings, username)
    try:
        shared_id = _ragflow_shared_dataset_id(settings)
    except ValueError:
        shared_id = ""
    data = _ragflow_request(user_settings, "GET", "/api/v1/datasets?page=1&page_size=200") or []
    kbs = []
    shared_item = None
    for item in data:
        entry = {
            "id": item.get("id", ""),
            "name": item.get("name", ""),
            "shared": item.get("id") == shared_id,
            "default": item.get("id") == default_id,
            "document_count": int(item.get("document_count", item.get("doc_num", 0)) or 0),
            "chunk_count": int(item.get("chunk_count", item.get("chunk_num", 0)) or 0),
        }
        if entry["shared"]:
            shared_item = entry
        else:
            kbs.append(entry)
    kbs.sort(key=lambda e: (not e["default"], e["name"]))
    if shared_item:
        kbs.append(shared_item)
    return {"ok": True, "provider": "ragflow", "kbs": kbs}


def kb_dataset_manage(payload: dict, username: str = "") -> dict:
    """创建/删除用户自己的知识库（RAGFlow 数据集）。"""
    payload = payload or {}
    settings = kb_embedding_settings()
    if settings["provider"] != "ragflow":
        raise ValueError("仅 RAGFlow 引擎支持多知识库")
    _ragflow_ready(settings)
    action = str(payload.get("action", "") or "").strip()
    user_settings = _ragflow_user_settings(settings, username)
    if action == "create":
        name = str(payload.get("name", "") or "").strip()[:50]
        if not name:
            raise ValueError("请输入知识库名称")
        chunk = kb_ragflow_chunk_settings()
        dataset = _ragflow_request(user_settings, "POST", "/api/v1/datasets", {
            "name": name,
            "chunk_method": chunk["chunk_method"],
            "parser_config": {"chunk_token_num": chunk["chunk_token_num"], "layout_recognize": chunk["layout_recognize"]},
        })
        return {"ok": True, "id": dataset["id"], "name": dataset.get("name", name)}
    if action == "delete":
        kb_id = str(payload.get("id", "") or "").strip()
        if not kb_id or _kb_normalize_scope(settings, kb_id) == "shared":
            raise ValueError("共享知识库不能删除")
        if kb_id == _ragflow_user_dataset_id(user_settings, username):
            # 删除默认库后清掉 ID 缓存，下次使用自动重建空库。
            mapping = _kb_mapping("ragflow_datasets.json")
            mapping.pop(f"tenant::{username or 'anonymous'}", None)
            _kb_mapping_save("ragflow_datasets.json", mapping)
        # RAGFlow 仅允许创建者删除数据集，越权由其把关。
        _ragflow_request(user_settings, "DELETE", "/api/v1/datasets", {"ids": [kb_id]})
        return {"ok": True}
    raise ValueError("不支持的操作")


def _ragflow_dataset_has_content(settings: dict, dataset_id: str) -> bool:
    try:
        data = _ragflow_request(settings, "GET", f"/api/v1/datasets/{dataset_id}/documents?page=1&page_size=100") or {}
    except ValueError:
        return False
    return any(int(item.get("chunk_count", item.get("chunk_num", 0)) or 0) > 0 for item in data.get("docs", []))


def _ragflow_user_chat_id(settings: dict, username: str, dataset_ids) -> str:
    """按用户和知识库集合创建专属助手，避免不同检索范围互相覆盖。"""
    user_key = username or "anonymous"
    mapping_key = _ragflow_chat_cache_key(user_key, dataset_ids)
    mapping = _kb_mapping("ragflow_tenant_chats.json")
    entry = mapping.get(mapping_key)
    if not entry and mapping_key != user_key:
        legacy = mapping.get(user_key)
        if isinstance(legacy, dict) and sorted(legacy.get("datasets", [])) == sorted(dataset_ids):
            entry = legacy
            mapping[mapping_key] = legacy
            _kb_mapping_save("ragflow_tenant_chats.json", mapping)
    if isinstance(entry, str):
        entry = {"id": entry, "datasets": []}
    if entry and entry.get("id"):
        body = {"prompt": _ragflow_chat_prompt_payload()}
        if list(entry.get("datasets", [])) != list(dataset_ids):
            body["dataset_ids"] = dataset_ids
            mapping[mapping_key] = {"id": entry["id"], "datasets": dataset_ids}
            _kb_mapping_save("ragflow_tenant_chats.json", mapping)
        _ragflow_request(settings, "PUT", f"/api/v1/chats/{entry['id']}", body)
        return entry["id"]
    suffix = "" if mapping_key == user_key else "-" + mapping_key.rsplit("::", 1)[-1]
    name = f"{KB_RAGFLOW_CHAT_NAME}-{user_key}{suffix}"
    try:
        data = _ragflow_request(settings, "GET", f"/api/v1/chats?page=1&page_size=100&name={urllib.parse.quote(name)}") or []
    except ValueError:
        data = []
    if isinstance(data, dict):
        data = data.get("chats") or data.get("items") or data.get("data") or []
    if not isinstance(data, list):
        raise ValueError("RAGFlow 助手列表格式异常，请检查服务版本")
    chat = next((item for item in data if isinstance(item, dict) and item.get("name") == name), None)
    if not chat:
        chat = _ragflow_request(settings, "POST", "/api/v1/chats", {
            "name": name,
            "dataset_ids": dataset_ids,
            "prompt": _ragflow_chat_prompt_payload(),
        })
    mapping[mapping_key] = {"id": chat["id"], "datasets": dataset_ids}
    if mapping_key == user_key:
        mapping[user_key] = mapping[mapping_key]
    _kb_mapping_save("ragflow_tenant_chats.json", mapping)
    return chat["id"]


def _ragflow_upload(filename: str, raw: bytes, settings: dict, dataset_id: str, username: str = "") -> dict:
    boundary = uuid.uuid4().hex
    body = (
        f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{filename}"\r\n'
        f"Content-Type: application/octet-stream\r\n\r\n".encode("utf-8")
        + raw
        + f"\r\n--{boundary}--\r\n".encode("utf-8")
    )
    docs = _ragflow_request(
        settings,
        "POST",
        f"/api/v1/datasets/{dataset_id}/documents",
        raw_body=body,
        content_type=f"multipart/form-data; boundary={boundary}",
    )
    doc = docs[0] if isinstance(docs, list) and docs else {}
    doc_id = doc.get("id", "")
    if doc_id:
        # 触发 DeepDoc 异步解析，前端通过列表轮询进度。
        _ragflow_request(settings, "POST", f"/api/v1/datasets/{dataset_id}/chunks", {"document_ids": [doc_id]})
    return {
        "id": doc_id,
        "name": doc.get("name", filename),
        "uploaded_by": username,
        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "chars": int(doc.get("size", 0) or 0),
        "chunk_count": 0,
        "embedding_model": "ragflow",
        "status": "parsing",
    }


def _ragflow_doc_status(item: dict) -> str:
    run = str(item.get("run", "") or "").upper()
    progress = float(item.get("progress", 0) or 0)
    if run in ("DONE", "3") or progress >= 1:
        return ""
    if run in ("FAIL", "4"):
        return "failed"
    return "parsing"


def _ragflow_list(settings: dict, dataset_id: str) -> dict:
    data = _ragflow_request(settings, "GET", f"/api/v1/datasets/{dataset_id}/documents?page=1&page_size=200") or {}
    docs = []
    for item in data.get("docs", []):
        created = item.get("create_time", 0)
        try:
            created_text = datetime.fromtimestamp(int(created) / 1000).strftime("%Y-%m-%d %H:%M:%S") if created else ""
        except Exception:
            created_text = str(created)
        progress_msg = str(item.get("progress_msg", "") or "").strip()
        docs.append(
            {
                "id": item.get("id", ""),
                "name": item.get("name", ""),
                "uploaded_by": "",
                "created_at": created_text,
                "chars": int(item.get("size", 0) or 0),
                "chunk_count": int(item.get("chunk_count", item.get("chunk_num", 0)) or 0),
                "embedding_model": "ragflow",
                "status": _ragflow_doc_status(item),
                "progress": round(float(item.get("progress", 0) or 0) * 100),
                # progress_msg 是累积日志，取最后一行即当前阶段。
                "progress_msg": progress_msg.splitlines()[-1][:120] if progress_msg else "",
            }
        )
    return {"ok": True, "docs": docs, "embedding_model": "RAGFlow（DeepDoc 解析 + 混合检索）", "stale_docs": []}


def _ragflow_search(query: str, top_k: int, settings: dict, dataset_ids) -> dict:
    # RAGFlow 0.21 may return an empty ``chunks`` array for ``page_size=1``
    # even when the same query has hits. Fetch at least two candidates and
    # apply the caller's requested limit after normalizing the response.
    fetch_size = max(2, top_k)
    data = _ragflow_request(
        settings,
        "POST",
        "/api/v1/retrieval",
        {"question": query, "dataset_ids": dataset_ids, "page": 1, "page_size": fetch_size},
    ) or {}
    results = []
    for index, chunk in enumerate(data.get("chunks", [])):
        results.append(
            {
                "doc_id": chunk.get("document_id", ""),
                "dataset_id": chunk.get("dataset_id", chunk.get("kb_id", "")),
                "doc_name": chunk.get("document_keyword", chunk.get("docnm_kwd", "")),
                "chunk_index": index,
                "text": chunk.get("content", ""),
                "score": round(float(chunk.get("similarity", 0) or 0), 4),
                "page": _ragflow_chunk_page(chunk),
            }
        )
    return {"ok": True, "results": results[:top_k]}


def _kb_is_generation_request(question: str) -> bool:
    text = str(question or "")
    return any(word in text for word in ("试卷", "考试", "考题", "题目", "培训题", "测试题", "生成一份"))


def _kb_sources_from_search_results(results, top_k: int) -> list:
    sources = []
    for index, item in enumerate((results or [])[:top_k]):
        sources.append({
            "index": index + 1,
            "doc_id": item.get("doc_id", ""),
            "dataset_id": item.get("dataset_id", ""),
            "doc_name": item.get("doc_name", ""),
            "score": round(float(item.get("score", 0) or 0), 4),
            "text": str(item.get("text", ""))[:4000],
            "page": item.get("page", 0),
        })
    return sources


def _kb_generate_answer_from_sources(question: str, sources: list) -> str:
    context = "\n\n".join(
        f"[{idx + 1}]（来自《{item.get('doc_name', '')}》）\n{item.get('text', '')}"
        for idx, item in enumerate(sources)
    )
    system = (
        "你是企业知识库问答助手。只依据提供的资料片段回答。"
        "如果用户要求生成试卷、考试题、培训题或清单，只要资料片段与主题相关，"
        "就基于现有资料生成可用草稿，并说明资料范围；不要因为资料不完整就回答未找到。"
        "只有资料片段全部无关时，才说明知识库中没有找到相关内容。引用时标注 [1]、[2]。"
    )
    return _meeting_llm_chat(system, f"资料片段：\n{context}\n\n用户问题：{question}", temperature=0.2)


def _ragflow_chunk_page(chunk: dict) -> int:
    """片段在 PDF 中的起始页码（positions 格式 [页, x1, x2, y1, y2]，页码从 1 起）；无定位返回 0。"""
    positions = chunk.get("positions") or []
    try:
        return int(positions[0][0]) if positions and positions[0] else 0
    except (TypeError, ValueError, IndexError):
        return 0


def _ragflow_reference_chunks(reference) -> list:
    if isinstance(reference, dict):
        chunks = reference.get("chunks") or []
    elif isinstance(reference, list):
        chunks = reference
    else:
        chunks = []
    return [chunk for chunk in chunks if isinstance(chunk, dict)]


def _ragflow_answer_reference_ids(answer: str) -> list[int]:
    ids = []
    for pattern in (r"\[ID:(\d+)\]", r"##(\d+)\$\$"):
        for match in re.finditer(pattern, str(answer or "")):
            try:
                idx = int(match.group(1))
            except (TypeError, ValueError):
                continue
            if idx not in ids:
                ids.append(idx)
    return ids


def _ragflow_source_from_chunk(chunk: dict, index: int) -> dict:
    return {
        "index": index + 1,
        "doc_id": chunk.get("document_id", ""),
        "dataset_id": chunk.get("dataset_id", chunk.get("kb_id", "")),
        "doc_name": chunk.get("document_name", chunk.get("docnm_kwd", chunk.get("document_keyword", ""))),
        "score": round(float(chunk.get("similarity", 0) or 0), 4),
        "text": str(chunk.get("content", chunk.get("content_with_weight", "")) or "")[:4000],
        "page": _ragflow_chunk_page(chunk),
    }


def _ragflow_normalize_answer_sources(answer: str, reference, top_k: int) -> tuple[str, list]:
    chunks = _ragflow_reference_chunks(reference)
    if not chunks:
        return str(answer or "").strip(), []
    max_sources = max(1, min(int(top_k or 5), 10))
    raw_ids = _ragflow_answer_reference_ids(answer)
    valid_ids = [idx for idx in raw_ids if 0 <= idx < len(chunks)]
    if raw_ids and len(valid_ids) != len(raw_ids):
        selected_ids = list(range(min(len(raw_ids), len(chunks), max_sources)))
        id_map = {old_idx: new_idx for new_idx, old_idx in enumerate(raw_ids[:len(selected_ids)])}
    else:
        selected_ids = (valid_ids if valid_ids else list(range(min(top_k, len(chunks)))))[:max_sources]
        id_map = {old_idx: new_idx for new_idx, old_idx in enumerate(selected_ids)}

    def replace_id(match):
        old_idx = int(match.group(1))
        return f"[ID:{id_map[old_idx]}]" if old_idx in id_map else ""

    normalized_answer = re.sub(r"\[ID:(\d+)\]", replace_id, str(answer or ""))
    normalized_answer = re.sub(r"##(\d+)\$\$", replace_id, normalized_answer).strip()
    sources = [_ragflow_source_from_chunk(chunks[old_idx], new_idx) for new_idx, old_idx in enumerate(selected_ids)]
    return normalized_answer, sources


def _ragflow_session_cache_key(username: str, office_session_id: str = "") -> str:
    username = username or "anonymous"
    office_session_id = str(office_session_id or "").strip()
    return f"{username}::{office_session_id}" if office_session_id else username


def _ragflow_session_id(settings: dict, chat_id: str, username: str, office_session_id: str = "") -> str:
    sessions = _kb_mapping("ragflow_tenant_sessions.json")
    cache_key = _ragflow_session_cache_key(username, office_session_id)
    if office_session_id and chat_id:
        cache_key = f"{cache_key}::chat::{chat_id}"
    cached = sessions.get(cache_key, "")
    if cached:
        return cached
    suffix = ("-" + office_session_id[-8:]) if office_session_id else ""
    data = _ragflow_request(settings, "POST", f"/api/v1/chats/{chat_id}/sessions", {"name": f"office-{username}{suffix}"})
    session_id = data["id"]
    sessions[cache_key] = session_id
    _kb_mapping_save("ragflow_tenant_sessions.json", sessions)
    return session_id


def _ragflow_ask(question: str, top_k: int, username: str, settings: dict, office_session_id: str = "", selected_kb_ids=None) -> dict:
    """用用户自己的租户问答：助手建在用户租户，绑定所选个人/共享知识库。"""
    user_settings, dataset_ids = _ragflow_selected_query_datasets(settings, username, selected_kb_ids)
    if not dataset_ids:
        return {"ok": True, "answer": "知识库中还没有可用文档，请先上传并等待解析完成。", "sources": []}
    if _kb_is_generation_request(question):
        try:
            search = _ragflow_search(question, max(top_k, 8), user_settings, dataset_ids)
            direct_sources = _kb_sources_from_search_results(search.get("results", []), top_k)
            if direct_sources:
                return {"ok": True, "answer": _kb_generate_answer_from_sources(question, direct_sources), "sources": direct_sources}
        except Exception:
            pass
    chat_id = _ragflow_user_chat_id(user_settings, username, dataset_ids)
    session_id = _ragflow_session_id(user_settings, chat_id, username or "anonymous", office_session_id)

    def _completion(sid):
        return _ragflow_request(
            user_settings,
            "POST",
            f"/api/v1/chats/{chat_id}/completions",
            {"question": question, "stream": False, "session_id": sid},
        )

    try:
        data = _completion(session_id)
    except ValueError:
        # 会话可能在 RAGFlow 侧被删除，重建一次。
        office_doc_delete("knowledge", _kb_dir() / "ragflow_tenant_sessions.json")
        session_id = _ragflow_session_id(user_settings, chat_id, username or "anonymous", office_session_id)
        data = _completion(session_id)
    data = data or {}
    reference = data.get("reference") or {}
    answer, sources = _ragflow_normalize_answer_sources(str(data.get("answer", "") or ""), reference, top_k)
    return {"ok": True, "answer": answer, "sources": sources}


# ---------------------------------------------------------------------------
# Qdrant 向量库（REST，stdlib 直连）
# ---------------------------------------------------------------------------

def _qdrant_request(settings: dict, method: str, path: str, body=None):
    request = urllib.request.Request(
        settings["qdrant_url"].rstrip("/") + path,
        data=json.dumps(body, ensure_ascii=False).encode("utf-8") if body is not None else None,
        headers={"Content-Type": "application/json"},
        method=method,
    )
    try:
        with urllib.request.urlopen(request, timeout=30) as response:
            return json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", "ignore")[:200]
        raise ValueError(f"Qdrant 返回 {exc.code}：{detail}") from exc
    except Exception as exc:
        raise ValueError(f"Qdrant 连接失败（{settings['qdrant_url']}）：{exc}") from exc


def _qdrant_collection(model: str) -> str:
    # 不同嵌入模型向量不兼容，按模型分集合。
    return "kb_" + re.sub(r"[^0-9A-Za-z]+", "_", model).strip("_").lower()


def _qdrant_ensure_collection(settings: dict, name: str, dims: int):
    try:
        _qdrant_request(settings, "GET", f"/collections/{name}")
        return
    except ValueError:
        pass
    _qdrant_request(
        settings,
        "PUT",
        f"/collections/{name}",
        {"vectors": {"size": dims, "distance": "Cosine"}},
    )


def _qdrant_upsert_chunks(settings: dict, record: dict, chunks, vectors):
    name = _qdrant_collection(record["embedding_model"])
    _qdrant_ensure_collection(settings, name, len(vectors[0]))
    points = [
        {
            "id": str(uuid.uuid4()),
            "vector": vector,
            "payload": {
                "doc_id": record["id"],
                "doc_name": record["name"],
                "chunk_index": index,
                "text": chunk,
            },
        }
        for index, (chunk, vector) in enumerate(zip(chunks, vectors))
    ]
    _qdrant_request(settings, "PUT", f"/collections/{name}/points?wait=true", {"points": points})


def _qdrant_delete_doc(settings: dict, model: str, doc_id: str):
    name = _qdrant_collection(model)
    try:
        _qdrant_request(
            settings,
            "POST",
            f"/collections/{name}/points/delete?wait=true",
            {"filter": {"must": [{"key": "doc_id", "match": {"value": doc_id}}]}},
        )
    except ValueError:
        # 集合不存在等情况不阻塞删除元数据。
        pass


def _qdrant_search(settings: dict, query_vector, limit: int):
    name = _qdrant_collection(settings["model"])
    try:
        data = _qdrant_request(
            settings,
            "POST",
            f"/collections/{name}/points/search",
            {"vector": query_vector, "limit": limit, "with_payload": True},
        )
    except ValueError as exc:
        if "404" in str(exc):
            return []
        raise
    hits = []
    for item in data.get("result", []):
        payload = item.get("payload", {})
        hits.append(
            {
                "doc_id": payload.get("doc_id", ""),
                "doc_name": payload.get("doc_name", ""),
                "chunk_index": int(payload.get("chunk_index", 0) or 0),
                "text": payload.get("text", ""),
                "score": round(float(item.get("score", 0) or 0), 4),
            }
        )
    return hits


def kb_embed_texts(texts, settings=None):
    """批量调 /v1/embeddings，返回与输入等长的向量列表。"""
    settings = settings or kb_embedding_settings()
    url, key = resolve_speech_api_target(settings["api_url"], settings["api_key"])
    if not url:
        raise ValueError("未配置嵌入模型 API 地址")
    vectors = []
    for start in range(0, len(texts), KB_EMBED_BATCH):
        batch = texts[start:start + KB_EMBED_BATCH]
        request = urllib.request.Request(
            url.rstrip("/") + "/v1/embeddings",
            data=json.dumps({"model": settings["model"], "input": batch}, ensure_ascii=False).encode("utf-8"),
            headers=speech_api_headers(key),
            method="POST",
        )
        try:
            with urllib.request.urlopen(request, timeout=120) as response:
                data = json.loads(response.read().decode("utf-8"))
        except urllib.error.HTTPError as exc:
            detail = exc.read().decode("utf-8", "ignore")[:300]
            raise ValueError(f"嵌入接口返回 {exc.code}：{detail}") from exc
        items = sorted(data.get("data", []), key=lambda item: item.get("index", 0))
        if len(items) != len(batch):
            raise ValueError("嵌入接口返回的向量数量与输入不一致")
        vectors.extend([item["embedding"] for item in items])
    return vectors


def kb_test_embedding(payload: dict) -> dict:
    """设置页"测试连接"：按引擎实测。RAGFlow 测数据集接口；内置引擎测嵌入/Qdrant/重排。"""
    payload = payload or {}
    settings = kb_embedding_settings()
    for key in ("provider", "api_url", "api_key", "model", "vector_store", "qdrant_url", "rerank_model", "ragflow_url", "ragflow_api_key"):
        override = str(payload.get(key, "") or "").strip()
        if override:
            settings[key] = override
    if settings["provider"] == "ragflow":
        _ragflow_ready(settings)
        datasets = _ragflow_request(settings, "GET", "/api/v1/datasets?page=1&page_size=50") or []
        return {"ok": True, "model": "RAGFlow", "dimensions": 0, "ragflow": f"连接正常，{len(datasets)} 个知识库"}
    vectors = kb_embed_texts(["知识库连接测试"], settings)
    result = {"ok": True, "model": settings["model"], "dimensions": len(vectors[0])}
    if settings["vector_store"] == "qdrant":
        _qdrant_request(settings, "GET", "/collections")
        result["qdrant"] = "连接正常"
    if settings["rerank_model"]:
        reranked = _kb_rerank("测试", [{"text": "知识库连接测试"}, {"text": "无关内容"}], 1, settings)
        result["rerank"] = f"{settings['rerank_model']} 正常（top1 分数 {reranked[0]['score'] if reranked else '无'}）"
    return result


# ---------------------------------------------------------------------------
# 文档解析与分块
# ---------------------------------------------------------------------------

KB_SUPPORTED_SUFFIXES = ".txt/.md/.csv/.log/.json/.html/.docx/.pdf/.xlsx/.pptx"


def _kb_extract_text(filename: str, raw: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in (".txt", ".md", ".csv", ".log", ".json"):
        return raw.decode("utf-8", "ignore")
    if suffix in (".html", ".htm"):
        text = raw.decode("utf-8", "ignore")
        text = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", text, flags=re.S | re.I)
        text = re.sub(r"<[^>]+>", "\n", text)
        import html as html_module

        return html_module.unescape(re.sub(r"\n{2,}", "\n", text))
    if suffix == ".docx":
        from docx import Document

        doc = Document(io.BytesIO(raw))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(part for part in parts if part and part.strip())
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
        if not text.strip():
            raise ValueError("PDF 没有可提取的文本层（可能是扫描件，暂不支持 OCR）")
        return text
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        parts = []
        for sheet in workbook.worksheets:
            parts.append(f"【工作表：{sheet.title}】")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                if cells:
                    parts.append(" | ".join(cells))
        workbook.close()
        return "\n".join(parts)
    if suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(raw))
        parts = []
        for index, slide in enumerate(presentation.slides, start=1):
            parts.append(f"【第 {index} 页】")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        line = "".join(run.text for run in para.runs).strip()
                        if line:
                            parts.append(line)
                if getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    if suffix in (".doc", ".xls", ".ppt"):
        raise ValueError(f"暂不支持旧版 Office 格式 {suffix}，请另存为 {suffix}x 后上传")
    raise ValueError(f"暂不支持的文件类型：{suffix or '未知'}（支持 {KB_SUPPORTED_SUFFIXES}）")


def _kb_chunk_text(text: str):
    """按段落聚合成 ~KB_CHUNK_SIZE 字的块，相邻块保留少量重叠保证语义连续。"""
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n|\n", text) if p.strip()]
    chunks = []
    current = ""
    for para in paragraphs:
        # 超长段落先硬切
        while len(para) > KB_CHUNK_SIZE:
            piece, para = para[:KB_CHUNK_SIZE], para[KB_CHUNK_SIZE - KB_CHUNK_OVERLAP:]
            if current:
                chunks.append(current)
                current = ""
            chunks.append(piece)
        if len(current) + len(para) + 1 > KB_CHUNK_SIZE and current:
            chunks.append(current)
            current = current[-KB_CHUNK_OVERLAP:] if KB_CHUNK_OVERLAP else ""
        current = (current + "\n" + para).strip()
    if current.strip():
        chunks.append(current.strip())
    return [c for c in chunks if len(c.strip()) >= 10][:KB_MAX_CHUNKS_PER_DOC]


# ---------------------------------------------------------------------------
# 存储
# ---------------------------------------------------------------------------

def _kb_dir():
    return USER_DATA_DIR / "knowledge"


def _kb_doc_path(doc_id: str):
    safe = Path(str(doc_id or "")).name
    if not safe or not re.match(r"^[0-9A-Za-z_-]+$", safe):
        raise ValueError("文档 ID 无效")
    return _kb_dir() / f"{safe}.json"


def _kb_load_doc(path):
    record = office_doc_read("knowledge", path, None)
    return record if isinstance(record, dict) else None


def _kb_local_docs():
    """本地内置知识库的文档记录（不含 ragflow_*.json 映射），按 id 倒序：[(id, record)]。"""
    entries = office_doc_list("knowledge", _kb_dir())
    docs = []
    for name, record, _ in sorted(entries, key=lambda entry: entry[0], reverse=True):
        stem = name[:-5]
        if not re.match(r"^[0-9A-Za-z_-]+$", stem) or stem.startswith(("ragflow_", "knowledge_legacy_")):
            continue
        if isinstance(record, dict) and record.get("id"):
            docs.append((stem, record))
    return docs


def kb_upload(payload: dict, username: str = "") -> dict:
    payload = payload or {}
    filename = Path(str(payload.get("filename", "") or "document.txt")).name
    data = str(payload.get("data", "") or "")
    if "," in data:
        data = data.split(",", 1)[1]
    raw = base64.b64decode(data) if data else b""
    if not raw:
        raise ValueError("文件内容为空")

    settings = kb_embedding_settings()
    limit = KB_RAGFLOW_MAX_FILE_BYTES if settings["provider"] == "ragflow" else KB_MAX_FILE_BYTES
    if len(raw) > limit:
        raise ValueError(f"文件超过 {limit // (1024 * 1024)}MB 限制")
    scope = _kb_normalize_scope(settings, payload.get("scope"))
    if settings["provider"] == "ragflow":
        _ragflow_ready(settings)
        if scope == "shared" and not _kb_is_admin(username):
            raise ValueError("共享知识库仅管理员可以上传")
        scope_settings, dataset_id = _ragflow_scope_settings(settings, scope, username)
        return {"ok": True, "doc": _ragflow_upload(filename, raw, scope_settings, dataset_id, username)}

    text = _kb_extract_text(filename, raw)
    chunks = _kb_chunk_text(text)
    if not chunks:
        raise ValueError("没有从文件中解析出有效文本")

    vectors = kb_embed_texts(chunks, settings)

    doc_id = "kb" + datetime.now().strftime("%Y%m%d%H%M%S") + uuid.uuid4().hex[:6]
    record = {
        "id": doc_id,
        "name": filename,
        "uploaded_by": username,
        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "chars": len(text),
        "chunk_count": len(chunks),
        "embedding_model": settings["model"],
        "store": settings["vector_store"],
    }
    if settings["vector_store"] == "qdrant":
        _qdrant_upsert_chunks(settings, record, chunks, vectors)
        record["chunks"] = []
    else:
        record["chunks"] = [{"text": chunk, "vector": vector} for chunk, vector in zip(chunks, vectors)]
    office_doc_write("knowledge", _kb_doc_path(doc_id), record)
    return {"ok": True, "doc": _kb_doc_meta(record)}


def _kb_doc_meta(record: dict) -> dict:
    return {
        "id": record.get("id", ""),
        "name": record.get("name", ""),
        "uploaded_by": record.get("uploaded_by", ""),
        "created_at": record.get("created_at", ""),
        "chars": int(record.get("chars", 0) or 0),
        "chunk_count": int(record.get("chunk_count", 0) or 0),
        "embedding_model": record.get("embedding_model", ""),
    }


def kb_list(scope: str = "mine", username: str = "") -> dict:
    settings = kb_embedding_settings()
    if settings["provider"] == "ragflow":
        _ragflow_ready(settings)
        scope = _kb_normalize_scope(settings, scope)
        scope_settings, dataset_id = _ragflow_scope_settings(settings, scope, username)
        result = _ragflow_list(scope_settings, dataset_id)
        result["scope"] = scope
        result["provider"] = "ragflow"
        result["can_manage"] = scope != "shared" or _kb_is_admin(username)
        return result
    docs = [_kb_doc_meta(record) for _, record in _kb_local_docs()]
    current_model = kb_embedding_settings()["model"]
    stale = [d["name"] for d in docs if d["embedding_model"] and d["embedding_model"] != current_model]
    return {"ok": True, "docs": docs, "embedding_model": current_model, "stale_docs": stale, "provider": "builtin"}


def kb_delete(payload: dict, username: str = "") -> dict:
    payload = payload or {}
    settings = kb_embedding_settings()
    if settings["provider"] == "ragflow":
        _ragflow_ready(settings)
        scope = _kb_normalize_scope(settings, payload.get("scope"))
        if scope == "shared" and not _kb_is_admin(username):
            raise ValueError("共享知识库仅管理员可以删除文档")
        scope_settings, dataset_id = _ragflow_scope_settings(settings, scope, username)
        _ragflow_request(scope_settings, "DELETE", f"/api/v1/datasets/{dataset_id}/documents", {"ids": [str(payload.get("id", "") or "")]})
        return {"ok": True}
    path = _kb_doc_path(str(payload.get("id", "") or ""))
    record = _kb_load_doc(path)
    if record:
        user = find_user(username)
        is_admin = bool(user and user.get("role") in ("admin", "superadmin"))
        if record.get("uploaded_by") != username and not is_admin:
            raise ValueError("只能删除自己上传的文档")
        if record.get("store") == "qdrant":
            _qdrant_delete_doc(kb_embedding_settings(), record.get("embedding_model", ""), record.get("id", ""))
        office_doc_delete("knowledge", path)
    return {"ok": True}


# ---------------------------------------------------------------------------
# 检索与问答
# ---------------------------------------------------------------------------

def _cosine(a, b) -> float:
    dot = sum(x * y for x, y in zip(a, b))
    norm_a = math.sqrt(sum(x * x for x in a))
    norm_b = math.sqrt(sum(y * y for y in b))
    if not norm_a or not norm_b:
        return 0.0
    return dot / (norm_a * norm_b)


def _kb_rerank(query: str, candidates, top_k: int, settings: dict):
    """可选重排：调 OpenAI 兼容 /v1/rerank（如 Xinference 的 bge-reranker-v2-m3）精排候选。"""
    url, key = resolve_speech_api_target(settings["api_url"], settings["api_key"])
    request = urllib.request.Request(
        url.rstrip("/") + "/v1/rerank",
        data=json.dumps(
            {
                "model": settings["rerank_model"],
                "query": query,
                "documents": [item["text"] for item in candidates],
                "top_n": top_k,
            },
            ensure_ascii=False,
        ).encode("utf-8"),
        headers=speech_api_headers(key),
        method="POST",
    )
    with urllib.request.urlopen(request, timeout=60) as response:
        data = json.loads(response.read().decode("utf-8"))
    reranked = []
    for item in data.get("results", [])[:top_k]:
        index = int(item.get("index", -1))
        if 0 <= index < len(candidates):
            candidate = dict(candidates[index])
            candidate["score"] = round(float(item.get("relevance_score", 0) or 0), 4)
            reranked.append(candidate)
    return reranked


# FastGPT ⇄ RAGFlow 打通：内网适配器，供 FastGPT 工作流的 HTTP 节点实时检索
# RAGFlow 知识库。用管理端（服务账号）RAGFlow Key，默认检索其可见的全部有内容
# 数据集（含共享企业知识库）；由 RAGFlow 出检索结果，FastGPT 无需再配嵌入模型。
FASTGPT_KB_BRIDGE_SECRET = os.environ.get("FASTGPT_KB_BRIDGE_SECRET") or read_config().get("fastgpt_kb_bridge_secret", "")


def fastgpt_kb_list() -> dict:
    settings = kb_embedding_settings()
    _ragflow_ready(settings)
    data = _ragflow_request(settings, "GET", "/api/v1/datasets?page=1&page_size=200") or []
    items = []
    for item in data:
        items.append({
            "id": str(item.get("id") or ""),
            "name": item.get("name", ""),
            "chunk_count": int(item.get("chunk_count", item.get("chunk_num", 0)) or 0),
        })
    return {"ok": True, "datasets": items}


def fastgpt_kb_retrieval(payload: dict) -> dict:
    payload = payload or {}
    query = str(payload.get("query") or payload.get("question") or "").strip()
    if not query:
        raise ValueError("query 不能为空")
    top_k = min(20, max(1, int(payload.get("top_k") or payload.get("topK") or 6)))
    settings = kb_embedding_settings()
    _ragflow_ready(settings)
    dataset_ids = [str(x).strip() for x in (payload.get("dataset_ids") or payload.get("datasetIds") or []) if str(x).strip()]
    if not dataset_ids:
        data = _ragflow_request(settings, "GET", "/api/v1/datasets?page=1&page_size=200") or []
        dataset_ids = [str(it.get("id")) for it in data if int(it.get("chunk_count", it.get("chunk_num", 0)) or 0) > 0 and it.get("id")]
    if not dataset_ids:
        return {"ok": True, "results": [], "context": ""}
    res = _ragflow_search(query, top_k, settings, dataset_ids)
    results = res.get("results", [])
    context = "\n\n".join(
        f"【资料{index + 1}：{item.get('doc_name') or '知识库'}】\n{item.get('text', '')}"
        for index, item in enumerate(results)
    )
    return {"ok": True, "results": results, "context": context, "count": len(results)}


# ---------------------------------------------------------------------------
# FastGPT「API 文件库」(apiDataset) 连接器：让 RAGFlow 企业知识库以 FastGPT 原生
# 知识库形式接入，可在「知识库检索」节点直接选用、带引用来源。FastGPT 按下述契约
# 调用本服务（Authorization: Bearer <密钥>），RAGFlow 为文档来源，改文档只需在
# RAGFlow 侧维护后在 FastGPT 重新同步即可。契约（v4.15.2 实测）：
#   POST /v1/file/list  {searchKey,parentId} -> {success,data:[{id,name,type,parentId,hasChild,updateTime,createTime}]}
#   GET  /v1/file/content?id=  -> {success,data:{title,content,previewUrl}}
#   GET  /v1/file/detail?id=   -> {success,data:{id,name,parentId,type,updateTime,createTime}}
#   GET  /v1/file/read?id=     -> {success,data:{url}}
# id 约定：数据集(文件夹)="ds::<datasetId>"；文档(文件)="doc::<datasetId>::<docId>"。
# 企业库使用共享凭据；个人库使用 HMAC(username) 凭据，并从
# /apidataset/personal/<username> 路由进入。这样即使 FastGPT 所有账户在同一
# team，也不能把 A 用户的数据源改成 B 用户的数据源来读取。


def _apidataset_iso(value) -> str:
    """RAGFlow 的毫秒时间戳/字符串 -> ISO8601，供 FastGPT 判断是否需要重新同步。"""
    if not value:
        return ""
    try:
        ts = float(value)
        if ts > 1e12:  # 毫秒
            ts /= 1000.0
        import datetime
        return datetime.datetime.utcfromtimestamp(ts).strftime("%Y-%m-%dT%H:%M:%SZ")
    except Exception:
        return str(value)


def fastgpt_apidataset_personal_token(username: str) -> str:
    """Return the identity-bound API 文件库 credential for one office user."""
    username = str(username or "").strip()
    if not username:
        return ""
    if not FASTGPT_KB_BRIDGE_SECRET:
        raise ValueError('请配置 FASTGPT_KB_BRIDGE_SECRET 后使用知识库桥接服务')
    return hmac.new(
        FASTGPT_KB_BRIDGE_SECRET.encode("utf-8"),
        ("apidataset:personal:" + username).encode("utf-8"),
        hashlib.sha256,
    ).hexdigest()


def _apidataset_personal_scope(settings: dict, username: str):
    """Return this user's RAGFlow credentials and the one private default KB."""
    if not username:
        return settings, None
    user_settings = _ragflow_user_settings(settings, username)
    return user_settings, _ragflow_user_dataset_id(user_settings, username)


def _apidataset_content_datasets(settings: dict, username: str = "") -> list:
    settings, personal_dataset_id = _apidataset_personal_scope(settings, username)
    data = _ragflow_request(settings, "GET", "/api/v1/datasets?page=1&page_size=200") or []
    out = []
    for it in data:
        if personal_dataset_id and str(it.get("id") or "") != str(personal_dataset_id):
            continue
        if int(it.get("chunk_count", it.get("chunk_num", 0)) or 0) <= 0:
            continue
        out.append(it)
    return out


def _apidataset_documents(settings: dict, dataset_id: str) -> list:
    docs = []
    page = 1
    while True:
        data = _ragflow_request(
            settings, "GET",
            f"/api/v1/datasets/{dataset_id}/documents?page={page}&page_size=100",
        ) or {}
        batch = data.get("docs", data if isinstance(data, list) else []) or []
        docs.extend(batch)
        if len(batch) < 100:
            break
        page += 1
    return docs


def fastgpt_apidataset_list(parent_id: str = "", search_key: str = "", username: str = "") -> dict:
    """FastGPT /v1/file/list：根级列 RAGFlow 数据集（文件夹），文件夹内列文档（文件）。"""
    settings = kb_embedding_settings()
    _ragflow_ready(settings)
    parent_id = str(parent_id or "").strip()
    search_key = str(search_key or "").strip().lower()
    items = []
    if not parent_id or parent_id in ("root", "/"):
        for ds in _apidataset_content_datasets(settings, username):
            name = ds.get("name", "")
            if search_key and search_key not in name.lower():
                continue
            items.append({
                "id": f"ds::{ds.get('id')}",
                "parentId": "",
                "name": name,
                "type": "folder",
                "hasChild": True,
                "updateTime": _apidataset_iso(ds.get("update_time")),
                "createTime": _apidataset_iso(ds.get("create_time")),
            })
    elif parent_id.startswith("ds::"):
        dataset_id = parent_id[len("ds::"):]
        user_settings, personal_dataset_id = _apidataset_personal_scope(settings, username)
        if personal_dataset_id and dataset_id != str(personal_dataset_id):
            raise ValueError("无权访问该个人知识库")
        for doc in _apidataset_documents(user_settings, dataset_id):
            name = doc.get("name", "")
            if search_key and search_key not in name.lower():
                continue
            if int(doc.get("chunk_count", 0) or 0) <= 0:
                continue
            items.append({
                "id": f"doc::{dataset_id}::{doc.get('id')}",
                "parentId": parent_id,
                "name": name,
                "type": "file",
                "hasChild": False,
                "updateTime": _apidataset_iso(doc.get("update_time")),
                "createTime": _apidataset_iso(doc.get("create_time")),
            })
    return {"success": True, "data": items}


def _apidataset_parse_file_id(file_id: str):
    file_id = str(file_id or "").strip()
    if file_id.startswith("doc::"):
        rest = file_id[len("doc::"):]
        dataset_id, _, doc_id = rest.partition("::")
        return "file", dataset_id, doc_id
    if file_id.startswith("ds::"):
        return "folder", file_id[len("ds::"):], ""
    return "unknown", "", ""


def fastgpt_apidataset_content(file_id: str, username: str = "") -> dict:
    """FastGPT /v1/file/content：拼接文档全部分块文本，供 FastGPT 切块入库。"""
    settings = kb_embedding_settings()
    _ragflow_ready(settings)
    kind, dataset_id, doc_id = _apidataset_parse_file_id(file_id)
    if kind != "file":
        raise ValueError("仅文件支持读取内容")
    settings, personal_dataset_id = _apidataset_personal_scope(settings, username)
    if personal_dataset_id and dataset_id != str(personal_dataset_id):
        raise ValueError("无权访问该个人知识库")
    title = ""
    parts = []
    page = 1
    while True:
        data = _ragflow_request(
            settings, "GET",
            f"/api/v1/datasets/{dataset_id}/documents/{doc_id}/chunks?page={page}&page_size=100",
        ) or {}
        doc_meta = data.get("doc") or data.get("document") or {}
        if not title:
            title = doc_meta.get("name", "")
        chunks = data.get("chunks", []) or []
        for ch in chunks:
            text = ch.get("content") or ch.get("content_with_weight") or ""
            if text:
                parts.append(text.strip())
        total = int(data.get("total", 0) or 0)
        if len(chunks) < 100 or page * 100 >= total:
            break
        page += 1
    content = "\n\n".join(p for p in parts if p)
    return {"success": True, "data": {"title": title, "content": content, "previewUrl": ""}}


def fastgpt_apidataset_detail(file_id: str, username: str = "") -> dict:
    """FastGPT /v1/file/detail：返回文件/文件夹元数据。"""
    settings = kb_embedding_settings()
    _ragflow_ready(settings)
    kind, dataset_id, doc_id = _apidataset_parse_file_id(file_id)
    settings, personal_dataset_id = _apidataset_personal_scope(settings, username)
    if personal_dataset_id and dataset_id != str(personal_dataset_id):
        raise ValueError("无权访问该个人知识库")
    if kind == "folder":
        data = _ragflow_request(settings, "GET", "/api/v1/datasets?page=1&page_size=200") or []
        ds = next((d for d in data if str(d.get("id")) == dataset_id), None)
        if not ds:
            raise ValueError("数据集不存在")
        return {"success": True, "data": {
            "id": file_id, "parentId": "", "name": ds.get("name", ""), "type": "folder",
            "updateTime": _apidataset_iso(ds.get("update_time")),
            "createTime": _apidataset_iso(ds.get("create_time")),
        }}
    if kind == "file":
        docs = _apidataset_documents(settings, dataset_id)
        doc = next((d for d in docs if str(d.get("id")) == doc_id), None)
        if not doc:
            raise ValueError("文档不存在")
        return {"success": True, "data": {
            "id": file_id, "parentId": f"ds::{dataset_id}", "name": doc.get("name", ""), "type": "file",
            "updateTime": _apidataset_iso(doc.get("update_time")),
            "createTime": _apidataset_iso(doc.get("create_time")),
        }}
    raise ValueError("无效的文件标识")


def fastgpt_apidataset_read(file_id: str, username: str = "") -> dict:
    """FastGPT /v1/file/read：返回原文查看链接（RAGFlow 公网地址，可为空）。"""
    settings = kb_embedding_settings()
    kind, dataset_id, doc_id = _apidataset_parse_file_id(file_id)
    _, personal_dataset_id = _apidataset_personal_scope(settings, username)
    if personal_dataset_id and dataset_id != str(personal_dataset_id):
        raise ValueError("无权访问该个人知识库")
    public = str(settings.get("ragflow_public_url") or "").strip().rstrip("/")
    url = ""
    if kind == "file" and public and doc_id:
        url = f"{public}/document/{doc_id}"
    return {"success": True, "data": {"url": url}}


def kb_search(payload: dict, username: str = "") -> dict:
    payload = payload or {}
    query = str(payload.get("query", "") or "").strip()
    if not query:
        raise ValueError("query is required")
    top_k = min(20, max(1, int(payload.get("top_k", 5) or 5)))
    settings = kb_embedding_settings()
    if settings["provider"] == "ragflow":
        _ragflow_ready(settings)
        # 用用户自己的 Key 检索：个人库在自己租户，共享库经 team 成员权限可读。
        user_settings, dataset_ids = _ragflow_selected_query_datasets(settings, username, payload.get("kb_ids"))
        if not dataset_ids:
            return {"ok": True, "results": []}
        return _ragflow_search(query, top_k, user_settings, dataset_ids)
    # 启用重排时先取更宽的粗排候选，再精排到 top_k。
    recall_k = min(40, top_k * 4) if settings["rerank_model"] else top_k

    query_vector = kb_embed_texts([query[:1000]], settings)[0]

    scored = []
    skipped = []
    # 路径一：Qdrant 向量库（配置启用时）。
    if settings["vector_store"] == "qdrant":
        scored.extend(_qdrant_search(settings, query_vector, recall_k))
    # 路径二：本地 JSON 向量（旧数据或 local 模式），与 Qdrant 同为余弦分数可直接合并。
    for stem, record in _kb_local_docs():
        # 不同嵌入模型的向量空间不兼容，跳过并提示重新入库。
        if record.get("embedding_model") and record["embedding_model"] != settings["model"]:
            skipped.append(record.get("name", stem))
            continue
        for index, chunk in enumerate(record.get("chunks", [])):
            vector = chunk.get("vector")
            if not vector or len(vector) != len(query_vector):
                continue
            scored.append(
                {
                    "doc_id": record.get("id", ""),
                    "doc_name": record.get("name", ""),
                    "chunk_index": index,
                    "text": chunk.get("text", ""),
                    "score": round(_cosine(query_vector, vector), 4),
                }
            )
    # 去重（同一块可能两路都命中）并粗排。
    unique = {}
    for item in scored:
        unique.setdefault((item["doc_id"], item["chunk_index"]), item)
    candidates = sorted(unique.values(), key=lambda item: item["score"], reverse=True)[:recall_k]

    warning_parts = []
    if settings["rerank_model"] and candidates:
        try:
            candidates = _kb_rerank(query, candidates, top_k, settings)
        except Exception as exc:
            warning_parts.append(f"重排失败已按向量分数返回：{exc}")
            candidates = candidates[:top_k]
    else:
        candidates = candidates[:top_k]

    if skipped:
        warning_parts.append("以下文档与当前嵌入模型不一致，未参与检索（需删除后重新上传）：" + "、".join(sorted(set(skipped))))
    result = {"ok": True, "results": candidates}
    if warning_parts:
        result["warning"] = "；".join(warning_parts)
    return result


def kb_ask(payload: dict, username: str = "") -> dict:
    payload = payload or {}
    question = str(payload.get("question", "") or "").strip()
    if not question:
        raise ValueError("question is required")
    settings = kb_embedding_settings()
    if settings["provider"] == "ragflow":
        _ragflow_ready(settings)
        return _ragflow_ask(
            question,
            int(payload.get("top_k", 5) or 5),
            username,
            settings,
            str(payload.get("session_id", "") or ""),
            payload.get("kb_ids"),
        )
    search = kb_search({"query": question, "top_k": int(payload.get("top_k", 5) or 5)}, username)
    results = search.get("results", [])
    if not results:
        return {"ok": True, "answer": "知识库中没有找到相关内容。", "sources": [], "warning": search.get("warning", "")}

    context = "\n\n".join(f"[{i + 1}]（来自《{item['doc_name']}》）\n{item['text']}" for i, item in enumerate(results))
    answer = _meeting_llm_chat(
        KB_ASK_SYSTEM,
        f"资料片段：\n{context}\n\n问题：{question}",
        temperature=0.2,
    )
    sources = [
        {"index": i + 1, "doc_name": item["doc_name"], "score": item["score"], "text": item["text"][:200]}
        for i, item in enumerate(results)
    ]
    return {"ok": True, "answer": answer, "sources": sources, "warning": search.get("warning", "")}


# ---------------------------------------------------------------------------
# RAGFlow 平台内管理：初始化、模型同步、解析配置、总览（管理员无需登 RAGFlow 后台）
# ---------------------------------------------------------------------------

# RAGFlow 前端内置的固定 RSA 公钥（登录/注册接口要求密码用它加密）。
RAGFLOW_LOGIN_PUBKEY = (
    "-----BEGIN PUBLIC KEY-----\n"
    "MIIBIjANBgkqhkiG9w0BAQEFAAOCAQ8AMIIBCgKCAQEArq9XTUSeYr2+N1h3Afl/"
    "z8Dse/2yD0ZGrKwx+EEEcdsBLca9Ynmx3nIB5obmLlSfmskLpBo0UACBmB5rEjBp"
    "2Q2f3AG3Hjd4B+gNCG6BDaawuDlgANIhGnaTLrIqWrrcm4EMzJOnAOI1fgzJRsOO"
    "UEfaS318Eq9OVO3apEyCCt0lOQK6PuksduOjVxtltDav+guVAA068NrPYmRNabVK"
    "RNLJpL8w4D44sfth5RvZ3q9t+6RTArpEtc5sh5ChzvqPOzKGMXW83C95TxmXqpbK"
    "6olN4RevSfVjEAgCydH6HN6OhtOQEcnrU97r9H0iZOWwbw3pVrZiUkuRD1R56Wzs"
    "2wIDAQAB\n"
    "-----END PUBLIC KEY-----"
)
RAGFLOW_DEFAULT_ADMIN_EMAIL = "admin@chencytech.com"

RAGFLOW_CHUNK_METHODS = [
    {"id": "naive", "label": "通用（默认）"},
    {"id": "qa", "label": "问答对"},
    {"id": "table", "label": "表格"},
    {"id": "paper", "label": "论文"},
    {"id": "book", "label": "书籍"},
    {"id": "laws", "label": "法规制度"},
    {"id": "manual", "label": "手册"},
    {"id": "presentation", "label": "PPT 演示"},
    {"id": "one", "label": "整篇不切分"},
]


def _ragflow_encrypt_password(password: str) -> str:
    from cryptography.hazmat.primitives import serialization
    from cryptography.hazmat.primitives.asymmetric import padding as rsa_padding

    key = serialization.load_pem_public_key(RAGFLOW_LOGIN_PUBKEY.encode())
    cipher = key.encrypt(base64.b64encode(str(password).encode()), rsa_padding.PKCS1v15())
    return base64.b64encode(cipher).decode()


def _ragflow_web_request(settings: dict, path: str, body=None, auth: str = "", method=None):
    """RAGFlow 前端 /v1 管理接口（登录态），返回 (json, headers)。"""
    headers = {"Content-Type": "application/json"}
    if auth:
        headers["Authorization"] = auth
    request = urllib.request.Request(
        settings["ragflow_url"].rstrip("/") + path,
        data=json.dumps(body, ensure_ascii=False).encode("utf-8") if body is not None else None,
        headers=headers,
        method=method or ("POST" if body is not None else "GET"),
    )
    opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
    try:
        response = opener.open(request, timeout=60)
    except urllib.error.HTTPError as error:
        routes = {
            "/v1/user/login": ("/api/v1/auth/login", "POST"),
            "/v1/user/register": ("/api/v1/users", "POST"),
            "/v1/system/new_token": ("/api/v1/system/tokens", "POST"),
            "/v1/system/token_list": ("/api/v1/system/tokens", "GET"),
            "/v1/system/version": ("/api/v1/system/version", "GET"),
            "/v1/user/tenant_info": ("/api/v1/users/me/models", "GET"),
            "/v1/user/set_tenant_info": ("/api/v1/users/me/models", "PATCH"),
        }
        invite = re.fullmatch(r"/v1/tenant/([^/]+)/user", path)
        agree = re.fullmatch(r"/v1/tenant/agree/([^/]+)", path)
        if invite:
            routes[path] = (f"/api/v1/tenants/{invite.group(1)}/users", "POST")
        if agree:
            routes[path] = (f"/api/v1/tenants/{agree.group(1)}", "PATCH")
        if path in routes and error.code in (404, 405):
            error.close()
            target, target_method = routes[path]
            return _ragflow_web_request(settings, target, body, auth, target_method)
        raise
    with response:
        return json.loads(response.read().decode("utf-8")), response.headers


def kb_ragflow_chunk_settings() -> dict:
    config = read_config()
    return {
        "chunk_method": str(config.get("knowledge_ragflow_chunk_method") or "").strip() or "naive",
        "chunk_token_num": int(config.get("knowledge_ragflow_chunk_size") or 512),
        # DeepDOC=OCR+版面识别（扫描件/复杂排版，慢）；Plain Text=直接抽文本层（快一到两个数量级）。
        "layout_recognize": str(config.get("knowledge_ragflow_layout") or "").strip() or "DeepDOC",
    }


def kb_ragflow_status() -> dict:
    """RAGFlow 服务状态 + 全部知识库总览（管理员视角）。"""
    settings = kb_embedding_settings()
    config = read_config()
    result = {"ok": True, "configured": bool(settings["ragflow_api_key"]), "reachable": False,
              "datasets": [], "chunk_methods": RAGFLOW_CHUNK_METHODS, **kb_ragflow_chunk_settings(),
              "admin_email": str(config.get("knowledge_ragflow_admin_email") or ""),
              "sso_available": bool(config.get("knowledge_ragflow_admin_password"))}
    try:
        data, _ = _ragflow_web_request(settings, "/v1/system/version")
        result["reachable"] = True
        result["version"] = str((data or {}).get("data", "") or "")[:40]
    except Exception as exc:
        result["error"] = f"RAGFlow 服务不可达：{exc}"
        return result
    if not settings["ragflow_api_key"]:
        return result
    try:
        datasets = _ragflow_request(settings, "GET", "/api/v1/datasets?page=1&page_size=200") or []
        for item in datasets:
            result["datasets"].append(
                {
                    "id": item.get("id", ""),
                    "name": item.get("name", ""),
                    "document_count": int(item.get("document_count", item.get("doc_num", 0)) or 0),
                    "chunk_count": int(item.get("chunk_count", item.get("chunk_num", 0)) or 0),
                    "chunk_method": item.get("chunk_method", item.get("parser_id", "")),
                }
            )
    except Exception as exc:
        result["error"] = str(exc)
    return result


def kb_ragflow_bootstrap(payload: dict) -> dict:
    """一键初始化 RAGFlow：注册/登录服务账号 → API Key → 同步大模型 → 设默认。幂等可重复执行。"""
    payload = payload or {}
    settings = kb_embedding_settings()
    config = read_config()
    email = str(payload.get("email", "") or config.get("knowledge_ragflow_admin_email") or RAGFLOW_DEFAULT_ADMIN_EMAIL).strip()
    password = str(payload.get("password", "") or config.get("knowledge_ragflow_admin_password") or "").strip()
    if not password:
        password = "Chency@Ragflow" + secrets.token_hex(4)
    steps = []
    encrypted = _ragflow_encrypt_password(password)

    # 1. 注册（已存在则忽略）
    try:
        data, _ = _ragflow_web_request(settings, "/v1/user/register", {"nickname": "office-admin", "email": email, "password": encrypted})
        steps.append("注册账号：" + ("成功" if data.get("code") == 0 else str(data.get("message", ""))[:60]))
    except Exception as exc:
        steps.append(f"注册账号：{str(exc)[:60]}")

    # 2. 登录
    data, headers = _ragflow_web_request(settings, "/v1/user/login", {"email": email, "password": encrypted})
    if data.get("code") != 0:
        raise ValueError(f"RAGFlow 登录失败：{data.get('message')}（如账号已存在但密码不同，请在设置里填写正确密码）")
    auth = headers.get("Authorization", "")
    steps.append("登录：成功")
    config["knowledge_ragflow_admin_email"] = email
    config["knowledge_ragflow_admin_password"] = password

    # 3. API Key（没有才生成）
    if not settings["ragflow_api_key"]:
        data, _ = _ragflow_web_request(settings, "/v1/system/new_token", {}, auth=auth)
        token = (data.get("data") or {}).get("token", "")
        if not token:
            raise ValueError(f"生成 API Key 失败：{data.get('message')}")
        config["knowledge_ragflow_api_key"] = token
        settings["ragflow_api_key"] = token
        steps.append("生成 API Key：成功")
    else:
        steps.append("API Key：已存在")

    # 4/5. 同步大模型并设为默认（与用户租户开通共用同一套逻辑）
    steps.extend(_ragflow_tenant_sync_llm(settings, auth))

    config["knowledge_provider"] = "ragflow"
    write_config(config)
    return {"ok": True, "steps": steps}


def kb_ragflow_sso(username: str = "") -> dict:
    """用当前用户自己的 RAGFlow 租户账号换取登录态，返回免登录直达控制台的链接。

    RAGFlow 前端的 /login?auth=<JWT> 是其 OAuth 回跳通道，会直接写入登录态。
    每个办公助手用户对应一个 RAGFlow 租户：控制台里只见自己的个人知识库，
    共享知识库以团队资源形式可读。账号按需自动开通，用户无需记 RAGFlow 密码。
    """
    settings = kb_embedding_settings()
    _ragflow_ready(settings)
    account = _ragflow_user_account(settings, username)
    data, headers = _ragflow_web_request(
        settings, "/v1/user/login",
        {"email": account["email"], "password": _ragflow_encrypt_password(account["password"])},
    )
    if data.get("code") != 0:
        raise ValueError(f"RAGFlow 登录失败：{data.get('message')}")
    auth = headers.get("Authorization", "")
    if not auth:
        raise ValueError("RAGFlow 未返回登录凭证")
    public = settings["ragflow_public_url"]
    return {
        "ok": True,
        "auth": auth,
        "email": account["email"],
        "url": f"{public}/login?auth={urllib.parse.quote(auth)}" if public else "",
    }


def kb_ragflow_save_chunk_config(payload: dict) -> dict:
    """保存默认解析配置并应用到已有的全部知识库数据集。"""
    payload = payload or {}
    method = str(payload.get("chunk_method", "") or "naive").strip()
    if method not in {item["id"] for item in RAGFLOW_CHUNK_METHODS}:
        raise ValueError(f"不支持的分块模板：{method}")
    token_num = max(64, min(2048, int(payload.get("chunk_token_num", 512) or 512)))
    layout = str(payload.get("layout_recognize", "") or "DeepDOC").strip()
    if layout not in ("DeepDOC", "Plain Text"):
        raise ValueError("版面识别只支持 DeepDOC 或 Plain Text")
    config = read_config()
    config["knowledge_ragflow_chunk_method"] = method
    config["knowledge_ragflow_chunk_size"] = token_num
    config["knowledge_ragflow_layout"] = layout
    write_config(config)

    settings = kb_embedding_settings()
    _ragflow_ready(settings)
    updated = []
    failed = []
    seen = set()
    # 服务租户（共享库）+ 各用户租户（个人库只能用本人 Key 修改）逐一应用。
    key_rounds = [settings] + [
        {**settings, "ragflow_api_key": entry.get("api_key", "")}
        for entry in _kb_mapping("ragflow_accounts.json").values()
        if entry.get("api_key")
    ]
    for round_settings in key_rounds:
        try:
            datasets = _ragflow_request(round_settings, "GET", "/api/v1/datasets?page=1&page_size=200") or []
        except ValueError:
            continue
        for item in datasets:
            if item["id"] in seen:
                continue
            seen.add(item["id"])
            try:
                _ragflow_request(round_settings, "PUT", f"/api/v1/datasets/{item['id']}",
                                 {"chunk_method": method, "parser_config": {"chunk_token_num": token_num, "layout_recognize": layout}})
                updated.append(item.get("name", item["id"]))
            except Exception:
                failed.append(item.get("name", item["id"]))
    result = {"ok": True, "updated": updated, **kb_ragflow_chunk_settings()}
    if failed:
        result["warning"] = "以下知识库更新失败（可能有文档正在解析）：" + "、".join(failed)
    result["hint"] = "新配置对之后上传的文档生效；已有文档需点\"重新解析\"。"
    return result



def _kb_agent_session_owned(session_id: str, username: str) -> dict:
    session_id = str(session_id or "").strip()
    if not session_id:
        return {}
    getter = globals().get("get_session")
    session = getter(session_id) if callable(getter) else None
    if not session or session.get("user_id") != (username or "anonymous") or session.get("kind") != "knowledge":
        return {}
    if session.get("status") and session.get("status") != "active":
        return {}
    return session


def _kb_format_session(session: dict) -> dict:
    getter = globals().get("get_messages")
    messages = getter(session["id"], limit=500) if callable(getter) and session.get("id") else []
    title = str(session.get("title") or "").strip() or "新会话"
    return {
        "id": session.get("id", ""),
        "title": title,
        "created_at": session.get("created_at", ""),
        "updated_at": session.get("updated_at", ""),
        "message_count": len(messages),
        "selected_kb_ids": list((session.get("metadata") or {}).get("selected_kb_ids") or []),
    }


def _kb_legacy_import_flags() -> dict:
    return _kb_mapping("knowledge_legacy_chat_imports.json")


def _kb_save_legacy_import_flags(flags: dict):
    _kb_mapping_save("knowledge_legacy_chat_imports.json", flags)


def _kb_ragflow_sources(reference, top_k: int = 10) -> list:
    chunks = []
    if isinstance(reference, dict):
        chunks = reference.get("chunks") or []
    elif isinstance(reference, list):
        chunks = reference
    sources = []
    for index, chunk in enumerate(chunks[:top_k]):
        if not isinstance(chunk, dict):
            continue
        sources.append({
            "index": index + 1,
            "doc_id": chunk.get("document_id", ""),
            "dataset_id": chunk.get("dataset_id", chunk.get("kb_id", "")),
            "doc_name": chunk.get("document_name", chunk.get("docnm_kwd", "")),
            "score": round(float(chunk.get("similarity", 0) or 0), 4),
            "text": str(chunk.get("content", ""))[:400],
            "page": _ragflow_chunk_page(chunk),
        })
    return sources


def _kb_load_legacy_ragflow_messages(username: str = "") -> list:
    settings = kb_embedding_settings()
    if settings["provider"] != "ragflow" or not settings["ragflow_api_key"]:
        return []
    chat_id, session_id = _kb_chat_session_ids(username, "")
    if not (chat_id and session_id):
        return []
    try:
        user_settings = _ragflow_user_settings(settings, username)
        data = _ragflow_request(
            user_settings, "GET",
            f"/api/v1/chats/{chat_id}/sessions?page=1&page_size=1&id={session_id}",
        ) or []
    except ValueError:
        return []
    raw_messages = (data[0].get("messages") if data else []) or []
    messages = []
    for i, msg in enumerate(raw_messages):
        role = msg.get("role")
        content = str(msg.get("content", "") or "")
        if role == "assistant" and i == 0:
            continue
        if role not in ("user", "assistant") or not content:
            continue
        item = {"role": role, "content": content, "metadata": {"feature": "knowledge", "legacy_ragflow": True}}
        if role == "assistant":
            item["metadata"]["sources"] = _kb_ragflow_sources(msg.get("reference") or [], top_k=10)
        messages.append(item)
    return messages


def _kb_import_legacy_session_if_needed(username: str) -> dict:
    key = username or "anonymous"
    flags = _kb_legacy_import_flags()
    if flags.get(key):
        return {}
    messages = _kb_load_legacy_ragflow_messages(username)
    flags[key] = True
    _kb_save_legacy_import_flags(flags)
    if not messages:
        return {}
    _, legacy_ragflow_session_id = _kb_chat_session_ids(username, "")
    session = _kb_create_session(username, "历史会话", metadata={"synced_from": "ragflow", "ragflow_session_id": legacy_ragflow_session_id})
    for msg in messages[-100:]:
        _kb_add_chat_message(session["id"], msg["role"], msg["content"], msg.get("metadata") or {"feature": "knowledge"})
    return globals().get("get_session")(session["id"]) if callable(globals().get("get_session")) else session


def _kb_ragflow_time_text(value) -> str:
    try:
        num = int(value or 0)
        if num > 10_000_000_000:
            num = num // 1000
        if num > 0:
            return datetime.fromtimestamp(num).strftime("%Y-%m-%d %H:%M")
    except Exception:
        pass
    return _kb_default_session_title()


def _kb_raw_ragflow_messages(raw_messages) -> list:
    messages = []
    for i, msg in enumerate(raw_messages or []):
        if not isinstance(msg, dict):
            continue
        role = msg.get("role")
        content = str(msg.get("content", "") or "")
        if role == "assistant" and i == 0:
            continue
        if role not in ("user", "assistant") or not content:
            continue
        metadata = {"feature": "knowledge", "synced_from": "ragflow"}
        if role == "assistant":
            metadata["sources"] = _kb_ragflow_sources(msg.get("reference") or [], top_k=10)
        messages.append({"role": role, "content": content, "metadata": metadata})
    return messages


def _kb_title_from_ragflow_session(item: dict, messages: list) -> str:
    date_text = _kb_ragflow_time_text(item.get("update_time") or item.get("create_time") or item.get("created_at"))
    topic = ""
    for msg in messages:
        if msg.get("role") == "user":
            topic = re.sub(r"\s+", " ", str(msg.get("content") or "")).strip()[:24]
            break
    if not topic:
        name = str(item.get("name") or "").strip()
        if name and not name.startswith("office-"):
            topic = name[:24]
    return f"{date_text} {topic or 'RAGFlow 会话'}"[:80]


def _kb_existing_ragflow_session_map(username: str) -> dict:
    db_connect = globals().get("_db_connect")
    if not callable(db_connect):
        return {}
    with db_connect() as conn:
        rows = conn.execute(
            """
            SELECT id, metadata_json FROM agent_sessions
            WHERE user_id = ? AND kind = 'knowledge' AND status = 'active'
            """,
            (username or "anonymous",),
        ).fetchall()
    result = {}
    for row in rows:
        meta_raw = dict(row).get("metadata_json", "{}")
        try:
            meta = json.loads(meta_raw or "{}")
        except Exception:
            meta = {}
        ragflow_session_id = str(meta.get("ragflow_session_id") or "").strip()
        if ragflow_session_id:
            result[ragflow_session_id] = dict(row).get("id", "")
    return result


def _kb_ragflow_sessions_for_user(username: str) -> tuple[str, list]:
    settings = kb_embedding_settings()
    if settings["provider"] != "ragflow" or not settings["ragflow_api_key"]:
        return "", []
    try:
        user_settings, dataset_ids = _ragflow_user_query_datasets(settings, username)
        if not dataset_ids:
            return "", []
        chat_id = _ragflow_user_chat_id(user_settings, username, dataset_ids)
        data = _ragflow_request(user_settings, "GET", f"/api/v1/chats/{chat_id}/sessions?page=1&page_size=100") or []
    except Exception as exc:
        print(f"[knowledge] 同步 RAGFlow 会话失败（{username}）：{exc}")
        return "", []
    if isinstance(data, dict):
        sessions = data.get("sessions") or data.get("data") or data.get("items") or []
    else:
        sessions = data
    return chat_id, [item for item in sessions if isinstance(item, dict) and item.get("id")]


def _kb_sync_ragflow_sessions(username: str):
    chat_id, ragflow_sessions = _kb_ragflow_sessions_for_user(username)
    if not chat_id or not ragflow_sessions:
        return
    existing = _kb_existing_ragflow_session_map(username)
    for item in ragflow_sessions:
        ragflow_session_id = str(item.get("id") or "").strip()
        if not ragflow_session_id or ragflow_session_id in existing:
            continue
        messages = _kb_raw_ragflow_messages(item.get("messages") or [])
        title = _kb_title_from_ragflow_session(item, messages)
        session = _kb_create_session(
            username,
            title,
            metadata={
                "synced_from": "ragflow",
                "ragflow_chat_id": chat_id,
                "ragflow_session_id": ragflow_session_id,
            },
        )
        for msg in messages[-100:]:
            _kb_add_chat_message(session["id"], msg["role"], msg["content"], msg.get("metadata") or {"feature": "knowledge"})
        existing[ragflow_session_id] = session["id"]


def _kb_list_sessions(username: str, limit: int = 30) -> list:
    init = globals().get("init_db")
    if callable(init):
        init()
    db_connect = globals().get("_db_connect")
    if not callable(db_connect):
        return []
    limit = max(1, min(int(limit or 30), 100))
    with db_connect() as conn:
        rows = conn.execute(
            """
            SELECT s.*, COUNT(m.id) AS message_count
            FROM agent_sessions s
            LEFT JOIN agent_messages m ON m.session_id = s.id
            WHERE s.user_id = ? AND s.kind = 'knowledge' AND s.status = 'active'
            GROUP BY s.id
            ORDER BY s.updated_at DESC
            LIMIT ?
            """,
            (username or "anonymous", limit),
        ).fetchall()
    sessions = []
    for row in rows:
        item = dict(row)
        title = str(item.get("title") or "").strip() or "新会话"
        try:
            metadata = json.loads(item.get("metadata_json") or "{}")
        except Exception:
            metadata = {}
        sessions.append({
            "id": item.get("id", ""),
            "title": title,
            "created_at": item.get("created_at", ""),
            "updated_at": item.get("updated_at", ""),
            "message_count": int(item.get("message_count", 0) or 0),
            "selected_kb_ids": list(metadata.get("selected_kb_ids") or []),
        })
    return sessions


def _kb_default_session_title() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M")


def _kb_create_session(username: str, title: str = "", metadata: dict = None) -> dict:
    creator = globals().get("create_session")
    if not callable(creator):
        raise ValueError("会话服务不可用")
    title = str(title or "").strip()[:80] or _kb_default_session_title()
    meta = {"feature": "knowledge"}
    if metadata:
        meta.update(metadata)
    session = creator(username or "anonymous", kind="knowledge", title=title, metadata=meta)
    getter = globals().get("get_session")
    return getter(session["id"]) if callable(getter) else session


def _kb_ensure_session(username: str, session_id: str = "", title: str = "") -> dict:
    session = _kb_agent_session_owned(session_id, username)
    if session:
        return session
    sessions = _kb_list_sessions(username, limit=1)
    if sessions and not session_id:
        getter = globals().get("get_session")
        existing = getter(sessions[0]["id"]) if callable(getter) else None
        if existing:
            return existing
    return _kb_create_session(username, title)


def _kb_message_to_chat(item: dict) -> dict:
    msg = {"role": item.get("role", ""), "content": item.get("content", "")}
    metadata = item.get("metadata") or {}
    if item.get("role") == "assistant":
        msg["sources"] = metadata.get("sources") or []
        if metadata.get("failed"):
            msg["failed"] = True
    return msg


def kb_chat_sessions(payload: dict, username: str = "") -> dict:
    payload = payload or {}
    action = str(payload.get("action", "list") or "list").strip()
    if action == "list":
        _kb_sync_ragflow_sessions(username)
        sessions = _kb_list_sessions(username, payload.get("limit", 30))
        if not sessions:
            imported = _kb_import_legacy_session_if_needed(username)
            if imported:
                sessions = _kb_list_sessions(username, payload.get("limit", 30))
        if not sessions:
            session = _kb_create_session(username)
            sessions = [_kb_format_session(session)]
        return {"ok": True, "sessions": sessions}
    if action == "create":
        selected_kb_ids = _kb_normalize_selected_kb_ids(payload.get("kb_ids"))
        metadata = {"selected_kb_ids": selected_kb_ids} if selected_kb_ids else None
        session = _kb_create_session(username, payload.get("title", ""), metadata=metadata)
        return {"ok": True, "session": _kb_format_session(session)}
    if action == "get":
        session = _kb_agent_session_owned(payload.get("session_id", ""), username)
        if not session:
            _kb_sync_ragflow_sessions(username)
            session = _kb_agent_session_owned(payload.get("session_id", ""), username)
        if not session:
            raise ValueError("会话不存在")
        getter = globals().get("get_messages")
        messages = getter(session["id"], limit=payload.get("limit", 100)) if callable(getter) else []
        return {"ok": True, "session": _kb_format_session(session), "messages": [_kb_message_to_chat(m) for m in messages]}
    if action in ("clear", "delete"):
        session = _kb_agent_session_owned(payload.get("session_id", ""), username)
        if not session:
            raise ValueError("会话不存在")
        updater = globals().get("update_session")
        if callable(updater):
            updater(session["id"], status="archived")
        _kb_delete_ragflow_session(username, session["id"])
        return {"ok": True}
    raise ValueError("不支持的会话操作")


def _kb_update_session_title_from_question(session_id: str, question: str):
    getter = globals().get("get_session")
    session = getter(session_id) if callable(getter) else None
    if not session:
        return
    current_title = str(session.get("title") or "").strip()
    date_match = re.match(r"^(\d{4}-\d{2}-\d{2} \d{2}:\d{2})(?:\s|$)", current_title)
    if current_title and current_title != "新会话" and not date_match:
        return
    summary = re.sub(r"\s+", " ", str(question or "")).strip()[:24]
    if not summary:
        return
    prefix = date_match.group(1) if date_match else _kb_default_session_title()
    title = f"{prefix} {summary}"[:80]
    updater = globals().get("update_session")
    if callable(updater):
        updater(session_id, title=title)


def _kb_bind_ragflow_session(office_session_id: str, chat_id: str, ragflow_session_id: str):
    if not (office_session_id and ragflow_session_id):
        return
    getter = globals().get("get_session")
    updater = globals().get("update_session")
    if not (callable(getter) and callable(updater)):
        return
    session = getter(office_session_id)
    if not session:
        return
    metadata = dict(session.get("metadata") or {})
    if metadata.get("ragflow_session_id") == ragflow_session_id and metadata.get("ragflow_chat_id") == chat_id:
        return
    metadata.update({"synced_from": "ragflow", "ragflow_chat_id": chat_id, "ragflow_session_id": ragflow_session_id})
    updater(office_session_id, metadata=metadata)


def _kb_add_chat_message(session_id: str, role: str, content: str, metadata: dict = None):
    adder = globals().get("add_message")
    if callable(adder) and session_id and content is not None:
        adder(session_id, role, str(content), metadata=metadata or {"feature": "knowledge"})


def _kb_delete_ragflow_session(username: str, office_session_id: str = ""):
    settings = kb_embedding_settings()
    if settings["provider"] != "ragflow":
        return
    chat_id, session_id = _kb_chat_session_ids(username, office_session_id)
    getter = globals().get("get_session")
    office_session = getter(office_session_id) if callable(getter) and office_session_id else None
    metadata = (office_session or {}).get("metadata") or {}
    chat_id = str(metadata.get("ragflow_chat_id") or chat_id or "").strip()
    session_id = str(metadata.get("ragflow_session_id") or session_id or "").strip()
    sessions = _kb_mapping("ragflow_tenant_sessions.json")
    if sessions:
        sessions.pop(_ragflow_session_cache_key(username, office_session_id), None)
        for key, value in list(sessions.items()):
            if session_id and value == session_id:
                sessions.pop(key, None)
        _kb_mapping_save("ragflow_tenant_sessions.json", sessions)
    if chat_id and session_id:
        try:
            user_settings = _ragflow_user_settings(settings, username)
            _ragflow_request(user_settings, "DELETE", f"/api/v1/chats/{chat_id}/sessions", {"ids": [session_id]})
        except ValueError as exc:
            print(f"[knowledge] 删除 RAGFlow 会话失败（{username}/{session_id}）：{exc}")

def kb_ask_stream(handler, payload: dict, username: str = ""):
    """SSE 流式问答：按办公平台知识库会话保存上下文和历史。"""
    payload = payload or {}
    question = str(payload.get("question", "") or "").strip()
    top_k = min(10, max(1, int(payload.get("top_k", 5) or 5)))
    session = _kb_ensure_session(username, str(payload.get("session_id", "") or ""), title="新会话")
    office_session_id = session["id"]
    selected_kb_ids = _kb_session_selected_kb_ids(session, payload)
    if _kb_normalize_selected_kb_ids(payload.get("kb_ids")):
        _kb_update_session_selected_kbs(office_session_id, selected_kb_ids)
        getter = globals().get("get_session")
        session = getter(office_session_id) if callable(getter) else session

    def send_event(obj):
        handler.wfile.write(f"data: {json.dumps(obj, ensure_ascii=False)}\n\n".encode("utf-8"))
        handler.wfile.flush()

    handler.send_response(200)
    handler.send_header("Content-Type", "text/event-stream; charset=utf-8")
    handler.send_header("Cache-Control", "no-cache")
    handler.send_header("X-Accel-Buffering", "no")
    handler.end_headers()
    send_event({"session_id": office_session_id, "session": _kb_format_session(session)})

    final_answer = ""
    sources = []
    try:
        if not question:
            raise ValueError("question is required")
        _kb_add_chat_message(office_session_id, "user", question, {"feature": "knowledge"})
        _kb_update_session_title_from_question(office_session_id, question)
        settings = kb_embedding_settings()
        if settings["provider"] != "ragflow":
            result = kb_ask({"question": question, "top_k": top_k, "session_id": office_session_id, "kb_ids": selected_kb_ids}, username)
            final_answer = result.get("answer", "")
            sources = result.get("sources", [])
            send_event({"answer": final_answer})
            send_event({"done": True, "sources": sources, "session_id": office_session_id})
            _kb_add_chat_message(office_session_id, "assistant", final_answer, {"feature": "knowledge", "sources": sources})
            return
        _ragflow_ready(settings)
        user_settings, dataset_ids = _ragflow_selected_query_datasets(settings, username, selected_kb_ids)
        if not dataset_ids:
            final_answer = "知识库中还没有可用文档，请先上传并等待解析完成。"
            send_event({"answer": final_answer})
            send_event({"done": True, "sources": [], "session_id": office_session_id})
            _kb_add_chat_message(office_session_id, "assistant", final_answer, {"feature": "knowledge", "sources": []})
            return
        if _kb_is_generation_request(question):
            try:
                search = _ragflow_search(question, max(top_k, 8), user_settings, dataset_ids)
                sources = _kb_sources_from_search_results(search.get("results", []), top_k)
                if sources:
                    final_answer = _kb_generate_answer_from_sources(question, sources)
                    send_event({"answer": final_answer})
                    send_event({"done": True, "sources": sources, "session_id": office_session_id})
                    _kb_add_chat_message(office_session_id, "assistant", final_answer, {"feature": "knowledge", "sources": sources})
                    return
            except Exception:
                sources = []
        chat_id = _ragflow_user_chat_id(user_settings, username, dataset_ids)
        ragflow_session_id = _ragflow_session_id(user_settings, chat_id, username or "anonymous", office_session_id)
        _kb_bind_ragflow_session(office_session_id, chat_id, ragflow_session_id)

        def open_stream(sid):
            request = urllib.request.Request(
                user_settings["ragflow_url"].rstrip("/") + f"/api/v1/chats/{chat_id}/completions",
                data=json.dumps({"question": question, "stream": True, "session_id": sid}, ensure_ascii=False).encode("utf-8"),
                headers={"Content-Type": "application/json", "Authorization": "Bearer " + user_settings["ragflow_api_key"]},
                method="POST",
            )
            return urllib.request.build_opener(urllib.request.ProxyHandler({})).open(request, timeout=300)

        try:
            response = open_stream(ragflow_session_id)
        except urllib.error.HTTPError:
            _kb_delete_ragflow_session(username, office_session_id)
            ragflow_session_id = _ragflow_session_id(user_settings, chat_id, username or "anonymous", office_session_id)
            _kb_bind_ragflow_session(office_session_id, chat_id, ragflow_session_id)
            response = open_stream(ragflow_session_id)

        reference = {}
        with response:
            for raw_line in response:
                line = raw_line.decode("utf-8", "ignore").strip()
                if not line.startswith("data:"):
                    continue
                try:
                    packet = json.loads(line[5:])
                except json.JSONDecodeError:
                    continue
                data = packet.get("data")
                if not isinstance(data, dict):
                    continue
                if data.get("reference"):
                    reference = data["reference"]
                answer = str(data.get("answer", "") or "")
                if answer:
                    final_answer = answer
                    send_event({"answer": answer})
        if final_answer:
            normalized_answer, sources = _ragflow_normalize_answer_sources(final_answer, reference, top_k)
            if normalized_answer != final_answer:
                final_answer = normalized_answer
                send_event({"answer": final_answer})
        if not final_answer:
            final_answer = "知识库中没有找到相关内容。"
            send_event({"answer": final_answer})
        _kb_add_chat_message(office_session_id, "assistant", final_answer, {"feature": "knowledge", "sources": sources})
        send_event({"done": True, "sources": sources, "session_id": office_session_id})
    except Exception as exc:
        error_text = str(exc)
        try:
            send_event({"error": error_text, "session_id": office_session_id})
        except Exception:
            pass
        if question:
            _kb_add_chat_message(office_session_id, "assistant", "回答失败：" + error_text, {"feature": "knowledge", "sources": [], "failed": True})


def _kb_chat_session_ids(username: str, office_session_id: str = ""):
    """当前用户的 (chat_id, ragflow_session_id)，任一不存在返回空串。"""
    entry = _kb_mapping("ragflow_tenant_chats.json").get(username or "anonymous")
    chat_id = (entry.get("id") if isinstance(entry, dict) else entry) or ""
    sessions = _kb_mapping("ragflow_tenant_sessions.json")
    key = _ragflow_session_cache_key(username or "anonymous", office_session_id)
    return chat_id, str(sessions.get(key, "") or "")


def kb_chat_history(username: str = "", session_id: str = "") -> dict:
    """兼容旧接口：返回指定知识库会话的消息；未指定时返回最近会话。"""
    session = _kb_ensure_session(username, session_id)
    getter = globals().get("get_messages")
    messages = getter(session["id"], limit=100) if callable(getter) else []
    return {
        "ok": True,
        "session": _kb_format_session(session),
        "messages": [_kb_message_to_chat(m) for m in messages],
    }


def kb_chat_clear(username: str = "", session_id: str = "") -> dict:
    """清空当前办公平台知识库会话，并删除对应 RAGFlow 上下文。"""
    session = _kb_agent_session_owned(session_id, username) if session_id else _kb_ensure_session(username)
    if session:
        updater = globals().get("update_session")
        if callable(updater):
            updater(session["id"], status="archived")
        _kb_delete_ragflow_session(username, session["id"])
    return {"ok": True}


def kb_chat_embed(username: str = "") -> dict:
    """返回当前用户的 RAGFlow 原生对话页嵌入链接（/chat/share）。

    RAGFlow 自带网站嵌入方案：share 页通过租户 API Token 的 beta 凭据免登录访问，
    自带流式输出、引用角标与文档预览。助手建在用户自己租户，绑定个人库 + 共享库。
    """
    settings = kb_embedding_settings()
    if settings["provider"] != "ragflow":
        raise ValueError("仅 RAGFlow 引擎支持原生对话嵌入")
    _ragflow_ready(settings)
    account = _ragflow_user_account(settings, username)
    user_settings = {**settings, "ragflow_api_key": account["api_key"]}

    # 绑定当前有内容的知识库（个人全部库 + 共享库），并确保助手存在。
    _, dataset_ids = _ragflow_user_query_datasets(settings, username)
    if not dataset_ids:
        return {"ok": True, "url": "", "empty": True,
                "hint": "知识库还没有可用文档，上传并解析完成后即可对话。"}
    chat_id = _ragflow_user_chat_id(user_settings, username, dataset_ids)

    # share 页的 beta 凭据：随租户 API Token 生成，缓存进账号映射避免每次查询。
    beta = str(account.get("beta") or "").strip()
    if not beta:
        data, headers = _ragflow_web_request(
            settings, "/v1/user/login",
            {"email": account["email"], "password": _ragflow_encrypt_password(account["password"])},
        )
        if data.get("code") != 0:
            raise ValueError(f"RAGFlow 登录失败：{data.get('message')}")
        data, _ = _ragflow_web_request(settings, "/v1/system/token_list", auth=headers.get("Authorization", ""))
        tokens = data.get("data") or []
        beta = str((tokens[0] if tokens else {}).get("beta") or "").strip()
        if not beta:
            raise ValueError("RAGFlow 未返回嵌入凭据（beta token）")
        accounts = _kb_mapping("ragflow_accounts.json")
        entry = accounts.get(username or "anonymous") or account
        entry["beta"] = beta
        accounts[username or "anonymous"] = entry
        _kb_mapping_save("ragflow_accounts.json", accounts)

    public = settings["ragflow_public_url"] or settings["ragflow_url"]
    url = f"{public.rstrip('/')}/chat/share?shared_id={chat_id}&from=chat&auth={urllib.parse.quote(beta)}"
    return {"ok": True, "url": url, "empty": False}


def kb_external_api_info(username: str = "") -> dict:
    """Return the current user's RAGFlow OpenAI-compatible chat endpoint.

    The API key itself is intentionally never returned to the browser. Users
    obtain or rotate it in their own RAGFlow console; the chat id is safe to
    show and is bound to the user's personal plus shared knowledge datasets.
    """
    settings = kb_embedding_settings()
    public = str(settings.get("ragflow_public_url") or settings.get("ragflow_url") or "").strip().rstrip("/")
    result = {"ok": True, "provider": settings.get("provider"), "public_url": public, "ready": False,
              "chat_id": "", "base_url": "", "completion_url": "", "dataset_count": 0}
    if settings.get("provider") != "ragflow":
        result["hint"] = "当前知识库未使用 RAGFlow 引擎，请在引擎设置中启用后再接入标准 RAGFlow API。"
        return result
    _ragflow_ready(settings)
    user_settings = _ragflow_user_settings(settings, username)
    _, dataset_ids = _ragflow_user_query_datasets(settings, username)
    result["dataset_count"] = len(dataset_ids)
    if not dataset_ids:
        result["hint"] = "当前没有解析完成的知识库文档；上传并解析完成后会自动创建可供外部调用的 Chat Assistant。"
        return result
    chat_id = _ragflow_user_chat_id(user_settings, username, dataset_ids)
    base_url = f"{public}/api/v1/chats_openai/{chat_id}"
    result.update({"ready": True, "chat_id": chat_id, "base_url": base_url,
                   "completion_url": base_url + "/chat/completions",
                   "api_key_masked": "已为当前用户开通，请在 RAGFlow 控制台的 API Key 页面查看或轮换。"})
    return result


def _kb_preview_ids(dataset_id: str, doc_id: str):
    if not re.match(r"^[0-9a-f]{32}$", str(dataset_id or "")) or not re.match(r"^[0-9a-f]{32}$", str(doc_id or "")):
        raise ValueError("文档定位参数无效")
    return dataset_id, doc_id


def _kb_preview_settings(dataset_id: str, username: str) -> dict:
    """预览用的 Key：RAGFlow 的原件下载/分块接口是 owner-only（不认 team 权限）。
    共享库归服务账号租户，用服务 Key（team 成员本就可读）；其余用用户自己的 Key，
    越权访问他人个人库由 RAGFlow 拒绝。"""
    settings = kb_embedding_settings()
    if settings["provider"] != "ragflow":
        raise ValueError("仅 RAGFlow 引擎支持文档预览")
    _ragflow_ready(settings)
    try:
        if dataset_id == _ragflow_shared_dataset_id(settings):
            return settings
    except ValueError:
        pass
    return _ragflow_user_settings(settings, username)


def kb_doc_file(dataset_id: str, doc_id: str, username: str = ""):
    """下载文档原件用于引用预览，返回 (文件名, 字节)。"""
    dataset_id, doc_id = _kb_preview_ids(dataset_id, doc_id)
    user_settings = _kb_preview_settings(dataset_id, username)
    data = _ragflow_request(user_settings, "GET", f"/api/v1/datasets/{dataset_id}/documents?page=1&page_size=1&id={doc_id}") or {}
    docs = data.get("docs", [])
    name = docs[0].get("name", "document") if docs else "document"
    url = user_settings["ragflow_url"].rstrip("/") + f"/api/v1/datasets/{dataset_id}/documents/{doc_id}"
    request = urllib.request.Request(url, headers={"Authorization": "Bearer " + user_settings["ragflow_api_key"]})
    opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
    try:
        with opener.open(request, timeout=120) as response:
            raw = response.read()
    except Exception as exc:
        raise ValueError(f"文档下载失败：{exc}") from exc
    # RAGFlow 出错时返回 JSON 体（如权限不足），不能当文件内容透传。
    if raw[:1] == b"{" and b'"code"' in raw[:200]:
        try:
            err = json.loads(raw.decode("utf-8", "ignore"))
            raise ValueError(f"文档下载失败：{str(err.get('message', ''))[:100]}")
        except (json.JSONDecodeError, AttributeError):
            pass
    return name, raw


def kb_doc_chunks(dataset_id: str, doc_id: str, username: str = "") -> dict:
    """文档的解析分块文本（非 PDF/图片格式的预览方式）。"""
    dataset_id, doc_id = _kb_preview_ids(dataset_id, doc_id)
    user_settings = _kb_preview_settings(dataset_id, username)
    data = _ragflow_request(user_settings, "GET", f"/api/v1/datasets/{dataset_id}/documents/{doc_id}/chunks?page=1&page_size=60") or {}
    doc = data.get("doc", {}) or {}
    chunks = [str(c.get("content", "") or "") for c in (data.get("chunks") or [])]
    return {"ok": True, "name": doc.get("name", ""), "total": int(data.get("total", len(chunks)) or 0), "chunks": chunks}


def kb_reparse(payload: dict, username: str = "") -> dict:
    """重新解析指定文档（模板调整或解析失败后使用）。"""
    payload = payload or {}
    settings = kb_embedding_settings()
    if settings["provider"] != "ragflow":
        raise ValueError("仅 RAGFlow 引擎支持重新解析")
    _ragflow_ready(settings)
    scope = _kb_normalize_scope(settings, payload.get("scope"))
    if scope == "shared" and not _kb_is_admin(username):
        raise ValueError("共享知识库仅管理员可以重新解析")
    scope_settings, dataset_id = _ragflow_scope_settings(settings, scope, username)
    doc_id = str(payload.get("id", "") or "")
    _ragflow_request(scope_settings, "POST", f"/api/v1/datasets/{dataset_id}/chunks", {"document_ids": [doc_id]})
    return {"ok": True}


# ---------------------------------------------------------------------------
# 统一模型服务（Xinference）：一处配置语音识别/语音合成/嵌入模型，
# 一键应用到会议助手、数字人与知识库（含 RAGFlow 全部租户）。
# ---------------------------------------------------------------------------

# 宿主 LAN IP 对办公助手容器和 ragflow-server 容器都可达（host.docker.internal 在 RAGFlow 容器内不一定能解析）。
XINFERENCE_DEFAULT_URL = "http://10.1.6.230:9997"


def xinference_unified_config() -> dict:
    config = read_config()
    asr = meeting_asr_settings()
    tts = digital_human_tts_settings()
    kb = kb_embedding_settings()
    return {
        "ok": True,
        "api_url": str(config.get("xinference_api_url") or "").strip() or XINFERENCE_DEFAULT_URL,
        "asr_model": str(config.get("xinference_asr_model") or "").strip(),
        "tts_model": str(config.get("xinference_tts_model") or "").strip(),
        "embedding_model": str(config.get("xinference_embedding_model") or "").strip(),
        "current": {
            "meeting": {"api_url": asr["api_url"], "model": asr["model"]},
            "tts": {"provider": tts["provider"], "api_url": tts["api_url"], "model": tts["api_model"]},
            "kb": {"provider": kb["provider"], "api_url": kb["api_url"], "model": kb["model"]},
        },
    }


def xinference_probe(payload: dict) -> dict:
    """探测 Xinference /v1/models，按能力分类返回可选模型。"""
    payload = payload or {}
    url = str(payload.get("api_url", "") or "").strip().rstrip("/") or xinference_unified_config()["api_url"]
    opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
    try:
        with opener.open(urllib.request.Request(url + "/v1/models"), timeout=15) as response:
            data = json.loads(response.read().decode("utf-8"))
    except Exception as exc:
        raise ValueError(f"Xinference 连接失败（{url}）：{exc}")
    models = {"embedding": [], "asr": [], "tts": [], "chat": []}
    for item in data.get("data", []):
        name = item.get("model_name") or item.get("id", "")
        ability = item.get("model_ability") or []
        if item.get("model_type") == "embedding":
            models["embedding"].append(name)
        elif "audio2text" in ability:
            models["asr"].append(name)
        elif "text2audio" in ability:
            models["tts"].append(name)
        elif item.get("model_type") == "LLM":
            models["chat"].append(name)
    return {"ok": True, "api_url": url, **models}


def _ragflow_tenant_apply_embedding(settings: dict, auth: str, api_url: str, model: str) -> str:
    """给一个 RAGFlow 租户挂载 Xinference 嵌入模型并设为默认。返回结果描述。"""
    try:
        data, _ = _ragflow_web_request(settings, "/v1/llm/add_llm", {
            "llm_factory": "Xinference",
            "llm_name": model,
            "model_type": "embedding",
            "api_base": api_url.rstrip("/") + "/v1",
            "max_tokens": 8192,
        }, auth=auth)
        if data.get("code") != 0 and "exist" not in str(data.get("message", "")).lower():
            return f"挂载失败：{str(data.get('message', ''))[:60]}"
        data, _ = _ragflow_web_request(settings, "/v1/user/tenant_info", auth=auth)
        info = data.get("data") or {}
        data, _ = _ragflow_web_request(settings, "/v1/user/set_tenant_info", {
            "tenant_id": info.get("tenant_id"),
            "llm_id": info.get("llm_id", ""),
            "embd_id": f"{model}@Xinference",
            "asr_id": info.get("asr_id", ""),
            "img2txt_id": info.get("img2txt_id", ""),
            "rerank_id": info.get("rerank_id", ""),
        }, auth=auth)
        return "成功" if data.get("code") == 0 else f"设默认失败：{str(data.get('message', ''))[:60]}"
    except Exception as exc:
        return f"异常：{str(exc)[:60]}"


def _ragflow_switch_datasets_embedding(settings: dict, model: str, owner: str, warnings: list, skip_ids=()):
    """把该 Key 名下的空数据集切到新嵌入模型；有文档的库无法切换，提示重新上传。"""
    embd_id = f"{model}@Xinference"
    try:
        data = _ragflow_request(settings, "GET", "/api/v1/datasets?page=1&page_size=200") or []
    except ValueError:
        return
    for item in data:
        if item.get("id") in skip_ids:
            continue
        current = str(item.get("embedding_model", "") or "")
        if current == embd_id:
            continue
        chunks = int(item.get("chunk_count", item.get("chunk_num", 0)) or 0)
        if chunks > 0:
            warnings.append(f"{owner}《{item.get('name', '')}》已有文档，需删除文档后重新上传才会使用新嵌入模型")
            continue
        try:
            _ragflow_request(settings, "PUT", f"/api/v1/datasets/{item['id']}", {"embedding_model": embd_id})
        except ValueError as exc:
            warnings.append(f"{owner}《{item.get('name', '')}》切换失败：{str(exc)[:60]}")


def xinference_apply(payload: dict) -> dict:
    """保存统一配置并应用到会议助手、数字人 TTS、知识库嵌入与 RAGFlow 全部租户。"""
    payload = payload or {}
    api_url = str(payload.get("api_url", "") or "").strip().rstrip("/")
    if not api_url:
        raise ValueError("请填写 Xinference 地址")
    asr_model = str(payload.get("asr_model", "") or "").strip()[:100]
    tts_model = str(payload.get("tts_model", "") or "").strip()[:100]
    embedding_model = str(payload.get("embedding_model", "") or "").strip()[:100]
    config = read_config()
    config["xinference_api_url"] = api_url
    config["xinference_asr_model"] = asr_model
    config["xinference_tts_model"] = tts_model
    config["xinference_embedding_model"] = embedding_model
    write_config(config)

    steps = []
    warnings = []
    if asr_model:
        save_meeting_config({"api_url": api_url, "model": asr_model, "api_key_clear": True})
        steps.append(f"会议助手识别 → {asr_model}")
    if tts_model:
        tts = digital_human_tts_settings()
        save_digital_human_tts_config({
            "provider": "api", "voice": tts["voice"], "rate": tts["rate"], "pitch": tts["pitch"],
            "api_url": api_url, "api_model": tts_model, "api_voice": tts["api_voice"], "api_key_clear": True,
        })
        steps.append(f"数字人语音合成 → {tts_model}")
    if embedding_model:
        kb = kb_embedding_settings()
        save_kb_config({
            "provider": kb["provider"], "api_url": api_url, "model": embedding_model,
            "vector_store": kb["vector_store"], "qdrant_url": kb["qdrant_url"],
            "rerank_model": kb["rerank_model"], "ragflow_url": kb["ragflow_url"],
            "api_key_clear": True,
        })
        steps.append(f"知识库嵌入（内置引擎）→ {embedding_model}")
        settings = kb_embedding_settings()
        if settings["provider"] == "ragflow" and settings["ragflow_api_key"]:
            try:
                shared_id = _ragflow_shared_dataset_id(settings)
            except ValueError:
                shared_id = ""
            try:
                auth = _ragflow_admin_login(settings)
                steps.append("RAGFlow 服务租户 → " + _ragflow_tenant_apply_embedding(settings, auth, api_url, embedding_model))
                _ragflow_switch_datasets_embedding(settings, embedding_model, "共享", warnings)
            except Exception as exc:
                warnings.append(f"RAGFlow 服务租户：{str(exc)[:80]}")
            for username, entry in _kb_mapping("ragflow_accounts.json").items():
                try:
                    data, headers = _ragflow_web_request(
                        settings, "/v1/user/login",
                        {"email": entry["email"], "password": _ragflow_encrypt_password(entry["password"])},
                    )
                    if data.get("code") != 0:
                        warnings.append(f"{username}：RAGFlow 登录失败")
                        continue
                    steps.append(f"RAGFlow 租户 {username} → " + _ragflow_tenant_apply_embedding(
                        settings, headers.get("Authorization", ""), api_url, embedding_model))
                    # 共享库已在服务租户轮处理，用户轮跳过避免重复警告。
                    _ragflow_switch_datasets_embedding(
                        {**settings, "ragflow_api_key": entry.get("api_key", "")}, embedding_model, username, warnings,
                        skip_ids=(shared_id,))
                except Exception as exc:
                    warnings.append(f"{username}：{str(exc)[:60]}")
    result = {"ok": True, "steps": steps}
    if warnings:
        result["warning"] = "；".join(warnings)
    return result
