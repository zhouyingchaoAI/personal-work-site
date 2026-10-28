import ast
import base64
import io
import json
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


class AnalysisBlockTemplates(unittest.TestCase):
    def setUp(self):
        source = Path('ai-office-assistant-web/packages/joint-command/backend/crm_excel.py').read_text()
        tree = ast.parse(source)
        wanted = {'crm_analysis_template_export', 'crm_business_import'}
        nodes = [node for node in tree.body if isinstance(node, ast.FunctionDef) and node.name in wanted or isinstance(node, ast.Assign) and any(isinstance(target, ast.Name) and target.id == 'CRM_ANALYSIS_BLOCKS' for target in node.targets)]
        self.data = {'version': 'current', 'sheets': [
            {'key': '01_核心指标总览', 'title': '核心', 'source': '真实来源', 'headers': ['指标', '实际完成', '全年目标'], 'rows': [['订货', 100, 200], ['毛利', 20, 40]]},
            {'key': '06_费用明细', 'title': '费用', 'source': '汇总', 'headers': ['区域', '费用'], 'rows': [['公司合计', 18]]},
            {'key': '实际回款明细', 'title': '回款', 'source': '', 'headers': ['项目名称', '实际回款金额（元）', '回款日期'], 'rows': [['=不是公式', 1234, '2026-10-01']]},
        ]}
        self.saved = {}
        self.env = dict(globals(),
            crm_capabilities=lambda user: {'read': True, 'write': True, 'data_scope': user.get('scope', 'all')},
            crm_business_analysis=lambda *args, **kwargs: {'ok': True, 'data': self.data},
            read_config=lambda: {'users': [{'username': 'admin'}]},
            crm_compute_dashboard_data=lambda sheets: {'count': len(sheets)},
            crm_now=lambda: '2026-10-21T12:00:00', crm_text=lambda value, *args: str(value or ''),
            crmdb_enabled=lambda: True, crmdb_blob_keys=lambda prefix: [],
            crmdb_blob_put=lambda key, value: self.saved.update(value),
        )
        exec(compile(ast.Module(body=nodes, type_ignores=[]), '<templates>', 'exec'), self.env)

    def test_all_data_batch_roundtrip(self):
        raw, name, mime = self.env['crm_analysis_template_export']({'username': 'admin'}, 'current')
        self.assertEqual(mime, 'application/zip')
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            full = archive.read('全部板块_完整当前数据.xlsx')
        result = self.env['crm_business_import']({'data': base64.b64encode(full).decode(), 'filename': 'full.xlsx'}, 'admin')
        self.assertTrue(result['ok'])
        self.assertEqual([sheet['rows'] for sheet in self.saved['sheets']], [sheet['rows'] for sheet in self.data['sheets']])

    def test_partial_update_preserves_other_blocks_and_two_column_headers(self):
        raw, name, mime = self.env['crm_analysis_template_export']({}, blocks=['structure'])
        book = load_workbook(io.BytesIO(raw))
        book['06_费用明细']['B4'] = 99
        output = io.BytesIO()
        book.save(output)
        result = self.env['crm_business_import']({'data': base64.b64encode(output.getvalue()).decode()}, 'admin')
        self.assertTrue(result['ok'])
        self.assertEqual(self.saved['sheets'][0], self.data['sheets'][0])
        self.assertEqual(self.saved['sheets'][1]['rows'], [['公司合计', 99]])
        self.assertEqual(self.saved['dashboard']['count'], 3)

    def test_scope_and_invalid_version_are_rejected(self):
        with self.assertRaises(ValueError):
            self.env['crm_analysis_template_export']({'scope': 'region'})
        with self.assertRaises(ValueError):
            self.env['crm_analysis_template_export']({}, version='missing')

    def test_company_cash_flow_exports_and_roundtrips(self):
        self.data['sheets'].append({'key': '公司月度现金流', 'headers': ['年份', '月份', '经营现金流（万元）'], 'rows': [[2026, 1, -12.5]]})
        raw, _, _ = self.env['crm_analysis_template_export']({}, blocks=['overall'])
        book = load_workbook(io.BytesIO(raw))
        self.assertIn('公司月度现金流', book.sheetnames)
        self.assertEqual(book['公司月度现金流']['C4'].value, -12.5)

    def test_current_page_guidance_and_formatting(self):
        raw, _, _ = self.env['crm_analysis_template_export']({}, blocks=['overall'])
        book = load_workbook(io.BytesIO(raw))
        guide = str(list(book['_板块导入配置'].values))
        for text in ('整表替换', '管理订货', '应收存量', '实际回款明细', 'TOP10', '不使用公式'):
            self.assertIn(text, guide)
        sheet = book['01_核心指标总览']
        self.assertEqual(sheet.freeze_panes, 'A4')
        self.assertEqual(sheet['B4'].value, 100)
        self.assertEqual(sheet['B4'].number_format, '#,##0.00;[Red]-#,##0.00')
        self.assertFalse(sheet.sheet_view.showGridLines)


if __name__ == '__main__':
    unittest.main()
