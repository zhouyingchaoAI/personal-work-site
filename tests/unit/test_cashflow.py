import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import backend.runtime as r
from tests.unit.officedb_testing import db_variant


class CashflowAssistantTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.data_dir = Path(self.temp.name) / "data"
        self.patch = patch.object(r, "USER_DATA_DIR", self.data_dir)
        self.patch.start()

    def tearDown(self):
        self.patch.stop()
        self.temp.cleanup()

    def test_save_calculates_cash_gap_and_receivable_summary(self):
        result = r.cashflow_save({
            "project_name": "示例项目", "region": "广东区", "customer": "示例客户",
            "annual_rate": 0.06,
            "entries": [
                {"date": "2026-10-01", "milestone": "首付款", "receivable": 100000},
                {"date": "2026-09-01", "milestone": "采购", "procurement": 160000},
                {"date": "2026-11-01", "milestone": "服务", "personnel": 10000},
            ],
        }, "tester")
        item = result["item"]
        self.assertEqual(item["region"], "广东区")
        self.assertEqual(item["summary"]["total_receivable"], 100000)
        self.assertEqual(item["summary"]["funding_gap"], 160000)
        self.assertGreater(item["summary"]["funding_cost"], 0)
        overview = r.cashflow_list("tester")
        self.assertEqual(overview["summary"]["total"], 1)
        self.assertEqual(overview["by_region"][0]["name"], "广东区")
        self.assertEqual(overview["by_customer"][0]["outstanding"], 100000)

    def test_rejects_empty_project_or_unordered_invalid_rows(self):
        with self.assertRaises(ValueError):
            r.cashflow_save({"region": "广东区"}, "tester")
        with self.assertRaises(ValueError):
            r.cashflow_save({"project_name": "x", "entries": [{"date": "bad-date"}]}, "tester")

    def test_history_import_summarizes_region_customer_type_and_scale(self):
        result = r.cashflow_history_import({"source_name": "历史回款.xlsx", "records": [
            {"region": "广东", "customer": "客户甲", "project_name": "轨道维保项目", "project_type": "维保款", "contract_amount": 3000000, "planned_amount": 500000, "received_amount": 500000, "status": "已回款", "contract_date": "2025-01-01", "planned_date": "2025-03-01", "received_date": "2025-03-02"},
            {"region": "广东", "customer": "客户甲", "project_name": "轨道维保项目", "project_type": "验收款", "contract_amount": 3000000, "planned_amount": 300000, "received_amount": 0, "status": "未回款", "contract_date": "2025-01-01"},
        ]}, "tester")
        self.assertEqual(result["imported"], 2)
        history = r.cashflow_history_list("tester")
        self.assertEqual(history["summary"]["records"], 2)
        self.assertEqual(history["summary"]["unpaid_amount"], 300000)
        self.assertEqual(history["by_region"][0]["name"], "广东区")
        self.assertEqual(history["by_customer"][0]["received"], 500000)
        self.assertEqual({row["name"] for row in history["by_project_type"]}, {"维保款", "验收款"})
        self.assertEqual(history["by_status"][0]["name"], "已回款")
        self.assertEqual(history["by_year"][0]["name"], "2025")
        profile = history["customer_profiles"][0]
        self.assertEqual(profile["name"], "客户甲")
        self.assertEqual(profile["average_collection_days"], 60)
        self.assertIn("测算", profile["recommendation"])
        self.assertEqual(profile["payment_patterns"][0]["name"], "维保款")
        self.assertEqual(profile["median_collection_days"], 60)
        self.assertEqual(profile["p75_collection_days"], 60)
        self.assertGreater(profile["cashflow_risk_score"], 0)
        self.assertEqual(profile["on_time_collection_rate"], 0)
        self.assertEqual(profile["median_schedule_variance_days"], 1)
        self.assertEqual(history["project_risks"][0]["name"], "轨道维保项目")
        self.assertEqual(history["project_risks"][0]["unpaid"], 300000)
        self.assertEqual(history["project_risks"][0]["customer"], "客户甲")

    def test_history_records_can_be_saved_and_deleted(self):
        imported = r.cashflow_history_import({"records": [{"project_name": "项目A", "customer": "客户A", "planned_amount": 100, "status": "未回款"}]}, "tester")
        self.assertEqual(imported["imported"], 1)
        row = r.cashflow_history_records("tester")["items"][0]
        saved = r.cashflow_history_save({**row, "region": "广东区", "received_amount": 60, "status": "部分回款"}, "tester")
        self.assertEqual(saved["item"]["received_amount"], 60)
        self.assertEqual(r.cashflow_history_list("tester")["summary"]["received_amount"], 60)
        r.cashflow_history_delete(saved["item"]["id"], "tester")
        self.assertEqual(r.cashflow_history_records("tester")["total"], 0)

    def test_workbook_overlay_exports_template_formulas_and_user_cells(self):
        from openpyxl import Workbook, load_workbook
        self.data_dir.mkdir(parents=True, exist_ok=True)
        template = self.data_dir / "project-template.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "项目概算表模板"
        sheet["A1"] = "项目名称"
        sheet["B2"] = 100
        sheet["C2"] = "=B2*2"
        sheet["D2"] = __import__("datetime").date(2026, 9, 9)
        book.save(template)
        with patch.object(r, "CASHFLOW_WORKBOOK_TEMPLATE", template):
            workbook = r.cashflow_workbook_get("tester")
            self.assertEqual(workbook["sheets"][0]["name"], "项目概算表模板")
            __import__("json").dumps(workbook)
            r.cashflow_workbook_save({"project_name": "示例项目", "cells": {"项目概算表模板": {"B2": "250"}}}, "tester")
            exported_name, exported_raw = r.cashflow_workbook_export("tester")
        self.assertTrue(exported_name.startswith("示例项目-") and exported_name.endswith(".xlsx"))
        exported_book = load_workbook(__import__("io").BytesIO(exported_raw), data_only=False)
        self.assertEqual(exported_book["项目概算表模板"]["B2"].value, "250")
        self.assertEqual(exported_book["项目概算表模板"]["C2"].value, "=B2*2")

    def test_frontend_registers_cashflow_assistant(self):
        root = Path(__file__).resolve().parents[2]
        menu = (root / "ai-office-assistant-web/src/layout/menu.ts").read_text(encoding="utf-8")
        workspace = (root / "ai-office-assistant-web/src/views/WorkspaceView/index.vue").read_text(encoding="utf-8")
        self.assertIn("'cashflow'", menu)
        self.assertIn("资金测算助手", menu)
        self.assertIn("CashflowAssistantView", workspace)

    def test_workbook_frontend_has_sheet_specific_interactions(self):
        root = Path(__file__).resolve().parents[2]
        view = (root / "ai-office-assistant-web/src/views/CashflowAssistantView/WorkbookWorkbench.vue").read_text(encoding="utf-8")
        self.assertIn("sheetInteractions", view)
        for name in ("项目情况确认", "项目概算表模板--成本", "项目概算表模板-现金流", "假设清单风险清单", "人力成本明细表", "会议记录"):
            self.assertIn(name, view)
        self.assertIn("cashflow-timeline", view)


CashflowAssistantDbModeTests = db_variant(CashflowAssistantTests)
