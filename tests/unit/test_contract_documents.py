import base64
import io
import tempfile
import unittest
from pathlib import Path
from docx import Document
import backend.runtime as runtime
from tests.unit.officedb_testing import db_variant


class ContractDocumentsTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.original = runtime.USER_DATA_DIR
        runtime.USER_DATA_DIR = Path(self.temporary.name)
        document = Document()
        document.add_paragraph("测试合同")
        table = document.add_table(rows=2, cols=4)
        for cell, value in zip(table.rows[0].cells, ["名称", "型号", "单位", "数量"]):
            cell.text = value
        for cell, value in zip(table.rows[1].cells, ["交换机", "ABC", "台", "10"]):
            cell.text = value
        output = io.BytesIO()
        document.save(output)
        self.original_bytes = output.getvalue()
        self.record = runtime.contract_upload({"name": "contract.docx", "content_base64": base64.b64encode(output.getvalue()).decode()}, "alice")["contract"]

    def tearDown(self):
        runtime.USER_DATA_DIR = self.original
        self.temporary.cleanup()

    def test_extract_save_generate_download(self):
        self.assertEqual(self.record["lines"][0]["quantity"], "10")
        self.record.update(contract_no="C001", customer="客户")
        record = runtime.contract_save(self.record, "alice")["contract"]
        for kind in ["送货单", "到货签收单", "验收单"]:
            record = runtime.contract_generate({"id": record["id"], "revision": record["revision"], "kind": kind, "confirmed": True, "selections": [{"id": record["lines"][0]["id"], "quantity": "3", "sn": "SN001"}]}, "alice")["contract"]
            _, raw = runtime.contract_download("alice", record["id"], record["documents"][-1]["id"])
            document = Document(io.BytesIO(raw))
            self.assertEqual(document.tables[0].rows[1].cells[4].text, "3")
            self.assertIn("草稿", document.paragraphs[0].text)
        self.assertEqual(len(record["documents"]), 3)

    def test_isolation_and_identity(self):
        self.assertEqual(runtime.contract_list("bob")["contracts"], [])
        with self.assertRaises(ValueError):
            runtime.contract_read("bob", self.record["id"])
        with self.assertRaises(ValueError):
            runtime.contract_save(self.record, "bob")
        for username in ["", "../alice", "a/b"]:
            with self.assertRaises(PermissionError):
                runtime.contract_list(username)
        with self.assertRaises(ValueError):
            runtime.contract_download("bob", self.record["id"], "unknown")

    def test_original_and_archive(self):
        record = self.record
        name, raw = runtime.contract_original("alice", record["id"])
        self.assertEqual(name, "contract.docx")
        self.assertEqual(raw, self.original_bytes)
        for owner, identifier in [("bob", record["id"]), ("alice", "../other")]:
            with self.assertRaises(ValueError):
                runtime.contract_original(owner, identifier)
        self.assertFalse(runtime.contract_list("alice")["contracts"][0]["archived"])
        for state in ["true", None, 1]:
            with self.assertRaises(ValueError):
                runtime.contract_archive({"id": record["id"], "revision": 1, "archived": state}, "alice")
        archived = runtime.contract_archive({"id": record["id"], "revision": 1, "archived": True}, "alice")["contract"]
        self.assertEqual(archived["revision"], 2)
        self.assertTrue(runtime.contract_list("alice")["contracts"][0]["archived"])
        with self.assertRaises(ValueError):
            runtime.contract_save(archived, "alice")
        with self.assertRaises(ValueError):
            runtime.contract_generate({"id": record["id"]}, "alice")
        with self.assertRaises(ValueError):
            runtime.contract_archive({"id": record["id"], "revision": 1, "archived": False}, "alice")
        self.assertEqual(runtime.contract_original("alice", record["id"])[1], raw)
        restored = runtime.contract_archive({"id": record["id"], "revision": 2, "archived": False}, "alice")["contract"]
        restored.update(contract_no="C001", customer="客户")
        runtime.contract_save(restored, "alice")
        summary = runtime.contract_list("alice")["contracts"][0]
        self.assertEqual(summary["contract_no"], "C001")
        self.assertEqual(summary["document_count"], 0)

    def test_structured_candidates_and_preview(self):
        text = "合同编号：C2026\n甲 方：客户公司\n乙 方：服务公司\n含税总价为￥70000元，税率为6%。\n预付款 | 50 | 35000 | 7 | 先开票\n验收款 |  |  |  | \n"
        structured = runtime.contract_structure(text)
        self.assertEqual(structured["contract_no"], "C2026")
        self.assertEqual(structured["customer"], "客户公司")
        self.assertEqual(structured["supplier"], "服务公司")
        self.assertEqual(structured["total_amount"], "70000")
        self.assertEqual(len(structured["payments"]), 1)
        self.assertEqual(structured["payments"][0]["days"], "7")
        before = runtime.contract_read("alice", self.record["id"])
        preview = runtime.contract_extract_preview(self.record, "alice")["contract"]
        self.assertEqual(runtime.contract_read("alice", self.record["id"]), before)
        with self.assertRaises(ValueError):
            runtime.contract_extract_preview(self.record, "bob")
        record = {**self.record, **structured}
        saved = runtime.contract_save(record, "alice")["contract"]
        self.assertEqual(saved["lines"][0]["id"], self.record["lines"][0]["id"])
        self.assertEqual(saved["payments"], structured["payments"])
        self.assertEqual(saved["supplier"], "服务公司")

    def test_service_and_combined_quantity(self):
        document = Document()
        for header, name, quantity in [("服务项目", "安装调试", "1项"), ("设备名称", "防火墙", "2台")]:
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = header
            table.cell(0, 1).text = "数量"
            table.cell(1, 0).text = name
            table.cell(1, 1).text = quantity
        output = io.BytesIO()
        document.save(output)
        _, lines = runtime.contract_extract(output.getvalue(), ".docx")
        self.assertEqual([(line["quantity"], line["unit"]) for line in lines], [("1", "项"), ("2", "台")])

    def test_template_routing(self):
        for identifier, name, title, markers in runtime.CONTRACT_TEMPLATE_PROFILES:
            text = title + "\n" + "\n".join(markers)
            route = runtime.contract_detect_template(text)
            self.assertEqual(route["template_id"], identifier)
            self.assertEqual(route["mode"], "template")
            self.assertEqual(runtime.contract_detect_template(title)["mode"], "general")
        self.assertEqual(runtime.contract_detect_template("第三方技术服务合同")["mode"], "general")
        text = "服务销售合同 设备销售合同 " + " ".join(runtime.CONTRACT_TEMPLATE_PROFILES[0][3] + runtime.CONTRACT_TEMPLATE_PROFILES[1][3])
        self.assertEqual(runtime.contract_detect_template(text)["mode"], "general")
        document = Document()
        document.add_paragraph("服务销售合同 合同服务清单（不含设备软硬件供货） 服务有效期 服务交付方式")
        document.add_paragraph("合同总金额（含税）：¥70000元，增值税税率为6%")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "服务项目"
        table.cell(0, 1).text = "服务次数"
        table.cell(1, 0).text = "调试"
        table.cell(1, 1).text = "2次"
        output = io.BytesIO()
        document.save(output)
        record = runtime.contract_upload({"name": "模板合同.docx", "content_base64": base64.b64encode(output.getvalue()).decode()}, "alice")["contract"]
        self.assertEqual(record["extraction"]["mode"], "template")
        self.assertEqual(record["total_amount"], "70000")
        self.assertEqual(record["lines"][0]["quantity"], "2")
        self.assertEqual(runtime.contract_read("alice", record["id"])["extraction"], record["extraction"])
        self.assertEqual(self.record["extraction"]["mode"], "general")

    def test_equipment_table_parties_prices_and_blank_number(self):
        document = Document()
        document.add_paragraph('设备销售合同\n合同编号：\n为保护供需双方的合法权益')
        document.add_paragraph('设备清单 到货签收约定 保修责任')
        document.add_paragraph('第四条：付款方式\n乙方设备送达现场并经甲方签收确认后 2个月内，甲方支付已到货设备全额款项。\n第五条：开票方式')
        parties = document.add_table(rows=1, cols=4)
        for c, v in zip(parties.rows[0].cells, ['甲 方：', '客户公司', '乙 方：', '供应商公司']):
            c.text = v
        table = document.add_table(rows=3, cols=8)
        for row, values in zip(table.rows, [
            ['产品名称','品牌','规格型号/参数','单位','数量','单价(元)','合计(元)','备注'],
            ['跳纤','品牌A','LC-10','条','8','69','552','OLT1'],
            ['合计总金额（大写）：伍佰伍拾贰元'] * 8]):
            for c, v in zip(row.cells, values): c.text = v
        output = io.BytesIO(); document.save(output)
        record = runtime.contract_upload({'name':'设备合同.docx','content_base64':base64.b64encode(output.getvalue()).decode()}, 'alice')['contract']
        self.assertEqual(record['contract_no'], '')
        self.assertEqual(record['customer'], '客户公司')
        self.assertEqual(record['supplier'], '供应商公司')
        self.assertEqual(len(record['lines']), 1)
        line = record['lines'][0]
        self.assertEqual((line['spec'], line['brand'], line['unit_price'], line['amount'], line['remark']), ('LC-10','品牌A','69','552','OLT1'))
        self.assertIn('2个月内', record['payment_terms'])
        saved = runtime.contract_save(record, 'alice')['contract']
        self.assertEqual(saved['lines'][0]['unit_price'], '69')

    def test_stale_revision_and_excess_quantity(self):
        self.record.update(contract_no="C001", customer="客户")
        record = runtime.contract_save(self.record, "alice")["contract"]
        with self.assertRaises(ValueError):
            runtime.contract_save(self.record, "alice")
        for quantity in ["11", "-1", "NaN", "0"]:
            with self.assertRaises(ValueError):
                runtime.contract_generate({"id": record["id"], "revision": record["revision"], "kind": "送货单", "confirmed": True, "selections": [{"id": record["lines"][0]["id"], "quantity": quantity}]}, "alice")


ContractDocumentsDbModeTests = db_variant(ContractDocumentsTests)
