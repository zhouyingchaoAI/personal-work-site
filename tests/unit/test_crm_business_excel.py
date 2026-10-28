import base64
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
import backend.runtime as r

class BusinessExcelTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.patches=[patch.object(r,'USER_DATA_DIR',Path(self.tmp.name)),patch.object(r,'crm_on_change')]
        for p in self.patches:p.start()
    def tearDown(self):
        for p in reversed(self.patches):p.stop()
        self.tmp.cleanup()
    def payload(self,values=None):
        b=load_workbook(io.BytesIO(r.crm_business_template()))
        if values is not None:
            data={'公司全称':'示例集团有限公司','单位*':'示例集团','姓名*':'张工','单位类型':'业主单位',**values}
            headers=[cell.value for cell in b['经营收集表'][1]]
            b['经营收集表'].append([data.get(h,'') for h in headers])
        out=io.BytesIO();b.save(out);b.close()
        return {'filename':'客户经营.xlsx','data':base64.b64encode(out.getvalue()).decode()}

    def test_task_progress_summary_is_split_into_company_region_and_owner_datasets(self):
        book = Workbook(); sheet = book.active; sheet.title = '10月任务进度总表'
        sheet.append(['2026年10月冲刺任务令简报 | 数据截止：2026-10-18'])
        sheet.append([]); sheet.append(['一、公司整体-经营指标达成进度（保底+冲刺）'])
        sheet.append(['指标', '', '2026年度目标', '1-10月已完成', '年度目标完成进度', '年度目标缺口', '10月冲刺任务令', '10月任务令已完成', '10月任务令完成进度', '10月进度排名', '10月任务令缺口'])
        sheet.append(['回款', '', 38000, 13433.46, .3535, -24566.54, 3960.11, 138.42, .03495, 2, -3821.69])
        sheet.append([]); sheet.append(['二、区域 - 10月回款任务'])
        sheet.append(['区域', '责任人', '回款 任务令', '第一周', '第二周', '第三周', '第四周', '10月合计完成', '完成率', '排名', '缺口'])
        sheet.append(['广东', '黄宇', 414.53, 70, 4.92, 0, 0, 74.92, .1807, 1, -339.61])
        sheet.append([]); sheet.append(['五、项目经理 - 10月回款任务'])
        sheet.append(['区域', '责任人', '回款 任务令', '第一周', '第二周', '第三周', '第四周', '10月合计完成', '完成率', '排名', '缺口'])
        sheet.append(['广东', '孙良珑', 414.53, 70, 4.92, 0, 0, 74.92, .1807, 1, -339.61])

        sheet.append(['合计', '', 414.53, 70, 4.92, 0, 0, 74.92, .1807, '', -339.61])
        parsed = r._crm_task_progress_sheets(book)

        self.assertEqual([item['key'] for item in parsed], ['任务令总表说明', '任务令公司进度', '任务令区域进度', '任务令责任人进度', '任务令进度总表原表'])
        self.assertEqual(parsed[0]['rows'][0][0], '2026年10月冲刺任务令简报 | 数据截止：2026-10-18')
        self.assertEqual(parsed[1]['rows'][0][0:3], ['回款', 38000, 13433.46])
        self.assertEqual(parsed[2]['rows'][0][0:4], ['区域', '回款', '广东', '黄宇'])
        self.assertEqual(parsed[3]['rows'][0][0:4], ['项目经理', '回款', '广东', '孙良珑'])
        self.assertTrue(any(row[2] == '合计' for row in parsed[3]['rows']))
        self.assertTrue(any(row[1] == '二、区域 - 10月回款任务' for row in parsed[4]['rows']))

    def external_collection_payload(self):
        headers = ['区域', '公司全称', '单位*', '部门', '职务', '姓名*', '手机', '办公电话', '邮箱', '微信 / 其他联系', '重要程度', '支持程度', '跟进轨迹', '上次互动时间', '本周互动计划', '计划跟进日期', '单位类型', '省份', '城市', '维系周期（天）', '我方姓名']
        row = ['东区', '华为技术有限公司', '华为', '政企部', '经理', '张工', '', '', '', '', '重要', '支持', '已沟通方案', '/', '', '', '厂商', '浙江省', '杭州', '', '李经理']
        book = Workbook(); sheet = book.active; sheet.title = '经营收集表'; sheet.append(headers); sheet.append(row)
        output = io.BytesIO(); book.save(output); book.close()
        return {'filename':'外部客户经营收集表.xlsx','data':base64.b64encode(output.getvalue()).decode()}

    def external_completion_payload(self):
        headers = ['区域', '公司全称', '单位*', '部门', '职务', '姓名*', '手机', '办公电话', '邮箱', '微信 / 其他联系', '重要程度', '支持程度', '跟进轨迹', '上次互动时间', '本周互动计划', '计划跟进日期', '单位类型', '省份', '城市', '维系周期（天）', '我方姓名']
        direct = ['东区', '直入客户有限公司', '直入客户', '', '', '完整联系人', '', '', '', '', '', '', '', '', '', '', '业主单位', '', '杭州', '', '董涛']
        pending = ['东区', '', '待补公司简称', '', '', '', '', '', '', '', '', '', '', '', '', '', '', '', '杭州', '', '董涛']
        book = Workbook(); sheet = book.active; sheet.title = '经营收集表'; sheet.append(headers); sheet.append(direct); sheet.append(pending)
        output = io.BytesIO(); book.save(output); book.close()
        return {'filename':'待补全客户经营收集表.xlsx','data':base64.b64encode(output.getvalue()).decode()}

    def test_opportunity_import_keeps_cancelled_rows_as_archive(self):
        headers = ['机会点编码', '机会点名称', '业主单位', '机会点阶段', '赢率', '业务状态', '作废状态']
        book = Workbook(); sheet = book.active; sheet.append(headers)
        sheet.append(['OPP-001', '地铁信号改造', '城市地铁集团', '招标', '60%', '正常', ''])
        sheet.append(['OPP-002', '历史项目', '城市地铁集团', '立项', '10%', '丢单', '已作废'])
        output = io.BytesIO(); book.save(output); book.close()
        payload = {'filename': '机会点.xlsx', 'data': base64.b64encode(output.getvalue()).decode()}

        result = r.crm_opportunity_import(payload, 'alice')

        self.assertEqual((result['imported'], result['active'], result['archived']), (2, 1, 1))
        stored = r.crm_load('alice', 'opportunities')
        self.assertFalse(stored['OPP-001']['cancelled'])
        self.assertTrue(stored['OPP-002']['cancelled'])

    def test_sales_contract_import_marks_matched_opportunity_as_won(self):
        r.crm_store('alice', 'opportunities', {'OPP-001': {'id': 'OPP-001', 'code': 'OPP-001', 'name': '测试商机', 'outcome': 'open', 'cancelled': True, 'version': 1}})
        headers = ['销售项目编码', '销售项目名称', '销售子项目编码', '销售合同编码', '销售子项目名称', '客户', '合同签订日期', '合同金额', '业主单位', '项目经理', '客户经理', '归档']
        book = Workbook(); sheet = book.active; sheet.append(headers)
        sheet.append(['OPP-001', '测试商机', 'OPP-001-1', 'HT-001', '子项目一', '客户甲', '2026/10/01', 100000, '业主甲', '项目经理甲', '客户经理甲', '是'])
        sheet.append(['OPP-001', '测试商机', 'OPP-001-2', 'HT-002', '子项目二', '客户乙', '2026/10/03', 200000, '业主甲', '项目经理乙', '客户经理甲', '是'])
        sheet.append(['NO-MATCH', '未匹配项目', 'NO-MATCH-1', 'HT-003', '未匹配', '客户丙', '2026/10/05', 50000, '业主丙', '项目经理丙', '客户经理丙', '否'])
        output = io.BytesIO(); book.save(output); book.close()

        result = r.crm_sales_contract_import({'filename': '销售合同.xlsx', 'data': base64.b64encode(output.getvalue()).decode()}, 'alice')

        self.assertEqual((result['contracts'], result['matched_opportunities'], result['unmatched_projects']), (3, 1, 1))
        opportunity = r.crm_load('alice', 'opportunities')['OPP-001']
        self.assertEqual(opportunity['outcome'], 'won')
        self.assertEqual(opportunity['outcome_source'], 'sales_contract')
        self.assertFalse(opportunity['cancelled'])
        self.assertEqual(opportunity['contract_count'], 2)
        self.assertEqual(opportunity['contract_amount_yuan'], 300000)
        self.assertEqual(opportunity['contract_date'], '2026/10/01')
        self.assertEqual(len(r.crm_load('alice', 'sales_contracts')), 3)

    def test_opportunity_export_keeps_business_fields_and_current_filters(self):
        r.crm_store('alice', 'opportunities', {
            'A': {'id':'A', 'code':'A', 'name':'有效机会', 'owner_unit':'地铁集团', 'stage':'招标',
                  'region':'西区', 'manager':'张工', 'amount':1200, 'cancelled':False},
            'B': {'id':'B', 'code':'B', 'name':'历史机会', 'owner_unit':'地铁集团', 'stage':'立项',
                  'cancelled':True, 'cancel_reason':'业主取消'},
        })
        raw, filename = r.crm_opportunity_export({'username':'alice', 'role':'member', 'assistant_permissions':{'crm':True}}, 'active', '招标')
        book = load_workbook(io.BytesIO(raw), data_only=True)
        sheet = book['机会点清单']
        self.assertEqual(filename, '机会点清单.xlsx')
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual([cell.value for cell in sheet[1]][:4], ['机会点编码', '机会点名称', '业主单位', '项目归属'])
        self.assertEqual(sheet['A2'].value, 'A')
        amount_col = [c.value for c in sheet[1]].index('可参与份额-总额(万元)') + 1
        self.assertEqual(sheet.cell(2, amount_col).value, 1200)
        book.close()

    def test_opportunity_save_updates_business_fields_and_keeps_project_code_stable(self):
        r.crm_store('alice', 'opportunities', {'OPP-1': {
            'id':'OPP-1', 'code':'OPP-1', 'name':'原机会点', 'owner_unit':'原业主',
            'amount':0, 'cancelled':False, 'version':1,
        }})
        saved = r.crm_opportunity_save({
            'id':'OPP-1', 'code':'OPP-1', 'name':'更新后的机会点', 'owner_unit':'城市地铁集团',
            'stage':'招标', 'amount':'850.5', 'cancelled':True, 'cancel_reason':'业主暂缓', 'version':1,
        }, 'alice')
        self.assertTrue(saved['ok'])
        self.assertEqual(saved['opportunity']['amount'], 850.5)
        self.assertTrue(saved['opportunity']['cancelled'])
        self.assertEqual(saved['opportunity']['cancel_reason'], '业主暂缓')
        changed_code = r.crm_opportunity_save({**saved['opportunity'], 'code':'OTHER'}, 'alice')
        self.assertFalse(changed_code['ok'])
    def test_empty_template_does_not_import_example(self):
        result=r.crm_excel_import(self.payload(),'alice')
        self.assertEqual(result['created'],0)
        self.assertFalse(r.crm_load('alice','people'))

    def test_bulk_import_flushes_collections_and_pending_queue_once(self):
        book = load_workbook(io.BytesIO(base64.b64decode(self.external_completion_payload()['data'])))
        sheet = book['经营收集表']
        complete = [cell.value for cell in sheet[2]]
        pending = [cell.value for cell in sheet[3]]
        for index in range(20):
            row = list(complete)
            row[5] = f'联系人{index}'
            sheet.append(row)
            row = list(pending)
            row[2] = f'待补单位{index}'
            sheet.append(row)
        output = io.BytesIO()
        book.save(output)
        book.close()
        with patch.object(r, '_crm_store_disk', wraps=r._crm_store_disk) as store, \
                patch.object(r, '_crm_import_completion_write', wraps=r._crm_import_completion_write) as queue:
            result = r.crm_excel_import({'filename': '批量.xlsx', 'data': base64.b64encode(output.getvalue()).decode()}, 'alice')
        self.assertEqual(result['errors'], [])
        self.assertEqual((result['created'], result['pending']), (21, 21))
        keys = [(call.args[0], call.args[1]) for call in store.call_args_list]
        self.assertEqual(len(keys), len(set(keys)))
        self.assertEqual(queue.call_count, 1)
        self.assertEqual(len(r.crm_load('alice', 'people')), 21)
        self.assertEqual(len(r._crm_import_completion_load()), 21)
        self.assertIsNone(r.CRM_IMPORT_CONTEXT.batch)
        self.assertFalse(r.CRM_IMPORT_CONTEXT.active)

    def test_retry_flushes_pending_queue_and_clears_context_after_write_failure(self):
        rows = [{'row': index + 2, 'values': {'公司全称': '', '单位*': f'待补{index}'}}
                for index in range(10)]
        with patch.object(r, '_crm_import_completion_write', wraps=r._crm_import_completion_write) as queue:
            result = r.crm_excel_retry_rows({'rows': rows}, 'alice')
        self.assertEqual(result['pending'], 10)
        self.assertEqual(queue.call_count, 1)
        self.assertEqual(len(r._crm_import_completion_load()), 10)
        rows = [{'row': 99, 'values': {'公司全称': '', '单位*': '新增待补'}}]
        with patch.object(r, '_crm_import_completion_write', side_effect=OSError('disk full')):
            with self.assertRaises(OSError):
                r.crm_excel_retry_rows({'rows': rows}, 'alice')
        self.assertIsNone(r.CRM_IMPORT_CONTEXT.batch)
        self.assertFalse(r.CRM_IMPORT_CONTEXT.active)

    def test_external_collection_layout_uses_company_unit_path_and_accepts_blank_slash_date(self):
        try:
            result = r.crm_excel_import(self.external_collection_payload(), 'alice')
        except ValueError as exc:
            self.fail(f'外部经营收集表应可直接导入：{exc}')
        self.assertEqual(result['errors'], [], result)
        units = {unit['name']: unit for unit in r.crm_load('alice', 'units').values()}
        company = units['华为技术有限公司']
        unit = units['华为']
        self.assertEqual(company['parent_unit_id'], '')
        self.assertEqual(unit['parent_unit_id'], company['id'])
        self.assertEqual(company['region'], '')
        self.assertEqual(unit['region'], '杭州')
        self.assertEqual(unit['organization_level'], 'unit')
        self.assertEqual(unit['company_full_name'], '华为技术有限公司')
        self.assertEqual(unit['subsidiary_full_name'], '')
        self.assertEqual(unit['type'], '厂商')
        person = r.crm_read_people({'username': 'alice', 'role': 'member'})['people'][0]
        self.assertEqual(person['company_full_name'], '华为技术有限公司')
        self.assertEqual(person['subsidiary_full_name'], '')
        self.assertEqual(person['organization_unit_name'], '华为')
        self.assertEqual(person['organization_path'], '华为技术有限公司 / 华为')
        relation = next(iter(r.crm_load('alice', 'relations').values()))
        self.assertEqual(relation['action_count'], 0)

    def test_external_rows_with_conflicting_unit_type_or_region_keep_all_assigned_regions(self):
        headers = r.CRM_BUSINESS_EXTERNAL_HEADERS
        first = ['东区', '示例集团有限公司', '运营公司', '', '', '张工', '', '', '', '', '', '', '', '', '', '', '业主单位', '', '杭州', '', '']
        later = ['中区', '示例集团有限公司', '运营公司', '', '', '李工', '', '', '', '', '', '', '', '', '', '', '主管单位', '', '杭州', '', '']
        book = Workbook(); sheet = book.active; sheet.title = '经营收集表'; sheet.append(headers); sheet.append(first); sheet.append(later)
        output = io.BytesIO(); book.save(output); book.close()
        result = r.crm_excel_import({'filename': '跨区组织.xlsx', 'data': base64.b64encode(output.getvalue()).decode()}, 'alice')
        self.assertEqual(result['errors'], [], result)
        self.assertEqual(result['created'], 2, result)
        unit = next(unit for unit in r.crm_load('alice', 'units').values() if unit['name'] == '运营公司')
        self.assertEqual(unit['type'], '业主单位')
        self.assertEqual(unit['business_region'], '东区')
        self.assertEqual(unit['business_regions'], ['东区', '中区'])

    def test_external_invalid_optional_interaction_date_does_not_block_contact_import(self):
        headers = r.CRM_BUSINESS_EXTERNAL_HEADERS
        row = ['东区', '示例集团有限公司', '运营公司', '', '', '张工', '', '', '', '', '', '', '', '2026/3/11\n2026/6/4', '', '', '业主单位', '', '杭州', '', '']
        book = Workbook(); sheet = book.active; sheet.title = '经营收集表'; sheet.append(headers); sheet.append(row)
        output = io.BytesIO(); book.save(output); book.close()
        result = r.crm_excel_import({'filename': '日期待完善.xlsx', 'data': base64.b64encode(output.getvalue()).decode()}, 'alice')
        self.assertEqual(result['errors'], [], result)
        self.assertEqual(result['created'], 1, result)
        self.assertFalse(r.crm_load('alice', 'actions'))

    def test_incomplete_external_row_becomes_assignee_completion_task_before_entering_crm(self):
        with patch.object(r, 'read_config', return_value={'users': [{'username': 'dongtao', 'name': '董涛', 'role': 'member'}]}):
            result = r.crm_excel_import(self.external_completion_payload(), 'uploader')
        self.assertEqual(result['created'], 1, result)
        self.assertEqual(result['pending'], 1, result)
        self.assertFalse(any(person['name'] == '待补联系人' for person in r.crm_load('uploader', 'people').values()))
        self.assertTrue(hasattr(r, 'crm_import_completion_tasks'))
        tasks = r.crm_import_completion_tasks({'username': 'dongtao', 'role': 'member'})['tasks']
        self.assertEqual(len(tasks), 1)
        self.assertEqual(tasks[0]['assignee_username'], 'dongtao')
        self.assertEqual(tasks[0]['missing_fields'], ['公司全称', '姓名*', '单位类型'])

    def test_legacy_name_only_completion_task_is_visible_and_completable_by_assignee(self):
        users = {'users': [
            {'username': 'chenzhengyun', 'name': '陈正云', 'employee_id': '240512', 'role': 'member'},
            {'username': 'other', 'name': '其他人员', 'employee_id': '999999', 'role': 'member'},
        ]}
        with patch.object(r, 'read_config', return_value=users), \
                patch.object(r, 'crm_capabilities', return_value={'read': True, 'data_scope': 'self'}):
            pending_values = {header: '' for header in r.CRM_BUSINESS_HEADERS}
            pending_values.update({'区域': '广东区', '单位*': '待补单位'})
            task = r._crm_import_completion_create(
                {'我方姓名': '陈正云'},
                pending_values,
                r.CRM_BUSINESS_EXTERNAL_HEADERS,
                ['公司全称', '姓名*', '单位类型'],
                'uploader',
                2,
            )
            stored = r._crm_import_completion_load()
            stored[task['id']]['assignee_username'] = ''
            r._crm_import_completion_store(stored)

            visible = r.crm_import_completion_tasks(users['users'][0])['tasks']
            self.assertEqual([item['id'] for item in visible], [task['id']])
            self.assertEqual(r.crm_import_completion_tasks(users['users'][1])['tasks'], [])
            completed = r.crm_import_completion_complete({
                'id': task['id'],
                'values': {'公司全称': '客户集团有限公司', '姓名*': '客户联系人', '单位类型': '业主单位'},
            }, users['users'][0])
            self.assertTrue(completed['ok'])
            self.assertEqual(completed['task']['assignee_username'], 'chenzhengyun')
            self.assertEqual(len(r.crm_load('chenzhengyun', 'people')), 1)

    def test_admin_can_complete_an_assigned_import_task_into_the_assignee_crm(self):
        users = {'users': [
            {'username': 'dongtao', 'name': '董涛', 'role': 'member'},
            {'username': 'admin', 'name': '管理员', 'role': 'superadmin'},
        ]}
        with patch.object(r, 'read_config', return_value=users):
            r.crm_excel_import(self.external_completion_payload(), 'uploader')
            task = r.crm_import_completion_tasks({'username': 'admin', 'role': 'superadmin'})['tasks'][0]
            completed = r.crm_import_completion_complete({'id': task['id'], 'values': {'公司全称': '待补公司科技有限公司', '姓名*': '待补联系人', '单位类型': '厂商'}}, {'username': 'admin', 'role': 'superadmin'})
        self.assertTrue(completed['ok'])
        self.assertEqual(completed['task']['completed_by'], 'admin')
        self.assertEqual(len(r.crm_load('dongtao', 'people')), 1)
        self.assertFalse(r.crm_import_completion_tasks({'username': 'dongtao', 'role': 'member'})['pending'])

    def test_completion_export_then_upload_completes_and_removes_pending_task(self):
        users = {'users': [
            {'username': 'dongtao', 'name': '董涛', 'role': 'member'},
            {'username': 'admin', 'name': '管理员', 'role': 'superadmin'},
        ]}
        with patch.object(r, 'read_config', return_value=users):
            r.crm_excel_import(self.external_completion_payload(), 'uploader')
            raw, _ = r.crm_import_completion_excel_export({'username': 'admin', 'role': 'superadmin'}, '东区', 'dongtao')
            book = load_workbook(io.BytesIO(raw))
            sheet = book['资料补充']
            headers = [cell.value for cell in sheet[1]]
            self.assertEqual(sheet.max_row, 2)
            sheet.cell(2, headers.index('公司全称') + 1, '待补公司科技有限公司')
            sheet.cell(2, headers.index('姓名*') + 1, '待补联系人')
            sheet.cell(2, headers.index('单位类型') + 1, '厂商')
            output = io.BytesIO(); book.save(output); book.close()
            result = r.crm_import_completion_excel_import({
                'filename': '资料补充-东区.xlsx', 'data': base64.b64encode(output.getvalue()).decode(),
            }, {'username': 'admin', 'role': 'superadmin'})
        self.assertEqual(result['completed'], 1, result)
        self.assertEqual(result['errors'], [], result)
        self.assertFalse(r.crm_import_completion_tasks({'username': 'admin', 'role': 'superadmin'})['pending'])

    def test_completion_export_marks_missing_business_fields_by_color(self):
        users = {'users': [{'username': 'admin', 'name': '管理员', 'role': 'superadmin'}]}
        with patch.object(r, 'read_config', return_value=users):
            values = {header: '' for header in r.CRM_BUSINESS_HEADERS}
            values.update({'区域': '东区', '单位*': '待补单位', '部门': '物资部', '省份': '广东省'})
            r._crm_import_completion_create(
                {'区域': '东区', '单位*': '待补单位', '部门': '物资部', '省份': '广东省'}, values,
                r.CRM_BUSINESS_EXTERNAL_HEADERS, ['公司全称', '姓名*', '单位类型'], 'admin', 2,
            )
            raw, _ = r.crm_import_completion_excel_export({'username': 'admin', 'role': 'superadmin'}, '东区')
        book = load_workbook(io.BytesIO(raw))
        sheet = book['资料补充']
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(headers[-4:], ['手机', '办公电话', '邮箱', '微信 / 其他联系'])
        missing_fill = sheet.cell(2, headers.index('公司全称') + 1).fill.fgColor.rgb
        complete_fill = sheet.cell(2, headers.index('部门') + 1).fill.fgColor.rgb
        self.assertEqual(missing_fill, '00FFF2CC')
        self.assertEqual(complete_fill, '00E2F0D9')
        book.close()
    def test_template_collects_our_name_once_without_organization_columns(self):
        book=load_workbook(io.BytesIO(r.crm_business_template()), read_only=True)
        headers=[cell.value for cell in next(book['经营收集表'].iter_rows(max_row=1))]
        self.assertEqual(headers.count('我方姓名'),1)
        self.assertNotIn('我方对接人',headers)
        self.assertFalse({'我方公司（选填）','我方部门（选填）','我方岗位（选填）'} & set(headers))
        self.assertNotIn('我方组织目录',book.sheetnames)
        book.close()

    def test_downloaded_template_uses_the_three_field_import_gate_and_explains_completion_tasks(self):
        book = load_workbook(io.BytesIO(r.crm_business_template()), read_only=True)
        headers = [cell.value for cell in next(book['经营收集表'].iter_rows(max_row=1))]
        instructions = [row[0] for row in book['填写说明'].iter_rows(values_only=True) if row[0]]
        self.assertIn('公司全称', headers)
        self.assertNotIn('子公司全称', headers)
        self.assertEqual(headers, r.CRM_BUSINESS_EXTERNAL_HEADERS)
        self.assertTrue(any('资料补充' in line and '公司全称、姓名、单位类型' in line for line in instructions))
        book.close()

    def test_legacy_external_sheet_with_subsidiary_column_remains_importable(self):
        headers = r.CRM_BUSINESS_EXTERNAL_LEGACY_SUBSIDIARY_HEADERS
        row = ['东区', '历史集团有限公司', '历史集团杭州分公司', '历史集团', '', '', '历史联系人', '', '', '', '', '', '', '', '', '', '', '业主单位', '', '杭州', '', '董涛']
        book = Workbook(); sheet = book.active; sheet.title = '经营收集表'; sheet.append(headers); sheet.append(row)
        output = io.BytesIO(); book.save(output); book.close()
        result = r.crm_excel_import({'filename': '旧版.xlsx', 'data': base64.b64encode(output.getvalue()).decode()}, 'alice')
        self.assertEqual(result['errors'], [], result)
        units = {unit['name']: unit for unit in r.crm_load('alice', 'units').values()}
        self.assertEqual(units['历史集团杭州分公司']['organization_level'], 'subsidiary')
        self.assertEqual(units['历史集团']['parent_unit_id'], units['历史集团杭州分公司']['id'])

    def test_same_unit_short_name_under_different_companies_imports_without_conflict(self):
        headers = r.CRM_BUSINESS_EXTERNAL_HEADERS
        rows = [
            ['东区', '绍兴市轨道交通集团有限公司', '轨道集团', '建设部', '经理', '张工', '', '', '', '', '重要', '支持', '', '', '', '', '业主单位', '', '绍兴', '', '董涛'],
            ['东区', '宁波市轨道交通集团有限公司', '轨道集团', '运营部', '经理', '李工', '', '', '', '', '重要', '支持', '', '', '', '', '厂商', '', '宁波', '', '董涛'],
        ]
        book = Workbook(); sheet = book.active; sheet.title = '经营收集表'; sheet.append(headers)
        for row in rows: sheet.append(row)
        output = io.BytesIO(); book.save(output); book.close()
        result = r.crm_excel_import({'filename': '同名单位.xlsx', 'data': base64.b64encode(output.getvalue()).decode()}, 'alice')
        self.assertEqual(result['errors'], [], result)
        self.assertEqual(result['created'], 2, result)
        units = list(r.crm_load('alice', 'units').values())
        self.assertEqual(len([unit for unit in units if unit['name'] == '轨道集团' and unit.get('organization_level') == 'unit']), 2)

    def test_collection_template_offers_only_the_five_supported_unit_types(self):
        book = load_workbook(io.BytesIO(r.crm_business_template()))
        sheet = book['经营收集表']
        unit_type_column = next(cell.column_letter for cell in sheet[1] if cell.value == '单位类型')
        validation = next(item for item in sheet.data_validations.dataValidation if f'{unit_type_column}2' in str(item.sqref))
        self.assertEqual(validation.formula1, '=CRM_单位类型')
        region = next(item for item in sheet.data_validations.dataValidation if 'A2' in str(item.sqref))
        self.assertEqual(region.formula1, '=CRM_经营区域')
        self.assertEqual(r.CRM_UNIT_TYPES, ['主管单位', '业主单位', '总包单位', '设计单位', '厂商'])
        book.close()
    def test_full_row_and_repeat_preserves_scoring(self):
        data={'区域':'东区','部门':'设备部','重要程度':'重要','支持程度':'0','跟进轨迹':'历史原文，不猜次数','上次互动时间':'2026/1/26','本周互动计划':'确认接口','计划跟进日期':'2026-09-30'}
        payload=self.payload(data);res=r.crm_excel_import(payload,'alice')
        self.assertEqual(res['errors'],[],res)
        person=next(iter(r.crm_load('alice','people').values()))
        self.assertEqual(person['department'],'设备部');self.assertEqual(person['support_level'],'中立');self.assertEqual(person['maintenance_days'],14)
        relation=next(iter(r.crm_load('alice','relations').values()))
        self.assertEqual(relation['action_count'],1);self.assertEqual(relation['last_touch'],'2026-01-26')
        self.assertEqual(len(r.crm_followups('alice')),1)
        self.assertEqual(r.crm_excel_import(payload,'alice')['skipped'],1)
        self.assertEqual(len(r.crm_load('alice','actions')),2)
    def test_plan_only_never_counts_as_touch_even_after_completion(self):
        res=r.crm_excel_import(self.payload({'本周互动计划':'下周交流'}),'alice')
        self.assertEqual(res['errors'],[])
        relation=next(iter(r.crm_load('alice','relations').values()))
        self.assertEqual(relation['action_count'],0);self.assertEqual(relation['strength'],0);self.assertEqual(relation['last_touch'],'')
        plan=r.crm_followups('alice')[0]
        r.crm_followup_complete({'id':plan['id']},'alice')
        r.crm_recompute_relation('alice',relation['id'])
        self.assertEqual(r.crm_load('alice','relations')[relation['id']]['last_touch'],'')
        self.assertFalse(r.crm_followups('alice'))
    def test_invalid_required_business_values_leave_no_data_while_optional_bad_dates_are_ignored(self):
        for index, fields in enumerate([{'支持程度':'3'},{'上次互动时间':'2099-01-01'},{'计划跟进日期':'2026-09-30'},{'维系周期（天）':'0'}]):
            username = f'invalid-{index}'
            result=r.crm_excel_import(self.payload(fields), username)
            self.assertEqual(len(result['errors']),1,fields)
            self.assertFalse(r.crm_load(username,'units'))
        result = r.crm_excel_import(self.payload({'上次互动时间':'26/1/26'}), 'optional-date')
        self.assertEqual(result['errors'], [], result)
        self.assertEqual(result['created'], 1)
        self.assertFalse(r.crm_load('optional-date', 'actions'))

    def test_import_returns_editable_failed_row_and_maps_legacy_south_region(self):
        invalid = r.crm_excel_import(self.payload({'支持程度':'不确定选项'}), 'alice')
        self.assertEqual(invalid['errors'][0]['values']['支持程度'], '不确定选项')
        self.assertEqual(invalid['errors'][0]['headers'][0], '区域')
        imported = r.crm_excel_import(self.payload({'区域':'南区', '姓名*':'南区联系人'}), 'legacy_user')
        self.assertEqual(imported['errors'], [], imported)
        unit = next(iter(r.crm_load('legacy_user', 'units').values()))
        self.assertEqual(unit['business_region'], '广东区')

    def test_failed_business_row_can_be_corrected_and_retried_in_place(self):
        failed = r.crm_excel_import(self.payload({'支持程度':'错误选项'}), 'alice')['errors'][0]
        failed['values']['支持程度'] = '支持'
        result = r.crm_excel_retry_rows({'rows':[failed]}, 'alice')
        self.assertEqual(result['errors'], [], result)
        self.assertEqual(result['created'], 1)
        self.assertEqual(next(iter(r.crm_load('alice', 'people').values()))['support_level'], '支持')
    def test_only_fill_blank_profile(self):
        r.crm_excel_import(self.payload({'重要程度':'重要','支持程度':'1'}),'alice')
        r.crm_excel_import(self.payload({'重要程度':'一般','支持程度':'2','部门':'运营部'}),'alice')
        person=next(iter(r.crm_load('alice','people').values()))
        self.assertEqual(person['importance'],'重要');self.assertEqual(person['support_level'],'支持');self.assertEqual(person['department'],'运营部');self.assertEqual(person['maintenance_days'],14)
    def test_page_and_excel_share_business_rules(self):
        row={'区域':'东区','部门':'设备部','重要程度':'重要','支持程度':'0','跟进轨迹':'业务历史','上次互动时间':'2026-01-26','本周互动计划':'下周确认资料','计划跟进日期':'2026-09-30','我方对接人':'李经理'}
        imported=r.crm_excel_import(self.payload(row),'excel_user')
        self.assertEqual(imported['errors'],[])
        payload={'format':'business','unit_name':'示例集团','name':'张工','business_region':'东区','department':'设备部','importance':'重要','support_level':'0','history':'业务历史','last_interaction':'2026-01-26','weekly_plan':'下周确认资料','due_date':'2026-09-30','our_member':'李经理'}
        created=r.crm_quick_create(payload,'page_user');self.assertTrue(created['ok'],created)
        excel=next(iter(r.crm_load('excel_user','people').values()));page=next(iter(r.crm_load('page_user','people').values()))
        for key in ('department','importance','support_level','maintenance_days','imported_history','name'):
            self.assertEqual(excel[key],page[key],key)
        self.assertEqual(r.crm_quick_create(payload,'page_user')['status'],'skipped')
        relation=next(iter(r.crm_load('page_user','relations').values()))
        self.assertEqual(relation['action_count'],1);self.assertEqual(len(r.crm_followups('page_user')),1)
        failed=r.crm_quick_create({**payload,'name':'错误客户','last_interaction':'2099-01-01'},'page_user')
        self.assertFalse(failed['ok']);self.assertEqual(len(r.crm_load('page_user','people')),1)
    def test_missing_our_member_stays_blank_instead_of_account(self):
        result = r.crm_excel_import(self.payload({'区域':'东区'}), 'alice')  # 我方对接人/我方姓名 留空
        self.assertEqual(result['errors'], [], result)
        relation = next(iter(r.crm_load('alice', 'relations').values()))
        self.assertEqual(relation['our_member'], '')  # 不再自动填当前账号
        with patch.object(r,'read_config',return_value={'users':[{'username':'alice','role':'superadmin'}]}):
            self.assertEqual(r.crm_read_relations({'username':'alice','role':'superadmin'})['relations'][0]['our_member_label'], '—')
    def _person(self, user='alice'):
        r.crm_unit_save({'name':'集团','type':'业主单位'}, user)
        unit = next(iter(r.crm_load(user,'units').values()))
        r.crm_person_save({'name':'张三','unit_id':unit['id']}, user)
        return next(iter(r.crm_load(user,'people').values()))
    def test_migration_merges_account_autofill_into_blank(self):
        person = self._person()
        r.crm_relation_save({'person_id':person['id'],'our_member':'alice','source':'x'}, 'alice')  # 旧自动填账号
        r.crm_relation_save({'person_id':person['id'],'our_member':'','source':'x'}, 'alice')        # 复传产生的待指派重复
        rels = list(r.crm_load('alice','relations').values())
        acc = next(x for x in rels if x['our_member']=='alice'); blank = next(x for x in rels if x['our_member']=='')
        r.crm_action_add({'relation_id':acc['id'],'type':'拜访','happened_at':'2026-01-01','effective':True}, 'alice')
        r.crm_action_add({'relation_id':blank['id'],'type':'拜访','happened_at':'2026-01-01','effective':True}, 'alice')  # 内容重复
        res = r.crm_migrate_account_autofill_our_member('alice')
        self.assertEqual(res['removed'], 1)
        remaining = list(r.crm_load('alice','relations').values())
        self.assertEqual(len(remaining), 1)
        self.assertEqual(remaining[0]['our_member'], '')             # 只剩待指派
        self.assertEqual(len(r.crm_load('alice','actions')), 1)      # 重复动作已合并
        self.assertEqual(r.crm_migrate_account_autofill_our_member('alice')['removed'], 0)  # 幂等
    def test_migration_converts_lone_account_relation_to_blank(self):
        person = self._person()
        r.crm_relation_save({'person_id':person['id'],'our_member':'alice','source':'x'}, 'alice')
        res = r.crm_migrate_account_autofill_our_member('alice')
        self.assertEqual(res['cleared'], 1)
        self.assertEqual(next(iter(r.crm_load('alice','relations').values()))['our_member'], '')

    def test_removing_assignee_keeps_the_relation_and_its_actions(self):
        person = self._person()
        saved = r.crm_relation_save({'person_id': person['id'], 'our_member': '张工', 'source': '技术交流'}, 'alice')['relation']
        r.crm_action_add({'relation_id': saved['id'], 'type': '拜访', 'happened_at': '2026-09-10', 'effective': True}, 'alice')
        result = r.crm_relation_save({
            'id': saved['id'], 'version': saved['version'], 'person_id': person['id'],
            'our_member': '', 'source': '技术交流', 'intro_path': '', 'single_point_risk': False,
        }, 'alice')
        self.assertTrue(result['ok'], result)
        self.assertEqual(result['relation']['our_member'], '')
        self.assertEqual(len(r.crm_load('alice', 'actions')), 1)
    def test_page_optional_manager_validation_before_writes(self):
        result=r.crm_quick_create({'format':'business','unit_name':'新单位','name':'客户','manager_id':'not-owned'},'alice')
        self.assertFalse(result['ok']);self.assertFalse(r.crm_load('alice','units'))

    def test_optional_contact_channels_are_stored_by_page_and_excel(self):
        contacts = {'手机':'13800000000','办公电话':'020-12345678','邮箱':'zhang@example.com','微信 / 其他联系':'zhang-work'}
        result = r.crm_excel_import(self.payload(contacts), 'excel_user')
        self.assertEqual(result['errors'], [], result)
        excel_person = next(iter(r.crm_load('excel_user', 'people').values()))
        self.assertEqual({key: excel_person[key] for key in ('mobile','office_phone','email','wechat')}, {
            'mobile': '13800000000', 'office_phone': '020-12345678', 'email': 'zhang@example.com', 'wechat': 'zhang-work',
        })
        page = r.crm_quick_create({'format':'business','unit_name':'页面单位','name':'李工','mobile':'13900000000'}, 'page_user')
        self.assertTrue(page['ok'], page)
        page_person = next(iter(r.crm_load('page_user', 'people').values()))
        self.assertEqual(page_person['mobile'], '13900000000')
        self.assertEqual(page_person.get('email', ''), '')

    def test_old_business_template_without_contact_columns_remains_importable(self):
        book = Workbook()
        sheet = book.active
        sheet.title = '经营收集表'
        sheet.append(r.CRM_BUSINESS_CONTACT_LEGACY_HEADERS)
        sheet.append(['东区', '旧模板单位', '设备部', '经理', '王工', '重要', '支持', '', '', '', '', '业主单位', '上海', '', '李经理'])
        output = io.BytesIO(); book.save(output); book.close()
        result = r.crm_excel_import({'filename':'旧经营收集表.xlsx','data':base64.b64encode(output.getvalue()).decode()}, 'legacy_user')
        self.assertEqual(result['errors'], [], result)
        person = next(iter(r.crm_load('legacy_user', 'people').values()))
        self.assertEqual(person.get('mobile', ''), '')

    def test_chinese_support_and_legacy_read(self):
        for i,label in enumerate(r.CRM_SUPPORT_LABELS):
            result=r.crm_excel_import(self.payload({'姓名*':'客户'+str(i),'支持程度':label}),'alice')
            self.assertEqual(result['errors'],[])
        self.assertEqual({p['support_level'] for p in r.crm_load('alice','people').values()},set(r.CRM_SUPPORT_LABELS))
        people=r.crm_load('alice','people')
        for person,code in zip(people.values(),['-1',0,'1','2']): person['support_level']=code
        r.crm_store('alice','people',people)
        self.assertEqual({p['support_level'] for p in r.crm_people('alice')['people']},set(r.CRM_SUPPORT_LABELS))
        book=load_workbook(io.BytesIO(r.crm_business_template()))
        support_column = next(cell.column_letter for cell in book['经营收集表'][1] if cell.value == '支持程度')
        options=[d.formula1 for d in book['经营收集表'].data_validations.dataValidation if f'{support_column}2' in str(d.sqref)]
        self.assertEqual(options,['=CRM_支持程度']);book.close()
