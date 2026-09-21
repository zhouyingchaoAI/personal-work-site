import base64
import io
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from openpyxl import load_workbook
import backend.runtime as r

class ExcelTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.p1=patch.object(r,'USER_DATA_DIR',Path(self.tmp.name));self.p1.start()
        self.p2=patch.object(r,'crm_on_change');self.p2.start()
    def tearDown(self):
        self.p2.stop();self.p1.stop();self.tmp.cleanup()
    def payload(self,invalid=False):
        b=load_workbook(io.BytesIO(r.crm_excel_template(advanced=True)))
        b['单位'].append(['轨道集团','业主单位','上海'])
        b['联系人'].append(['轨道集团','张工','经理','定技术','','60','中立'])
        b['关系'].append(['轨道集团','张工','alice','拜访','','否'])
        b['动作'].append(['轨道集团','张工','alice','技术交流','2026-01-01','设计','确认接口','','是','否','P','提交清单','alice','2026-09-30'])
        b['项目决策链'].append(['P','设计','轨道集团','张工','alice'])
        if invalid: b['联系人'].append(['不存在的单位','李工'])
        raw=io.BytesIO();b.save(raw);b.close()
        return {'filename':'收集.xlsx','data':base64.b64encode(raw.getvalue()).decode()}
    def test_full_import_and_repeat(self):
        p=self.payload();first=r.crm_excel_import(p,'alice')
        self.assertEqual(first['created'],5,first)
        self.assertEqual(first['errors'],[])
        second=r.crm_excel_import(p,'alice')
        self.assertEqual(second['created'],0)
        self.assertEqual(second['skipped'],5)
        self.assertFalse(r.crm_load('bob','units'))
        self.assertEqual(r.crm_load('alice','actions').popitem()[1]['bidirectional'],False)
    def test_partial_error_location(self):
        result=r.crm_excel_import(self.payload(True),'alice')
        self.assertEqual(result['created'],5)
        self.assertEqual(result['errors'][0]['sheet'],'联系人')
        self.assertEqual(result['errors'][0]['row'],3)
    def test_blank_template(self):
        p={'filename':'template.xlsx','data':base64.b64encode(r.crm_excel_template(advanced=True)).decode()}
        result=r.crm_excel_import(p,'alice')
        self.assertEqual(result['created'],0)
        self.assertFalse(result['errors'])

    def test_missing_worksheet_dimension_metadata_is_treated_as_empty(self):
        # Some WPS-generated workbooks expose a missing worksheet dimension as
        # None.  The upload preflight must not compare that value with an int.
        self.assertEqual(r._crm_sheet_dimensions(SimpleNamespace(max_row=None, max_column=None)), (0, 0))
        self.assertEqual(r._crm_sheet_dimensions(SimpleNamespace(max_row='8', max_column='16')), (8, 16))

    def test_formula_and_invalid_schema(self):
        book=load_workbook(io.BytesIO(r.crm_excel_template(advanced=True)));book['单位'].append(['=1+1','业主单位'])
        buf=io.BytesIO();book.save(buf)
        result=r.crm_excel_import({'filename':'a.xlsx','data':base64.b64encode(buf.getvalue()).decode()},'alice')
        self.assertEqual(result['errors'][0]['row'],2)
        self.assertEqual(result['created'],0)
        book['单位']['A1']='错误表头';buf=io.BytesIO();book.save(buf);book.close()
        with self.assertRaises(ValueError): r.crm_excel_import({'filename':'a.xlsx','data':base64.b64encode(buf.getvalue()).decode()},'alice')


class SimpleExcelTests(ExcelTests):
    def test_simple_minimum_and_repeat(self):
        book=load_workbook(io.BytesIO(r.crm_excel_template()))
        self.assertEqual(book.sheetnames,['客户收集表'])
        book['客户收集表'].append(['业主A','业主单位','李工'])
        book['客户收集表'].append(['业主A','业主单位','王工'])
        raw=io.BytesIO();book.save(raw);book.close()
        payload={'filename':'simple.xlsx','data':base64.b64encode(raw.getvalue()).decode()}
        result=r.crm_excel_import(payload,'alice')
        self.assertEqual(result['created'],2,result)
        self.assertEqual(len(r.crm_load('alice','units')),1)
        self.assertEqual(len(r.crm_load('alice','people')),2)
        self.assertEqual(len(r.crm_load('alice','relations')),2)
        self.assertEqual(r.crm_excel_import(payload,'alice')['skipped'],2)

    def test_invalid_role_no_partial_writes(self):
        book=load_workbook(io.BytesIO(r.crm_excel_template()))
        book['客户收集表'].append(['业主A','业主单位','李工','','','错误角色'])
        raw=io.BytesIO();book.save(raw);book.close()
        result=r.crm_excel_import({'filename':'a.xlsx','data':base64.b64encode(raw.getvalue()).decode()},'alice')
        self.assertEqual(len(result['errors']),1)
        self.assertFalse(r.crm_load('alice','units'))

class QuickCaptureTests(ExcelTests):
    def test_quick_capture_validation_and_repeat(self):
        self.assertFalse(r.crm_quick_create({'name':'李工'},'alice')['ok'])
        self.assertFalse(r.crm_load('alice','units'))
        payload={'unit_name':'业主A','unit_type':'业主单位','name':'李工'}
        self.assertEqual(r.crm_quick_create(payload,'alice')['status'],'created')
        self.assertEqual(r.crm_quick_create(payload,'alice')['status'],'skipped')
        self.assertEqual(len(r.crm_load('alice','relations')),1)
        self.assertFalse(r.crm_load('bob','people'))
        self.assertFalse(r.crm_quick_create({**payload,'unit_type':'设计单位'},'alice')['ok'])


class ImportLimitTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.data_patch = patch.object(r, 'USER_DATA_DIR', Path(self.tmp.name))
        self.change_patch = patch.object(r, 'crm_on_change')
        self.data_patch.start(); self.change_patch.start()

    def tearDown(self):
        self.change_patch.stop(); self.data_patch.stop(); self.tmp.cleanup()

    def test_import_is_not_capped_at_500_rows(self):
        book = load_workbook(io.BytesIO(r.crm_excel_template()))
        sheet = book['客户收集表']
        for index in range(501):
            sheet.append([f'业主{index}', '业主单位', f'联系人{index}'])
        raw = io.BytesIO(); book.save(raw); book.close()
        result = r.crm_excel_import({'filename':'large.xlsx','data':base64.b64encode(raw.getvalue()).decode()}, 'alice')
        self.assertEqual(result['created'], 501, result)
        self.assertEqual(result['errors'], [], result)
