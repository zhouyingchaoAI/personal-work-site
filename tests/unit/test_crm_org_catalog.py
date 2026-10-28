import io
import unittest
from openpyxl import load_workbook
import backend.runtime as r

class OrgCatalogTests(unittest.TestCase):
    def test_template_has_shared_company_and_department_choices(self):
        b=load_workbook(io.BytesIO(r.crm_business_template()))
        # 我方组织目录改由「设置 → 我方组织」维护，模板只保留共享下拉选项。
        self.assertNotIn('我方组织目录',b.sheetnames)
        self.assertIn('下拉选项',b.sheetnames)
        headers=[c.value for c in b['下拉选项'][1]]
        for name in ('经营区域','重要程度','支持程度','单位类型'):
            self.assertIn(name,headers)
        dvs=list(b['经营收集表'].data_validations.dataValidation)
        self.assertTrue(dvs,'收集表应带下拉校验')
        self.assertTrue(any('CRM_单位类型' in str(v.formula1) and str(v.sqref).startswith('Q2') for v in dvs))
        # 区域 / 重要程度 / 支持程度同样走共享命名区域的下拉校验。
        for name in ('CRM_经营区域','CRM_重要程度','CRM_支持程度'):
            self.assertTrue(any(name in str(v.formula1) for v in dvs), name)
        b.close()
