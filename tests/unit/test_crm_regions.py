import base64
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
import backend.runtime as r

class RegionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.patches = [patch.object(r, 'USER_DATA_DIR', Path(self.tmp.name)), patch.object(r, 'crm_on_change')]
        for p in self.patches: p.start()
    def tearDown(self):
        for p in reversed(self.patches): p.stop()
        self.tmp.cleanup()
    def test_save_preserve_clear_invalid(self):
        u = r.crm_unit_save({'name':'示例业主','type':'业主单位','region':'上海','business_region':'东区'},'alice')['unit']
        self.assertEqual(u['region'],'上海')
        legacy = {k:v for k,v in u.items() if k != 'business_region'}
        u = r.crm_unit_save(legacy,'alice')['unit']
        self.assertEqual(u['business_region'],'东区')
        self.assertFalse(r.crm_unit_save({**u,'business_region':'华东'},'alice')['ok'])
        # 多区域模型：单字段留空但 business_regions 仍有值时保留；两者都清空才真正清除。
        kept = r.crm_unit_save({**u,'business_region':''},'alice')['unit']
        self.assertEqual(kept['business_region'],'东区')
        cleared = r.crm_unit_save({**kept,'business_region':'','business_regions':[]},'alice')['unit']
        self.assertEqual(cleared['business_region'],'')
        self.assertEqual(cleared['business_regions'],[])
    def workbook(self,area,legacy=False):
        book=load_workbook(io.BytesIO(r.crm_excel_template()))
        sheet=book['客户收集表']
        self.assertEqual(sheet.cell(1,9).value,'经营区域')
        if legacy: sheet.delete_cols(9)
        row=['示例业主','业主单位','张工','上海','','','','']
        sheet.append(row if legacy else row+[area])
        out=io.BytesIO();book.save(out);book.close()
        return {'filename':'客户.xlsx','data':base64.b64encode(out.getvalue()).decode()}
    def test_new_and_legacy_excel(self):
        result=r.crm_excel_import(self.workbook('广东区'),'alice')
        self.assertEqual(result['errors'],[])
        self.assertEqual(next(iter(r.crm_load('alice','units').values()))['business_region'],'广东区')
        result=r.crm_excel_import(self.workbook('',True),'bob')
        self.assertEqual(result['errors'],[])
        self.assertEqual(next(iter(r.crm_load('bob','units').values()))['business_region'],'')
    def test_invalid_or_conflicting_region_no_new_records(self):
        result=r.crm_excel_import(self.workbook('华东'),'alice')
        self.assertEqual(len(result['errors']),1)
        self.assertFalse(r.crm_load('alice','units'))
        r.crm_excel_import(self.workbook('东区'),'alice')
        # 同一单位再次以另一区域导入：并入多区域而不是报错，首个区域保持为主区域。
        result=r.crm_excel_import(self.workbook('西区'),'alice')
        self.assertEqual(result['errors'],[])
        unit=next(iter(r.crm_load('alice','units').values()))
        self.assertEqual(unit['business_region'],'东区')
        self.assertEqual(unit['business_regions'],['东区','西区'])
    def test_quick_capture_region(self):
        result=r.crm_quick_create({'unit_name':'示例单位','unit_type':'业主单位','name':'张工','business_region':'广东区'},'alice')
        self.assertTrue(result['ok'],result)
        self.assertEqual(next(iter(r.crm_load('alice','units').values()))['business_region'],'广东区')

    def test_discontinued_regions_are_rejected(self):
        for region in ('南区', '北区'):
            with self.subTest(region=region):
                result = r.crm_quick_create({'unit_name': f'{region}示例单位', 'unit_type': '业主单位', 'name': '张工', 'business_region': region}, 'alice')
                self.assertFalse(result['ok'])
