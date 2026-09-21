import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import backend.runtime as r


class RelationshipExportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.users = [
            {'username': 'root', 'role': 'superadmin'},
            {'username': 'manager', 'name': '经营管理员', 'role': 'member'},
            {'username': 'liaison', 'name': '李经理', 'employee_id': 'L001', 'role': 'member'},
            {'username': 'other', 'name': '王经理', 'employee_id': 'L002', 'role': 'member'},
            {'username': 'uploader', 'name': '上传人', 'role': 'member'},
        ]
        patches = [
            patch.object(r, 'USER_DATA_DIR', Path(self.tmp.name)),
            patch.object(r, 'read_config', return_value={'users': self.users}),
            patch.object(r, 'crm_on_change'),
        ]
        for item in patches:
            item.start()
            self.addCleanup(item.stop)
        r.crm_security_grant({'username': 'manager', 'role': 'disabled', 'manager': True}, self.users[0], mode='appoint')
        r.crm_store('uploader', 'units', {
            'east': {'id': 'east', 'name': '东区客户', 'type': '业主单位', 'business_regions': ['东区']},
            'west': {'id': 'west', 'name': '西区客户', 'type': '总包单位', 'business_regions': ['西区']},
        })
        r.crm_store('uploader', 'people', {
            'pe': {'id': 'pe', 'name': '东区联系人', 'unit_id': 'east', 'importance': '重要', 'support_level': '积极支持', 'title': '总经理'},
            'pw': {'id': 'pw', 'name': '西区联系人', 'unit_id': 'west', 'importance': '一般', 'support_level': '', 'title': '采购'},
        })
        r.crm_store('uploader', 'relations', {
            're': {'id': 're', 'person_id': 'pe', 'our_member': 'L001', 'version': 1},
            'rw': {'id': 'rw', 'person_id': 'pw', 'our_member': 'L002', 'version': 1},
        })

    def _rows(self, raw):
        book = load_workbook(io.BytesIO(raw))
        sheet = book.active
        return [[c.value for c in row] for row in sheet.iter_rows()]

    def test_only_manager_can_export(self):
        with self.assertRaises(PermissionError):
            r.crm_relationship_export(self.users[2], 'all', '')

    def test_delivered_export_has_attribution_and_exact_file_fingerprint(self):
        import hashlib
        raw, name = r.crm_relationship_export(self.users[1], 'all', '')
        delivered = r.crm_trace_export(raw, self.users[1], '/api/crm/excel/export', name, {'level': 'all'})
        book = load_workbook(io.BytesIO(delivered))
        self.assertIn('manager', book.active.oddFooter.center.text)
        self.assertEqual(book.properties.creator, 'manager')
        self.assertEqual(self._rows(raw), self._rows(delivered))
        book.close()
        event = r.crm_security_audit(self.users[0])['events'][0]
        self.assertEqual(event['operation'], 'export')
        self.assertEqual(event['grant']['sha256'], hashlib.sha256(delivered).hexdigest())

    def test_all_level_carries_importance_and_support(self):
        raw, name = r.crm_relationship_export(self.users[1], 'all', '')
        rows = self._rows(raw)
        header = rows[1]
        self.assertIn('重要程度', header)
        self.assertIn('支持程度', header)
        self.assertEqual(header[-4:], ['手机', '办公电话', '邮箱', '微信 / 其他联系'])
        body = rows[2:]
        self.assertEqual(len(body), 2)
        east = next(row for row in body if row[3] == '东区联系人')
        self.assertEqual(east[header.index('重要程度')], '重要')
        self.assertEqual(east[header.index('支持程度')], '积极支持')
        west = next(row for row in body if row[3] == '西区联系人')
        self.assertEqual(west[header.index('支持程度')], '待补充')
        self.assertIn('联合指挥', name)

    def test_export_separates_our_contact_organization_fields(self):
        """A contact name must not be concatenated with their organization."""
        r.crm_staff_save({
            'name': '李经理', 'member_key': 'L001', 'company': '联合指挥中心',
            'department': '东区作战平台', 'title': '泛交通BD',
        }, 'uploader')
        raw, _ = r.crm_relationship_export(self.users[1], 'all', '')
        rows = self._rows(raw)
        header = rows[1]
        self.assertEqual(
            header[header.index('我方对接人'):header.index('我方对接人') + 4],
            ['我方对接人', '我方公司', '我方部门', '我方岗位'],
        )
        east = next(row for row in rows[2:] if row[3] == '东区联系人')
        self.assertEqual(east[header.index('我方对接人')], '李经理')
        self.assertEqual(east[header.index('我方公司')], '联合指挥中心')
        self.assertEqual(east[header.index('我方部门')], '东区作战平台')
        self.assertEqual(east[header.index('我方岗位')], '泛交通BD')

    def test_customer_organization_export_has_hierarchy_and_contacts(self):
        raw, name = r.crm_customer_organization_export(self.users[1])
        book = load_workbook(io.BytesIO(raw))
        self.assertEqual(book.sheetnames, ['客户组织架构', '客户联系人'])
        organization = [[cell.value for cell in row] for row in book['客户组织架构'].iter_rows()]
        contacts = [[cell.value for cell in row] for row in book['客户联系人'].iter_rows()]
        self.assertIn('组织路径', organization[1])
        self.assertIn('联系人数量', organization[1])
        self.assertIn('客户单位层级', contacts[1])
        east = next(row for row in organization[2:] if row[1] == '东区客户')
        self.assertEqual(east[organization[1].index('经营区域')], '东区')
        self.assertEqual(east[organization[1].index('联系人数量')], 1)
        self.assertEqual(contacts[2][contacts[1].index('联系人')], '东区联系人')
        self.assertIn('客户组织架构', name)
        book.close()

    def test_region_level_filters_by_region(self):
        raw, name = r.crm_relationship_export(self.users[1], 'region', '东区')
        body = self._rows(raw)[2:]
        self.assertEqual([row[3] for row in body], ['东区联系人'])
        self.assertIn('东区', name)

    def test_region_level_rejects_unknown_region(self):
        with self.assertRaises(ValueError):
            r.crm_relationship_export(self.users[1], 'region', '南区')

    def test_manager_level_filters_by_account(self):
        raw, name = r.crm_relationship_export(self.users[1], 'manager', 'liaison')
        body = self._rows(raw)[2:]
        self.assertEqual([row[3] for row in body], ['东区联系人'])
        self.assertIn('李经理', name)

    def test_manager_export_can_be_narrowed_to_a_region(self):
        raw, name = r.crm_relationship_export(self.users[1], 'manager', 'liaison', '东区')
        body = self._rows(raw)[2:]
        self.assertEqual([row[3] for row in body], ['东区联系人'])
        self.assertIn('东区', name)


if __name__ == '__main__':
    unittest.main()
