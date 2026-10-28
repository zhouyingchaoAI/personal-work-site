import tempfile
import base64
import io
import unittest
from openpyxl import Workbook, load_workbook
from pathlib import Path
from unittest.mock import patch
import backend.runtime as r


class CrmSimplificationTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.patches = [patch.object(r, 'USER_DATA_DIR', Path(self.tmp.name)), patch.object(r, 'crm_on_change'),
                        patch.object(r, 'read_config', return_value={'users': [{'username': 'alice', 'role': 'superadmin'}]})]
        for p in self.patches: p.start()
        self.unit = r.crm_unit_save({'name': '同名客户', 'type': '业主单位'}, 'alice')['unit']
        self.other = r.crm_unit_save({'name': '同名客户分公司', 'type': '业主单位'}, 'alice')['unit']
        self.payload = {'code': 'OPP-TEST', 'name': '测试商机', 'owner_unit': '同名客户', 'owner_unit_id': self.unit['id'], 'sales_stage': '需求确认', 'outcome': 'open'}

    def tearDown(self):
        for p in reversed(self.patches): p.stop()
        self.tmp.cleanup()

    def test_opportunity_template_matches_full_source_layout(self):
        expected = [
            '机会点编码', '机会点名称', '业务日期', '业务组织', '业主单位', '机会点阶段', '数据状态',
            '业务状态', '作废状态', '项目归属', '投资规模(万元)', '可参与份额-总额(万元)',
            '可参与份额-集成产业(万元)', '可参与份额-服务1.0(含配套件)(万元)', '项目预计招标时间',
            '合同预计签订时间', '创建人', '创建日期', '修改人', '修改日期', '客户经理', '备注',
            '作废原因', '审核人', '审核日期', '总包or二代单位', '涉及产品', '赢率',
            '生成销售项目次数', '作废人', '作废日期', '品牌',
        ]
        book = load_workbook(io.BytesIO(r.crm_opportunity_template()), read_only=True, data_only=True)
        headers = [str(cell.value or '').strip() for cell in book['机会点导入'][1]]
        book.close()
        self.assertEqual(headers, expected)

    def test_explicit_customer_link_excludes_similar_name(self):
        saved = r.crm_opportunity_save(self.payload, 'alice')
        self.assertTrue(saved['ok'], saved)
        context = r.crm_opportunity_context('alice', 'OPP-TEST')
        self.assertEqual(context['unit_ids'], [self.unit['id']])
        self.assertEqual(saved['opportunity']['sales_stage'], '需求确认')

    def test_hidden_customer_and_invalid_stage_rejected(self):
        self.assertFalse(r.crm_opportunity_save({**self.payload, 'owner_unit_id': 'invisible'}, 'alice')['ok'])
        self.assertFalse(r.crm_opportunity_save({**self.payload, 'sales_stage': '随意阶段'}, 'alice')['ok'])

    def test_missing_explicit_id_does_not_fall_back_to_similar_name(self):
        r.crm_opportunity_save(self.payload, 'alice')
        context = r.crm_opportunity_context('alice', 'OPP-TEST', units={self.other['id']: self.other})
        self.assertEqual(context['unit_ids'], [])

    def test_battle_rejects_contact_outside_linked_customer(self):
        r.crm_opportunity_save(self.payload, 'alice')
        person = r.crm_person_save({'name': '其他客户联系人', 'unit_id': self.other['id']}, 'alice')['person']
        result = r.crm_battle_map_generate({'project_id': 'OPP-TEST', 'target_person_id': person['id']}, 'alice')
        self.assertFalse(result['ok'])

    def test_reimport_preserves_sales_fields_and_releases_changed_customer_link(self):
        r.crm_opportunity_save(self.payload, 'alice')
        def upload(owner):
            book = Workbook(); sheet = book.active
            sheet.append(['机会点编码', '机会点名称', '业主单位', '机会点阶段'])
            sheet.append(['OPP-TEST', '测试商机', owner, '招标'])
            buf = io.BytesIO(); book.save(buf); book.close()
            return r.crm_opportunity_import({'filename': 'test.xlsx', 'data': base64.b64encode(buf.getvalue()).decode()}, 'alice')
        self.assertTrue(upload('同名客户')['ok'])
        item = r.crm_load('alice', 'opportunities')['OPP-TEST']
        self.assertEqual(item['owner_unit_id'], self.unit['id'])
        self.assertEqual(item['sales_stage'], '需求确认')
        self.assertTrue(upload('新的客户')['ok'])
        self.assertEqual(r.crm_load('alice', 'opportunities')['OPP-TEST']['owner_unit_id'], '')

    def test_opportunity_import_previews_duplicates_without_writing(self):
        r.crm_opportunity_save(self.payload, 'alice')
        book = Workbook(); sheet = book.active
        sheet.append(['机会点编码', '机会点名称', '业主单位'])
        sheet.append(['OPP-TEST', '重复商机', '同名客户'])
        sheet.append(['OPP-NEW', '新增商机', '新客户'])
        buf = io.BytesIO(); book.save(buf); book.close()
        payload = {'filename': 'test.xlsx', 'data': base64.b64encode(buf.getvalue()).decode(), 'mode': 'preview'}
        result = r.crm_opportunity_import(payload, 'alice')
        self.assertEqual(result['duplicate_count'], 1)
        self.assertEqual(result['new_count'], 1)
        self.assertEqual(result['duplicates'][0]['code'], 'OPP-TEST')
        self.assertNotIn('OPP-NEW', r.crm_load('alice', 'opportunities'))

    def test_opportunity_import_skip_keeps_existing_and_imports_only_new(self):
        r.crm_opportunity_save(self.payload, 'alice')
        book = Workbook(); sheet = book.active
        sheet.append(['机会点编码', '机会点名称', '业主单位'])
        sheet.append(['OPP-TEST', '不应覆盖', '同名客户'])
        sheet.append(['OPP-NEW', '新增商机', '新客户'])
        buf = io.BytesIO(); book.save(buf); book.close()
        payload = {'filename': 'test.xlsx', 'data': base64.b64encode(buf.getvalue()).decode(), 'mode': 'skip'}
        result = r.crm_opportunity_import(payload, 'alice')
        stored = r.crm_load('alice', 'opportunities')
        self.assertEqual(result['imported'], 1)
        self.assertEqual(result['skipped'], 1)
        self.assertEqual(stored['OPP-TEST']['name'], '测试商机')
        self.assertEqual(stored['OPP-NEW']['name'], '新增商机')

    def test_reimport_keeps_confirmed_duplicate_archived(self):
        records = {
            'OPP-OLD': {'id': 'OPP-OLD', 'code': 'OPP-OLD', 'name': '重复商机', 'owner_unit': '客户甲',
                        'amount': 10, 'cancelled': True, 'outcome': 'cancelled', 'duplicate_of': 'OPP-KEEP',
                        'cancel_reason': '重复商机，已归并至 OPP-KEEP', 'version': 2}
        }
        r.crm_store('alice', 'opportunities', records)
        book = Workbook(); sheet = book.active
        sheet.append(['机会点编码', '机会点名称', '业主单位', '作废状态'])
        sheet.append(['OPP-OLD', '重复商机', '客户甲', '未作废'])
        buf = io.BytesIO(); book.save(buf); book.close()
        result = r.crm_opportunity_import({'filename': 'test.xlsx', 'data': base64.b64encode(buf.getvalue()).decode()}, 'alice')
        item = r.crm_load('alice', 'opportunities')['OPP-OLD']
        self.assertTrue(result['ok'])
        self.assertTrue(item['cancelled'])
        self.assertEqual(item['outcome'], 'cancelled')
        self.assertEqual(item['duplicate_of'], 'OPP-KEEP')

    def test_outcome_and_version_preserved(self):
        first = r.crm_opportunity_save(self.payload, 'alice')['opportunity']
        second = r.crm_opportunity_save({**first, 'outcome': 'won'}, 'alice')['opportunity']
        self.assertEqual(second['outcome'], 'won')
        self.assertTrue(second['outcome_changed_at'])
        self.assertFalse(second['cancelled'])
        self.assertFalse(r.crm_opportunity_save(first, 'alice')['ok'])
        cancelled = r.crm_opportunity_save({**second, 'outcome': 'cancelled'}, 'alice')['opportunity']
        self.assertTrue(cancelled['cancelled'])

    def test_export_includes_customer_link_and_sales_result(self):
        r.crm_opportunity_save({**self.payload, 'outcome': 'won'}, 'alice')
        raw, _ = r.crm_opportunity_export({'username': 'alice', 'role': 'superadmin'})
        book = load_workbook(io.BytesIO(raw)); rows = list(book.active.values); book.close()
        item = dict(zip(rows[0], rows[1]))
        self.assertEqual(item['关联客户ID'], self.unit['id'])
        self.assertEqual(item['销售阶段'], '需求确认')
        self.assertEqual(item['成交结果'], '赢单')

    def test_plan_creates_followup_without_increasing_relationship_score(self):
        person = r.crm_person_save({'name': '测试联系人', 'unit_id': self.unit['id']}, 'alice')['person']
        relation = r.crm_relation_save({'person_id': person['id'], 'our_member': 'alice'}, 'alice')['relation']
        score = relation['score']
        result = r.crm_action_add({'relation_id': relation['id'], 'record_kind': 'plan', 'next_step': '确认需求', 'due_date': '2026-10-01', 'owner': 'alice', 'project_id': 'OPP-TEST', 'effective': True, 'bidirectional': True}, 'alice')
        self.assertTrue(result['ok'])
        self.assertEqual(r.relations_get('alice', relation['id'])['score'], score)
        self.assertEqual(result['action']['happened_at'], '')
        self.assertEqual(len(r.crm_followups('alice')), 1)
