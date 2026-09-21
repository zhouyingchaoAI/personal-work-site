import io
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
import backend.runtime as r

class StaffTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.old = r.USER_DATA_DIR
        r.USER_DATA_DIR = Path(self.tmp.name)
    def tearDown(self):
        r.USER_DATA_DIR = self.old
        self.tmp.cleanup()
    def test_hierarchy_isolation_version_and_stable_key(self):
        boss = r.crm_staff_save({'name':'王经理','department':'营销中心'},'alice')['staff']
        worker = r.crm_staff_save({'name':'张工','manager_id':boss['id'],'department':'东区'},'alice')['staff']
        self.assertFalse(r.crm_staff_save(dict(boss,manager_id=worker['id']),'alice')['ok'])
        self.assertFalse(r.crm_staff_save({'name':'其他','manager_id':boss['id']},'bob')['ok'])
        self.assertEqual(r.crm_load('bob','staff'),{})
        updated = r.crm_staff_save(dict(worker,name='张工程师',member_key='新标识'),'alice')['staff']
        self.assertEqual(updated['member_key'],'张工')
        self.assertFalse(r.crm_staff_save(worker,'alice')['ok'])
        self.assertFalse(r.crm_staff_save({'name':'张工'},'alice')['ok'])
        self.assertTrue(r.crm_staff_save({'name':'张工','member_key':'E002'},'alice')['ok'])
    def test_template_opens_example_without_polluting_collection(self):
        book=load_workbook(io.BytesIO(r.crm_business_template()))
        self.assertEqual(book.active.title,'示例（不导入）')
        headers = [cell.value for cell in book.active[1]]
        self.assertEqual(book.active.cell(2, headers.index('单位*') + 1).value,'示例轨道集团')
        self.assertEqual(book.active.cell(2, headers.index('支持程度') + 1).value,'支持')
        self.assertIsNone(book['经营收集表']['B2'].value)
        book.close()

    def test_template_import_saves_separate_our_organization_columns(self):
        import base64
        book = load_workbook(io.BytesIO(r.crm_business_template()))
        sheet = book['经营收集表']
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(headers[-4:], ['我方姓名', '我方公司', '我方部门', '我方岗位'])
        row = {header: '' for header in headers}
        row.update({
            '区域': '东区', '公司全称': '测试客户有限公司', '单位*': '测试客户',
            '单位类型': '业主单位', '姓名*': '客户甲', '我方姓名': '李工',
            '我方公司': '联合指挥中心', '我方部门': '东区作战平台', '我方岗位': '泛交通BD',
        })
        sheet.append([row[header] for header in headers])
        raw = io.BytesIO(); book.save(raw); book.close()
        result = r.crm_excel_import({
            'filename': '经营收集表.xlsx',
            'data': base64.b64encode(raw.getvalue()).decode(),
        }, 'alice')
        self.assertEqual(result['errors'], [], result)
        staff = r.crm_member_profile('alice', '李工')
        self.assertEqual(staff['company'], '联合指挥中心')
        self.assertEqual(staff['department'], '东区作战平台')
        self.assertEqual(staff['title'], '泛交通BD')
    def test_retry_overwrites_conflicting_unit_type_only_on_confirmation(self):
        created = r.crm_excel_retry_rows({'rows':[{'row':8,'values':{'区域':'东区','单位*':'测试业主','姓名*':'客户甲','手机':'','单位类型':'业主单位','城市':'长沙'}}]},'alice')
        self.assertEqual(created['errors'],[])
        unit_id = next(u['id'] for u in r.crm_load('alice','units').values() if u['name']=='测试业主')
        self.assertEqual(r.crm_load('alice','units')[unit_id]['type'],'业主单位')
        row = {'row':8,'values':{'区域':'东区','单位*':'测试业主','姓名*':'客户甲','手机':'','单位类型':'厂商','城市':'长沙'}}
        blocked = r.crm_excel_retry_rows({'rows':[dict(row)]},'alice')
        self.assertEqual(blocked['created'],0)
        self.assertTrue(any('单位类型' in e['error'] for e in blocked['errors']))
        self.assertEqual(r.crm_load('alice','units')[unit_id]['type'],'业主单位')
        confirmed = r.crm_excel_retry_rows({'rows':[dict(row)],'overwrite':True},'alice')
        self.assertEqual(confirmed['errors'],[])
        self.assertEqual(confirmed['created'],1)
        self.assertEqual(r.crm_load('alice','units')[unit_id]['type'],'厂商')
        history = r.crm_record_versions('alice','units',unit_id)
        self.assertEqual(history[0]['record']['type'],'业主单位')
    def test_staff_access_uses_crm_permissions(self):
        self.assertEqual(r.crm_security_operation('/api/crm/staff','GET'),'read')
        self.assertEqual(r.crm_security_operation('/api/crm/staff/save','POST'),'write')

    def test_staff_snapshot_exposes_login_account_for_picker_search(self):
        old_config, old_capabilities = r.read_config, r.crm_capabilities
        try:
            r.read_config = lambda: {'users': [{'username': 'zhangsan', 'employee_id': 'E100'}]}
            r.crm_capabilities = lambda _user: {'data_scope': 'self'}
            r.crm_store('zhangsan', 'staff', {'staff-1': {
                'id': 'staff-1', 'name': '张三', 'member_key': 'E100', 'version': 1,
            }})
            staff = r.crm_read_snapshot({'username': 'zhangsan'})['staff']
            self.assertEqual(staff[0]['username'], 'zhangsan')
        finally:
            r.read_config, r.crm_capabilities = old_config, old_capabilities

    def test_capture_integrates_staff_relations_followups_and_battle(self):
        payload={'format':'business','unit_name':'测试业主','name':'客户甲','our_member':'E100','our_name':'李工','our_department':'东区营销部','our_position':'客户经理','weekly_plan':'确认需求','due_date':'2026-09-30'}
        result=r.crm_quick_create(payload,'alice')
        self.assertTrue(result['ok'],result)
        staff=next(iter(r.crm_load('alice','staff').values()))
        self.assertEqual(staff['department'],'东区营销部')
        rel=r.crm_relations('alice')['relations'][0]
        self.assertEqual(rel['our_staff']['id'],staff['id'])
        self.assertIn('客户经理',rel['our_member_label'])
        self.assertIn('东区营销部',r.crm_followups('alice')[0]['owner_label'])
        battle=r.crm_battle_map_generate({'project_id':'P','target_person_id':rel['person_id'],'our_members':['E100']},'alice')
        self.assertIn('李工',battle['battle_map']['who_goes'])
        excluded=r.crm_battle_map_generate({'project_id':'P','target_person_id':rel['person_id'],'our_members':[]},'alice')
        self.assertEqual(excluded['battle_map']['who_goes'],'待指派')
        r.crm_staff_save(dict(staff,title='区域负责人'),'alice')
        self.assertIn('区域负责人',r.crm_relations('alice')['relations'][0]['our_member_label'])
        conflict=r.crm_quick_create(dict(payload,name='不应新增',our_department='错误部门'),'alice')
        self.assertFalse(conflict['ok'])
        self.assertEqual(len(r.crm_load('alice','people')),1)

    def test_excel_new_staff_columns_and_old_template_compatible(self):
        import base64
        from openpyxl import Workbook
        for headers, name in [(r.CRM_BUSINESS_LEGACY_HEADERS,'客户一'),(r.CRM_BUSINESS_HEADERS,'客户二')]:
            data={'区域':'东区','单位*':'测试业主','部门':'技术部','职务':'主任','姓名*':name,'重要程度':'重要','支持程度':'支持','本周互动计划':'确认资料','单位类型':'业主单位','城市':'上海','维系周期（天）':'14','我方姓名':'王工'}
            if headers == r.CRM_BUSINESS_LEGACY_HEADERS:
                data.update({'我方对接人':'E200','我方公司':'测试公司','我方部门':'营销部','我方岗位':'售前工程师'})
            values=[data.get(header,'') for header in headers]
            b=Workbook();s=b.active;s.title='经营收集表';s.append(headers);s.append(values)
            raw=io.BytesIO();b.save(raw);b.close()
            result=r.crm_excel_import({'filename':'test.xlsx','data':base64.b64encode(raw.getvalue()).decode()},'alice')
            self.assertEqual(result['errors'],[])
        self.assertEqual(len(r.crm_load('alice','people')),2)
        self.assertEqual(r.crm_member_profile('alice','E200')['title'],'售前工程师')
