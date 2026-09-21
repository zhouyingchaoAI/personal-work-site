import base64
import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
import backend.runtime as runtime


class OrganizationExcelTests(unittest.TestCase):
    def test_pinyin_username_and_collision(self):
        self.assertEqual(runtime.organization_default_username('杨光耀', {}), 'yangguangyao')
        self.assertEqual(runtime.organization_default_username('杨光耀', {'yangguangyao', 'yangguangyao2'}), 'yangguangyao3')

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.config = Path(self.tmp.name) / 'config.json'
        self.config.write_text(json.dumps({'users': [
            {'username': 'root', 'role': 'superadmin', 'name': '管理员'},
            {'username': 'zhangsan', 'role': 'member', 'name': '张三', 'employee_id': '120001', 'company': '旧公司', 'department': '旧部门', 'position': '旧岗位'},
        ]}, ensure_ascii=False), encoding='utf-8')
        self.patch = patch.object(runtime, 'CONFIG_PATH', self.config)
        self.patch.start()
        self.directory_patch = patch.object(runtime, 'USER_DATA_DIR', Path(self.tmp.name))
        self.directory_patch.start()

    def tearDown(self):
        self.directory_patch.stop()
        self.patch.stop()
        self.tmp.cleanup()

    def test_original_position_sheet_and_unmatched_staff(self):
        book = Workbook(); sheet = book.active; sheet.title = 'sheet1'
        sheet.append(['姓名', '在职状态', '一级部门', '二级部门', '岗位', '成本中心', '职位族', '职位类', '职位子类', None])
        sheet.append(['张三', '在职', '服务中心', '交付部', '工程师', '新公司', '技术族', '支援类', '交付'])
        sheet.append(['未开户人员', '在职', '服务中心', '', '经理', '新公司'])
        sheet.append(['张三', '离职', '错误部门', '', '错误岗位', '错误公司'])
        raw = io.BytesIO(); book.save(raw); book.close()
        with patch.object(runtime, 'openclaw_sync_user'):
            result = runtime.organization_excel_import({'filename': '岗位信息.xlsx', 'data': base64.b64encode(raw.getvalue()).decode()}, 'root')
        user = runtime.read_config()['users'][1]
        self.assertEqual((user['company'], user['department'], user['position']), ('新公司', '服务中心 / 交付部', '工程师'))
        self.assertEqual(user['role'], 'member')
        self.assertEqual(result['updated'], 1)
        self.assertEqual(result['errors'], [])
        self.assertEqual(result['imported'], 3)
        self.assertEqual([e['row'] for e in result['notices']], [3, 4])
        self.assertIn('未开户人员', result['notices'][0]['error'])
        self.assertEqual(len(runtime.read_config()['organization_roster']), 3)
        exported = load_workbook(io.BytesIO(runtime.organization_excel_export()))
        self.assertEqual(exported['岗位名册'].max_row, 4)
        self.assertIn('技术族', list(exported['岗位名册'].values)[1])
        self.assertEqual(len(runtime.read_config()['users']), 2)

    def test_auto_create_is_opt_in_member_only_and_repeat_safe(self):
        book = Workbook(); sheet = book.active
        sheet.append(['姓名', '在职状态', '一级部门', '二级部门', '岗位', '成本中心'])
        sheet.append(['新同事', '在职', '服务中心', '交付部', '总经理', '公司'])
        sheet.append(['离职同事', '离职', '服务中心', '', '经理', '公司'])
        raw = io.BytesIO(); book.save(raw)
        payload = {'filename': '岗位.xlsx', 'data': base64.b64encode(raw.getvalue()).decode(), 'create_users': True}
        with patch.object(runtime, 'openclaw_sync_user'):
            first = runtime.organization_excel_import(payload, 'root')
            second = runtime.organization_excel_import(payload, 'root')
        self.assertEqual(first['created'], 1)
        self.assertEqual(second['created'], 0)
        self.assertEqual(first['pending'], 1)
        users = runtime.read_config()['users']
        self.assertEqual(len(users), 3)
        created = users[-1]
        self.assertEqual(created['role'], 'member')
        self.assertEqual(created['department'], '服务中心 / 交付部')
        self.assertTrue(runtime._is_hashed(created['password']))
        self.assertNotIn('password', first['created_users'][0])

    def test_import_updates_matching_employee_organization_without_changing_role(self):
        book = Workbook()
        sheet = book.active
        sheet.title = '组织架构'
        sheet.append(['工号', '姓名', '账号', '公司', '部门', '岗位', '系统角色'])
        sheet.append(['120001', '张三', 'zhangsan', '新公司', '服务中心 / 交付部', '交付工程师', 'member'])
        raw = io.BytesIO(); book.save(raw); book.close()
        payload = {'filename': '组织架构.xlsx', 'data': base64.b64encode(raw.getvalue()).decode()}
        with patch.object(runtime, 'openclaw_sync_user'):
            result = runtime.organization_excel_import(payload, 'root')
        updated = next(item for item in runtime.read_config()['users'] if item['username'] == 'zhangsan')
        self.assertEqual(result['updated'], 1)
        self.assertEqual(updated['company'], '新公司')
        self.assertEqual(updated['department'], '服务中心 / 交付部')
        self.assertEqual(updated['position'], '交付工程师')
        self.assertEqual(updated['role'], 'member')

    def test_import_uses_username_as_primary_organization_match(self):
        config = runtime.read_config()
        config['users'].append({'username': 'other', 'role': 'member', 'name': '李四', 'employee_id': '120002', 'company': '原公司'})
        runtime.write_config(config)
        book = Workbook(); sheet = book.active; sheet.title = '组织架构'
        sheet.append(['工号', '姓名', '账号', '公司', '部门', '岗位', '系统角色'])
        sheet.append(['120002', '张三', 'zhangsan', '用户名公司', '用户名部门', '用户名岗位', 'member'])
        raw = io.BytesIO(); book.save(raw); book.close()
        payload = {'filename': '组织架构.xlsx', 'data': base64.b64encode(raw.getvalue()).decode()}
        with patch.object(runtime, 'openclaw_sync_user'):
            runtime.organization_excel_import(payload, 'root')
        users = {item['username']: item for item in runtime.read_config()['users']}
        self.assertEqual(users['zhangsan']['company'], '旧公司')
        self.assertEqual(users['other']['company'], '原公司')

    def import_positions(self, rows):
        book = Workbook(); sheet = book.active
        sheet.append(['姓名', '在职状态', '一级部门', '二级部门', '岗位', '成本中心', '工号'])
        for row in rows:
            sheet.append(row)
        raw = io.BytesIO(); book.save(raw); book.close()
        with patch.object(runtime, 'openclaw_sync_user'):
            return runtime.organization_excel_import({'filename': '岗位.xlsx', 'data': base64.b64encode(raw.getvalue()).decode(), 'create_users': True}, 'root')

    def test_new_employee_binds_existing_name_and_repeat_is_safe(self):
        config = runtime.read_config()
        config['users'][1]['employee_id'] = ''
        runtime.write_config(config)
        base = ['张三', '在职', '服务中心', '交付部', '工程师', '公司']
        self.import_positions([base + ['']])
        first = self.import_positions([base + ['00123']])
        second = self.import_positions([base + ['00123']])
        config = runtime.read_config()
        self.assertEqual(config['users'][1]['employee_id'], '00123')
        self.assertEqual(len(config['users']), 2)
        self.assertEqual(len(config['organization_roster']), 1)
        self.assertEqual(first['updated'], 1)
        self.assertEqual(second['updated'], 0)

    def test_complete_roster_import_marks_employee_accounts_missing_from_latest_sheet(self):
        config = runtime.read_config()
        config['users'].append({
            'username': 'legacy', 'role': 'member', 'name': '历史成员',
            'employee_id': '990001', 'company': '旧公司', 'department': '旧部门', 'position': '旧岗位',
        })
        runtime.write_config(config)

        result = self.import_positions([
            ['张三', '在职', '服务中心', '交付部', '工程师', '公司', '120001'],
        ])

        users = {item['username']: item for item in runtime.read_config()['users']}
        self.assertEqual(users['zhangsan']['organization_check']['status'], 'present')
        self.assertEqual(users['legacy']['organization_check']['status'], 'missing')
        self.assertNotIn('organization_check', users['root'])
        self.assertEqual(result['roster_check']['missing'], 1)
        self.assertEqual(result['roster_check']['missing_users'][0]['username'], 'legacy')

    def test_later_complete_roster_import_clears_missing_marker(self):
        config = runtime.read_config()
        config['users'][1]['organization_check'] = {
            'status': 'missing', 'checked_at': '2026-09-01T08:00:00', 'source': '旧岗位表.xlsx',
        }
        runtime.write_config(config)

        self.import_positions([
            ['张三', '在职', '服务中心', '交付部', '工程师', '公司', '120001'],
        ])

        user = runtime.read_config()['users'][1]
        self.assertEqual(user['organization_check']['status'], 'present')
        self.assertNotEqual(user['organization_check']['checked_at'], '2026-09-01T08:00:00')
        self.assertEqual(user['organization_check']['source'], '岗位.xlsx')

    def test_new_user_with_employee_id_and_conflicts(self):
        result = self.import_positions([
            ['新同事', '在职', '部门', '', '工程师', '公司', '0007'],
            ['张三', '在职', '部门', '', '工程师', '公司', 'different'],
            ['冒名人员', '在职', '部门', '', '工程师', '公司', '120001'],
            ['重复甲', '在职', '部门', '', '工程师', '公司', '0099'],
            ['重复乙', '在职', '部门', '', '工程师', '公司', '0099'],
            ['离职人员', '离职', '部门', '', '工程师', '公司', '0088'],
        ])
        self.assertEqual(result['created'], 1)
        self.assertEqual(result['pending'], 5)
        users = runtime.read_config()['users']
        self.assertEqual(users[-1]['employee_id'], '0007')
        self.assertEqual(users[1]['employee_id'], '120001')
        self.assertEqual(users[1]['department'], '旧部门')

    def test_employee_cell_formats(self):
        sheet = Workbook().active
        sheet['A1'] = 12; sheet['A1'].number_format = '00000'
        self.assertEqual(runtime.organization_cell_text(sheet['A1']), '00012')
        sheet['A1'] = 0; sheet['A1'].number_format = 'General'
        self.assertEqual(runtime.organization_cell_text(sheet['A1']), '0')

    def test_employee_formula_uses_saved_result_only(self):
        book = Workbook(); sheet = book.active
        sheet.append(['姓名', '在职状态', '一级部门', '二级部门', '岗位', '成本中心', '工号'])
        sheet.append(['张三', '在职', '部门', '', '工程师', '公司', '=VLOOKUP(A2,[1]sheet1!A:E,5,0)'])
        raw = io.BytesIO(); book.save(raw); book.close()
        payload = {'filename': '岗位.xlsx', 'data': base64.b64encode(raw.getvalue()).decode()}
        with patch.object(runtime, 'openclaw_sync_user'):
            result = runtime.organization_excel_import(payload, 'root')
        self.assertIn('没有有效的已保存结果', result['errors'][0]['error'])
        saved = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as source, zipfile.ZipFile(saved, 'w') as output:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == 'xl/worksheets/sheet1.xml':
                    data = data.replace(b'</f><v></v>', b'</f><v>120001</v>')
                output.writestr(item, data)
        payload['data'] = base64.b64encode(saved.getvalue()).decode()
        with patch.object(runtime, 'openclaw_sync_user'):
            result = runtime.organization_excel_import(payload, 'root')
        self.assertEqual(result['errors'], [])
        self.assertEqual(result['updated'], 1)
        self.assertEqual(runtime.read_config()['users'][1]['employee_id'], '120001')
