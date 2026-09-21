import base64
import json
import shutil
import unittest
from pathlib import Path
from unittest.mock import patch

import backend.runtime as rt


class WritingAssistantTest(unittest.TestCase):
    def setUp(self):
        self.username = "unit_writing_user"
        self.root = rt.user_root(self.username)
        if self.root.exists():
            shutil.rmtree(self.root)

    def tearDown(self):
        if self.root.exists():
            shutil.rmtree(self.root)

    def _file_payload(self, name, text):
        encoded = base64.b64encode(text.encode("utf-8")).decode("ascii")
        return {"name": name, "data": "data:text/plain;base64," + encoded}

    def test_upload_extracts_text_and_lists_material(self):
        result = rt.writing_upload({"files": [self._file_payload("素材.txt", "这是上传素材")]} , self.username)
        self.assertTrue(result["ok"])
        self.assertEqual(len(result["files"]), 1)
        self.assertEqual(result["files"][0]["name"], "素材.txt")
        self.assertIn("上传素材", result["files"][0]["text_preview"])

    def test_agent_attachment_context_uses_the_uploaded_file_text(self):
        upload = rt.writing_upload({"files": [self._file_payload("会议纪要.txt", "项目计划在周五评审")]} , self.username)

        context = rt.writing_agent_attachment_context([upload["files"][0]["id"]], self.username)

        self.assertIn("会议纪要.txt", context)
        self.assertIn("项目计划在周五评审", context)

    def test_selected_personal_report_is_added_as_writing_reference(self):
        report = self.root / "reports" / "weekly" / "工作周报-2026.8.14.xlsx"
        report.parent.mkdir(parents=True, exist_ok=True)
        report.write_bytes(b"placeholder")
        with patch.object(rt, "all_files", return_value=[{
            "kind": "weekly", "name": report.name, "path": str(report),
        }]), patch.object(rt, "_kb_extract_text", return_value="历史周报中的项目进展"):
            materials = rt._writing_collect_materials({
                "report_refs": [{"kind": "weekly", "name": report.name}],
            }, self.username)

        self.assertEqual(len(materials), 1)
        self.assertIn("参考周报", materials[0]["name"])
        self.assertEqual(materials[0]["text"], "历史周报中的项目进展")

    def test_weekly_compose_falls_back_when_model_returns_malformed_json(self):
        with patch.object(rt, "assistant_settings", lambda: {"url": "http://model", "key": "key", "model": "test"}), patch.object(
            rt, "request_json", lambda *args, **kwargs: {"choices": [{"message": {"content": '{"weekly_summary":[{"content":"未闭合}'}}]}
        ):
            result = rt.compose_weekly_skill({"raw_work": "完成平台联调并修复超时问题"})

        self.assertTrue(result["ok"])
        self.assertEqual(result["mode"], "local_fallback")
        self.assertTrue(result["draft"]["weekly_summary"])

    def test_generate_uses_uploaded_files_and_knowledge_sources(self):
        upload = rt.writing_upload({"files": [self._file_payload("制度.txt", "素材里有项目背景")]} , self.username)
        captured = {}

        def fake_kb_search(payload, username):
            captured["payload"] = payload
            captured["username"] = username
            return {"ok": True, "results": [{"doc_name": "知识库文档", "text": "知识库里的参考内容", "score": 0.91}]}

        with patch.object(rt, "kb_search", fake_kb_search), patch.object(rt, "assistant_settings", lambda: {"url": "", "key": "", "model": "local"}):
            result = rt.writing_generate({
                "title": "写一份方案",
                "writing_type": "方案文档",
                "requirements": "正式、结构完整",
                "file_ids": [upload["files"][0]["id"]],
                "use_knowledge": True,
                "kb_ids": ["kb1"],
            }, self.username)

        self.assertTrue(result["ok"])
        self.assertIn("写一份方案", result["item"]["title"])
        self.assertIn("素材里有项目背景", result["item"]["content"])
        self.assertEqual(result["item"]["sources"][0]["doc_name"], "知识库文档")
        self.assertEqual(captured["payload"]["kb_ids"], ["kb1"])
        self.assertEqual(captured["username"], self.username)

    def test_multi_file_summary_prompt_includes_every_file_with_a_bounded_context(self):
        materials = [
            {"id": f"mat{index}", "name": f"材料{index}.txt", "text": f"文件{index}内容" + "x" * 12000}
            for index in range(5)
        ]

        prompt = rt._writing_prompt({"title": "本周总结", "writing_type": "工作总结"}, materials, [])

        for index in range(5):
            self.assertIn(f"【素材：材料{index}.txt】", prompt)
            self.assertIn(f"文件{index}内容", prompt)
        material_section = prompt.split("上传素材：\n", 1)[1].split("\n\n知识库参考：", 1)[0]
        self.assertLessEqual(len(material_section), rt.WRITING_MAX_MATERIAL_TOTAL_CHARS + 300)

    def test_rewrite_prompt_includes_the_latest_user_instruction(self):
        prompt = rt._writing_prompt(
            {"title": "本周工作", "rewrite_instruction": "内容太简单，请扩写但不要加入下周计划"},
            [],
            [],
            "rewrite",
        )

        self.assertIn("内容太简单，请扩写但不要加入下周计划", prompt)

    def test_templates_history_and_txt_export_are_persisted(self):
        template = rt.writing_save_template({"name": "我的模板", "writing_type": "汇报材料", "prompt": "按三段式写"}, self.username)
        self.assertTrue(template["ok"])
        templates = rt.writing_templates(self.username)
        self.assertTrue(any(item["name"] == "我的模板" for item in templates["templates"]))

        result = rt.writing_generate({"title": "工作总结", "writing_type": "工作总结", "requirements": "简洁"}, self.username)
        history = rt.writing_history(self.username)
        self.assertEqual(history["items"][0]["id"], result["item"]["id"])

        export = rt.writing_export({"id": result["item"]["id"], "format": "txt"}, self.username)
        self.assertTrue(export["ok"])
        path = Path(export["path"])
        self.assertTrue(path.exists())
        self.assertIn("工作总结", path.read_text(encoding="utf-8"))

    def test_docx_export_converts_markdown_table_to_native_word_table(self):
        from docx import Document

        doc = Document()
        rt._writing_markdown_to_docx(doc, """三、归档后续行动计划

| 序号 | 事项内容 | 责任人 | 完成时间 |。
| --- | --- | --- | --- |。
| 1 | 完成 ASC 项目归档 | 余鸿慧 | 8月31日 |。
| 2 | 整理资料分类清单 | 田凯 | 9月30日 |。
""")

        self.assertEqual(len(doc.tables), 1)
        table = doc.tables[0]
        self.assertEqual(table.cell(0, 0).text, "序号")
        self.assertEqual(table.cell(1, 1).text, "完成 ASC 项目归档")

    def test_docx_export_sets_explicit_chinese_and_latin_fonts(self):
        from docx import Document
        from docx.oxml.ns import qn

        doc = Document()
        rt._writing_markdown_to_docx(doc, "会议纪要 Meeting minutes")

        run = doc.paragraphs[0].runs[0]
        self.assertEqual(run._element.rPr.rFonts.get(qn("w:eastAsia")), "微软雅黑")
        self.assertEqual(run._element.rPr.rFonts.get(qn("w:ascii")), "Arial")
        normal_fonts = doc.styles["Normal"]._element.rPr.rFonts
        self.assertEqual(normal_fonts.get(qn("w:eastAsia")), "微软雅黑")
        self.assertEqual(normal_fonts.get(qn("w:ascii")), "Arial")

    def test_xlsx_export_converts_markdown_table_to_native_sheet(self):
        from openpyxl import load_workbook

        export = rt.writing_export({
            "title": "归档行动计划",
            "format": "xlsx",
            "content": """三、归档后续行动计划

| 序号 | 事项内容 | 责任人 | 完成时间 |。
| --- | --- | --- | --- |。
| 1 | 完成 ASC 项目归档 | 余鸿慧 | 8月31日 |。
| 2 | 整理资料分类清单 | 田凯 | 9月30日 |。
""",
        }, self.username)

        self.assertEqual(export["format"], "xlsx")
        workbook = load_workbook(export["path"])
        sheet = workbook["表格1"]
        self.assertEqual(sheet["A1"].value, "序号")
        self.assertEqual(sheet["B2"].value, "完成 ASC 项目归档")

    def test_xlsx_export_requires_a_markdown_table(self):
        with self.assertRaisesRegex(ValueError, "未识别到可导出的表格"):
            rt.writing_export({"title": "工作总结", "format": "xlsx", "content": "这是一段普通正文。"}, self.username)


if __name__ == "__main__":
    unittest.main()
